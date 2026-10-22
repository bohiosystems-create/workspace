"""Function sweep (synthetic): the Excel functions and edge cases real estate and fund models use, saved as Excel saves them (original
formula text, _xlfn. prefixes, each result stored). Expected values were checked by hand against Excel; the rest come from a LibreOffice
recalculation (pass its values as a JSON file). Usage: python3 gen_fnsweep.py out.xlsx [lo_values.json]
Function sweep: Excel functions and edge cases real estate and fund models use. Sheet F: label, formula, Excel's expected value
(hand-checked, where certain). LibreOffice and the app's engine are then compared against it."""
import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName

wb = Workbook()
d = wb.active; d.title = 'Data'
# A: periods 1..12, B: cash flows, C: dates (month ends), D: category, E: rent, F: text numbers
cf = [-1000, 50, 60, 70, 80, 90, 100, 110, 120, 130, 140, 1200]
cats = ['Retail', 'Office', 'Retail', 'Hotel', 'Office', 'Retail', 'Hotel', 'Office', 'Retail', 'Hotel', 'Office', 'Retail']
for i in range(12):
    d.cell(i + 2, 1, i + 1)
    d.cell(i + 2, 2, cf[i])
    d.cell(i + 2, 3, dt.date(2026, 1, 31) + dt.timedelta(days=0)) if i == 0 else None
    d.cell(i + 2, 4, cats[i])
    d.cell(i + 2, 5, 100 + 7 * i)
for i in range(12):
    d.cell(i + 2, 3).value = f'=EOMONTH(DATE(2026,1,31),{i})' if i else dt.date(2026, 1, 31)
    d.cell(i + 2, 3).number_format = 'yyyy-mm-dd'
d['A1'], d['B1'], d['C1'], d['D1'], d['E1'] = 'Period', 'CF', 'Date', 'Cat', 'Rent'
d['G1'] = 'Rate'; d['G2'] = 0.08
d['H1'] = 'Len'; d['H2'] = 6
d['I1'] = 'Txt'; d['I2'] = '5'; d['I3'] = '=""'; d['I4'] = 'abc'
d['J2'] = True
wb.defined_names['Rate'] = DefinedName('Rate', attr_text='Data!$G$2')
wb.defined_names['CFs'] = DefinedName('CFs', attr_text='Data!$B$2:$B$13')
# a second and third sheet for 3D references
for n, v in (('S1', 10), ('S2', 20), ('S3', 30)):
    s = wb.create_sheet(n); s['A1'] = v
# an Excel table (structured references)
t = wb.create_sheet('Units')
t.append(['Unit', 'Area', 'Price'])
for r in (('A', 100, 1.5), ('B', 120, 1.6), ('C', 90, 1.4)):
    t.append(list(r))
tab = Table(displayName='UnitTbl', ref='A1:C4'); t.add_table(tab)

f = wb.create_sheet('F')
f.append(['Label', 'Formula', 'Excel expected'])
R = []
def add(label, formula, expect=None):
    R.append((label, formula, expect))

D = 'Data!'
# --- financial
add('NPV', f'=NPV(0.08,{D}B3:B13)+{D}B2')
add('IRR', f'=IRR({D}B2:B13)')
add('IRR guess', f'=IRR({D}B2:B13,0.2)')
add('XNPV', f'=XNPV(0.08,{D}B2:B13,{D}C2:C13)')
add('XIRR', f'=XIRR({D}B2:B13,{D}C2:C13)')
add('MIRR', f'=MIRR({D}B2:B13,0.06,0.1)')
add('PMT', '=PMT(0.07/12,240,-1000000)')
add('IPMT', '=IPMT(0.07/12,5,240,-1000000)')
add('PPMT', '=PPMT(0.07/12,5,240,-1000000)')
add('PV', '=PV(0.06,10,-50000,0)')
add('FV', '=FV(0.05/12,120,-1000)')
add('NPER', '=NPER(0.06/12,-2000,200000)')
add('RATE', '=RATE(240,-7752.99,1000000)')
add('CUMIPMT', '=CUMIPMT(0.07/12,240,1000000,1,12,0)')
add('CUMPRINC', '=CUMPRINC(0.07/12,240,1000000,1,12,0)')
add('EFFECT', '=EFFECT(0.06,12)')
add('NOMINAL', '=NOMINAL(0.0617,12)')
add('SLN', '=SLN(100000,10000,10)')
add('DB', '=DB(100000,10000,10,1)')
add('DDB', '=DDB(100000,10000,10,1)')
add('Named range NPV', '=NPV(Rate,CFs)')
# --- dates
add('DATE', '=DATE(2026,2,30)', 46083)
add('EDATE end-month', '=EDATE(DATE(2026,1,31),1)', 46081)
add('EOMONTH negative', '=EOMONTH(DATE(2026,3,15),-1)', 46081)
add('YEARFRAC basis 0', '=YEARFRAC(DATE(2026,1,31),DATE(2026,3,31),0)')
add('YEARFRAC basis 1', '=YEARFRAC(DATE(2026,1,1),DATE(2027,7,1),1)')
add('YEARFRAC basis 3', '=YEARFRAC(DATE(2026,1,1),DATE(2027,7,1),3)')
add('DAYS360', '=DAYS360(DATE(2026,1,31),DATE(2026,3,31))', 60)
add('DAYS', '=DAYS(DATE(2026,12,31),DATE(2026,1,1))', 364)
add('DATEDIF m', '=DATEDIF(DATE(2026,1,31),DATE(2026,7,30),"m")', 5)
add('DATEDIF y', '=DATEDIF(DATE(2020,2,29),DATE(2026,2,28),"y")', 5)
add('NETWORKDAYS', '=NETWORKDAYS(DATE(2026,1,1),DATE(2026,1,31))', 22)
add('WORKDAY', '=WORKDAY(DATE(2026,1,1),10)', 46037)
add('WEEKDAY 2', '=WEEKDAY(DATE(2026,1,1),2)', 4)
add('YEAR of text date', '=YEAR("2026-03-15")', 2026)
add('1900 leap serial', '=DATE(1900,3,1)', 61)
add('Month count from dates', f'=(YEAR({D}C13)-YEAR({D}C2))*12+MONTH({D}C13)-MONTH({D}C2)', 11)
# --- lookup
add('INDEX/MATCH exact', f'=INDEX({D}E2:E13,MATCH("Hotel",{D}D2:D13,0))', 121)
add('MATCH approx', f'=MATCH(105,{D}E2:E13,1)', 1)
add('MATCH descending', '=MATCH(25,{50,40,30,20,10},-1)', 3)
add('VLOOKUP exact', f'=VLOOKUP(5,{D}A2:E13,5,FALSE)', 128)
add('VLOOKUP approx', f'=VLOOKUP(5.5,{D}A2:E13,2,TRUE)', 80)
add('HLOOKUP', '=HLOOKUP("b",{"a","b";1,2},2,FALSE)', 2)
add('LOOKUP', f'=LOOKUP(2,1/({D}D2:D13="Office"),{D}E2:E13)', 170)
add('OFFSET variable width', f'=SUM(OFFSET({D}B2,0,0,{D}H2,1))', -650)
add('OFFSET in NPV', f'=NPV(0.08,OFFSET({D}B3,0,0,{D}H2-1,1))')
add('INDIRECT text', f'=INDIRECT("Data!B"&(1+{D}H2))', 90)
add('INDIRECT sheet name', '=INDIRECT("\'S2\'!A1")', 20)
add('CHOOSE', '=CHOOSE(2,"Base","Upside","Downside")', 'Upside')
add('XLOOKUP', f'=XLOOKUP("Hotel",{D}D2:D13,{D}E2:E13)', 121)
add('XLOOKUP not found', f'=XLOOKUP("Spa",{D}D2:D13,{D}E2:E13,"none")', 'none')
add('XMATCH', f'=XMATCH("Office",{D}D2:D13)', 2)
add('ROWS/COLUMNS', f'=ROWS({D}A2:E13)*COLUMNS({D}A2:E13)', 60)
add('Structured ref sum', '=SUM(UnitTbl[Area])', 310)
add('Structured ref sumproduct', '=SUMPRODUCT(UnitTbl[Area],UnitTbl[Price])', 468)
add('3D sum', "=SUM(S1:S3!A1)", 60)
add('Whole column sum', f'=SUM({D}E:E)', 1662)
# --- conditional aggregation
add('SUMIF', f'=SUMIF({D}D2:D13,"Retail",{D}B2:B13)', 470)
add('SUMIFS two', f'=SUMIFS({D}B2:B13,{D}D2:D13,"Office",{D}A2:A13,">4")', 330)
add('SUMIFS wildcard', f'=SUMIFS({D}E2:E13,{D}D2:D13,"Ho*")', 426)
add('COUNTIFS', f'=COUNTIFS({D}D2:D13,"<>Retail",{D}E2:E13,">=120")', 6)
add('AVERAGEIFS', f'=AVERAGEIFS({D}E2:E13,{D}D2:D13,"Retail")', 136.4)
add('MAXIFS', f'=MAXIFS({D}E2:E13,{D}D2:D13,"Office")', 170)
add('MINIFS', f'=MINIFS({D}E2:E13,{D}D2:D13,"Hotel")', 121)
add('SUMPRODUCT condition', f'=SUMPRODUCT(({D}D2:D13="Retail")*{D}E2:E13)', 682)
add('SUMPRODUCT --', f'=SUMPRODUCT(--({D}A2:A13>6),{D}B2:B13)', 1800)
add('COUNTA', f'=COUNTA({D}D2:D13)', 12)
add('COUNTBLANK', f'=COUNTBLANK({D}I2:I5)', 2)
# --- math and rounding
add('ROUND half', '=ROUND(2.5,0)', 3)
add('ROUND negative half', '=ROUND(-2.5,0)', -3)
add('ROUND float', '=ROUND(1.005,2)', 1.01)
add('ROUNDUP', '=ROUNDUP(3.14159,3)', 3.142)
add('ROUNDDOWN', '=ROUNDDOWN(-3.14159,2)', -3.14)
add('MROUND', '=MROUND(17,5)', 15)
add('CEILING', '=CEILING(4.3,0.5)', 4.5)
add('CEILING.MATH negative', '=CEILING.MATH(-4.3,1)', -4)
add('FLOOR.MATH', '=FLOOR.MATH(4.7,2)', 4)
add('INT negative', '=INT(-1.25)', -2)
add('TRUNC', '=TRUNC(-1.25)', -1)
add('MOD negative', '=MOD(-7,3)', 2)
add('QUOTIENT', '=QUOTIENT(-7,3)', -2)
add('POWER fraction', '=POWER(1.08,1/12)')
add('SQRT', '=SQRT(2)')
add('LN', '=LN(10)')
add('PRODUCT growth', '=PRODUCT(1.03,1.03,1.03)', 1.092727)
# --- logic
add('IFS', '=IFS(5>6,"a",5>4,"b",TRUE,"c")', 'b')
add('SWITCH', '=SWITCH(2,1,"one",2,"two","other")', 'two')
add('IFERROR div0', '=IFERROR(1/0,0)', 0)
add('IFNA', '=IFNA(MATCH("x",{"a","b"},0),-1)', -1)
add('XOR', '=XOR(TRUE,FALSE)', True)
add('AND with range', f'=AND({D}E2:E13>0)', True)
add('Empty cell arithmetic', f'={D}I5+1', 1)
add('Empty string arithmetic', f'={D}I3+1', '#VALUE!')
add('Text number arithmetic', f'={D}I2+1', 6)
add('Boolean arithmetic', f'={D}J2+1', 2)
add('Text compare case', '=IF("abc"="ABC",1,0)', 1)
add('Error propagation in SUM', f'=IFERROR(SUM({D}B2,1/0),-1)', -1)
# --- text
add('TEXT percent', '=TEXT(0.1234,"0.0%")', '12.3%')
add('TEXT thousands', '=TEXT(1234567.891,"#,##0.00")', '1,234,567.89')
add('TEXT date', '=TEXT(DATE(2026,3,1),"mmm-yy")', 'Mar-26')
add('TEXT millions', '=TEXT(1234567,"#,##0.0,,\\M")', '1.2M')
add('CONCAT', '=CONCAT("Case ","Base")', 'Case Base')
add('TEXTJOIN', '=TEXTJOIN(", ",TRUE,"a","","b")', 'a, b')
add('ampersand number', '="IRR "&ROUND(0.12345*100,1)&"%"', 'IRR 12.3%')
add('VALUE', '=VALUE("1,234.5")', 1234.5)
add('NUMBERVALUE', '=NUMBERVALUE("1.234,5",",",".")', 1234.5)
add('SUBSTITUTE', '=SUBSTITUTE("SAR 1,000",",","")', 'SAR 1000')
add('FIND/MID', '=MID("Riyadh-North",FIND("-","Riyadh-North")+1,5)', 'North')
add('PROPER', '=PROPER("al waha park")', 'Al Waha Park')
add('FIXED', '=FIXED(1234.567,1)', '1,234.6')
add('LEFT number', '=LEFT(2026,2)', '20')
# --- statistics (Monte Carlo, comps)
add('MEDIAN', f'=MEDIAN({D}E2:E13)', 138.5)
add('STDEV.S', f'=STDEV.S({D}E2:E13)')
add('PERCENTILE.INC', f'=PERCENTILE.INC({D}E2:E13,0.9)')
add('QUARTILE', f'=QUARTILE({D}E2:E13,1)')
add('LARGE', f'=LARGE({D}E2:E13,2)', 170)
add('RANK.EQ', f'=RANK.EQ(121,{D}E2:E13)', 9)
add('CORREL', f'=CORREL({D}A2:A13,{D}E2:E13)', 1)
add('FORECAST.LINEAR', f'=FORECAST.LINEAR(13,{D}E2:E13,{D}A2:A13)', 184)
add('NORM.INV', '=NORM.INV(0.95,0,1)')
add('NORM.S.DIST', '=NORM.S.DIST(1.645,TRUE)')
add('NORM.DIST', '=NORM.DIST(1,0,1,TRUE)')
add('LOGNORM.INV', '=LOGNORM.INV(0.5,0,1)', 1)
# --- array and new
add('MMULT', '=SUMPRODUCT(MMULT({1,2;3,4},{1;1}))', 10)
add('FREQUENCY first', f'=INDEX(FREQUENCY({D}E2:E13,{{120,150}}),1)', 3)
add('LET', '=LET(x,2,y,3,x*y+1)', 7)
add('SUM of SEQUENCE', '=SUM(SEQUENCE(10))', 55)
add('SUM of FILTER', f'=SUM(FILTER({D}E2:E13,{D}D2:D13="Hotel"))', 426)
add('COUNT of UNIQUE', f'=COUNTA(UNIQUE({D}D2:D13))', 3)
add('MAX of SORT', f'=INDEX(SORT({D}E2:E13,1,-1),1)', 177)
add('LAMBDA', '=LAMBDA(x,x*2)(21)', 42)
# --- Excel 365 regular expressions
add('REGEXTEST', '=REGEXTEST("Unit A-102","^Unit [A-Z]-\\d{3}$")', True)
add('REGEXEXTRACT', '=REGEXEXTRACT("Lease 2026-03-15 signed","\\d{4}-\\d{2}-\\d{2}")', '2026-03-15')
add('REGEXEXTRACT group', '=REGEXEXTRACT("Tenant: Northstar LLC","Tenant: (\\w+)",2)', 'Northstar')
add('REGEXREPLACE', '=REGEXREPLACE("SAR 1,250,000","[^0-9]","")', '1250000')
add('REGEXREPLACE nth', '=REGEXREPLACE("a-b-c","-","+",2)', 'a-b+c')

import re
XLFN = ['REGEXTEST','REGEXEXTRACT','REGEXREPLACE','MAXIFS','MINIFS','IFS','SWITCH','XOR','CONCAT','TEXTJOIN','XLOOKUP','XMATCH','LET','SEQUENCE','UNIQUE','LAMBDA','DAYS','STDEV.S',
        'PERCENTILE.INC','RANK.EQ','FORECAST.LINEAR','NORM.INV','NORM.S.DIST','NORM.DIST','LOGNORM.INV','CEILING.MATH','FLOOR.MATH','NUMBERVALUE','IFNA']
XLWS = ['FILTER','SORT']
def excelise(f):
    for n in XLWS:
        f = re.sub(r'(?<![A-Z._])'+re.escape(n)+r'\(', '_xlfn._xlws.'+n+'(', f)
    for n in XLFN:
        f = re.sub(r'(?<![A-Z._])'+re.escape(n)+r'\(', '_xlfn.'+n+'(', f)
    return f
R = [(l, excelise(fm), e) for l, fm, e in R]
for i, (label, formula, expect) in enumerate(R, start=2):
    f.cell(i, 1, label)
    f.cell(i, 2, formula)
    if expect is not None:
        f.cell(i, 3, expect)
import json, zipfile, tempfile, os, re as _re, datetime as _dt
LO = json.load(open(sys.argv[2])) if len(sys.argv) > 2 else {}
tmp = tempfile.mktemp(suffix='.xlsx'); wb.save(tmp)
def xv(v):
    if isinstance(v, bool): return 'b', '1' if v else '0'
    if isinstance(v, (int, float)): return None, repr(float(v)) if isinstance(v, float) else str(v)
    if isinstance(v, str) and v.startswith('#'): return 'e', v
    return 'str', str(v).replace('&', '&amp;').replace('<', '&lt;')
serial = lambda d: (d - _dt.date(1899, 12, 30)).days
import calendar
eom = [serial(_dt.date(2026 + (i // 12), (i % 12) + 1, calendar.monthrange(2026 + (i // 12), (i % 12) + 1)[1])) for i in range(12)]
names = {n: i for i, n in enumerate(wb.sheetnames, start=1)}
zin = zipfile.ZipFile(tmp); zout = zipfile.ZipFile(sys.argv[1], 'w', zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    data = zin.read(it.filename)
    def fill(x, ref, v):
        t, txt = xv(v)
        def rep(m):
            head = m.group(1)
            head = _re.sub(r'\st="[^"]*"', '', head)
            if t: head = head[:-1] + ' t="%s">' % t
            return head + m.group(2) + '<v>%s</v>' % txt
        return _re.sub(r'(<c r="%s"[^>]*>)(<f>[^<]*</f>)(?:<v>[^<]*</v>|<v\s*/>)?' % ref, rep, x, 1)
    if it.filename == 'xl/worksheets/sheet%d.xml' % names['F']:
        x = data.decode()
        for i, (label, formula, expect) in enumerate(R, start=2):
            v = expect if expect is not None else LO.get(str(i))
            if v is None: continue
            x = fill(x, 'B%d' % i, v)
        data = x.encode()
    if it.filename == 'xl/worksheets/sheet%d.xml' % names['Data']:
        x = data.decode()
        for i in range(1, 12): x = fill(x, 'C%d' % (i + 2), eom[i])
        x = fill(x, 'I3', '')
        data = x.encode()
    zout.writestr(it, data)
zout.close(); os.remove(tmp)
print(len(R), 'cases')
