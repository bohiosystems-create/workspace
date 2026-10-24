# Round 44 (synthetic): two models for a CIO day of highly complex tasks. python3 -I gen/gen_c44.py → models/c44_*.xlsx, golden/c44.json
#  c44_phased_sar   a phased mixed-use development in SAR '000 on a 60-month timeline: a scenario switch (Base / Downside / Upside read by
#                   INDEX/MATCH), a hidden Drivers sheet (S-curve, sales phasing, lease-up), a cash-sweep construction loan with capitalised
#                   interest, an annual roll-up by SUMIF on a year key, returns annualised from monthly IRRs over an OFFSET-defined name,
#                   a hold input in months. One planted fault: the last month's interest typed as 0 over its formula row.
#  c44_fund_usd     a three-deal value-add fund in USD: deals in SAR and AED on quarterly dates (EDATE), converted on an FX sheet, consolidated
#                   with INDIRECT from a list of sheet names, fees on committed capital, a waterfall (8% pref compounding quarterly, 100% GP
#                   catch-up to 20%, then 80/20), LP net IRR by XIRR, a Checks sheet. One planted fault: the Jeddah deal's sale reads the
#                   Riyadh deal's exit yield.
# Every golden figure is computed here independently of any spreadsheet engine.
import openpyxl, os, json, math, datetime, calendar
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
D = os.path.join(os.path.dirname(__file__), '..', 'models'); GD = os.path.join(os.path.dirname(__file__), '..', 'golden'); G = {}

def irr(c, lo=-0.99, hi=10.0):
    f = lambda r: sum(x / (1 + r) ** t for t, x in enumerate(c))
    if f(lo) * f(hi) > 0: return None
    for _ in range(400):
        m = (lo + hi) / 2
        if f(lo) * f(m) <= 0: hi = m
        else: lo = m
    return (lo + hi) / 2

def xirr(c, d):
    d0 = d[0]; f = lambda r: sum(x / (1 + r) ** ((dd - d0).days / 365) for x, dd in zip(c, d))
    lo, hi = -0.99, 10.0
    if f(lo) * f(hi) > 0: return None
    for _ in range(400):
        m = (lo + hi) / 2
        if f(lo) * f(m) <= 0: hi = m
        else: lo = m
    return (lo + hi) / 2

# ---------------------------------------------------------------- 1. the phased development (SAR '000, 60 months)
N = 60
def phased(case='Base', price=None, rent=None, cost=None, ey=None, hold=60, ltc=0.55, rate=0.07, land=120000, nsa=30000, nla=8000, gfa=50000,
           cstart=1, cmonths=24, launch=6, dep=0.3, speriod=24, lstart=25, lmonths=12, sc=0.02, months=N):
    SC = {'Base': (12, 1.2, 4.5, 0.07), 'Downside': (10.8, 1.0, 4.9, 0.0775), 'Upside': (13, 1.3, 4.3, 0.065)}[case]
    price = SC[0] if price is None else price; rent = SC[1] if rent is None else rent; cost = SC[2] if cost is None else cost; ey = SC[3] if ey is None else ey
    w = [math.sin(math.pi * (m - cstart + 0.5) / cmonths) if cstart <= m < cstart + cmonths else 0.0 for m in range(1, months + 1)]; sw = sum(w)
    U = []; E = []; bal = 0.0; peak = 0.0; tdc = 0.0; ints = 0.0
    for m in range(1, months + 1):
        ld = land if m == 1 else 0.0; con = cost * gfa * w[m - 1] / sw
        sales = price * nsa * (dep if m == launch else ((1 - dep) / speriod if launch < m <= launch + speriod else 0.0))
        occ = 0.0 if m < lstart else min(1.0, (m - lstart + 1) / lmonths); rr = rent * nla * occ / 12 if m <= hold else 0.0
        sale = rent * nla / ey * (1 - sc) if m == hold else 0.0
        u = -ld - con + sales + rr + sale
        draw = ltc * (ld + con); intr = rate / 12 * bal; avail = bal + draw + intr
        rep = avail if m == hold else min(avail, sales + rr)
        bal = avail - rep; peak = max(peak, bal); tdc += ld + con + intr; ints += intr
        U.append(u); E.append(u + draw - rep)
    li, ui = irr(E, -0.5, 1.0), irr(U, -0.5, 1.0)
    return {'levered_irr': (1 + li) ** 12 - 1 if li is not None else None, 'unlevered_irr': (1 + ui) ** 12 - 1 if ui is not None else None,
            'net_profit': sum(E), 'total_cost': tdc, 'profit_on_cost': sum(E) / tdc, 'peak_debt': peak, 'interest': ints, 'equity': E}

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Inputs'
ws['A1'] = "AL BAHR MIXED-USE – PHASED DEVELOPMENT (SAR '000)"; ws['A2'] = 'Case'; ws['B2'] = 'Base'
dv = DataValidation(type='list', formula1='"Base,Downside,Upside"', allow_blank=False); ws.add_data_validation(dv); dv.add('B2')
ws.append([]); ws.append(['Scenario inputs', 'Base', 'Downside', 'Upside', 'Live'])
for i, (lab, b, d, u) in enumerate([('Residential sale price per sqm', 12, 10.8, 13), ('Retail rent per sqm pa', 1.2, 1.0, 1.3), ('Construction cost per sqm', 4.5, 4.9, 4.3), ('Exit yield', 0.07, 0.0775, 0.065)]):
    r = 5 + i; ws.append([lab, b, d, u, f'=INDEX(B{r}:D{r},MATCH($B$2,$B$4:$D$4,0))'])
    if lab == 'Exit yield':
        for c in 'BCDE': ws[f'{c}{r}'].number_format = '0.00%'
ws.append([])
single = [('Residential NSA (sqm)', 30000), ('Retail NLA (sqm)', 8000), ('GFA (sqm)', 50000), ("Land cost (SAR '000)", 120000), ('Construction start (month)', 1), ('Construction months', 24),
          ('Launch month', 6), ('Deposit share at launch', 0.3), ('Sales period (months after launch)', 24), ('Lease-up start (month)', 25), ('Lease-up months', 12), ('Hold (months)', 60),
          ('Loan to cost', 0.55), ('Interest rate', 0.07), ('Selling costs', 0.02)]
for lab, v in single: ws.append([lab, v])
R = {lab: 10 + i for i, (lab, _) in enumerate(single)}   # row of each single input
for lab in ['Deposit share at launch', 'Loan to cost', 'Interest rate', 'Selling costs']: ws[f'B{R[lab]}'].number_format = '0.00%'
I = lambda lab: f"Inputs!$B${R[lab]}"; LAND = "Land cost (SAR '000)"
PRICE, RENT, COST, EY = 'Inputs!$E$5', 'Inputs!$E$6', 'Inputs!$E$7', 'Inputs!$E$8'
dr = wb.create_sheet('Drivers'); dr.sheet_state = 'hidden'
dr['A1'] = 'Drivers (monthly)'; dr['A2'] = 'Month'; dr['A3'] = 'Construction S-curve weight'; dr['A4'] = 'Construction share'; dr['A5'] = 'Sales share'; dr['A6'] = 'Retail occupancy'; dr['A7'] = 'Year'
last = L(N + 2)
for m in range(1, N + 1):
    k = L(m + 2); dr[f'{k}2'] = m
    dr[f'{k}3'] = f"=IF(AND({k}$2>={I('Construction start (month)')},{k}$2<{I('Construction start (month)')}+{I('Construction months')}),SIN(PI()*({k}$2-{I('Construction start (month)')}+0.5)/{I('Construction months')}),0)"
    dr[f'{k}4'] = f'={k}3/SUM($C$3:${last}$3)'
    dr[f'{k}5'] = f"=IF({k}$2={I('Launch month')},{I('Deposit share at launch')},IF(AND({k}$2>{I('Launch month')},{k}$2<={I('Launch month')}+{I('Sales period (months after launch)')}),(1-{I('Deposit share at launch')})/{I('Sales period (months after launch)')},0))"
    dr[f'{k}6'] = f"=IF({k}$2<{I('Lease-up start (month)')},0,MIN(1,({k}$2-{I('Lease-up start (month)')}+1)/{I('Lease-up months')}))"
    dr[f'{k}7'] = f'=INT(({k}$2-1)/12)+1'
mo = wb.create_sheet('Monthly'); mo['A1'] = "Monthly cash flow (SAR '000)"
labs = ['Month', 'Land', 'Construction', 'Residential sales', 'Retail rent', 'Retail sale', 'Unlevered cash flow', 'Loan draw', 'Interest (capitalised)', 'Loan repayment', 'Loan balance', 'Equity cash flow']
for i, t in enumerate(labs): mo.cell(i + 2, 1, t)
for m in range(1, N + 1):
    k = L(m + 2); p = L(m + 1); prev = f'{p}12' if m > 1 else '0'
    mo[f'{k}2'] = f'=Drivers!{k}2'
    mo[f'{k}3'] = f"=IF({k}$2=1,{I(LAND)},0)"
    mo[f'{k}4'] = f"={COST}*{I('GFA (sqm)')}*Drivers!{k}4"
    mo[f'{k}5'] = f"={PRICE}*{I('Residential NSA (sqm)')}*Drivers!{k}5"
    mo[f'{k}6'] = f"=IF({k}$2<={I('Hold (months)')},{RENT}*{I('Retail NLA (sqm)')}*Drivers!{k}6/12,0)"
    mo[f'{k}7'] = f"=IF({k}$2={I('Hold (months)')},{RENT}*{I('Retail NLA (sqm)')}/{EY}*(1-{I('Selling costs')}),0)"
    mo[f'{k}8'] = f'=-{k}3-{k}4+{k}5+{k}6+{k}7'
    mo[f'{k}9'] = f"={I('Loan to cost')}*({k}3+{k}4)"
    mo[f'{k}10'] = f"={I('Interest rate')}/12*{prev}" if m < N else 0   # planted: the last month's interest typed as 0 over the formula row
    mo[f'{k}11'] = f"=IF({k}$2={I('Hold (months)')},{prev}+{k}9+{k}10,MIN({prev}+{k}9+{k}10,{k}5+{k}6))"
    mo[f'{k}12'] = f'={prev}+{k}9+{k}10-{k}11'
    mo[f'{k}13'] = f'={k}8+{k}9-{k}11'
an = wb.create_sheet('Annual'); an['A1'] = "Annual roll-up (SAR '000)"; an['A2'] = 'Year'
for i, t in enumerate(['Unlevered cash flow', 'Equity cash flow', 'Residential sales', 'Interest (capitalised)', 'Loan balance (year end)']): an.cell(3 + i, 1, t)
for y in range(1, 6):
    k = L(y + 2); an[f'{k}2'] = y
    for i, row in enumerate([8, 13, 5, 10]): an[f'{k}{3 + i}'] = f'=SUMIF(Drivers!$C$7:${last}$7,{k}$2,Monthly!$C${row}:${last}${row})'
    an[f'{k}7'] = f'=INDEX(Monthly!$C$12:${last}$12,1,{k}$2*12)'
wb.defined_names['EQCF'] = DefinedName('EQCF', attr_text=f'OFFSET(Monthly!$C$13,0,0,1,COUNT(Monthly!$C$2:${last}$2))')
wb.defined_names['UCF'] = DefinedName('UCF', attr_text=f'OFFSET(Monthly!$C$8,0,0,1,COUNT(Monthly!$C$2:${last}$2))')
rt = wb.create_sheet('Returns'); rt['A1'] = "Returns (SAR '000)"
for t, f, nf in [('Levered IRR', '=(1+IRR(EQCF,0.02))^12-1', '0.00%'), ('Unlevered IRR', '=(1+IRR(UCF))^12-1', '0.00%'), ('Net profit', f'=SUM(Monthly!C13:{last}13)', '#,##0'),
                 ('Total development cost', f'=SUM(Monthly!C3:{last}3)+SUM(Monthly!C4:{last}4)+SUM(Monthly!C10:{last}10)', '#,##0'), ('Profit on cost', '=B4/B5', '0.00%'),
                 ('Peak debt', f'=MAX(Monthly!C12:{last}12)', '#,##0'), ('Case run', '=Inputs!B2', None)]:
    rt.append([t, f])
    if nf: rt[f'B{rt.max_row}'].number_format = nf
wb.save(os.path.join(D, 'c44_phased_sar.xlsx'))
base = phased(); G['phased'] = {k: {a: b for a, b in v.items() if a != 'equity'} for k, v in {
    'base': base, 'downside': phased('Downside'), 'upside': phased('Upside'), 'price_minus_10': phased(price=10.8), 'hold_48': phased(hold=48), 'ltc_50': phased(ltc=0.50),
    'rate_8': phased(rate=0.08), 'ltc50_rate8': phased(ltc=0.50, rate=0.08), 'hold_72_past_timeline': None}.items() if v is not None}
lo, hi = 5.0, 30.0   # the sale price that gives a 20% levered IRR
for _ in range(200):
    m = (lo + hi) / 2
    if phased(price=m)['levered_irr'] < 0.20: lo = m
    else: hi = m
G['phased']['price_for_lev_20'] = (lo + hi) / 2
G['phased']['units'] = "SAR '000"; G['phased']['planted'] = "Monthly!" + L(N + 2) + "10 typed 0 over the interest formula row (harmless: the loan is repaid by then)"

# ---------------------------------------------------------------- 2. the three-deal fund (USD, 16 quarters)
Q = 16; START = datetime.date(2026, 3, 31)
def eom_add(d, months):
    y, m = divmod(d.month - 1 + months, 12); y += d.year; m += 1
    return datetime.date(y, m, min(d.day, calendar.monthrange(y, m)[1]))
DATES = [eom_add(START, 3 * (q - 1)) for q in range(1, Q + 1)]
DEALS = [  # name, currency, price (local m), acquisition quarter, NOI yield on cost, capex share, exit quarter, exit yield, rent growth pa
    ('Riyadh Office', 'SAR', 300.0, 1, 0.075, 0.10, 12, 0.065, 0.03), ('Jeddah Logistics', 'SAR', 180.0, 1, 0.080, 0.10, 14, 0.070, 0.04), ('Dubai Residential', 'AED', 220.0, 3, 0.065, 0.10, 14, 0.060, 0.02)]
FX = {'USD': 1.0, 'SAR': 3.75, 'AED': 3.6725}
def deal_cf(name, cur, price, acq, ny, capex, ex, ey, g, fx=None, sc=0.015):
    fx = fx or FX; out = []
    for q in range(1, Q + 1):
        buy = -price if q == acq else 0.0; cx = -price * capex / 4 if acq < q <= acq + 4 else 0.0
        noi = price * ny / 4 * (1 + g) ** ((q - acq - 1) // 4) if acq < q <= ex else 0.0
        sale = (price * ny * (1 + g) ** ((ex - acq - 1) // 4)) / ey * (1 - sc) if q == ex else 0.0
        out.append((buy + cx + noi + sale) / fx[cur])
    return out
def fund(deals=DEALS, fx=None, fee=0.015, committed=250.0, pref=0.08, carry=0.20, bug=True):
    rows = []
    for i, d in enumerate(deals):
        d = list(d)
        if bug and d[0] == 'Jeddah Logistics': d[7] = deals[0][7]   # the planted fault: Jeddah's sale reads Riyadh's exit yield
        rows.append(deal_cf(*d, fx=fx))
    gross = [sum(r[q] for r in rows) for q in range(Q)]; fees = [committed * fee / 4] * Q; net = [g - f for g, f in zip(gross, fees)]
    hb = 0.0; cumpref = 0.0; cumcu = 0.0; LP = []; GP = []; contrib = 0.0; dist_lp = 0.0
    for q in range(Q):
        c = max(0.0, -net[q]); dist = max(0.0, net[q]); hb = hb * (1 + pref) ** 0.25 + c
        topref = min(dist, hb); hb -= topref; cumpref += topref; rem = dist - topref
        cu = min(rem, max(0.0, carry / (1 - carry) * cumpref - cumcu)); cumcu += cu; rest = rem - cu
        lp = -c + topref + (1 - carry) * rest; gp = cu + carry * rest; LP.append(lp); GP.append(gp); contrib += c; dist_lp += topref + (1 - carry) * rest
    return {'lp_net_irr': xirr(LP, DATES), 'gross_irr': xirr(gross, DATES), 'tvpi': dist_lp / contrib, 'total_carry': sum(GP), 'lp_profit': sum(LP), 'peak_drawn': max(-min(0.0, x) for x in [sum(net[:q + 1]) for q in range(Q)]),
            'gross': gross, 'lp': LP}

def fund_book(indirect):
  wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Fund'
  ws['A1'] = 'GULF VALUE-ADD FUND II (USD m)'; ws['A2'] = 'Quarter'; ws['A3'] = 'Date'
  for q in range(1, Q + 1):
      k = L(q + 2); ws[f'{k}2'] = q; ws[f'{k}3'] = f'=EDATE(Inputs!$B$2,3*({k}2-1))'; ws[f'{k}3'].number_format = 'dd/mm/yyyy'
  ws['A5'] = 'Deal'; ws['B5'] = 'Currency'
  for i, d in enumerate(DEALS):
      r = 6 + i; ws[f'A{r}'] = d[0]; ws[f'B{r}'] = f"=INDIRECT(\"'\"&$A{r}&\"'!$B$3\")" if indirect else f"='{d[0]}'!$B$3"
      for q in range(1, Q + 1): k = L(q + 2); ws[f'{k}{r}'] = f"=INDEX(INDIRECT(\"'\"&$A{r}&\"'!$C$22:$R$22\"),1,{k}$2)" if indirect else f"='{d[0]}'!{k}22"
  labs = ['Gross cash flow (USD m)', 'Management fee', 'Net cash flow', 'LP contribution', 'Distributable', 'Hurdle balance', 'LP preferred return', 'Cumulative pref paid', 'GP catch-up', 'Cumulative catch-up',
          'Residual', 'LP net cash flow', 'GP carry']
  for i, t in enumerate(labs): ws.cell(10 + i, 1, t)
  for q in range(1, Q + 1):
      k = L(q + 2); p = L(q + 1); first = q == 1
      ws[f'{k}10'] = f'=SUM({k}6:{k}8)'; ws[f'{k}11'] = '=Inputs!$B$3*Inputs!$B$4/4'; ws[f'{k}12'] = f'={k}10-{k}11'
      ws[f'{k}13'] = f'=MAX(0,-{k}12)'; ws[f'{k}14'] = f'=MAX(0,{k}12)'
      ws[f'{k}15'] = f"={'0' if first else p + '15'}*(1+Inputs!$B$5)^(1/4)+{k}13-{k}16"
      ws[f'{k}16'] = f"=MIN({k}14,{'0' if first else p + '15'}*(1+Inputs!$B$5)^(1/4)+{k}13)"
      ws[f'{k}17'] = f"={'0' if first else p + '17'}+{k}16"
      ws[f'{k}18'] = f"=MIN({k}14-{k}16,MAX(0,Inputs!$B$6/(1-Inputs!$B$6)*{k}17-{'0' if first else p + '19'}))"
      ws[f'{k}19'] = f"={'0' if first else p + '19'}+{k}18"
      ws[f'{k}20'] = f'={k}14-{k}16-{k}18'
      ws[f'{k}21'] = f'=-{k}13+{k}16+(1-Inputs!$B$6)*{k}20'
      ws[f'{k}22'] = f'={k}18+Inputs!$B$6*{k}20'
  inp = wb.create_sheet('Inputs'); inp['A1'] = 'GULF VALUE-ADD FUND II – FUND TERMS (USD m)'
  for r, (t, v, nf) in enumerate([('First quarter end', START, 'dd/mm/yyyy'), ('Committed capital (USD m)', 250, '#,##0.0'), ('Management fee (pa, on committed)', 0.015, '0.00%'), ('Preferred return', 0.08, '0.00%'), ('Carried interest', 0.20, '0.00%')], start=2):
      inp[f'A{r}'] = t; inp[f'B{r}'] = v; inp[f'B{r}'].number_format = nf
  fx = wb.create_sheet('FX'); fx.append(['Currency', 'Per USD']);
  for c, v in [('USD', 1), ('SAR', 3.75), ('AED', 3.6725)]: fx.append([c, v])
  for name, cur, price, acq, ny, capex, ex, ey, g in DEALS:
      s = wb.create_sheet(name); s['A1'] = f'{name} – deal cash flow ({cur} m)'
      rows = [('Deal', name), ('Currency', cur), ('Purchase price (m)', price), ('Acquisition quarter', acq), ('NOI yield on cost', ny), ('Capex (share of price, over 4 quarters)', capex), ('Exit quarter', ex), ('Exit yield', ey), ('Rent growth pa', g), ('Selling costs', 0.015), ('FX per USD', '=VLOOKUP(B3,FX!$A$2:$B$4,2,FALSE)')]
      for i, (t, v) in enumerate(rows): s[f'A{2 + i}'] = t; s[f'B{2 + i}'] = v
      for r in (6, 7, 9, 10, 11): s[f'B{r}'].number_format = '0.00%'
      s['A14'] = 'Quarter'; s['A15'] = 'Purchase'; s['A16'] = 'Capex'; s['A17'] = 'NOI'; s['A18'] = 'Sale proceeds'; s['A20'] = f'Net cash flow ({cur} m)'; s['A22'] = 'Net cash flow (USD m)'
      EYREF = "'Riyadh Office'!$B$9" if name == 'Jeddah Logistics' else '$B$9'   # the planted fault
      for q in range(1, Q + 1):
          k = L(q + 2); s[f'{k}14'] = f'=Fund!{k}2'
          s[f'{k}15'] = f'=IF({k}$14=$B$5,-$B$4,0)'; s[f'{k}16'] = f'=IF(AND({k}$14>$B$5,{k}$14<=$B$5+4),-$B$4*$B$7/4,0)'
          s[f'{k}17'] = f'=IF(AND({k}$14>$B$5,{k}$14<=$B$8),$B$4*$B$6/4*(1+$B$10)^INT(({k}$14-$B$5-1)/4),0)'
          s[f'{k}18'] = f'=IF({k}$14=$B$8,$B$4*$B$6*(1+$B$10)^INT(({k}$14-$B$5-1)/4)/{EYREF}*(1-$B$11),0)'
          s[f'{k}20'] = f'=SUM({k}15:{k}18)'; s[f'{k}22'] = f'={k}20/$B$12'
  rt = wb.create_sheet('Returns'); rt['A1'] = 'Fund returns (USD m)'
  for t, f, nf in [('LP net IRR', '=XIRR(Fund!C21:R21,Fund!C3:R3)', '0.00%'), ('Gross IRR', '=XIRR(Fund!C10:R10,Fund!C3:R3)', '0.00%'),
                   ('TVPI', '=(SUM(Fund!C16:R16)+(1-Inputs!B6)*SUM(Fund!C20:R20))/SUM(Fund!C13:R13)', '0.00"x"'), ('Total carry', '=SUM(Fund!C22:R22)', '#,##0.00'),
                   ('LP profit', '=SUM(Fund!C21:R21)', '#,##0.00'), ('Committed capital', '=Inputs!B3', '#,##0.0')]:
      rt.append([t, f]); rt[f'B{rt.max_row}'].number_format = nf
  ck = wb.create_sheet('Checks'); ck['A1'] = 'Checks (all should be 0)'
  ck['A2'] = 'Deals sum to the fund gross'; ck['B2'] = "=ROUND(SUM(Fund!C10:R10)-(SUM('Riyadh Office'!C22:R22)+SUM('Jeddah Logistics'!C22:R22)+SUM('Dubai Residential'!C22:R22)),6)"
  ck['A3'] = 'Waterfall distributes everything'; ck['B3'] = '=ROUND(SUM(Fund!C14:R14)-SUM(Fund!C16:R16)-SUM(Fund!C18:R18)-SUM(Fund!C20:R20),6)'
  wb.save(os.path.join(D, 'c44_fund_usd.xlsx' if indirect else 'c44_fund_direct_usd.xlsx'))
fund_book(True); fund_book(False)

slip = [list(d) for d in DEALS]; slip[2][3] = 5; slip[2][6] = 16
fixed_fx = dict(FX); fixed_fx['AED'] = 4.0
G['fund'] = {k: {a: b for a, b in v.items() if a not in ('gross', 'lp')} for k, v in {
    'base': fund(), 'dubai_slips_2q': fund(deals=[tuple(d) for d in slip]), 'aed_4': fund(fx=fixed_fx), 'fee_125': fund(fee=0.0125), 'jeddah_fixed': fund(bug=False)}.items()}
G['fund']['dates'] = [d.isoformat() for d in DATES]; G['fund']['planted'] = "'Jeddah Logistics'!C18:R18 read 'Riyadh Office'!$B$9 (6.5%) for the exit yield instead of their own $B$9 (7%)"
json.dump(G, open(os.path.join(GD, 'c44.json'), 'w'), indent=1)
for k, v in G['phased'].items(): print('phased', k, v if not isinstance(v, dict) else {a: (round(b, 6) if isinstance(b, float) else b) for a, b in v.items()})
for k, v in G['fund'].items(): print('fund', k, v if not isinstance(v, dict) else {a: (round(b, 6) if isinstance(b, float) else b) for a, b in v.items()})
