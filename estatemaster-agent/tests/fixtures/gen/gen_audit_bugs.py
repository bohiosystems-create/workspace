"""A synthetic development model with nine planted defects, for the structural audit test (t_audit2). Nothing from a client.
Planted: (1) an input nothing reads, (2) a total whose SUM stops one row short of its block, (3) a row whose last period skips the
growth, (4) a typed number at the end of a formula row, (5) a dashboard figure typed over where the column holds formulas,
(6) an S-curve that adds to 95%, (7) the model's own check row not zero, (8) a cost wired so that more cost raises profit,
(9) a year header that jumps a year. Saved through LibreOffice so every formula carries its result. Run: python3 gen_audit_bugs.py <out.xlsx>"""
import sys, os, shutil, subprocess, tempfile
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L

wb = Workbook()
inp = wb.active; inp.title = 'Inputs'
rows = [('Land cost', 120000), ('Construction cost', 400000), ('Professional fees %', 0.08), ('Contingency %', 0.05), ('Marketing %', 0.03),
        ('Sale price per unit', 1800), ('Units', 400), ('Rent growth', 0.03), ('Exit yield', 0.07), ('Interest rate', 0.06), ('Debt %', 0.55)]
inp['A1'], inp['B1'] = 'Input', 'Value'
for i, (k, v) in enumerate(rows, start=2):
    inp.cell(i, 1, k); inp.cell(i, 2, v)
    if isinstance(v, float) and v < 1: inp.cell(i, 2).number_format = '0.0%'
I = {k: f'Inputs!$B${i}' for i, (k, v) in enumerate(rows, start=2)}

m = wb.create_sheet('Model')
years = [2026, 2027, 2028, 2030, 2031, 2032, 2033, 2034]          # planted 9: 2029 is missing
cols = [L(5 + i) for i in range(8)]
m['A1'] = 'Model (SAR 000)'
m['A3'] = 'Year'
for c, y in zip(cols, years): m[f'{c}3'] = y
m['A4'] = 'S-curve'
scurve = [0.30, 0.40, 0.25, 0, 0, 0, 0, 0]                           # planted 6: adds to 95%
for c, v in zip(cols, scurve): m[f'{c}4'] = v; m[f'{c}4'].number_format = '0%'
m['A5'] = 'Sales phasing'
for c, v in zip(cols, [0, 0.2, 0.4, 0.4, 0, 0, 0, 0]): m[f'{c}5'] = v; m[f'{c}5'].number_format = '0%'
labels = {10: 'Land', 11: 'Construction', 12: 'Professional fees', 13: 'Contingency', 14: 'Marketing', 15: 'Total cost',
          17: 'Sales revenue', 18: 'Sales price index', 20: 'Net profit', 22: 'Debt drawn', 23: 'Interest', 25: 'Net cash flow'}
for r, t in labels.items(): m[f'A{r}'] = t
for i, c in enumerate(cols):
    p = cols[i - 1] if i else None
    m[f'{c}10'] = f'={I["Land cost"]}*IF({c}3=$E$3,1,0)'
    m[f'{c}11'] = f'={I["Construction cost"]}*{c}4'
    m[f'{c}12'] = f'={c}11*{I["Professional fees %"]}'
    m[f'{c}13'] = f'={c}11*{I["Contingency %"]}'
    m[f'{c}14'] = f'={c}17*{I["Marketing %"]}'
    m[f'{c}15'] = f'=SUM({c}10:{c}13)'                               # planted 2: marketing (row 14) left out
    m[f'{c}18'] = '=1' if not i else (f'={p}18*1.02' if i < 7 else f'={p}18')   # planted 3: last period skips the growth
    m[f'{c}17'] = f'={I["Sale price per unit"]}*{I["Units"]}*{c}5*{c}18'
    m[f'{c}20'] = f'={c}17+{c}15'                                      # planted 8: profit = revenue + cost (sign wrong)
    m[f'{c}22'] = f'={c}15*{I["Debt %"]}'
    m[f'{c}23'] = f'={c}22*{I["Interest rate"]}'
    m[f'{c}25'] = f'={c}17-{c}15-{c}23'
m['L25'] = 150000                                                     # planted 4: a typed number at the end of the cash-flow row
m['N3'] = 'Total'
for r in (10, 11, 12, 13, 14, 15, 17, 20, 22, 23, 25): m[f'N{r}'] = f'=SUM(E{r}:L{r})'

d = wb.create_sheet('Dashboard')
d['A1'] = 'Dashboard'
d['A3'], d['B3'] = 'Levered IRR', '=IRR(Model!E25:L25)'
d['A4'], d['B4'] = 'Total development cost', '=Model!N15'
d['A5'], d['B5'] = 'Gross revenue', '=Model!N17'
d['A6'], d['B6'] = 'Net profit', 310000                               # planted 5: typed over (B5 and B7 hold =Model!N17 / =Model!N19-pattern)
d['A7'], d['B7'] = 'Net profit check', '=Model!N19'
d['A8'], d['B8'] = 'Peak debt', '=MAX(Model!E22:L22)'
d['A9'], d['B9'] = 'Profit on cost', '=B6/B4'
d['B3'].number_format = '0.00%'; d['B9'].number_format = '0.0%'
m['A19'] = 'Net profit (sum)'
for c in cols: m[f'{c}19'] = f'={c}20'
m['N19'] = '=SUM(E19:L19)'

k = wb.create_sheet('Checks')
k['A1'] = 'Checks'
k['A3'], k['B3'] = 'Check: sources - uses', '=Model!N22-Model!N15*Inputs!B12+1'   # planted 7: never zero
k['A4'], k['B4'] = 'Check: phasing', '=SUM(Model!E5:L5)-1'

tmp = tempfile.mkdtemp(); src = os.path.join(tmp, 'audit_bugs.xlsx'); wb.save(src)
out = os.path.join(tmp, 'o'); os.makedirs(out)
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', out, src], check=True, capture_output=True, timeout=600)
shutil.copy(os.path.join(out, 'audit_bugs.xlsx'), sys.argv[1]); shutil.rmtree(tmp, ignore_errors=True)
print('written', sys.argv[1])
