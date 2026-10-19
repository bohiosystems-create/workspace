# Round 47 fixture (synthetic): a hyperscale data-centre campus built to suit (SAR, quarterly), with made-up figures.
#  r47_dc_campus_sar: two phases delivered by MW of IT load, each with its own start quarter, build length and capex per MW; leased per MW
#  a month with a pre-let share at delivery and the rest taking up over a few quarters; power passed through at the tariff times the PUE
#  (a margin kept); Murabaha financing at a share of each quarter's cost with its profit capitalised until the exit; a typed exit quarter;
#  the returns in a table with the columns "Unlevered" and "Levered" and the rows IRR · Equity multiple · Profit (the other way round from
#  the usual matrix); the IRRs annualised from quarterly IRRs. One Campus sheet holds the inputs and the timeline from column H.
import openpyxl,sys
from openpyxl.utils import get_column_letter as L
out=sys.argv[1].rstrip('/')
wb=openpyxl.Workbook();K=wb.active;K.title='Returns';C=wb.create_sheet('Campus')
N=40;C0=8   # quarters 0..N, quarter 0 in column H
C['A1']='HYPERSCALE CAMPUS | INPUTS AND QUARTERLY CASH FLOW (SAR)'
C['A3']='PHASES';C['B3']='IT load (MW)';C['C3']='Start quarter';C['D3']='Build (quarters)';C['E3']='Capex per MW (SAR m)';C['F3']='Pre-let at delivery'
for r,(n,mw,s,b,cx,pl) in enumerate((('Phase 1',36,1,6,38,0.7),('Phase 2',24,7,5,40,0.5)),4):
    C[f'A{r}']=n;C[f'B{r}']=mw;C[f'C{r}']=s;C[f'D{r}']=b;C[f'E{r}']=cx;C[f'F{r}']=pl;C[f'F{r}'].number_format='0%'
for r,l,v,f in ((8,'Land (SAR)',95000000,'#,##0'),(9,'Rent per kW a month (SAR)',330,'#,##0'),(10,'Rent escalation a year',0.025,'0.00%'),(11,'Take-up of the rest (quarters)',4,'0'),
                (12,'PUE',1.3,'0.00'),(13,'Power tariff (SAR per MWh)',180,'#,##0'),(14,'Power margin kept',0.1,'0.00%'),(15,'Operating costs (% of rent)',0.12,'0.00%'),
                (16,'Murabaha share of cost',0.6,'0.00%'),(17,'Murabaha profit rate',0.068,'0.00%'),(18,'Exit quarter',32,'0'),(19,'Exit cap rate',0.0725,'0.00%'),(20,'Disposal costs',0.01,'0.00%')):
    C[f'A{r}']=l;C[f'B{r}']=v;C[f'B{r}'].number_format=f
lab={22:'Quarter',23:'Phase 1 capex',24:'Phase 2 capex',25:'Land',26:'Total cost',28:'Phase 1 leased MW',29:'Phase 2 leased MW',30:'Rent',31:'Power margin',32:'Operating costs',
     33:'Net operating income',35:'Exit value',36:'Murabaha drawn',37:'Murabaha profit (capitalised)',38:'Murabaha balance',39:'Murabaha repaid',41:'Unlevered cash flow',42:'Levered cash flow'}
for r,l in lab.items():C[f'A{r}']=l
for k in range(N+1):
    c=L(C0+k);p=L(C0+k-1);live=f'({c}$22<=$B$18)'
    C[f'{c}22']=k
    for i,r in ((0,23),(1,24)):
        pr=4+i;C[f'{c}{r}']=f'=IF(AND({c}$22>=$C${pr},{c}$22<$C${pr}+$D${pr}),$B${pr}*$E${pr}*1000000/$D${pr},0)'
        dl=f'({c}$22-$C${pr}-$D${pr}+1)'   # quarters since delivery, 1 in the delivery quarter
        C[f'{c}{28+i}']=f'=IF({dl}>=1,$B${pr}*MIN(1,$F${pr}+(1-$F${pr})*({dl}-1)/$B$11),0)'
    C[f'{c}25']='=$B$8' if k==0 else 0;C[f'{c}26']=f'=SUM({c}23:{c}25)'
    C[f'{c}30']=f'=({c}28+{c}29)*1000*$B$9*3*(1+$B$10)^INT({c}$22/4)'
    C[f'{c}31']=f'=({c}28+{c}29)*$B$12*24*91.25*$B$13*$B$14';C[f'{c}32']=f'=-{c}30*$B$15';C[f'{c}33']=f'={c}30+{c}31+{c}32'
    C[f'{c}35']=f'=IF({c}$22=$B$18,4*{L(C0+k+1) if k<N else c}33/$B$19*(1-$B$20),0)'
    C[f'{c}36']=f'={c}26*$B$16'
    C[f'{c}37']='=0' if k==0 else f'=IF({c}$22<=$B$18,{p}38*$B$17/4,0)'
    C[f'{c}39']=f'=IF({c}$22=$B$18,{"0" if k==0 else p+"38"}+{c}36+{c}37,0)'
    C[f'{c}38']=f'={"0" if k==0 else p+"38"}+{c}36+{c}37-{c}39'
    C[f'{c}41']=f'={live}*(-{c}26+{c}33+{c}35)';C[f'{c}42']=f'={live}*({c}41+{c}36-{c}39)'
    for r in list(range(23,27))+list(range(30,34))+[35,36,37,38,39,41,42]:C[f'{c}{r}'].number_format='#,##0;(#,##0)'
E=L(C0+N)
K['A1']='HYPERSCALE CAMPUS | RETURNS';K['B3']='Unlevered';K['C3']='Levered'
K['A4']='IRR';K['B4']=f"=(1+IRR(Campus!H41:{E}41,0.03))^4-1";K['C4']=f"=(1+IRR(Campus!H42:{E}42,0.03))^4-1"
K['A5']='Equity multiple';K['B5']=f'=SUMIF(Campus!H41:{E}41,">0")/-SUMIF(Campus!H41:{E}41,"<0")';K['C5']=f'=SUMIF(Campus!H42:{E}42,">0")/-SUMIF(Campus!H42:{E}42,"<0")'
K['A6']='Profit';K['B6']=f'=SUM(Campus!H41:{E}41)';K['C6']=f'=SUM(Campus!H42:{E}42)'
K['A8']='Total development cost';K['B8']=f'=SUM(Campus!H26:{E}26)';K['A9']='Peak Murabaha balance';K['B9']=f'=MAX(Campus!H38:{E}38)'
for a in ('B4','C4'):K[a].number_format='0.00%'
for a in ('B5','C5'):K[a].number_format='0.00"x"'
for a in ('B6','C6','B8','B9'):K[a].number_format='#,##0'
wb.save(f'{out}/r47_dc_campus_sar.xlsx')
