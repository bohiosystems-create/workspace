# Round 38 fixture (synthetic): a single-tenant net lease valuation laid out the way STNL templates often are, with made-up figures.
#  r38_stnl_usd: defined names throughout (SF_NRA, Rent_Base, Residual_Month, Debt_Payment, ...); the analysis period set by the lease
#  end and a renewal option; the option's rent weighted by its exercise probability (credit income) with the rest let at market after
#  downtime (spec income); the purchase price a formula (year 1 NOI over a 7.5% cap written inside it), not a typed number; returns by
#  XIRR over OFFSET ranges that end at the residual month; an "Unlevered Returns" and a "Levered Returns" block with rows labelled just
#  "IRR" and "EMx"; a Summary returns matrix (Profit · IRR · EMx by Unlevered / Levered row); a valuation by XNPV split into credit,
#  spec and residual cash flows at their own discount rates. Month counts are rounded where the timeline reads them.
import openpyxl,sys,datetime
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
out=sys.argv[1].rstrip('/')
wb=openpyxl.Workbook();S=wb.active;S.title='Summary';I=wb.create_sheet('Inputs');R=wb.create_sheet('Investment Returns');V=wb.create_sheet('Valuation')
def name(n,ref):wb.defined_names[n]=DefinedName(n,attr_text=ref)
N=240   # months on the timeline (0..N)
# ---- Inputs ----
I['B2']='STNL VALUATION MODEL | INPUTS';I['B6']='ANALYSIS DETAILS';I['I6']='PROJECT INVESTMENT';I['K6']='%';I['M6']='Amount'
I['B7']='Analysis Start';I['G7']=datetime.datetime(2026,1,1);I['G7'].number_format='mmm-yy'
I['B8']='Analysis Period';I['G8']='=F9'
I['B9']='Analysis End';I['E9']='Residual Month';I['F9']='=MIN(IF($M$33="Yes",$L$35,$L$28),240)';I['G9']='=EOMONTH(G7,G8-1)';I['G9'].number_format='mmm-yy'
I['B10']='General Vacancy';I['G10']=0;I['G10'].number_format='0.00%'
I['B12']='OTHER INCOME ASSUMPTIONS';I['B13']='Other Income - Billboard Rent';I['G13']=1500;I['B15']='Income Growth (% /Yr.)';I['G15']=0.02;I['G15'].number_format='0.00%'
I['B17']='EXPENSE ASSUMPTIONS';I['B18']='Tenant Expense Recovery %';I['G18']=1;I['G18'].number_format='0%'
I['B20']='Year 1 Operating Expenses';I['G20']='=SF_NRA*7.25';I['B21']='Year 1 Capital Expenditures';I['G21']='=SF_NRA*0.4'
I['B22']='Expense Growth (% /Yr.)';I['G22']=0.025;I['G22'].number_format='0.00%'
I['B24']='Year 1 Net Operating Income';I['G24']="='Investment Returns'!D23"
I['B25']='RESIDUAL ASSUMPTIONS';I['B26']='Residual Discount Rate';I['G26']='=E39+0.02';I['G26'].number_format='0.00%'
I['B28']='Gross Rent';I['F28']='=$M$43*SF_NRA*(1+$G$15)^(F9/12)';I['B29']='Expense Recovery Income';I['F29']='=$G$20*$G$18*(1+$G$22)^(F9/12)'
I['B30']='Other Income';I['F30']='=$G$13*(1+$G$15)^(F9/12)';I['B34']='- Operating Expenses';I['F34']='=$G$20*(1+$G$22)^(F9/12)'
I['B35']='Net Operating Income';I['F35']='=F28+F29+F30-F34'
I['B38']='Net Residual Value';I['E38']='Input';I['G38']='Amount';I['B39']='Gross Residual Value';I['D39']='Cap Rate';I['E39']=0.0725;I['E39'].number_format='0.00%';I['G39']='=F35/E39'
I['B42']='- Selling Costs';I['D42']='%';I['E42']=0.02;I['E42'].number_format='0.00%';I['G42']='=E42*G39';I['B43']='Net Residual Value';I['G43']='=G39-G42'
I['I7']='Purchase Price';I['M7']="='Investment Returns'!D23/0.075";I['I8']='Acquisition Fees';I['K8']=0.01;I['M8']='=M7*K8';I['I9']='Due Diligence Costs';I['K9']=0.005;I['M9']='=M7*K9'
I['I10']='Construction Costs';I['M10']=0;I['I11']='Total Project Costs';I['M11']='=SUM(M7:M10)'
I['I13']='DEBT';I['K13']='% of Price';I['I14']='Loan Amount';I['K14']=0.65;I['M14']='=M7*K14';I['I15']='Fixed Interest Rate';I['M15']=0.0615;I['I16']='Amortization';I['M16']=360
for a in ('K8','K9','K14','M15'):I[a].number_format='0.00%'
I['I18']='TENANT DETAILS';I['I19']='Tenant Name';I['M19']='Harbor Pharmacy, Inc';I['I21']='Credit Rating';I['M21']='BBB+';I['I22']='Credit Income Discount Rate';I['M22']=0.0675;I['I23']='Spec Income Discount Rate';I['M23']='=G26'
I['I25']='LEASE RENT DETAILS';I['I27']='Remaining Lease Term (months)';I['M27']=96;I['I28']='Lease End';I['K28']='End Month';I['L28']='=ROUND(M27,0)'
I['I30']='Rent Increase Type';I['M30']='Annual';I['I31']='Base Rent /SF/Yr (as of Month 1)';I['L31']='=M31*SF_NRA';I['M31']=38.5;I['I32']='Rent Increase';I['M32']=0.02;I['M32'].number_format='0.00%'
I['I33']='OPTION #1 - RENEWAL INCOME';I['M33']='Yes';I['I34']='Exercise Probability';I['K34']='Start Month';I['L34']='=L28+1';I['M34']=0.7;I['M34'].number_format='0%'
I['I35']='Length';I['K35']='End Month';I['L35']='=L34+ROUND(M35,0)-1';I['M35']=60
I['I36']='Option Rent /SF/Yr';I['M36']='=M31*(1+M32)^(ROUND(M27,0)/12)*(1+M37)';I['I37']='Option Rent Increase (% /Yr.)';I['M37']=0.0;I['M37'].number_format='0.00%'
I['I43']='Market Rent /SF/Yr (if not renewed)';I['M43']=36;I['I44']='Downtime (months)';I['M44']=9
for n,ref in (('SF_NRA','Summary!$G$11'),('Rent_Base','Inputs!$M$31'),('Vacancy','Inputs!$G$10'),('Residual_Month','Inputs!$F$9'),('Residual_NetValue','Inputs!$G$43'),
 ('Project_TotalCost','Inputs!$M$11'),('Debt_Amount','Inputs!$M$14'),('Debt_Payment',"'Investment Returns'!$N$7"),('Debt_Balloon',"'Investment Returns'!$N$8"),
 ('Discount_Rate_Credit','Inputs!$M$22'),('Discount_Rate_Spec','Inputs!$M$23'),('Discount_Rate_Residual','Inputs!$G$26')):name(n,ref)
# rent schedule, month 0 in column Q
for r,l in ((78,'Year'),(79,'Month'),(80,'Month Ending'),(81,'Base Credit Income'),(82,'Option 1 Credit Income'),(83,'Option 1 Spec Income'),(91,'Expense Recovery'),(96,'Other Income'),(102,'OpEx'),(107,'CapEx')):I.cell(r,16,l)
for k in range(N+1):
    c=L(17+k);I[f'{c}79']=k;I[f'{c}78']=f'=ROUNDUP({c}79/12,0)';I[f'{c}80']=f'=EOMONTH($G$7,{c}79-2)+1';I[f'{c}80'].number_format='mmm-yy'
    if k==0:
        for r in (81,82,83,91,96,102,107):I[f'{c}{r}']=0
        continue
    I[f'{c}81']=f'=IF(AND({c}79>=1,{c}79<=$L$28),Rent_Base*SF_NRA/12*(1+$M$32)^({c}78-1),0)'
    I[f'{c}82']=f'=IF(AND($M$33="Yes",{c}79>=$L$34,{c}79<=$L$35),$M$34*$M$36*SF_NRA/12*(1+$M$37)^(ROUNDUP(({c}79-$L$34+1)/12,0)-1),0)'
    I[f'{c}83']=f'=IF(AND($M$33="Yes",{c}79>=$L$34+ROUND($M$44,0),{c}79<=$L$35),(1-$M$34)*$M$43*SF_NRA/12,0)'
    occ=f'(SUM({c}81:{c}83)>0)'
    I[f'{c}102']=f'={occ}*$G$20/12*(1+$G$22)^({c}78-1)';I[f'{c}91']=f'={c}102*$G$18';I[f'{c}96']=f'={occ}*$G$13/12*(1+$G$15)^({c}78-1)';I[f'{c}107']=f'={occ}*$G$21/12*(1+$G$22)^({c}78-1)'
# ---- Investment Returns: monthly cash flow, month 0 in column C ----
R['B2']='STNL VALUATION MODEL | INVESTMENT RETURNS';R['B6']='Unlevered Returns';R['B7']='IRR';R['E7']='=C61';R['B8']='EMx';R['E8']='=C62'
R['B11']='Levered Returns';R['B12']='IRR';R['E12']='=C66';R['B13']='EMx';R['E13']='=C67'
for a in ('E7','E12'):R[a].number_format='0.00%'
for a in ('E8','E13'):R[a].number_format='0.00"x"'
R['K6']='Debt';R['K7']='Payment (Monthly)';R['N7']='=PMT(Inputs!M15/12,Inputs!M16,-Inputs!M14)';R['K8']='Balloon Balance';R['N8']="=PV(Inputs!M15/12,Inputs!M16-'Investment Returns'!N9,-Debt_Payment)";R['K9']='Term';R['N9']='=Residual_Month'
lab={37:'Year',38:'Month',39:'Month Ending',40:'Gross Rents',41:'Expense Recovery',42:'Other Income',43:'General Vacancy',44:'Effective Gross Revenue',45:'Operating Expenses',
 46:'Net Operating Income',47:'Capital Expenditures',48:'Cash Flow from Operations',49:'Debt Service',50:'Cash Flow after Financing',51:'Initial Investment',52:'Initial Loan Draw',
 53:'Net Residual Value',54:'Loan Payoff',56:'Unlevered Cash Flow',57:'Levered Cash Flow',59:'Unlevered Returns',61:'IRR',62:'EMx',64:'Levered Returns',66:'IRR',67:'EMx'}
for r,l in lab.items():R.cell(r,2,l)
R['B36']='MONTHLY CASH FLOW'
for k in range(N+1):
    c=L(3+k);p=L(2+k);ic=L(17+k);bl=f'IF({c}38="","",'
    if k==0:
        R['C37']=0;R['C38']=0;R['C39']='=EOMONTH(D39,-1)'
    else:
        R[f'{c}38']=f'=IF({p}38+1>Residual_Month,"",{p}38+1)';R[f'{c}37']=f'={bl}ROUNDUP({c}38/12,0))';R[f'{c}39']=f'={bl}EOMONTH(Inputs!$G$7,{c}38-1))'
    R[f'{c}39'].number_format='mmm-yy'
    R[f'{c}40']=f'={bl}SUM(Inputs!{ic}81:{ic}83))';R[f'{c}41']=f'={bl}Inputs!{ic}91)';R[f'{c}42']=f'={bl}Inputs!{ic}96)';R[f'{c}43']=f'={bl}SUM({c}40:{c}42)*-Vacancy)'
    R[f'{c}44']=f'={bl}SUM({c}40:{c}43))';R[f'{c}45']=f'={bl}Inputs!{ic}102)';R[f'{c}46']=f'={bl}{c}44-{c}45)';R[f'{c}47']=f'={bl}Inputs!{ic}107)';R[f'{c}48']=f'={bl}{c}46-{c}47)'
    R[f'{c}49']=0 if k==0 else f'={bl}Debt_Payment)';R[f'{c}50']=f'={bl}{c}48-{c}49)'
    R[f'{c}51']=f'={bl}IF({c}38>0,0,Project_TotalCost))';R[f'{c}52']=f'={bl}IF({c}38=0,Debt_Amount,0))';R[f'{c}53']=f'={bl}IF({c}38<Residual_Month,0,Residual_NetValue))'
    R[f'{c}54']=f'={bl}IF({c}38=Residual_Month,Debt_Balloon,0))';R[f'{c}56']=f'={bl}{c}48-{c}51+{c}53)';R[f'{c}57']=f'={bl}{c}50-{c}51+{c}52+{c}53-{c}54)'
E=L(3+N)
R['C61']='=XIRR(OFFSET($C$56,0,0,1,Residual_Month+1),OFFSET($C$39,0,0,1,Residual_Month+1))';R['C62']=f'=SUMIF($C$56:${E}$56,">0")/-SUMIF($C$56:${E}$56,"<0")'
R['C66']='=XIRR(OFFSET($C$57,0,0,1,Residual_Month+1),OFFSET($C$39,0,0,1,Residual_Month+1))';R['C67']=f'=SUMIF($C$57:${E}$57,">0")/-SUMIF($C$57:${E}$57,"<0")'
# annual cash flow, year 0 in column C
R['B16']='ANNUAL CASH FLOW'
for y in range(21):
    c=L(3+y);R[f'{c}16']=y
    for r,src in ((17,40),(23,46),(25,48),(26,49),(33,56),(34,57)):R[f'{c}{r}']=f'=SUMIF($C$37:${E}$37,{c}$16,$C{src}:${E}{src})'
for r,l in ((17,'Gross Rents'),(23,'Net Operating Income'),(25,'Cash Flow from Operations'),(26,'Debt Service'),(33,'Unlevered Cash Flow'),(34,'Levered Cash Flow')):R.cell(r,2,l)
# ---- Valuation: XNPV by partition ----
V['B2']='STNL VALUATION MODEL | VALUATION';V['B6']='VALUATION SUMMARY';V['B7']='Year 1 NOI';V['C7']="='Investment Returns'!D23";V['B8']='PV of Project Cash Flows';V['C8']='=C42';V['B9']='Yield-on PV';V['C9']='=C7/C8'
for r,l in ((14,'Month'),(15,'Month Ending'),(34,'Credit Cash Flow'),(35,'Spec Cash Flow'),(36,'Residual Cash Flow'),(39,'Credit Cash Flow'),(40,'Spec Cash Flow'),(41,'Residual Cash Flow'),(42,'Present Value')):V.cell(r,2,l)
V['C33']='Discount Rate';V['C34']='=Discount_Rate_Credit';V['C35']='=Discount_Rate_Spec';V['C36']='=Discount_Rate_Residual';V['C38']='PV of CF'
for k in range(1,N+1):
    c=L(3+k);ic=L(17+k);rc=L(3+k)
    V[f'{c}14']=k;V[f'{c}15']=f"='Investment Returns'!{rc}39"
    occ=f'(Inputs!{ic}91+Inputs!{ic}96-Inputs!{ic}102-Inputs!{ic}107)'
    V[f'{c}34']=f'=IF({c}14>Residual_Month,0,Inputs!{ic}81+Inputs!{ic}82+IF(SUM(Inputs!{ic}81:{ic}83)=0,0,(Inputs!{ic}81+Inputs!{ic}82)/SUM(Inputs!{ic}81:{ic}83))*{occ})'
    V[f'{c}35']=f'=IF({c}14>Residual_Month,0,Inputs!{ic}83+IF(SUM(Inputs!{ic}81:{ic}83)=0,0,Inputs!{ic}83/SUM(Inputs!{ic}81:{ic}83))*{occ})'
    V[f'{c}36']=f'=({c}14=Residual_Month)*Residual_NetValue'
V['C39']='=XNPV(C34,OFFSET(D34,0,0,1,Residual_Month),OFFSET($D$15,0,0,1,Residual_Month))';V['C40']='=XNPV(C35,OFFSET(D35,0,0,1,Residual_Month),OFFSET($D$15,0,0,1,Residual_Month))'
V['C41']='=XNPV(C36,OFFSET(D36,0,0,1,Residual_Month),OFFSET($D$15,0,0,1,Residual_Month))';V['C42']='=SUM(C39:C41)'
# ---- Summary ----
S['B2']='STNL VALUATION MODEL | INVESTMENT SUMMARY';S['B6']='PROPERTY DESCRIPTION';S['B7']='Name';S['G7']='Harbor Pharmacy';S['B11']='Net Rentable Area (SF)';S['G11']=14500
S['I6']='LEASE SUMMARY';S['I11']='Remaining Term (months)';S['L11']='=Inputs!M27';S['I12']='Lease Type';S['L12']='NNN'
S['B16']='RETURN SUMMARY';S['F16']='Analysis Period:';S['G16']='=Residual_Month';S['E17']='Profit';S['F17']='IRR';S['G17']='EMx '
S['B18']='Unlevered';S['E18']=f"=SUM('Investment Returns'!C33:W33)";S['F18']="='Investment Returns'!E7";S['G18']="='Investment Returns'!E8"
S['B19']='Levered';S['E19']=f"=SUM('Investment Returns'!C34:W34)";S['F19']="='Investment Returns'!E12";S['G19']="='Investment Returns'!E13"
for r in (18,19):S[f'E{r}'].number_format='#,##0';S[f'F{r}'].number_format='0.00%';S[f'G{r}'].number_format='0.00"x"'
S['I14']='PRICING/VALUATION';S['L14']='Cap Rate';S['N14']='Price'
for r,l,v in ((15,'Asking Price',6600000),(16,'Acquisition Price','=Inputs!M7'),(18,'DCF Value','=Valuation!C42')):
    S[f'I{r}']=l;S[f'N{r}']=v;S[f'L{r}']=f"='Investment Returns'!$D$23/N{r}";S[f'L{r}'].number_format='0.00%';S[f'N{r}'].number_format='#,##0'
wb.save(f'{out}/r38_stnl_usd.xlsx')
