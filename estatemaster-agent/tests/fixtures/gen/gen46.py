# Round 46 fixture (synthetic): a single-tenant net lease (a drive-thru restaurant, AED) laid out the way small net-lease models often are,
# with made-up figures. r46_stnl_qsr_aed: the years run down the rows of a "Lease CF" sheet (year 0 in row 6), not across; the returns are
# labelled "Unleveraged IRR" / "Leveraged IRR" on a Deal Summary next to a going-in cap rate and a year-1 cash-on-cash; the rent steps up
# 10% every five lease years and the lease ends before the timeline does (a renewal at market rent after a vacancy); the purchase price is
# typed; the loan is interest-only for three years, then amortising, its balance at the exit from CUMPRINC; a typed hold period ends the
# cash flow (sale at the next year's NOI over the exit cap).
import openpyxl,sys
out=sys.argv[1].rstrip('/')
wb=openpyxl.Workbook();S=wb.active;S.title='Deal Summary';A=wb.create_sheet('Assumptions');C=wb.create_sheet('Lease CF')
# ---- assumptions ----
A['A1']='ASSUMPTIONS (AED)'
rows=[(3,'Purchase price',14200000,'#,##0'),(4,'Closing costs (% of price)',0.025,'0.00%'),(5,'Year-1 base rent',1010000,'#,##0'),(6,'Rent step every 5 lease years',0.10,'0.00%'),
      (7,'Lease years remaining',8,'0'),(8,'Vacancy before re-letting (years)',1,'0'),(9,'Market rent at re-letting (% of last rent)',0.95,'0.00%'),(10,'Market rent growth a year',0.02,'0.00%'),
      (11,'Non-recoverable costs (% of rent)',0.015,'0.00%'),(12,'Hold period (years)',10,'0'),(13,'Exit cap rate',0.0725,'0.00%'),(14,'Selling costs',0.015,'0.00%'),
      (16,'Loan to value',0.6,'0.00%'),(17,'Interest rate',0.0595,'0.00%'),(18,'Interest-only years',3,'0'),(19,'Amortisation (years)',25,'0'),(20,'Lease years elapsed at purchase',7,'0')]
for r,l,v,f in rows:A[f'A{r}']=l;A[f'B{r}']=v;A[f'B{r}'].number_format=f
A['A22']='Loan amount';A['B22']='=B3*B16';A['B22'].number_format='#,##0'
A['A23']='Annual debt service (amortising)';A['B23']='=PMT(B17,B19,-B22)';A['B23'].number_format='#,##0'
# ---- lease cash flow: year 0 in row 6, years running down ----
C['A1']='LEASE CASH FLOW (AED)';hdr=['Year','Lease year','Base rent','Non-recoverable costs','Net operating income','Interest','Principal','Debt service','Sale proceeds','Loan payoff','Unleveraged cash flow','Leveraged cash flow']
for j,h in enumerate(hdr):C.cell(5,1+j,h)
N=16
for y in range(N+1):
    r=6+y;C[f'A{r}']=y;C[f'B{r}']=f'=Assumptions!$B$20+A{r}'
    live=f'(A{r}<=Assumptions!$B$12)'
    rent=(f'IF(A{r}<=Assumptions!$B$7,Assumptions!$B$5*(1+Assumptions!$B$6)^(INT((B{r}-1)/5)-INT(Assumptions!$B$20/5)),'
          f'IF(A{r}<=Assumptions!$B$7+Assumptions!$B$8,0,Assumptions!$B$5*(1+Assumptions!$B$6)^(INT((Assumptions!$B$20+Assumptions!$B$7-1)/5)-INT(Assumptions!$B$20/5))*Assumptions!$B$9*(1+Assumptions!$B$10)^(A{r}-Assumptions!$B$7-Assumptions!$B$8-1)))')
    if y==0:
        for c in 'CDEFGHIJ':C[f'{c}{r}']=0
        C[f'K{r}']='=-Assumptions!B3*(1+Assumptions!B4)';C[f'L{r}']='=K6+Assumptions!B22'
    else:
        C[f'C{r}']=f'={rent}';C[f'D{r}']=f'=-C{r}*Assumptions!$B$11';C[f'E{r}']=f'=C{r}+D{r}'
        C[f'F{r}']=f'=IF(A{r}<=Assumptions!$B$18,Assumptions!$B$22*Assumptions!$B$17,-CUMIPMT(Assumptions!$B$17,Assumptions!$B$19,Assumptions!$B$22,A{r}-Assumptions!$B$18,A{r}-Assumptions!$B$18,0))'
        C[f'G{r}']=f'=IF(A{r}<=Assumptions!$B$18,0,-CUMPRINC(Assumptions!$B$17,Assumptions!$B$19,Assumptions!$B$22,A{r}-Assumptions!$B$18,A{r}-Assumptions!$B$18,0))'
        C[f'H{r}']=f'=F{r}+G{r}'
        nxt=f'E{r+1}' if y<N else f'E{r}*(1+Assumptions!$B$10)'
        C[f'I{r}']=f'=IF(A{r}=Assumptions!$B$12,{nxt}/Assumptions!$B$13*(1-Assumptions!$B$14),0)'
        C[f'J{r}']=(f'=IF(A{r}=Assumptions!$B$12,Assumptions!$B$22+IF(A{r}<=Assumptions!$B$18,0,CUMPRINC(Assumptions!$B$17,Assumptions!$B$19,Assumptions!$B$22,1,A{r}-Assumptions!$B$18,0)),0)')
        C[f'K{r}']=f'={live}*(E{r}+I{r})';C[f'L{r}']=f'={live}*(E{r}-H{r}+I{r}-J{r})'
    for c in 'CDEFGHIJKL':C[f'{c}{r}'].number_format='#,##0;(#,##0)'
E=6+N
# ---- deal summary ----
S['A1']='NET LEASE ACQUISITION | DEAL SUMMARY';S['A3']='Tenant';S['B3']='Drive-thru restaurant (corporate guarantee)';S['A4']='Location';S['B4']='Dubai, Al Barsha'
S['A6']='Going-in cap rate';S['B6']="='Lease CF'!E7/Assumptions!B3";S['A7']='Year-1 cash-on-cash';S['B7']="='Lease CF'!L7/-'Lease CF'!L6"
S['A9']='Unleveraged IRR';S['B9']=f"=IRR('Lease CF'!K6:K{E})";S['A10']='Leveraged IRR';S['B10']=f"=IRR('Lease CF'!L6:L{E})"
S['A11']='Equity multiple';S['B11']=f"=SUM('Lease CF'!L7:L{E})/-'Lease CF'!L6";S['A12']='Net profit (leveraged)';S['B12']=f"=SUM('Lease CF'!L6:L{E})"
for a in ('B6','B7','B9','B10'):S[a].number_format='0.00%'
S['B11'].number_format='0.00"x"';S['B12'].number_format='#,##0'
wb.save(f'{out}/r46_stnl_qsr_aed.xlsx')
