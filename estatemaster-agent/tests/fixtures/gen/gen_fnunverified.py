# a small development model that uses two functions the test models never use (GCD, FISHER): the upload check names them
import openpyxl,sys
wb=openpyxl.Workbook();I=wb.active;I.title='Inputs'
for i,(l,v) in enumerate([('Units',120),('Price per unit',450000),('Build cost per unit',300000),('Phasing step',24),('Growth',0.03)],1):I.cell(i,1,l);I.cell(i,2,v)
R=wb.create_sheet('Returns')
rows=[('Gross revenue','=Inputs!B1*Inputs!B2*(1+Inputs!B5)'),('Total cost','=Inputs!B1*Inputs!B3'),('Net profit','=B1-B2'),('Profit on cost','=B3/B2'),('Phase size (units)','=GCD(Inputs!B1,Inputs!B4)'),('Risk score','=FISHER(MIN(0.9,B4))')]
for i,(l,f) in enumerate(rows,1):R.cell(i,1,l);R.cell(i,2,f)
wb.save(sys.argv[1])
