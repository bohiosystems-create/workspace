"""Gulf fixtures (t_gulf), synthetic: which regime a project or a contract falls under.
  uae_dubai.xlsx     a Business Bay, Dubai acquisition in AED with the Dubai Land Department transfer fee and 5% VAT: UAE onshore
  gulf_neutral.xlsx  the same deal in USD with no location, tax or regulator anywhere: the workbook does not say
  difc_fund.docx     a DIFC public property fund prospectus (DFSA), borrowing up to 60% of gross asset value (the reported DFSA cap
                     for a public property fund is 50%); it also names "the CMA" once, which is no evidence either way
Nothing from a client. Run: python3 gen_gulf.py <out_dir_docs> <out_dir_models>"""
import sys, os, shutil, subprocess, tempfile, zipfile
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L

DOCS, MODELS = sys.argv[1], sys.argv[2]

def model(path, ccy, dubai):
    wb = Workbook(); i = wb.active; i.title = 'Inputs'
    i['A1'] = f'BUSINESS BAY OFFICE ACQUISITION ({ccy})' if dubai else f'OFFICE ACQUISITION ({ccy})'
    rows = ([('Location', 'Business Bay, Dubai')] if dubai else []) + [('Purchase price', 50000000)] + \
        ([('Dubai Land Department transfer fee', 0.04), ('VAT on fees', 0.05)] if dubai else []) + \
        [('Agency fee', 0.02), ('Rent year 1', 4000000), ('Rent growth', 0.03), ('Exit yield', 0.07), ('Hold (years)', 5)]
    for r, (k, v) in enumerate(rows, start=2):
        i.cell(r, 1, k); i.cell(r, 2, v)
        if isinstance(v, float) and v < 1: i.cell(r, 2).number_format = '0.00%'
    I = {k: f'Inputs!$B${r}' for r, (k, v) in enumerate(rows, start=2)}
    fees = f"{I['Purchase price']}*({I['Agency fee']}" + (f"*(1+{I['VAT on fees']})+{I['Dubai Land Department transfer fee']})" if dubai else ')')
    c = wb.create_sheet('Cash Flow'); c['A1'] = 'CASH FLOW'; c['A3'] = 'Year'
    for r, t in {5: 'Acquisition', 6: 'Acquisition costs', 7: 'Rent', 8: 'Sale', 9: 'Net cash flow'}.items(): c.cell(r, 1, t)
    for k in range(6):
        col = L(3 + k); c[f'{col}3'] = 2026 + k
        c[f'{col}5'] = f"=-{I['Purchase price']}" if k == 0 else 0
        c[f'{col}6'] = f'=-{fees}' if k == 0 else 0
        c[f'{col}7'] = 0 if k == 0 else f"={I['Rent year 1']}*(1+{I['Rent growth']})^{k - 1}"
        c[f'{col}8'] = f"=IF({k}={I['Hold (years)']},{col}7*(1+{I['Rent growth']})/{I['Exit yield']},0)"
        c[f'{col}9'] = f'=SUM({col}5:{col}8)'
    s = wb.create_sheet('Returns'); s['A3'] = 'Unlevered IRR'; s['B3'] = "=IRR('Cash Flow'!C9:H9)"; s['B3'].number_format = '0.00%'
    s['A4'] = 'Net profit'; s['B4'] = "=SUM('Cash Flow'!C9:H9)"
    tmp = tempfile.mkdtemp(); src = os.path.join(tmp, os.path.basename(path)); wb.save(src)
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', os.path.join(tmp, 'o'), src], check=True, capture_output=True, timeout=600)
    shutil.copy(os.path.join(tmp, 'o', os.path.basename(path)), path); shutil.rmtree(tmp, ignore_errors=True)

os.makedirs(MODELS, exist_ok=True)
model(os.path.join(MODELS, 'uae_dubai.xlsx'), 'AED', True)
model(os.path.join(MODELS, 'gulf_neutral.xlsx'), 'USD', False)

P = [(None, 'GULF HORIZON PROPERTY FUND'), (None, 'PROSPECTUS'), (None, '(synthetic fixture for testing; not a real fund)'),
     (0, 'THE FUND.'),
     (1, 'Domicile and regulator. The Fund is a public property fund domiciled in the Dubai International Financial Centre and regulated by the Dubai Financial Services Authority under the Collective Investment Rules.'),
     (1, 'Other authorities. Nothing in this Prospectus has been approved by the CMA or by any other authority outside the Dubai International Financial Centre.'),
     (0, 'INVESTMENT AND BORROWING.'),
     (1, 'Borrowing. The Fund may borrow up to sixty per cent (60%) of its gross asset value to finance acquisitions.'),
     (1, 'Structure. The Fund is a closed-ended fund and its units will be listed within three years of the first offering.'),
     (0, 'FEES.'),
     (1, 'Management fee. The Fund Manager is entitled to a management fee of one and a quarter per cent (1.25%) of net asset value a year, payable quarterly.')]
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
def para(lv, t):
    num = '' if lv is None else '<w:pPr><w:numPr><w:ilvl w:val="%d"/><w:numId w:val="1"/></w:numPr></w:pPr>' % lv
    return f'<w:p>{num}<w:r><w:t xml:space="preserve">{escape(t)}</w:t></w:r></w:p>'
body = ''.join(para(lv, t) for lv, t in P)
doc = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body>{body}<w:sectPr/></w:body></w:document>'
num = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:numbering {W}><w:abstractNum w:abstractNumId="0">'
       '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl><w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2"/></w:lvl>'
       '</w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>')
ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
      '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
      '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
      '<Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/></Types>')
rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
drels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
         '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/></Relationships>')
os.makedirs(DOCS, exist_ok=True)
with zipfile.ZipFile(os.path.join(DOCS, 'difc_fund.docx'), 'w', zipfile.ZIP_DEFLATED) as z:
    z.writestr('[Content_Types].xml', ct); z.writestr('_rels/.rels', rels); z.writestr('word/document.xml', doc)
    z.writestr('word/_rels/document.xml.rels', drels); z.writestr('word/numbering.xml', num)
print('written uae_dubai.xlsx, gulf_neutral.xlsx, difc_fund.docx')
