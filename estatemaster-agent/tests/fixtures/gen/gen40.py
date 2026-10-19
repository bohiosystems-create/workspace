# Round 40 fixture (synthetic): a Saudi off-plan residential development, with made-up figures, laid out the way developers' models often are.
#  r40_offplan_sar: amounts in SAR '000 stated once in the sheet heading; units sold off-plan on a sales curve, each sale paid in
#  instalments over construction (a deposit, then progress payments, the balance at handover); buyers' payments held in an escrow
#  (Wafi) account released in step with construction progress; a Murabaha facility drawn for the cost the released cash does not
#  cover, its profit rate SAIBOR plus a margin; a quarterly timeline with dated quarter ends; returns on a Results sheet whose
#  rows read "Project IRR (unlevered)", "Equity IRR" and "Developer profit" with "Profit on cost" beside them.
import openpyxl,sys,datetime,calendar
from openpyxl.utils import get_column_letter as L
out=sys.argv[1].rstrip('/')
wb=openpyxl.Workbook();R=wb.active;R.title='Results';I=wb.create_sheet('Inputs');Q=wb.create_sheet('Quarterly CF')
N=16   # quarters 0..N
I['B2']="NAKHEEL GARDENS · INPUTS (SAR '000)"
rows=[(4,'Units',240,'0'),(5,'Average unit price',1450,'#,##0'),(6,'Land cost',62000,'#,##0'),(7,'Construction cost',168000,'#,##0'),
 (8,'Construction start (quarter)',1,'0'),(9,'Construction length (quarters)',10,'0'),(10,'Soft costs (% of construction)',0.12,'0.00%'),
 (11,'Sales and marketing (% of sales)',0.03,'0.00%'),(12,'Deposit at sale',0.10,'0.00%'),(13,'Paid at handover',0.30,'0.00%'),
 (14,'Escrow release lag (quarters)',1,'0'),(15,'SAIBOR',0.055,'0.00%'),(16,'Murabaha margin',0.025,'0.00%'),(17,'Murabaha upfront fee',0.0075,'0.00%'),
 (18,'Equity (share of land and first costs)',1.0,'0%')]
for r,l,v,f in rows:I.cell(r,2,l);I.cell(r,3,v).number_format=f
I['B20']='Sales curve (share of units sold, by quarter from launch)';I['B21']='Quarter'
curve=[0.15,0.15,0.12,0.10,0.10,0.08,0.08,0.07,0.06,0.05,0.04]
for k in range(N+1):
    c=L(3+k);I[f'{c}21']=k;I[f'{c}22']=curve[k-1] if 1<=k<=len(curve) else 0;I[f'{c}22'].number_format='0.0%'
I['B22']='Share sold'
Q['B2']="QUARTERLY CASH FLOW (SAR '000)";Q['B3']='Quarter ending';Q['B4']='Quarter'
lab={6:'Units sold',7:'Sales contracted',8:'Construction progress',9:'Cumulative progress',11:'Buyer payments into escrow',12:'Escrow balance',13:'Escrow release',
 15:'Land',16:'Construction',17:'Soft costs',18:'Sales and marketing',19:'Total costs',21:'Project cash flow (unlevered)',
 23:'Funding need after escrow',24:'Equity',25:'Murabaha draw',26:'Murabaha profit',27:'Murabaha fee',28:'Murabaha repayment',29:'Murabaha balance',31:'Equity cash flow'}
for r,l in lab.items():Q.cell(r,2,l)
def qend(k):y,m=divmod(3*k+2,12);y+=2026;m+=1;return datetime.datetime(y,m,calendar.monthrange(y,m)[1])
for k in range(N+1):
    c=L(3+k);p=L(2+k);Q[f'{c}3']=qend(k);Q[f'{c}3'].number_format='mmm-yy';Q[f'{c}4']=k;qq=f'{c}$4'
    Q[f'{c}6']=f'=Inputs!$C$4*Inputs!{c}22';Q[f'{c}7']=f'={c}6*Inputs!$C$5'
    Q[f'{c}8']=f'=IF(AND({qq}>=Inputs!$C$8,{qq}<Inputs!$C$8+ROUND(Inputs!$C$9,0)),1/ROUND(Inputs!$C$9,0),0)';Q[f'{c}9']=f'=SUM($C$8:{c}8)'
    hand=f'Inputs!$C$8+ROUND(Inputs!$C$9,0)'
    # a sale pays its deposit now; the middle share follows construction progress; the balance comes at handover (sales after handover pay in full)
    mid='(1-Inputs!$C$12-Inputs!$C$13)';prevS='0' if k==0 else f'SUM($C$7:{p}7)'
    Q[f'{c}11']=(f'=IF({qq}<{hand},{c}7*(Inputs!$C$12+{mid}*{c}9)+{prevS}*{mid}*{c}8,0)'
                 f'+IF({qq}={hand},{prevS}*Inputs!$C$13+{c}7,0)+IF({qq}>{hand},{c}7,0)')
    Q[f'{c}13']=(f'=IF({qq}>={hand},{p}12+{c}11,MIN({p}12+{c}11,MAX(0,{c}9*SUM($C$11:{c}11)-SUM($C$13:{p}13))))' if k else '=0')
    Q[f'{c}12']=f'={"0" if k==0 else p+"12"}+{c}11-{c}13'
    Q[f'{c}15']=f'=IF({qq}=0,-Inputs!$C$6,0)';Q[f'{c}16']=f'=-Inputs!$C$7*{c}8';Q[f'{c}17']=f'={c}16*Inputs!$C$10';Q[f'{c}18']=f'=-{c}7*Inputs!$C$11'
    Q[f'{c}19']=f'=SUM({c}15:{c}18)'
    Q[f'{c}21']=f'={c}11+{c}19'
    Q[f'{c}23']=f'=-MIN(0,{c}13+{c}19)'
    prevbal='0' if k==0 else f'{p}29'
    Q[f'{c}26']='=0' if k==0 else f'=-{p}29*(Inputs!$C$15+Inputs!$C$16)/4'
    Q[f'{c}24']=f'=IF({qq}=0,{c}23,0)';Q[f'{c}25']=f'=IF({qq}=0,0,{c}23-{c}26)'
    Q[f'{c}27']=f'=-{c}25*Inputs!$C$17'
    Q[f'{c}28']=f'=MIN({prevbal}+{c}25,MAX(0,{c}13+{c}19+{c}26+{c}27))'
    Q[f'{c}29']=f'={prevbal}+{c}25-{c}28'
    Q[f'{c}31']=f'={c}13+{c}19+{c}25+{c}26+{c}27-{c}28+IF({qq}={N},{c}12,0)'
E=L(3+N)
R['B2']="NAKHEEL GARDENS · RESULTS (SAR '000)"
res=[(5,'Gross development value',f"=SUM('Quarterly CF'!C7:{E}7)",'#,##0'),(6,'Total development cost',f"=-SUM('Quarterly CF'!C19:{E}19)",'#,##0'),
 (7,'Developer profit',f"=C5-C6+SUM('Quarterly CF'!C26:{E}27)",'#,##0'),(8,'Profit on cost','=C7/C6','0.00%'),
 (10,'Project IRR (unlevered)',f"=(1+IRR('Quarterly CF'!C21:{E}21))^4-1",'0.00%'),(11,'Equity IRR',f"=(1+IRR('Quarterly CF'!C31:{E}31))^4-1",'0.00%'),
 (12,'Peak Murabaha balance',f"=MAX('Quarterly CF'!C29:{E}29)",'#,##0'),(13,'Equity multiple',f"=SUMIF('Quarterly CF'!C31:{E}31,\">0\")/-SUMIF('Quarterly CF'!C31:{E}31,\"<0\")",'0.00"x"')]
for r,l,f,fm in res:R.cell(r,2,l);R.cell(r,3,f).number_format=fm
wb.save(f'{out}/r40_offplan_sar.xlsx')
