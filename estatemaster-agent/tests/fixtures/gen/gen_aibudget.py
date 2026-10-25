"""Golden answers for the budgeted live-AI test (tests/t_aibudget.mjs) that props.json and sens2.json do not hold: the Al Waha model
(r33_waha_sar.xlsx) with a term from its lease or term sheet applied (rent growth 4%, interest rate 6.5%), recalculated by LibreOffice.
Run: python3 gen_aibudget.py <out.json>"""
import sys, os, json, shutil, subprocess, tempfile
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__)); M = os.path.join(HERE, '..', 'models', 'r33_waha_sar.xlsx')
READ = ['Returns!B2', 'Returns!B3']   # unlevered, levered IRR

def recalc(changes):
    tmp = tempfile.mkdtemp(); src = os.path.join(tmp, 'w.xlsx'); wb = load_workbook(M)
    for ref, v in changes:
        s, c = ref.split('!'); wb[s][c] = v
    wb.save(src); out = os.path.join(tmp, 'o'); os.makedirs(out)
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', out, src], check=True, capture_output=True, timeout=600)
    wb2 = load_workbook(os.path.join(out, 'w.xlsx'), data_only=True); r = [wb2[x.split('!')[0]][x.split('!')[1]].value for x in READ]
    shutil.rmtree(tmp, ignore_errors=True); return r

G = {'_about': 'LibreOffice recalculations of r33_waha_sar.xlsx for t_aibudget: unlevered and levered IRR (Returns!B2, Returns!B3), fractions.',
     'base': recalc([]), 'escalation_4': recalc([('Inputs!B5', 0.04)]), 'margin_65': recalc([('Inputs!B10', 0.065)])}
json.dump(G, open(sys.argv[1], 'w'), indent=1); print(json.dumps(G))
