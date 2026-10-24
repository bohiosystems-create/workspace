# The meaning test (synthetic): one small workbook per kind of meaning the app's picture must carry, beside the cells themselves: each
# figure computed here independently. python3 -I gen/gen_meaning.py → models/mean_*.xlsx and golden/meaning.json
#  mean_scale_sar   a model in SAR '000 (said in the sheet heading): a profit, a peak debt and a loan balance row, read in thousands
#  mean_period_qtr  a quarterly model: acquisition and exit by quarter number, dates by EDATE(start, 3*(q-1)), a pref compounded ^(1/4)
#  mean_case_txt    a case switch that holds a word (Base / Downside) read by INDEX/MATCH over a scenario table, single inputs below it
#  mean_roles       a price that also drives the income (NOI as a yield on cost), a GFA that drives only the cost, and a rent wired with
#                   the wrong sign (the one direction error the debug must still find)
#  mean_fund_vocab  a fund's “Gross IRR” and “LP net IRR”, an FX table keyed by currency codes
#  mean_engine      engine gaps: an IRR over a name defined with OFFSET, a row picked by INDIRECT from a list of sheet names, a function the
#                   browser engine does not have (TREND, which Excel computes: 130), and an IRR Excel's 10% start guess does not reach
import openpyxl, os, json, datetime, calendar, math
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
D = os.path.join(os.path.dirname(__file__), '..', 'models'); G = {}

def irr(c, lo=-0.99, hi=10.0):
    f = lambda r: sum(x / (1 + r) ** t for t, x in enumerate(c))
    for _ in range(400):
        m = (lo + hi) / 2
        if f(lo) * f(m) <= 0: hi = m
        else: lo = m
    return (lo + hi) / 2

def xirr(c, d):
    f = lambda r: sum(x / (1 + r) ** ((dd - d[0]).days / 365) for x, dd in zip(c, d))
    lo, hi = -0.99, 10.0
    for _ in range(400):
        m = (lo + hi) / 2
        if f(lo) * f(m) <= 0: hi = m
        else: lo = m
    return (lo + hi) / 2

def eom_add(d, months):
    y, m = divmod(d.month - 1 + months, 12); y += d.year; m += 1
    return datetime.date(y, m, min(d.day, calendar.monthrange(y, m)[1]))

# 1. scale: SAR '000
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Inputs'; ws['A1'] = "AL RAWDA – DEVELOPMENT (SAR '000)"
for r in [('Land', 40000), ('Construction', 120000), ('Sales', 210000), ('Loan to cost', 0.5)]: ws.append(list(r))
ws['B5'].number_format = '0.00%'
cf = wb.create_sheet('Cash Flow'); cf['A1'] = "Cash flow (SAR '000)"; cf['A2'] = 'Year'; cf['A3'] = 'Cost'; cf['A4'] = 'Sales'; cf['A5'] = 'Loan balance'; cf['A6'] = 'Equity cash flow'
ph = [(0.6, 0, 0.5), (0.4, 0.4, 1.0), (0, 0.6, 0), (0, 0, 0), (0, 0, 0)]   # cost share, sales share, loan balance as a share of the cumulative cost (five years)
cum = 0
for i, (c, sl, lb) in enumerate(ph):
    k = L(i + 2); cf[f'{k}2'] = i + 1
    cf[f'{k}3'] = f'=-(Inputs!$B$2*{1 if i == 0 else 0}+Inputs!$B$3*{c})'; cf[f'{k}4'] = f'=Inputs!$B$4*{sl}'
    cf[f'{k}5'] = f'=-SUM($B$3:{k}3)*Inputs!$B$5*{lb}'
    prev = f'{L(i + 1)}5' if i else '0'
    cf[f'{k}6'] = f'={k}3+{k}4+({k}5-{prev})'
rt = wb.create_sheet('Returns'); rt['A1'] = "Returns (SAR '000)"
for t, f in [('Net profit', "=SUM('Cash Flow'!B3:F4)"), ('Peak debt', "=MAX('Cash Flow'!B5:F5)"), ('Levered IRR', "=IRR('Cash Flow'!B6:F6)")]: rt.append([t, f])
rt['B4'].number_format = '0.00%'
wb.save(os.path.join(D, 'mean_scale_sar.xlsx'))
cost = [40000 + 120000 * 0.6, 120000 * 0.4, 0, 0, 0]; sales = [0, 210000 * 0.4, 210000 * 0.6, 0, 0]; cumc = [sum(cost[:i + 1]) for i in range(5)]
bal = [cumc[i] * 0.5 * ph[i][2] for i in range(5)]; eq = [-cost[i] + sales[i] + (bal[i] - (bal[i - 1] if i else 0)) for i in range(5)]
G['scale'] = {'net_profit_sar': (sum(sales) - sum(cost)) * 1000, 'peak_debt_sar': max(bal) * 1000, 'balance_sar': [b * 1000 for b in bal], 'levered_irr': irr(eq), 'scale': 1000}

# 2. periods: a quarterly model
Q = 12; start = datetime.date(2027, 3, 31)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Inputs'; ws['A1'] = 'NORTH QUAY – QUARTERLY (USD)'
for r in [('First quarter end', start), ('Price', 40000000), ('Acquisition quarter', 2), ('Exit quarter', 10), ('NOI yield', 0.07), ('Exit yield', 0.065), ('Preferred return', 0.08)]: ws.append(list(r))
ws['B2'].number_format = 'dd/mm/yyyy'
cf = wb.create_sheet('Cash Flow'); cf['A1'] = 'Quarter'; cf['A2'] = 'Date'; cf['A3'] = 'Net cash flow'; cf['A4'] = 'Hurdle balance'
for q in range(1, Q + 1):
    k = L(q + 1); p = L(q)
    cf[f'{k}1'] = q; cf[f'{k}2'] = f'=EDATE(Inputs!$B$2,3*({k}1-1))'; cf[f'{k}2'].number_format = 'dd/mm/yyyy'
    cf[f'{k}3'] = f'=IF({k}1=Inputs!$B$4,-Inputs!$B$3,0)+IF(AND({k}1>Inputs!$B$4,{k}1<=Inputs!$B$5),Inputs!$B$3*Inputs!$B$6/4,0)+IF({k}1=Inputs!$B$5,Inputs!$B$3*Inputs!$B$6/Inputs!$B$7,0)'
    cf[f'{k}4'] = f"={'0' if q == 1 else p + '4'}*(1+Inputs!$B$8)^(1/4)-{k}3"
rt = wb.create_sheet('Returns'); rt.append(['Unlevered IRR', "=XIRR('Cash Flow'!B3:M3,'Cash Flow'!B2:M2)"]); rt.append(['Net profit', "=SUM('Cash Flow'!B3:M3)"])
rt['B1'].number_format = '0.00%'
wb.save(os.path.join(D, 'mean_period_qtr.xlsx'))
def qmodel(acq=2, ex=10):
    d = [eom_add(start, 3 * (q - 1)) for q in range(1, Q + 1)]; c = []
    for q in range(1, Q + 1):
        v = (-40e6 if q == acq else 0) + (40e6 * 0.07 / 4 if acq < q <= ex else 0) + (40e6 * 0.07 / 0.065 if q == ex else 0); c.append(v)
    return xirr(c, d), sum(c)
G['period'] = {'irr': qmodel()[0], 'profit': qmodel()[1], 'irr_exit_8': qmodel(ex=8)[0], 'irr_exit_12': qmodel(ex=12)[0], 'first': 1, 'last': Q}

# 3. a case switch that holds a word
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Inputs'; ws['A1'] = 'PALM COURT – CASES (USD)'; ws['A2'] = 'Case'; ws['B2'] = 'Base'
ws.append([]); ws.append(['Case inputs', 'Base', 'Downside', 'Live'])
for i, (lab, b, d) in enumerate([('Rent', 5000000, 4500000), ('Exit yield', 0.06, 0.07)]):
    r = 5 + i; ws.append([lab, b, d, f'=INDEX(B{r}:C{r},MATCH($B$2,$B$4:$C$4,0))'])
ws.append([])
for t, v in [('Price', 60000000), ('Hold (years)', 5)]: ws.append([t, v])
cf = wb.create_sheet('Cash Flow'); cf['A1'] = 'Year'; cf['A2'] = 'Net cash flow'
for y in range(6):
    k = L(y + 2); cf[f'{k}1'] = y
    cf[f'{k}2'] = '=-Inputs!$B$8' if y == 0 else f'=IF({k}1<=Inputs!$B$9,Inputs!$D$5,0)+IF({k}1=Inputs!$B$9,Inputs!$D$5/Inputs!$D$6,0)'
rt = wb.create_sheet('Returns'); rt.append(['Unlevered IRR', "=IRR('Cash Flow'!B2:G2)"]); rt.append(['Net profit', "=SUM('Cash Flow'!B2:G2)"])
wb.save(os.path.join(D, 'mean_case_txt.xlsx'))
def case(rent, ey): c = [-60e6] + [rent] * 5; c[5] += rent / ey; return irr(c), sum(c)
G['case'] = {'base_irr': case(5e6, 0.06)[0], 'down_irr': case(4.5e6, 0.07)[0], 'down_profit': case(4.5e6, 0.07)[1], 'names': ['Base', 'Downside']}

# 4. roles: directions that are and are not certain
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Inputs'; ws['A1'] = 'CEDAR PARK – ROLES (USD)'
for r in [('Purchase price', 50000000), ('NOI yield on cost', 0.08), ('GFA', 10000), ('Fit-out cost per sqm', 500), ('Rent uplift', 300000), ('Exit yield', 0.07)]: ws.append(list(r))
cf = wb.create_sheet('Cash Flow'); cf['A1'] = 'Year'; cf['A2'] = 'NOI'; cf['A3'] = 'Fit-out'; cf['A4'] = 'Rent uplift'; cf['A5'] = 'Net cash flow'
for y in range(6):
    k = L(y + 2); cf[f'{k}1'] = y
    cf[f'{k}2'] = '=0' if y == 0 else '=Inputs!$B$2*Inputs!$B$3'
    cf[f'{k}3'] = '=-Inputs!$B$4*Inputs!$B$5' if y == 0 else '=0'
    cf[f'{k}4'] = '=0' if y == 0 else '=-Inputs!$B$6'   # the planted fault: an uplift in rent taken off the cash flow
    cf[f'{k}5'] = f'=IF({k}1=0,-Inputs!$B$2,0)+{k}2+{k}3+{k}4+IF({k}1=5,{k}2/Inputs!$B$7,0)'
rt = wb.create_sheet('Returns'); rt.append(['Unlevered IRR', "=IRR('Cash Flow'!B5:G5)"]); rt.append(['Net profit', "=SUM('Cash Flow'!B5:G5)"])
wb.save(os.path.join(D, 'mean_roles.xlsx'))
G['roles'] = {'planted': 'Inputs!B6', 'not_direction': ['Inputs!B2', 'Inputs!B4']}

# 5. a fund's words
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Inputs'; ws['A1'] = 'GULF CORE FUND (USD m)'
for r in [('Committed capital (USD m)', 200), ('Management fee', 0.015), ('Deal price (SAR m)', 375)]: ws.append(list(r))
fx = wb.create_sheet('FX'); fx.append(['Currency', 'Per USD'])
for c, v in [('USD', 1), ('SAR', 3.75), ('AED', 3.6725)]: fx.append([c, v])
cf = wb.create_sheet('Cash Flow'); cf['A1'] = 'Year'; cf['A2'] = 'Gross cash flow (USD m)'; cf['A3'] = 'Fees'; cf['A4'] = 'LP net cash flow'
for y in range(6):
    k = L(y + 2); cf[f'{k}1'] = y
    cf[f'{k}2'] = '=-Inputs!$B$4/FX!$B$3' if y == 0 else ('=Inputs!$B$4/FX!$B$3*0.08' if y < 5 else '=Inputs!$B$4/FX!$B$3*(0.08+1.25)')
    cf[f'{k}3'] = '=0' if y == 0 else '=-Inputs!$B$2*Inputs!$B$3'
    cf[f'{k}4'] = f'={k}2+{k}3'
rt = wb.create_sheet('Returns'); rt.append(['Gross IRR', "=IRR('Cash Flow'!B2:G2)"]); rt.append(['LP net IRR', "=IRR('Cash Flow'!B4:G4)"]); rt.append(['Price in USD (m)', '=Inputs!B4/FX!B3'])
for r in (1, 2): rt[f'B{r}'].number_format = '0.00%'
wb.save(os.path.join(D, 'mean_fund_vocab.xlsx'))
g = [-100] + [8] * 4 + [8 + 125]; n = [g[0]] + [v - 3 for v in g[1:]]
G['fund'] = {'gross': irr(g), 'lp': irr(n)}

# 6. engine gaps
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Inputs'; ws['A1'] = 'ENGINE EDGES (USD)'
for r in [('Price', 1000), ('Rent', 120), ('Deals', None)]: ws.append(list(r))
ws['A8'] = 'Rent history'; ws['B8'] = 100; ws['C8'] = 110; ws['D8'] = 120; ws['A9'] = 'Year'; ws['B9'] = 1; ws['C9'] = 2; ws['D9'] = 3
ws['A5'] = 'Deal A'; ws['A6'] = 'Deal B'
for nmv in ['Deal A', 'Deal B']:
    s = wb.create_sheet(nmv); s['A1'] = 'Net cash flow'
    for y in range(4): s[f'{L(y + 2)}1'] = f"=IF({y}=0,-Inputs!$B$2*{1 if nmv == 'Deal A' else 0.5},Inputs!$B$3*{1 if nmv == 'Deal A' else 0.5})+IF({y}=3,Inputs!$B$2*{1 if nmv == 'Deal A' else 0.5},0)"
cf = wb.create_sheet('Cash Flow'); cf['A1'] = 'Year'; cf['A2'] = 'Deal A'; cf['A3'] = 'Deal B'; cf['A4'] = 'Total'; cf['A6'] = 'Monthly flow'
for y in range(4):
    k = L(y + 2); cf[f'{k}1'] = y
    for r, src in ((2, 'Inputs!$A$5'), (3, 'Inputs!$A$6')): cf[f'{k}{r}'] = f'=INDEX(INDIRECT("\'"&{src}&"\'!$B$1:$E$1"),1,{k}$1+1)'
    cf[f'{k}4'] = f'={k}2+{k}3'
mf = [-54433.103893, -1291.9011512, -2128.5936225, -2928.8652668, -3679.0232143, 17396.713355, -563.85212001, -1735.6607635, -2698.1138589, -3434.743567, -3932.9459493, -4184.1966248, -4184.1966248, -3932.9459493, -3434.743567, -2698.1138589, -1735.6607635, -563.85212001, 797.26210072, 2324.392857, 3991.4105182, 5769.79195, 7629.108553, 9537.5469046, 10566.666667, 10633.333333, 10700, 10766.666667, 10833.333333, 10900, 466.66666667, 533.33333333, 600, 666.66666667, 733.33333333] + [800] * 24 + [135200]
# a monthly equity flow (round 44's Al Bahr): a rate exists (3.12% a month), but Excel's IRR, from its 10% guess, needs 22 tries and stops at 20
for i, v in enumerate(mf): cf.cell(6, i + 2, v)
wb.defined_names['TOTCF'] = DefinedName('TOTCF', attr_text="OFFSET('Cash Flow'!$B$4,0,0,1,COUNT('Cash Flow'!$B$1:$E$1))")
rt = wb.create_sheet('Returns')
for t, f in [('Deal IRR', '=IRR(TOTCF)'), ('Total profit', "=SUM('Cash Flow'!B4:E4)"), ('Rent forecast', '=TREND(Inputs!B8:D8,Inputs!B9:D9,4)'), ('Monthly IRR', f"=IRR('Cash Flow'!B6:{L(len(mf) + 1)}6)")]: rt.append([t, f])
wb.save(os.path.join(D, 'mean_engine.xlsx'))
tot = [-1500 + 0, 180, 180, 180 + 1500]
G['engine'] = {'irr': irr(tot), 'profit': sum(tot), 'monthly_rate': irr(mf, -0.5, 1.0), 'trend': 130}
json.dump(G, open(os.path.join(os.path.dirname(__file__), '..', 'golden', 'meaning.json'), 'w'), indent=1); print(json.dumps(G, indent=1))
