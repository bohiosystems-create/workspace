# Round 33 fixture (synthetic): a SAR development held in a fund. Annual 2026-2035 across columns, a construction cost typed year by
# year (hardcoded schedule), an occupancy ramp typed by year, interest-only senior debt sized on cost, DSCR and LTV rows with covenant
# inputs, a hold-period input that times the sale, LP / GP inputs, returns on their own sheet.
import openpyxl,sys
from openpyxl.utils import get_column_letter as L
out=sys.argv[1]
wb=openpyxl.Workbook();i=wb.active;i.title='Inputs';i['A1']='AL WAHA BUSINESS PARK (SAR)'
rows=[('Land cost',50000000),('Construction cost (total, phased on Cash Flow)',200000000),('Stabilised rent',30000000),('Rent growth',0.03),('Operating costs (share of rent)',0.20),
 ('Exit cap rate',0.07),('Hold period (years)',10),('Loan to cost',0.55),('Interest rate',0.06),('Minimum DSCR covenant',1.25),('Maximum LTV covenant',0.60),
 ('LP share of equity',0.90),('Preferred return',0.08),('GP promote',0.20),('Discount rate',0.09),('Selling costs',0.02)]
for r,(l,v) in enumerate(rows,2):i.cell(r,1,l);i.cell(r,2,v)
for r in (5,6,7,9,10,12,13,14,15,16,17):i.cell(r,2).number_format='0.00%'
c=wb.create_sheet('Cash Flow');c['A1']='AL WAHA BUSINESS PARK – CASH FLOW (SAR)'
c['A3']='Year';c['A4']='Period';c['M3']='Total'
cons={0:-60000000,1:-90000000,2:-50000000};occ={3:0.40,4:0.75}
labels={6:'Land',7:'Construction cost',8:'Occupancy',9:'Rent',10:'Operating costs',11:'Net operating income',12:'Sale proceeds',13:'Unlevered cash flow',
 14:'Loan drawdown',15:'Opening debt',16:'Interest',17:'Loan repayment',18:'Levered cash flow to equity',19:'Closing debt',20:'DSCR',21:'Property value',22:'LTV',23:'Cumulative equity cash flow'}
for r,l in labels.items():c.cell(r,1,l)
for k in range(10):
    col=L(3+k);p=L(2+k)
    c[f'{col}3']=2026+k;c[f'{col}4']=1 if k==0 else f'={p}4+1'
    c[f'{col}6']=f'=IF({col}$4=1,-Inputs!$B$2,0)'
    c[f'{col}7']=cons.get(k,0)
    c[f'{col}8']=occ.get(k,0.95 if k>=5 else 0)
    c[f'{col}9']=f'=IF({col}$4<=Inputs!$B$8,Inputs!$B$4*{col}8*(1+Inputs!$B$5)^MAX(0,{col}$4-4),0)'
    c[f'{col}10']=f'=-{col}9*Inputs!$B$6'
    c[f'{col}11']=f'={col}9+{col}10'
    c[f'{col}12']=f'=IF({col}$4=Inputs!$B$8,{col}11*(1+Inputs!$B$5)/Inputs!$B$7*(1-Inputs!$B$17),0)'
    c[f'{col}13']=f'={col}6+{col}7+{col}11+{col}12'
    c[f'{col}14']=f'=-({col}6+{col}7)*Inputs!$B$9'
    c[f'{col}15']='=0' if k==0 else f'={p}19'
    c[f'{col}16']=f'=-{col}15*Inputs!$B$10'
    c[f'{col}17']=f'=IF({col}$4=Inputs!$B$8,-({col}15+{col}14),0)'
    c[f'{col}18']=f'={col}13+{col}14+{col}16+{col}17'
    c[f'{col}19']=f'={col}15+{col}14+{col}17'
    c[f'{col}20']=f'=IF(AND({col}11>0,{col}16<0),{col}11/-{col}16,"")'
    c[f'{col}21']=f'=IF({col}11>0,{col}11/Inputs!$B$7,"")'
    c[f'{col}22']=f'=IF(AND(ISNUMBER({col}21),{col}19>0),{col}19/{col}21,"")'
    c[f'{col}23']=f'=SUM($C18:{col}18)'
for r in (6,7,9,10,11,12,13,14,16,17,18):c[f'M{r}']=f'=SUM(C{r}:L{r})'
s=wb.create_sheet('Returns')
for r,(l,f) in enumerate([('Unlevered IRR',"=IRR('Cash Flow'!C13:L13)"),('Levered IRR',"=IRR('Cash Flow'!C18:L18)"),
    ('Equity multiple','=SUMIF(\'Cash Flow\'!C18:L18,">0")/-SUMIF(\'Cash Flow\'!C18:L18,"<0")'),('Net profit',"=SUM('Cash Flow'!C18:L18)"),
    ('Peak equity',"=-MIN('Cash Flow'!C23:L23)"),('Minimum DSCR',"=MIN('Cash Flow'!C20:L20)"),('Maximum LTV',"=MAX('Cash Flow'!C22:L22)"),
    ('NPV at the discount rate',"=NPV(Inputs!B16,'Cash Flow'!C18:L18)"),('Total development cost',"=-('Cash Flow'!M6+'Cash Flow'!M7)"),('Profit on cost','=B5/B10')],2):
    s.cell(r,1,l);s.cell(r,2,f)
wb.save(out)
