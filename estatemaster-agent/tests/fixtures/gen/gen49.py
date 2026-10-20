# Round 49 fixture (synthetic): a neighbourhood retail centre development laid out the way US retail development templates often are, with
# made-up figures. r49_retail_dev_usd: an Underwriting sheet with a "Levered IRR" banner at the top (the same figure as the returns table);
# a uses table (START · END · METHOD · /RENTABLE SF · AMOUNT) phased Straight-Line or S-Curve; a tenant roll (SUITE · TENANT · LEASE START ·
# SF · RENT START · RENT GROWTH % · RENT/SF/YR · ANNUAL RENT · STEP) where rent steps every five years or every year; recoveries pro rata to
# each suite's share of the centre; a construction loan and a permanent loan at stabilisation; a typed sale month; a Summary whose PROPERTY
# RETURNS table puts ROI, EMx and IRR side by side (the IRR is the third figure, not the first) for Unlevered and Levered rows, next to a
# PARTNERSHIP RETURNS table (LP and GP: a preferred return and a promote paid at the sale) and DEVELOPMENT RETURNS (yield on cost, spread).
import openpyxl,sys,datetime
from openpyxl.utils import get_column_letter as L
out=sys.argv[1].rstrip('/')
wb=openpyxl.Workbook();S=wb.active;S.title='Summary';U=wb.create_sheet('Underwriting')
N=108;C0=12   # months 0..N on the timeline, month 0 in column L
U['B2']='RETAIL DEVELOPMENT MODEL';U['B5']='Levered IRR';U['C5']='=H212';U['D5']='Levered EMx:';U['E5']='=G212';U['H5']='Development Spread (Untrended)';U['I5']='=H217'
U['C5'].number_format='0.00%';U['E5'].number_format='0.00"x"';U['I5'].number_format='0" bps"'
U['B13']='Analysis Begin Date';U['E13']=datetime.datetime(2027,3,1);U['E13'].number_format='mmm-yy'
U['F15']='Sale Month';U['I15']=96;U['B17']='Land Area (SF)';U['E17']=520000;U['B18']='Net Rentable Area (SF)';U['E18']='=E105'
# ---- uses ----
U['B25']='USES';U['B26']='Land Costs';U['B29']='Hard Costs';U['D25']='START';U['E25']='END';U['F25']='METHOD';U['H25']='/RENTABLE SF';U['I25']='AMOUNT'
uses=((27,'Land Purchase Price',0,0,'Straight-Line',None,9800000),(28,'Closing / Legal Costs',0,0,'Straight-Line',None,'=I27*0.01'),
      (30,'Site Work',2,8,'Straight-Line',None,'=H30*E18'),(31,'Building Shell',4,16,'S-Curve',None,'=H31*E18'),(32,'Tenant Improvements',12,20,'S-Curve',None,'=H32*E18'),
      (33,'Soft Costs',1,20,'Straight-Line',None,'=0.09*SUM(I30:I32)'),(34,'Developer Fee',1,20,'Straight-Line',None,'=0.04*SUM(I27:I33)'),(35,'Contingency',4,20,'S-Curve',None,'=0.05*SUM(I30:I32)'))
for r,l,s,e,m,_,amt in uses:
    U[f'B{r}']=l;U[f'D{r}']=s;U[f'E{r}']=e;U[f'F{r}']=m;U[f'I{r}']=amt;U[f'I{r}'].number_format='#,##0'
    if not str(amt).startswith('=H'):U[f'H{r}']=f'=I{r}/$E$18'
for r,v in ((30,14),(31,118),(32,42)):U[f'H{r}']=v
U['B36']='Total Project Cost';U['I36']='=SUM(I27:I35)';U['I36'].number_format='#,##0'
# ---- financing ----
U['B40']='Construction Loan to Cost';U['E40']=0.6;U['B41']='Construction Loan Rate';U['E41']=0.0775
U['B42']='Permanent Loan LTV';U['E42']=0.6;U['B43']='Permanent Loan Rate';U['E43']=0.0625;U['B44']='Permanent Loan Month';U['E44']=30
for a in ('E40','E41','E42','E43'):U[a].number_format='0.00%'
# ---- tenant roll ----
U['B91']='RETAIL BASE RENTAL INCOME'
for j,h in enumerate(['SUITE','TENANT','LEASE START','SF','RENT START','RENT GROWTH %','RENT/SF/YR','ANNUAL RENT','','STEP']):U.cell(92,2+j,h) if h else None
ten=(('Anchor-1','Greenfield Grocers',18,52000,0.10,21,'Every 5 Years'),('101','Spec',20,4200,0.03,35,'Annual'),('102','Spec',22,3600,0.03,35,'Annual'),
     ('103','Spec',24,3600,0.03,34,'Annual'),('104','Spec',26,2800,0.03,37,'Annual'),('201','Spec',22,6000,0.03,32,'Annual'),('202','Spec',28,5200,0.03,32,'Annual'),
     ('Outparcel-1','Burger Drive-Thru',16,3200,0.10,58,'Every 5 Years'),('Outparcel-2','Corner Coffee',18,2400,0.10,64,'Every 5 Years'))
for r,(su,tn,ls,sf,g,rent,step) in enumerate(ten,93):
    U[f'B{r}']=su;U[f'C{r}']=tn;U[f'D{r}']=ls;U[f'E{r}']=sf;U[f'F{r}']=f'=D{r}+3';U[f'G{r}']=g;U[f'H{r}']=rent;U[f'I{r}']=f'=H{r}*E{r}';U[f'K{r}']=step
    U[f'G{r}'].number_format='0.00%';U[f'E{r}'].number_format='#,##0';U[f'I{r}'].number_format='#,##0'
U['B105']='TOTAL RENTAL INCOME';U['E105']='=SUM(E93:E101)';U['I105']='=SUM(I93:I101)';U['H105']='=I105/E105'
U['B120']='Operating Expenses / SF / Yr';U['E120']=7.4;U['B121']='Expense Growth';U['E121']=0.03;U['B122']='Recovery %';U['E122']=0.9
U['B123']='General Vacancy';U['E123']=0.05;U['B124']='Exit Cap Rate';U['E124']=0.07;U['B125']='Selling Costs';U['E125']=0.02;U['B126']='Market Cap Rate';U['E126']=0.0675
for a in ('E121','E122','E123','E124','E125','E126'):U[a].number_format='0.00%'
# ---- partnership terms ----
U['B130']='LP Share of Equity';U['E130']=0.9;U['B131']='Preferred Return (simple, a year)';U['E131']=0.08;U['B132']='GP Promote above the Pref';U['E132']=0.2
for a in ('E130','E131','E132'):U[a].number_format='0.00%'
# ---- monthly timeline ----
lab={3:'Analysis Month',24:'Analysis Date',140:'Rental Income',141:'Recovery Income',142:'General Vacancy',143:'Effective Gross Revenue',144:'Operating Expenses',145:'Net Operating Income',
     150:'Development Cost',151:'Construction Loan Draw',152:'Construction Interest',153:'Construction Loan Balance',154:'Construction Loan Payoff',
     155:'Permanent Loan Funding',156:'Permanent Loan Debt Service',157:'Permanent Loan Payoff',160:'Gross Sale Price',161:'Selling Costs',
     211:'Net Unlevered Cash Flow',212:'Net Levered Cash Flow',228:'GP',229:'LP'}
for r,l in lab.items():U.cell(r,2,l)
E=L(C0+N);live=lambda c:f'({c}$3<=$I$15)'
for k in range(N+1):
    c=L(C0+k);p=L(C0+k-1)
    U[f'{c}3']=k;U[f'{c}24']=f'=EOMONTH($E$13,{c}3-1)';U[f'{c}24'].number_format='mmm-yy'
    rent='+'.join(f'(({c}$3>=$F${r})*$I${r}/12*(1+$G${r})^IF($K${r}="Annual",INT(MAX(0,{c}$3-$F${r})/12),5*INT(MAX(0,{c}$3-$F${r})/60)))' for r in range(93,102))
    U[f'{c}140']=f'={live(c)}*({rent})'
    occ='+'.join(f'({c}$3>=$F${r})*$E${r}' for r in range(93,102))
    U[f'{c}144']=f'={live(c)}*({c}$3>=MIN($F$93:$F$101))*$E$120*$E$105/12*(1+$E$121)^INT({c}$3/12)'
    U[f'{c}141']=f'={c}144*$E$122*({occ})/$E$105'
    U[f'{c}142']=f'=-$E$123*({c}140+{c}141)';U[f'{c}143']=f'=SUM({c}140:{c}142)';U[f'{c}145']=f'={c}143-{c}144'
    cost='+'.join(f'IF(AND({c}$3>=$D${r},{c}$3<=$E${r}),$I${r}*IF($F${r}="S-Curve",(COS(PI()*({c}$3-$D${r})/($E${r}-$D${r}+1))-COS(PI()*({c}$3-$D${r}+1)/($E${r}-$D${r}+1)))/2,1/($E${r}-$D${r}+1)),0)' for r in range(27,36) if r!=29)
    U[f'{c}150']=f'={cost}';U[f'{c}151']=f'={c}150*$E$40'
    pb='0' if k==0 else f'{p}153'
    U[f'{c}152']='=0' if k==0 else f'=IF({c}$3<=$E$44,{p}153*$E$41/12,0)'
    U[f'{c}154']=f'=IF({c}$3=$E$44,{pb}+{c}151,0)';U[f'{c}153']=f'={pb}+{c}151-{c}154'
    U[f'{c}155']=f'=IF({c}$3=$E$44,$E$42*12*SUMPRODUCT(($L$3:${E}$3>$E$44)*($L$3:${E}$3<=$E$44+12)*$L$145:${E}$145)/12/$E$126,0)'
    U[f'{c}156']=f'=IF(AND({c}$3>$E$44,{live(c)}),SUM($L$155:${E}$155)*$E$43/12,0)';U[f'{c}157']=f'=IF({c}$3=$I$15,SUM($L$155:${E}$155),0)'
    U[f'{c}160']=f'=IF({c}$3=$I$15,SUMPRODUCT(($L$3:${E}$3>$I$15-12)*($L$3:${E}$3<=$I$15)*$L$145:${E}$145)*(1+$E$121)/$E$124,0)';U[f'{c}161']=f'=-{c}160*$E$125'
    U[f'{c}211']=f'={live(c)}*(-{c}150+{c}145+{c}160+{c}161)'
    U[f'{c}212']=f'={live(c)}*({c}211+{c}151-{c}152-{c}154+{c}155-{c}156-{c}157)'
    # partners: equity in pro rata; distributions pro rata, the GP promote paid at the sale
    U[f'{c}229']=f'={c}212*$E$130-IF({c}$3=$I$15,$I$233,0)';U[f'{c}228']=f'={c}212*(1-$E$130)+IF({c}$3=$I$15,$I$233,0)'
    for r in (140,141,142,143,144,145,150,151,152,153,154,155,156,157,160,161,211,212,228,229):U[f'{c}{r}'].number_format='#,##0;(#,##0)'
for r in (211,212,228,229):
    U[f'G{r}']=f'=SUMIF($L${r}:${E}${r},">0")/-SUMIF($L${r}:${E}${r},"<0")';U[f'H{r}']=f'=XIRR(OFFSET($L${r},0,0,1,$I$15+1),OFFSET($L$24,0,0,1,$I$15+1))';U[f'I{r}']=f'=SUM($L${r}:${E}${r})'
    U[f'G{r}'].number_format='0.00"x"';U[f'H{r}'].number_format='0.00%';U[f'I{r}'].number_format='#,##0'
U['G210']='EMx';U['H210']='IRR';U['I210']='Profit'
U['B215']='Yield-on-Cost';U['H215']=f'=SUMPRODUCT(($L$3:${E}$3>$E$44)*($L$3:${E}$3<=$E$44+12)*$L$145:${E}$145)/I36';U['B216']='Market Cap Rate';U['H216']='=E126'
U['B217']='Development Spread';U['H217']='=(H215-H216)*10000';U['H215'].number_format='0.00%';U['H216'].number_format='0.00%'
U['B232']='Equity Contributed';U['I232']=f'=-SUMIF($L$212:${E}$212,"<0")';U['B233']='GP Promote at Sale'
U['I233']=f'=MAX(0,$E$132*($I$212-$E$131*$I$232*$I$15/12))';U['I232'].number_format='#,##0';U['I233'].number_format='#,##0'
# ---- summary ----
S['B2']='=Underwriting!B2';S['B4']='INVESTMENT DESCRIPTION';S['B6']='Net Rentable Area (SF)';S['F6']='=Underwriting!E18';S['B11']='Sale Month';S['F11']='=Underwriting!I15'
S['N20']='PARTNERSHIP RETURNS';S['Q20']='ROI';S['R20']='EMx ';S['S20']='IRR'
S['N21']='LP';S['Q21']='=Underwriting!I229/-SUMIF(Underwriting!L229:'+E+'229,"<0")';S['R21']='=Underwriting!G229';S['S21']='=Underwriting!H229'
S['N22']='GP';S['Q22']='=Underwriting!I228/-SUMIF(Underwriting!L228:'+E+'228,"<0")';S['R22']='=Underwriting!G228';S['S22']='=Underwriting!H228'
S['N24']='PROPERTY RETURNS';S['Q24']='ROI';S['R24']='EMx ';S['S24']='IRR'
S['N25']='Unlevered';S['Q25']='=Underwriting!I211/Underwriting!I36';S['R25']='=Underwriting!G211';S['S25']='=Underwriting!H211'
S['N26']='Levered';S['Q26']='=Underwriting!I212/Underwriting!I232';S['R26']='=Underwriting!G212';S['S26']='=Underwriting!H212'
S['N27']='DEVELOPMENT RETURNS';S['N28']='Yield-on-Cost';S['S28']='=Underwriting!H215';S['N29']='Market Cap Rate';S['S29']='=Underwriting!H216';S['N30']='Development Spread';S['S30']='=Underwriting!H217'
for r in (21,22,25,26):S[f'Q{r}'].number_format='0.00%';S[f'R{r}'].number_format='0.00"x"';S[f'S{r}'].number_format='0.00%'
S['S28'].number_format='0.00%';S['S29'].number_format='0.00%';S['S30'].number_format='0" bps"'
S['H31']='USES';S['M31']='Amount';S['H32']='Total Project Cost';S['M32']='=Underwriting!I36';S['M32'].number_format='#,##0'
wb.save(f'{out}/r49_retail_dev_usd.xlsx')
