"""Arabic fixtures (t_arabic), synthetic: a Riyadh residential development model written in Arabic (Arabic sheet names with spaces, Arabic
labels, SAR, a scenario switch whose cases are named in Arabic, Real Estate Transaction Tax on the land) and the murabaha facility agreement
that finances it, in Arabic, as a Word file and as a PDF. Nothing from a client. The agreement against the model, on purpose:
  match     loan to cost at most 50% (the model's input is 50%)
  differ    profit rate 7.5% a year (the model 7%), written with Arabic-Indic digits (٧٫٥٪); final repayment by 30 Jumada al-Akhirah 1452 AH
            (27 October 2030), where the model repays in 2031
  breach    the facility balance with its rolled-up profit must stay within 50% of the total cost every year: the model's own row goes above it
  absent    an arrangement fee of 1% of the facility (not in the model)
  itself    signed 4 Rabi al-Awwal 1448 AH, paired with 18 August 2026: by Umm al-Qura that day is 17 August 2026
Run: python3 gen_arabic.py <out_dir_docs> <out_dir_models>"""
import sys, os, shutil, subprocess, tempfile, zipfile
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L

DOCS, MODELS = sys.argv[1], sys.argv[2]

# ---- the model ----
wb = Workbook()
i = wb.active; i.title = 'المدخلات'
i['A1'] = 'مشروع حي النرجس السكني (ريال سعودي)'
rows = [('السيناريو (1 = أساسي، 2 = متحفظ)', 1), ('سعر البيع للمتر المربع · أساسي', 6500), ('سعر البيع للمتر المربع · متحفظ', 5800),
        ('المساحة القابلة للبيع (م²)', 40000), ('تكلفة البناء للمتر المربع', 3200), ('تكلفة الأرض', 60000000), ('ضريبة التصرفات العقارية', 0.05),
        ('رسوم إدارة التطوير (من تكلفة البناء)', 0.03), ('نسبة التمويل إلى التكلفة', 0.5), ('معدل ربح المرابحة', 0.07), ('سنة السداد', 5)]
for r, (k, v) in enumerate(rows, start=3):
    i.cell(r, 1, k); i.cell(r, 2, v)
    if isinstance(v, float) and v < 1: i.cell(r, 2).number_format = '0.00%'
I = {k: f"'المدخلات'!$B${r}" for r, (k, v) in enumerate(rows, start=3)}
SC, PB, PD, AREA, BC, LAND, RETT, DM, LTC, RATE, REPAY = [I[k] for k, _ in rows]

N = 5
c = wb.create_sheet('التدفقات النقدية'); c['A1'] = 'التدفقات النقدية (ريال سعودي)'; c['A3'] = 'السنة'; c['A4'] = 'الفترة'
lab = {6: 'تكلفة الأرض', 7: 'ضريبة التصرفات العقارية', 8: 'تكاليف البناء', 9: 'رسوم إدارة التطوير', 10: 'إجمالي التكاليف',
       11: 'إيرادات المبيعات', 12: 'صافي التدفق النقدي للمشروع', 13: 'سحب التمويل', 14: 'ربح المرابحة المرسمل', 15: 'سداد التمويل',
       16: 'رصيد التمويل', 17: 'رصيد التمويل إلى إجمالي التكلفة', 18: 'صافي التدفق النقدي لحقوق الملكية'}
for r, t in lab.items(): c.cell(r, 1, t)
build = [0.3, 0.45, 0.25, 0, 0]; sales = [0, 0, 0.3, 0.4, 0.3]
TOT = f"(-SUM($C$10:$G$10))"
for k in range(N):
    col = L(3 + k); p = L(2 + k)
    c[f'{col}3'] = 2027 + k; c[f'{col}4'] = 1 if k == 0 else f'={p}4+1'
    c[f'{col}6'] = f'=IF({col}$4=1,-{LAND},0)'
    c[f'{col}7'] = f'={col}6*{RETT}'
    c[f'{col}8'] = f'=-{AREA}*{BC}*{build[k]}'
    c[f'{col}9'] = f'={col}8*{DM}'
    c[f'{col}10'] = f'=SUM({col}6:{col}9)'
    c[f'{col}11'] = f'={AREA}*CHOOSE({SC},{PB},{PD})*{sales[k]}'
    c[f'{col}12'] = f'={col}10+{col}11'
    c[f'{col}13'] = f'=-{col}10*{LTC}'
    c[f'{col}14'] = ('=0' if k == 0 else f'={p}16*{RATE}')
    c[f'{col}15'] = f'=IF({col}$4={REPAY},-({p}16+{col}13+{col}14),0)' if k else f'=IF({col}$4={REPAY},-({col}13+{col}14),0)'
    c[f'{col}16'] = (f'={col}13+{col}14+{col}15' if k == 0 else f'={p}16+{col}13+{col}14+{col}15')
    c[f'{col}17'] = f'={col}16/{TOT}'
    c[f'{col}18'] = f'={col}12+{col}13+{col}15'
    c[f'{col}17'].number_format = '0.00%'
c['H3'] = 'المجموع'
for r in (6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 18): c[f'H{r}'] = f'=SUM(C{r}:G{r})'
s = wb.create_sheet('العوائد'); s['A1'] = 'العوائد'
for r, (k, f) in enumerate([('معدل العائد الداخلي بعد التمويل', "=IRR('التدفقات النقدية'!C18:G18)"), ('معدل العائد الداخلي قبل التمويل', "=IRR('التدفقات النقدية'!C12:G12)"),
                            ('صافي الربح', "='التدفقات النقدية'!H18"), ('العائد على التكلفة', "='التدفقات النقدية'!H12/-'التدفقات النقدية'!H10"),
                            ('مضاعف حقوق الملكية', "=SUMIF('التدفقات النقدية'!C18:G18,\">0\")/-SUMIF('التدفقات النقدية'!C18:G18,\"<0\")")], start=3):
    s.cell(r, 1, k); s.cell(r, 2, f)
for r in (3, 4, 6): s.cell(r, 2).number_format = '0.00%'
tmp = tempfile.mkdtemp(); src = os.path.join(tmp, 'ar_dev.xlsx'); wb.save(src)
out = os.path.join(tmp, 'o'); os.makedirs(out)
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', out, src], check=True, capture_output=True, timeout=600)
os.makedirs(MODELS, exist_ok=True); shutil.copy(os.path.join(out, 'ar_dev.xlsx'), os.path.join(MODELS, 'ar_dev.xlsx')); shutil.rmtree(tmp, ignore_errors=True)

# ---- the agreement ----
P = [
 (None, 'اتفاقية تسهيلات مرابحة'), (None, 'بين بنك الريادة (المموّل) وشركة النرجس للتطوير العقاري (العميل)'),
 (None, '(نموذج اصطناعي لأغراض الاختبار، وليس اتفاقية حقيقية)'),
 (0, 'التعريفات. يكون للكلمات والعبارات الآتية المعاني المبينة أمام كل منها:'),
 (1, 'المشروع: مشروع حي النرجس السكني في مدينة الرياض بمساحة قابلة للبيع قدرها ٤٠٬٠٠٠ متر مربع.'),
 (1, 'إجمالي تكلفة المشروع: مجموع تكلفة الأرض وضريبة التصرفات العقارية وتكاليف البناء ورسوم إدارة التطوير.'),
 (0, 'التسهيل.'),
 (1, 'مبلغ التسهيل: يلتزم المموّل بتوفير تسهيل مرابحة لا يتجاوز مبلغه خمسين بالمائة (50%) من إجمالي تكلفة المشروع.'),
 (1, 'هامش الربح: يستحق للمموّل ربح المرابحة بمعدل ٧٫٥٪ سنوياً على الرصيد القائم، ويُرسمل الربح على رصيد التسهيل حتى السداد.'),
 (1, 'رسوم الترتيب: يدفع العميل للمموّل رسوم ترتيب بنسبة واحد بالمائة (1%) من مبلغ التسهيل عند التوقيع.'),
 (0, 'السداد والتعهدات.'),
 (1, 'تاريخ السداد النهائي: يلتزم العميل بسداد كامل رصيد التسهيل والأرباح المستحقة في موعد أقصاه ٣٠ جمادى الآخرة ١٤٥٢هـ.'),
 (1, 'تعهد الرصيد: يتعهد العميل بألا يتجاوز رصيد التسهيل شاملاً الأرباح المرسملة خمسين بالمائة (50%) من إجمالي تكلفة المشروع في أي وقت.'),
 (0, 'أحكام عامة.'),
 (1, 'القانون الواجب التطبيق: تخضع هذه الاتفاقية للأنظمة المعمول بها في المملكة العربية السعودية وأحكام الشريعة الإسلامية.'),
 (1, 'تاريخ الاتفاقية: حُررت هذه الاتفاقية بتاريخ ١٤٤٨/٠٣/٠٤هـ الموافق ٢٠٢٦/٠٨/١٨م من نسختين أصليتين.'),
]
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
def para(lv, t):
    num = '' if lv is None else f'<w:numPr><w:ilvl w:val="{lv}"/><w:numId w:val="1"/></w:numPr>'
    return f'<w:p><w:pPr>{num}<w:bidi/></w:pPr><w:r><w:rPr><w:rtl/></w:rPr><w:t xml:space="preserve">{escape(t)}</w:t></w:r></w:p>'
body = ''.join(para(lv, t) for lv, t in P)
document = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body>{body}<w:sectPr/></w:body></w:document>'
lvls = [('decimal', '%1.'), ('decimal', '%1.%2')]
numbering = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:numbering {W}><w:abstractNum w:abstractNumId="0">'
             + ''.join(f'<w:lvl w:ilvl="{k}"><w:start w:val="1"/><w:numFmt w:val="{f}"/><w:lvlText w:val="{t}"/></w:lvl>' for k, (f, t) in enumerate(lvls))
             + '</w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>')
ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
      '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
      '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
      '<Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/></Types>')
rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
drels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
         '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/></Relationships>')
os.makedirs(DOCS, exist_ok=True)
with zipfile.ZipFile(os.path.join(DOCS, 'murabaha_ar.docx'), 'w', zipfile.ZIP_DEFLATED) as z:
    z.writestr('[Content_Types].xml', ct); z.writestr('_rels/.rels', rels); z.writestr('word/document.xml', document)
    z.writestr('word/_rels/document.xml.rels', drels); z.writestr('word/numbering.xml', numbering)
# the same agreement as a PDF, laid out by LibreOffice (right to left, Arabic shaped by its own text engine): extracted text from such PDFs
# carries the usual Arabic artefacts (ligatures split, letters inside a word and digits inside a number in visual order)
from openpyxl.styles import Alignment
pw = Workbook(); ps = pw.active; ps.title = 'الاتفاقية'; ps.sheet_view.rightToLeft = True; ps.column_dimensions['A'].width = 120; n = [0, 0]
for r, (lv, t) in enumerate(P, start=1):
    if lv == 0: n[0] += 1; n[1] = 0; t = f'{n[0]}. {t}'
    elif lv == 1: n[1] += 1; t = f'{n[0]}.{n[1]} {t}'
    ps.cell(r, 1, t).alignment = Alignment(horizontal='right', wrap_text=True)
tmp = tempfile.mkdtemp(); src = os.path.join(tmp, 'murabaha_ar.xlsx'); pw.save(src)
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, src], check=True, capture_output=True, timeout=600)
shutil.copy(os.path.join(tmp, 'murabaha_ar.pdf'), os.path.join(DOCS, 'murabaha_ar.pdf')); shutil.rmtree(tmp, ignore_errors=True)
print('written ar_dev.xlsx, murabaha_ar.docx, murabaha_ar.pdf')
