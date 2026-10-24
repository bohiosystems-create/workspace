# The picture test (synthetic): one small workbook per Excel feature the app's picture of a workbook must hold in full, each with its
# figures computed independently here. python3 -I gen/gen_picture.py → models/pic_*.xlsx and golden/picture.json
#  pic_loop_sar     a loan whose interest is funded by the loan (a circular reference, iterative calculation on)
#  pic_table_usd    an Excel table read by structured references (two tenants with the same name)
#  pic_spill_usd    a timeline spilled by SEQUENCE, read by the rows under it
#  pic_dates_usd    dates typed as plain serial numbers (General format) read by XNPV
#  pic_xarr_sar     a cross-sheet array formula (SUM of a product of two ranges on another sheet)
#  pic_merged_sar   an inputs block under a merged two-row header
#  pic_names_usd    defined names for the inputs and for the cash-flow range the IRR reads
#  pic_extlink_usd  a headline that reads another workbook (an external link the file does not carry)
import openpyxl, os, json, datetime
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
D = os.path.join(os.path.dirname(__file__), '..', 'models'); G = {}

def irr(c):
    lo, hi = -0.99, 10.0
    f = lambda r: sum(x / (1 + r) ** t for t, x in enumerate(c))
    for _ in range(300):
        m = (lo + hi) / 2
        if f(lo) * f(m) <= 0: hi = m
        else: lo = m
    return (lo + hi) / 2

def book(title, rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Inputs'; ws.append([title, None])
    for r in rows: ws.append(list(r))
    return wb

def returns(wb, items):
    r = wb.create_sheet('Returns')
    for t, f in items: r.append([t, f])

# 1. the loop: interest on the average balance, capitalised into the loan it is funded by
wb = book('LOOP – CAPITALISED INTEREST (SAR)', [('Loan', 100000000), ('Interest rate', 0.08), ('Fee', 0.02)])
lp = wb.create_sheet('Loop'); lp['A1'] = 'Drawn'; lp['B1'] = '=Inputs!B2'; lp['A2'] = 'Interest'; lp['B2'] = '=Inputs!B3*(B1+B2/2)'
lp['A3'] = 'Fee'; lp['B3'] = '=Inputs!B4*B1'; lp['A4'] = 'Total debt'; lp['B4'] = '=B1+B2+B3'
lp['A5'] = 'Covenant'; lp['B5'] = '=IF(B4>105000000,"Breached","Within")'   # a text result downstream of the loop
returns(wb, [('Total debt', '=Loop!B4'), ('Interest cost', '=Loop!B2'), ('Covenant', '=Loop!B5')])
wb.calculation = CalcProperties(iterate=True, iterateCount=200, iterateDelta=1e-9, fullCalcOnLoad=True)
wb.save(os.path.join(D, 'pic_loop_sar.xlsx'))
I = 0.08 * 100e6 / (1 - 0.04); G['loop'] = {'interest': I, 'total_debt': 100e6 + I + 0.02 * 100e6, 'circ': ['Loop!B2'], 'covenant': 'Breached'}

# 2. the table
wb = book('TABLE – LEASE SCHEDULE (USD)', [('Rent growth', 0.03)])
ls = wb.create_sheet('Leases'); ls.append(['Tenant', 'Area', 'Rent psm'])
leases = [('Anchor', 12000, 180), ('Gym', 3000, 150), ('Anchor', 4000, 175), ('Cafe', 800, 320)]
for t in leases: ls.append(list(t))
tab = Table(displayName='Leases', ref='A1:C5'); tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2'); ls.add_table(tab)
returns(wb, [('Total rent', '=SUMPRODUCT(Leases[Area],Leases[Rent psm])'), ('Total area', '=SUM(Leases[Area])'), ('Rent next year', '=SUMPRODUCT(Leases[Area],Leases[Rent psm])*(1+Inputs!B2)')])
wb.save(os.path.join(D, 'pic_table_usd.xlsx'))
tr = sum(a * r for _, a, r in leases); G['table'] = {'total_rent': tr, 'total_area': sum(a for _, a, _ in leases), 'rent_next': tr * 1.03, 'rows': len(leases)}

# 3. the spill
wb = book('SPILL – TIMELINE BY SEQUENCE (USD)', [('Rent year 1', 100000), ('Rent growth', 0.03), ('Discount rate', 0.08)])
cf = wb.create_sheet('Cash Flow'); cf['A1'] = 'Year'; cf['B1'] = '=_xlfn.SEQUENCE(1,8)'; cf['A2'] = 'Rent'
for i in range(8): cf[f'{L(i + 2)}2'] = f'=Inputs!$B$2*(1+Inputs!$B$3)^({L(i + 2)}1-1)'
returns(wb, [('Total rent', "=SUM('Cash Flow'!B2:I2)"), ('NPV', "=NPV(Inputs!B4,'Cash Flow'!B2:I2)")])
wb.save(os.path.join(D, 'pic_spill_usd.xlsx'))
rent = [100000 * 1.03 ** i for i in range(8)]; G['spill'] = {'total_rent': sum(rent), 'npv': sum(x / 1.08 ** (t + 1) for t, x in enumerate(rent)), 'spilled': ['Cash Flow!' + L(i + 2) + '1' for i in range(1, 8)]}

# 4. dates by value: plain serial numbers, no date format
wb = book('DATES – SERIALS IN GENERAL FORMAT (USD)', [('Discount rate', 0.09), ('Price', 50000000), ('Rent', 4000000)])
cf = wb.create_sheet('Cash Flow'); cf['A1'] = 'Date'; cf['A2'] = 'Net cash flow'
ds = [datetime.date(2026 + y, 12, 31) for y in range(6)]; ser = lambda d: (d - datetime.date(1899, 12, 30)).days
for i, d in enumerate(ds):
    k = L(i + 2); cf[f'{k}1'] = ser(d); cf[f'{k}2'] = '=-Inputs!$B$3' if i == 0 else ('=Inputs!$B$4' if i < 5 else '=Inputs!$B$4+Inputs!$B$3*1.1')
returns(wb, [('NPV at the discount rate', "=XNPV(Inputs!B2,'Cash Flow'!B2:G2,'Cash Flow'!B1:G1)")])
wb.save(os.path.join(D, 'pic_dates_usd.xlsx'))
c = [-50e6] + [4e6] * 4 + [4e6 + 55e6]; G['dates'] = {'xnpv': sum(x / 1.09 ** ((d - ds[0]).days / 365) for x, d in zip(c, ds)), 'dateCells': ['Cash Flow!' + L(i + 2) + '1' for i in range(6)]}

# 5. a cross-sheet array formula
wb = book('ARRAYS – RENT ROLL ACROSS SHEETS (SAR)', [('Vacancy', 0.05)])
rr = wb.create_sheet('Rent roll'); rr.append(['Unit', 'Count', 'Rent per unit'])
units = [('Studio', 40, 30000), ('1 bed', 60, 42000), ('2 bed', 30, 58000), ('Penthouse', 4, 150000)]
for u in units: rr.append(list(u))
cf = wb.create_sheet('Cash Flow'); cf['A1'] = 'Gross rent'; cf['B1'] = ArrayFormula('B1', "=SUM('Rent roll'!B2:B5*'Rent roll'!C2:C5)")
cf['A2'] = 'Net rent'; cf['B2'] = '=B1*(1-Inputs!B2)'
returns(wb, [('Net rent', "='Cash Flow'!B2"), ('Units', "=SUM('Rent roll'!B2:B5)")])
wb.save(os.path.join(D, 'pic_xarr_sar.xlsx'))
gr = sum(n * r for _, n, r in units); G['xarr'] = {'gross': gr, 'net': gr * 0.95, 'units': sum(n for _, n, _ in units)}

# 6. a merged two-row header over the inputs
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Inputs'; ws.merge_cells('A1:C1'); ws['A1'] = 'MERGED – KEY ASSUMPTIONS (SAR)'
ws.append(['Item', 'Value', 'Unit'])
for r in [('GFA', 20000, 'sqm'), ('Cost per sqm', 4500, 'SAR'), ('Rent per sqm', 900, 'SAR'), ('Exit yield', 0.07, '%')]: ws.append(list(r))
returns(wb, [('Total cost', '=Inputs!B3*Inputs!B4'), ('Value on completion', '=Inputs!B3*Inputs!B5/Inputs!B6'), ('Profit on cost', '=Returns!B2/Returns!B1-1')])
wb.save(os.path.join(D, 'pic_merged_sar.xlsx'))
G['merged'] = {'cost': 20000 * 4500, 'value': 20000 * 900 / 0.07, 'poc': (20000 * 900 / 0.07) / (20000 * 4500) - 1}

# 7. defined names
wb = book('NAMES – HOLD AND SELL (USD)', [('Price', 30000000), ('Rent', 2400000), ('Growth', 0.02), ('Exit yield', 0.065)])
for n, ref in [('Price', 'Inputs!$B$2'), ('Rent', 'Inputs!$B$3'), ('Growth', 'Inputs!$B$4'), ('ExitYield', 'Inputs!$B$5'), ('CF', "'Cash Flow'!$B$2:$G$2")]:
    wb.defined_names[n] = DefinedName(n, attr_text=ref)
cf = wb.create_sheet('Cash Flow'); cf['A1'] = 'Year'; cf['A2'] = 'Net cash flow'
for i in range(6):
    k = L(i + 2); cf[f'{k}1'] = i
    cf[f'{k}2'] = '=-Price' if i == 0 else (f'=Rent*(1+Growth)^({k}1-1)' + ('+Rent*(1+Growth)^5/ExitYield' if i == 5 else ''))
returns(wb, [('Unlevered IRR', '=IRR(CF)'), ('Net profit', '=SUM(CF)')])
wb.save(os.path.join(D, 'pic_names_usd.xlsx'))
c = [-30e6] + [2.4e6 * 1.02 ** (y - 1) for y in range(1, 6)]; c[5] += 2.4e6 * 1.02 ** 5 / 0.065
G['names'] = {'irr': irr(c), 'profit': sum(c)}

# 8. an external link the file does not carry
wb = book('EXTERNAL – MARKET RENT FROM ANOTHER FILE (USD)', [('Area', 10000)])
returns(wb, [('Market rent', "=Inputs!B2*'[Market.xlsx]Rents'!$B$2"), ('Area', '=Inputs!B2')])
wb.save(os.path.join(D, 'pic_extlink_usd.xlsx'))
G['extlink'] = {'area': 10000}

json.dump(G, open(os.path.join(os.path.dirname(__file__), '..', 'golden', 'picture.json'), 'w'), indent=1); print(json.dumps(G, indent=1))
