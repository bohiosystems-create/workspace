"""Sensitivity test model (t_sens2), synthetic: a SAR logistics park built over two years and held, ten annual periods 2027-2036.
Planted for the sensitivity tools: a debt cap (Debt drawn = MIN(loan to cost x cost, cap)) that binds near today's construction cost; a cash
trap when the DSCR falls under its covenant (IF ... MAX(0, ...)) that switches on when rent falls; the levered IRR wrapped in
IFERROR(..., 0), so a deep downside shows 0; a scenario switch (1 = Base, 2 = Downside) choosing between two rent lines; an occupancy
ramp typed year by year; an exit year two periods before the model's end; professional fees written as a number inside a formula; rent
and exit yield multiplied together in the sale price (they interact). Saved through LibreOffice so every formula carries its result.
Run: python3 gen_sens2.py <out.xlsx>"""
import sys, os, shutil, subprocess, tempfile
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L

wb = Workbook()
i = wb.active; i.title = 'Inputs'
i['A1'] = 'RIVERSIDE LOGISTICS PARK (SAR)'
rows = [('Scenario (1 = Base, 2 = Downside)', 1), ('Rent per sqm · Base', 150), ('Rent per sqm · Downside', 120), ('Lettable area (sqm)', 400000),
        ('Rent growth', 0.03), ('Operating costs (share of rent)', 0.15), ('Construction cost', 600000000), ('Exit yield', 0.07), ('Exit year', 8),
        ('Interest rate', 0.06), ('Loan to cost', 0.55), ('Debt cap', 340000000), ('Minimum DSCR covenant', 1.30)]
for r, (k, v) in enumerate(rows, start=2):
    i.cell(r, 1, k); i.cell(r, 2, v)
for r in (6, 7, 9, 11, 12): i.cell(r, 2).number_format = '0.00%'
i['A15'] = 'Professional fees'; i['B15'] = '=B8*0.08'
I = {k: f'Inputs!$B${r}' for r, (k, v) in enumerate(rows, start=2)}
FEES = 'Inputs!$B$15'

c = wb.create_sheet('Cash Flow'); c['A1'] = 'RIVERSIDE LOGISTICS PARK - CASH FLOW (SAR)'
c['A3'] = 'Year'; c['A4'] = 'Period'; c['M3'] = 'Total'
labels = {6: 'Construction cost', 7: 'Occupancy', 8: 'Rent', 9: 'Operating costs', 10: 'Net operating income', 11: 'Sale proceeds',
          12: 'Unlevered cash flow', 13: 'Debt drawn', 14: 'Opening debt', 15: 'Interest', 16: 'Loan repayment', 17: 'Closing debt', 18: 'DSCR',
          19: 'Cash trapped', 20: 'Cash released', 21: 'Levered cash flow to equity', 22: 'Cumulative equity cash flow'}
for r, t in labels.items(): c.cell(r, 1, t)
occ = [0, 0, 0.60, 0.85, 0.95, 0.95, 0.95, 0.95, 0.95, 0.95]
for k in range(10):
    col = L(3 + k); p = L(2 + k)
    c[f'{col}3'] = 2027 + k; c[f'{col}4'] = 1 if k == 0 else f'={p}4+1'
    c[f'{col}6'] = f'=IF({col}$4<=2,-({I["Construction cost"]}+{FEES})/2,0)'
    c[f'{col}7'] = occ[k]; c[f'{col}7'].number_format = '0%'
    c[f'{col}8'] = f'=IF({col}$4<={I["Exit year"]},{I["Lettable area (sqm)"]}*CHOOSE({I["Scenario (1 = Base, 2 = Downside)"]},{I["Rent per sqm · Base"]},{I["Rent per sqm · Downside"]})*{col}7*(1+{I["Rent growth"]})^MAX(0,{col}$4-3),0)'
    c[f'{col}9'] = f'=-{col}8*{I["Operating costs (share of rent)"]}'
    c[f'{col}10'] = f'={col}8+{col}9'
    c[f'{col}11'] = f'=IF({col}$4={I["Exit year"]},{col}10*(1+{I["Rent growth"]})/{I["Exit yield"]}*0.98,0)'
    c[f'{col}12'] = f'={col}6+{col}10+{col}11'
    c[f'{col}13'] = f'=IF({col}$4<=2,MIN({I["Loan to cost"]}*({I["Construction cost"]}+{FEES}),{I["Debt cap"]})/2,0)'
    c[f'{col}14'] = '=0' if k == 0 else f'={p}17'
    c[f'{col}15'] = f'=-{col}14*{I["Interest rate"]}'
    c[f'{col}16'] = f'=IF({col}$4={I["Exit year"]},-({col}14+{col}13),0)'
    c[f'{col}17'] = f'={col}14+{col}13+{col}16'
    c[f'{col}18'] = f'=IF(AND({col}10>0,{col}15<0),{col}10/-{col}15,"")'
    c[f'{col}19'] = f'=IF(AND(ISNUMBER({col}18),{col}$4<{I["Exit year"]}),IF({col}18<{I["Minimum DSCR covenant"]},MAX(0,{col}10+{col}15),0),0)'
    c[f'{col}20'] = f'=IF({col}$4={I["Exit year"]},SUM($C19:{col}19),0)'
    c[f'{col}21'] = f'={col}12+{col}13+{col}15+{col}16-{col}19+{col}20'
    c[f'{col}22'] = f'=SUM($C21:{col}21)'
for r in (6, 8, 9, 10, 11, 12, 13, 15, 16, 19, 20, 21): c[f'M{r}'] = f'=SUM(C{r}:L{r})'

s = wb.create_sheet('Returns'); s['A1'] = 'RETURNS'
for r, (k, f) in enumerate([('Unlevered IRR', "=IRR('Cash Flow'!C12:L12)"), ('Levered IRR', "=IFERROR(IRR('Cash Flow'!C21:L21),0)"),
                            ('Equity multiple', "=SUMIF('Cash Flow'!C21:L21,\">0\")/-SUMIF('Cash Flow'!C21:L21,\"<0\")"), ('Net profit', "=SUM('Cash Flow'!C21:L21)"),
                            ('Peak debt', "=MAX('Cash Flow'!C17:L17)"), ('Minimum DSCR', "=MIN('Cash Flow'!C18:L18)"),
                            ('Total development cost', "=-'Cash Flow'!M6"), ('Profit on cost', '=B5/B8'), ('Cash trapped (total)', "='Cash Flow'!M19")], start=2):
    s.cell(r, 1, k); s.cell(r, 2, f)
for r in (2, 3, 9): s.cell(r, 2).number_format = '0.00%'

tmp = tempfile.mkdtemp(); src = os.path.join(tmp, 'sens2.xlsx'); wb.save(src)
out = os.path.join(tmp, 'o'); os.makedirs(out)
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', out, src], check=True, capture_output=True, timeout=600)
shutil.copy(os.path.join(out, 'sens2.xlsx'), sys.argv[1]); shutil.rmtree(tmp, ignore_errors=True)
print('written', sys.argv[1])
