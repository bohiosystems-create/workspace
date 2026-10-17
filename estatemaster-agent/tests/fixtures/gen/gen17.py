# round 17
#  pbsa   - Hong Kong student housing, HKD, nomination agreement share at a fixed rate, academic-year occupancy
#  dxb    - Dubai hotel, AED, management contract: base fee + incentive fee in tiers of GOP margin (IF ladder)
#  multi  - UK multi-let industrial, GBP, lease expiry profile, WAULT via SUMPRODUCT, ERV reversion on expiry
#  cold   - Norwegian cold store, NOK, pallet positions x price, power cost per pallet escalating
#  seoul  - Seoul presale apartments, KRW, presale schedule and VAT on non-housing parts
#  lender - mezzanine lender model, USD, PIK accrual, lender IRR and MOIC
import openpyxl,os
from openpyxl.utils import get_column_letter as L
OUT='/tmp/claude-0/-home-user-workspace/c7e8f42c-9ed0-5ba7-9852-d8a131185935/scratchpad/m17/';os.makedirs(OUT+'calc',exist_ok=True)
P='0.00%';N='#,##0';N2='#,##0.00'
def inputs(ws,rows,col=2,start=3):
    I={};r=start
    for k,lab,v,f in rows:
        ws.cell(r,1,lab);c=ws.cell(r,col,v);c.number_format=f
        I[k]=(f"'{ws.title}'" if ' ' in ws.title else ws.title)+f"!${L(col)}${r}";r+=1
    return I
def summary(wb,name,rows):
    s=wb.create_sheet(name)
    for i,(lab,f,fm) in enumerate(rows):s.cell(3+i,1,lab);c=s.cell(3+i,2,f);c.number_format=fm
def annual(wb,I,noi,years,price_key,exit_key,hold_key,extra=None):
    cf=wb.create_sheet('Cash Flow');cf['A3']='Year';cf['A4']='NOI';cf['A5']='Sale';cf['A6']='Purchase';cf['A7']='Net cash flow'
    for y in range(0,years+1):
        c=L(y+2);cf[f'{c}3']=y;on=f"({c}$3>=1)*({c}$3<={I[hold_key]})"
        cf[f'{c}4']=f"={on}*({noi})*1.025^MAX(0,{c}$3-1)";cf[f'{c}5']=f"=IF({c}$3={I[hold_key]},{c}4*1.025/{I[exit_key]},0)";cf[f'{c}6']=f"=IF({c}$3=0,-{I[price_key]},0)";cf[f'{c}7']=f"=SUM({c}4:{c}6)"
    return L(years+2)

# ---------- HK student housing ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='KOWLOON TONG STUDENT RESIDENCE (HKD)'
I=inputs(ws,[('beds','Beds',820,N),('price','Purchase price (HK$)',1450000000,N),('nom','Nomination agreement share',0.6,P),('nomrate','Nomination rate (HK$/bed/month)',7800,N),('dt','Direct-let rate (HK$/bed/month)',9200,N),
  ('occ','Direct-let occupancy',0.88,P),('months','Academic months let',10,'0'),('summer','Summer income (HK$)',9000000,N),('opex','Opex ratio',0.3,P),('exit','Exit yield',0.04,P),('hold','Hold (years)',7,'0')])
noi=f"({I['beds']}*({I['nom']}*{I['nomrate']}+(1-{I['nom']})*{I['dt']}*{I['occ']})*{I['months']}+{I['summer']})*(1-{I['opex']})"
last=annual(wb,I,noi,9,'price','exit','hold')
summary(wb,'Returns',[('Unlevered IRR',f"=IRR('Cash Flow'!B7:{last}7)",P),('Net profit',f"=SUM('Cash Flow'!B7:{last}7)",N),('Going-in yield',f"='Cash Flow'!C4/{I['price']}",P),('Income secured by nominations',f"={I['nom']}*{I['nomrate']}/({I['nom']}*{I['nomrate']}+(1-{I['nom']})*{I['dt']}*{I['occ']})",P)])
wb.save(OUT+'pbsa_hkd.xlsx')

# ---------- Dubai hotel with incentive fee ladder ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='DUBAI CREEK HOTEL – MANAGEMENT CONTRACT (AED)'
I=inputs(ws,[('keys','Keys',310,N),('adr','ADR (AED)',780,N),('occ','Occupancy',0.74,P),('other','Other revenue (% of rooms)',0.6,P),('gop','GOP margin',0.38,P),('base','Base fee (% of revenue)',0.03,P),
  ('t1','Incentive fee if GOP margin < 35%',0.06,P),('t2','Incentive fee if GOP margin 35–40%',0.08,P),('t3','Incentive fee if GOP margin > 40%',0.1,P),('ffe','FF&E reserve',0.04,P),('price','Purchase price (AED)',620000000,N),('exit','Exit cap',0.075,P),('hold','Hold (years)',6,'0')])
s=wb.create_sheet('Operations');rows=[('Revenue',f"={I['keys']}*365*{I['adr']}*{I['occ']}*(1+{I['other']})",N),('GOP',f"=B3*{I['gop']}",N),('Base fee',f"=-B3*{I['base']}",N),
 ('Incentive fee rate',f"=IF({I['gop']}<0.35,{I['t1']},IF({I['gop']}<=0.4,{I['t2']},{I['t3']}))",P),('Incentive fee',"=-B4*B6",N),('FF&E reserve',f"=-B3*{I['ffe']}",N),('Owner EBITDA','=B4+B5+B7+B8',N)]
for i,(lab,f,fm) in enumerate(rows):s.cell(3+i,1,lab);c=s.cell(3+i,2,f);c.number_format=fm
last=annual(wb,I,"Operations!$B$9",8,'price','exit','hold')
summary(wb,'Returns',[('Unlevered IRR',f"=IRR('Cash Flow'!B7:{last}7)",P),('Net profit',f"=SUM('Cash Flow'!B7:{last}7)",N),('Owner EBITDA (year 1)','=Operations!B9',N),('Fees as % of GOP','=-(Operations!B5+Operations!B7)/Operations!B4',P),('Price per key',f"={I['price']}/{I['keys']}",N)])
wb.save(OUT+'dxb_aed.xlsx')

# ---------- UK multi-let industrial with WAULT ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='TRAFFORD PARK MULTI-LET ESTATE (GBP)'
I=inputs(ws,[('price','Purchase price (£)',58000000,N),('erv','ERV (£/sq ft)',9.5,N2),('void','Void on expiry (months)',9,'0'),('renew','Renewal probability',0.65,P),('exit','Exit yield',0.0575,P),('hold','Hold (years)',5,'0'),('costs','Purchase costs',0.068,P)])
rr=wb.create_sheet('Tenancy');rr.append(['Unit','Tenant','Area (sq ft)','Passing rent (£/sq ft)','Years to expiry'])
for i,t in enumerate([('U1','Parcel Hub',85000,7.2,1.5),('U2','Northern Foods',62000,7.8,4),('U3','TradeCo',40000,8.1,2.5),('U4','BoxPack',55000,6.9,7),('U5','Vacant',30000,0,0)],start=2):
    for j,v in enumerate(t,start=1):rr.cell(i,j,v)
s=wb.create_sheet('Analysis');rows=[('Passing rent (£)',"=SUMPRODUCT(Tenancy!C2:C6,Tenancy!D2:D6)",N),('ERV (£)',f"=SUM(Tenancy!C2:C6)*{I['erv']}",N),('Reversion','=B4/B3-1',P),
 ('WAULT to expiry (years)',"=SUMPRODUCT(Tenancy!C2:C6,Tenancy!D2:D6,Tenancy!E2:E6)/B3",N2),('Leases expiring within hold (£)',f"=SUMPRODUCT((Tenancy!E2:E6<={I['hold']})*(Tenancy!E2:E6>0),Tenancy!C2:C6,Tenancy!D2:D6)",N),
 ('Expected void cost (£)',f"=(B7*(1-{I['renew']})+30000*{I['erv']})*{I['void']}/12",N),('Exit value (£)',f"=B4/{I['exit']}",N),('Total cost (£)',f"={I['price']}*(1+{I['costs']})",N),
 ('Net profit (£)',f"=B9-B10+B3*{I['hold']}-B8",N),('Unlevered IRR (approx.)',f"=((B9+B3*{I['hold']}-B8)/B10)^(1/{I['hold']})-1",P),('Net initial yield',"=B3/B10",P)]
for i,(lab,f,fm) in enumerate(rows):s.cell(3+i,1,lab);c=s.cell(3+i,2,f);c.number_format=fm
wb.save(OUT+'multi_gbp.xlsx')

# ---------- Norwegian cold store ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='DRAMMEN COLD STORE (NOK)'
I=inputs(ws,[('pallets','Pallet positions',42000,N),('util','Utilisation',0.86,P),('ppw','Price per pallet per week (kr)',38,N2),('power','Power cost per pallet per week (kr)',6.5,N2),('powinf','Power price escalation',0.06,P),
  ('staff','Staff and other costs (kr p.a.)',21000000,N),('price','Purchase price (kr)',720000000,N),('exit','Exit yield',0.0575,P),('hold','Hold (years)',8,'0')])
cf=wb.create_sheet('Cash Flow');cf['A3']='Year';cf['A4']='Storage revenue';cf['A5']='Power';cf['A6']='Staff and other';cf['A7']='EBITDA';cf['A8']='Sale';cf['A9']='Purchase';cf['A10']='Net cash flow'
for y in range(0,10):
    c=L(y+2);cf[f'{c}3']=y;on=f"({c}$3>=1)*({c}$3<={I['hold']})"
    cf[f'{c}4']=f"={on}*{I['pallets']}*{I['util']}*{I['ppw']}*52*1.03^MAX(0,{c}$3-1)";cf[f'{c}5']=f"=-{on}*{I['pallets']}*{I['util']}*{I['power']}*52*(1+{I['powinf']})^MAX(0,{c}$3-1)"
    cf[f'{c}6']=f"=-{on}*{I['staff']}*1.03^MAX(0,{c}$3-1)";cf[f'{c}7']=f"=SUM({c}4:{c}6)";cf[f'{c}8']=f"=IF({c}$3={I['hold']},{c}7*1.03/{I['exit']},0)";cf[f'{c}9']=f"=IF({c}$3=0,-{I['price']},0)";cf[f'{c}10']=f"=SUM({c}7:{c}9)"
summary(wb,'Returns',[('Unlevered IRR',"=IRR('Cash Flow'!B10:K10)",P),('Net profit',"=SUM('Cash Flow'!B10:K10)",N),('EBITDA margin (year 1)',"='Cash Flow'!C7/'Cash Flow'!C4",P),('Power as % of revenue (year 1)',"=-'Cash Flow'!C5/'Cash Flow'!C4",P)])
wb.save(OUT+'cold_nok.xlsx')

# ---------- Seoul presale apartments, KRW ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='MAPO RIVERSIDE – PRESALE APARTMENTS (KRW)'
I=inputs(ws,[('units','Units',480,N),('price','Average price (₩)',1250000000,N),('presale','Presold at launch',0.8,P),('deposit','Deposit at contract',0.1,P),('interim','Interim payments (before completion)',0.6,P),
  ('land','Land (₩)',210000000000,N),('build','Construction (₩)',265000000000,N),('vat','VAT on construction (non-housing share)',0.1*0.25,P),('pf','PF loan rate',0.065,P),('months','Construction (months)',33,'0')])
s=wb.create_sheet('Appraisal');rows=[('Sales revenue',f"={I['units']}*{I['price']}",N),('Collected before completion',f"=B3*{I['presale']}*({I['deposit']}+{I['interim']})",N),('Balance at completion','=B3-B4',N),
 ('VAT',f"={I['build']}*{I['vat']}",N),('PF interest',f"=({I['land']}+{I['build']}/2-B4/2)*{I['pf']}*{I['months']}/12",N),('Total cost',f"={I['land']}+{I['build']}+B6+B7",N),('Net profit','=B3-B8',N),('Profit on cost','=B9/B8',P),
 ('Margin on sales','=B9/B3',P)]
for i,(lab,f,fm) in enumerate(rows):s.cell(3+i,1,lab);c=s.cell(3+i,2,f);c.number_format=fm
wb.save(OUT+'seoul_krw.xlsx')

# ---------- mezzanine lender model ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Loan Terms';ws['A1']='CANAL STREET – MEZZANINE LOAN (LENDER MODEL, USD)'
I=inputs(ws,[('amount','Commitment ($)',35000000,N),('cash','Cash coupon',0.08,P),('pik','PIK coupon',0.04,P),('fee','Arrangement fee',0.015,P),('exitfee','Exit fee',0.01,P),('term','Term (years)',4,'0'),('value','Property value ($)',190000000,N),('senior','Senior loan ($)',105000000,N)])
cf=wb.create_sheet('Lender Cash Flow');cf['A3']='Year';cf['A4']='Balance (with PIK)';cf['A5']='Drawdown';cf['A6']='Fees';cf['A7']='Cash interest';cf['A8']='Repayment';cf['A9']='Lender cash flow'
for y in range(0,6):
    c=L(y+2);p=L(y+1);cf[f'{c}3']=y
    cf[f'{c}4']=(f"=IF({c}$3<={I['term']},{p}4*(1+{I['pik']}),0)" if y>0 else f"={I['amount']}")
    cf[f'{c}5']=f"=IF({c}$3=0,-{I['amount']},0)";cf[f'{c}6']=f"=IF({c}$3=0,{I['amount']}*{I['fee']},0)+IF({c}$3={I['term']},{c}4*{I['exitfee']},0)"
    cf[f'{c}7']=(f"=IF({c}$3<={I['term']},{p}4*{I['cash']},0)" if y>0 else '=0');cf[f'{c}8']=f"=IF({c}$3={I['term']},{c}4,0)";cf[f'{c}9']=f"=SUM({c}5:{c}8)"
summary(wb,'Lender Returns',[('Lender IRR',"=IRR('Lender Cash Flow'!B9:G9)",P),('Lender MOIC',"=SUMIF('Lender Cash Flow'!B9:G9,\">0\")/-SUMIF('Lender Cash Flow'!B9:G9,\"<0\")",'0.00"x"'),
  ('Attachment LTV',f"={I['senior']}/{I['value']}",P),('Detachment LTV',f"=({I['senior']}+{I['amount']})/{I['value']}",P),('Net profit',"=SUM('Lender Cash Flow'!B9:G9)",N)])
wb.save(OUT+'lender_usd.xlsx')
print('ok')
