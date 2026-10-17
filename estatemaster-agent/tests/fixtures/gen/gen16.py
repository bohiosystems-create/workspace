# round 16: five more model types with Excel features not tried yet
#  sdlt   - UK residential land deal, GBP, stamp duty by bands via approximate-match VLOOKUP
#  berlin - Berlin flats, EUR, rent cap: MIN(market, cap) per unit type, IF on modernisation surcharge
#  mall   - Sao Paulo mall, BRL, SELIC-linked debt, IPCA-indexed rents
#  wafi   - Riyadh off-plan tower, SAR, escrow (Wafi) collections by construction milestone via INDEX/MATCH
#  hall   - Virginia data hall, USD, monthly dates with EOMONTH, annual roll-up with SUMIFS by year
import openpyxl,datetime,os
from openpyxl.utils import get_column_letter as L
OUT='/tmp/claude-0/-home-user-workspace/c7e8f42c-9ed0-5ba7-9852-d8a131185935/scratchpad/m16/';os.makedirs(OUT+'calc',exist_ok=True)
P='0.00%';N='#,##0';N2='#,##0.00';D='dd-mmm-yy'
def inputs(ws,rows,col=2,start=3):
    I={};r=start
    for k,lab,v,f in rows:
        ws.cell(r,1,lab);c=ws.cell(r,col,v);c.number_format=f
        I[k]=(f"'{ws.title}'" if ' ' in ws.title else ws.title)+f"!${L(col)}${r}";r+=1
    return I
def summary(wb,name,rows):
    s=wb.create_sheet(name)
    for i,(lab,f,fm) in enumerate(rows):s.cell(3+i,1,lab);c=s.cell(3+i,2,f);c.number_format=fm

# ---------- UK land deal with SDLT bands ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='ASHFORD MEADOWS – RESIDENTIAL LAND DEAL (GBP)'
I=inputs(ws,[('land','Land price (£)',8400000,N),('homes','Homes',120,N),('gdv','Average sale price (£)',365000,N),('build','Build cost per home (£)',168000,N),('s106','S106 per home (£)',14000,N),
  ('fees','Professional fees',0.1,P),('fin','Finance rate',0.085,P),('months','Programme (months)',30,'0'),('mkt','Sales & marketing',0.03,P)])
bd=wb.create_sheet('SDLT');bd.append(['Band from (£)','Rate']);[bd.append(list(r)) for r in [(0,0),(150000,0.02),(250000,0.05)]]
for r in range(2,5):bd.cell(r,2).number_format=P
s=wb.create_sheet('Appraisal');rows=[('GDV',f"={I['homes']}*{I['gdv']}",N),('SDLT on land',f"=MAX(0,MIN({I['land']},SDLT!A4)-SDLT!A3)*SDLT!B3+MAX(0,{I['land']}-SDLT!A4)*SDLT!B4",N),
 ('Top SDLT band reached',f"=VLOOKUP({I['land']},SDLT!A2:B4,2,TRUE)",P),('Build',f"={I['homes']}*{I['build']}",N),('S106',f"={I['homes']}*{I['s106']}",N),('Fees',f"=B6*{I['fees']}",N),
 ('Finance',f"=({I['land']}+B4+(B6+B7+B8)/2)*{I['fin']}*{I['months']}/12",N),('Sales & marketing',f"=B3*{I['mkt']}",N),('Total development cost',f"={I['land']}+SUM(B4,B6:B10)",N),
 ('Net profit','=B3-B11',N),('Profit on cost','=B12/B11',P),('Margin on GDV','=B12/B3',P)]
for i,(lab,f,fm) in enumerate(rows):s.cell(3+i,1,lab);c=s.cell(3+i,2,f);c.number_format=fm
wb.save(OUT+'sdlt_gbp.xlsx')

# ---------- Berlin flats with a rent cap ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Annahmen';ws['A1']='PRENZLAUER BERG – WOHNUNGSBESTAND (EUR)'
I=inputs(ws,[('price','Kaufpreis (EUR)',42000000,N),('cap','Mietpreisbremse cap (EUR/sqm/month)',13.5,N2),('mod','Modernisation surcharge (EUR/sqm/month)',1.2,N2),('modon','Modernised (1 = yes)',1,'0'),
  ('vac','Vacancy',0.02,P),('opex','Non-recoverable costs (% of rent)',0.12,P),('growth','Rent growth',0.02,P),('exit','Exit yield',0.035,P),('hold','Hold (years)',10,'0'),('costs','Purchase costs',0.105,P)])
mx=wb.create_sheet('Units');mx.append(['Type','Units','Size (sqm)','Market rent (EUR/sqm/month)','Rent charged (EUR/sqm/month)'])
for i,(t,u,sz,mr) in enumerate([('1-room',24,38,17.5),('2-room',40,62,15.2),('3-room',22,84,13.9),('4-room',8,108,12.8)],start=2):
    mx.cell(i,1,t);mx.cell(i,2,u);mx.cell(i,3,sz);mx.cell(i,4,mr);mx.cell(i,5,f"=MIN(D{i},{I['cap']}+IF({I['modon']}=1,{I['mod']},0))")
cf=wb.create_sheet('Cash Flow');cf['A3']='Year';cf['A4']='Net rent';cf['A5']='Sale';cf['A6']='Purchase';cf['A7']='Net cash flow'
for y in range(0,12):
    c=L(y+2);cf[f'{c}3']=y;on=f"({c}$3>=1)*({c}$3<={I['hold']})"
    cf[f'{c}4']=f"={on}*SUMPRODUCT(Units!$B$2:$B$5,Units!$C$2:$C$5,Units!$E$2:$E$5)*12*(1-{I['vac']})*(1-{I['opex']})*(1+{I['growth']})^MAX(0,{c}$3-1)"
    cf[f'{c}5']=f"=IF({c}$3={I['hold']},{c}4*(1+{I['growth']})/{I['exit']},0)";cf[f'{c}6']=f"=IF({c}$3=0,-{I['price']}*(1+{I['costs']}),0)";cf[f'{c}7']=f"=SUM({c}4:{c}6)"
summary(wb,'Returns',[('Unlevered IRR',"=IRR('Cash Flow'!B7:M7)",P),('Net profit',"=SUM('Cash Flow'!B7:M7)",N),('Net initial yield',f"='Cash Flow'!C4/({I['price']}*(1+{I['costs']}))",P),('Reversion to market (%)',"=SUMPRODUCT(Units!B2:B5,Units!C2:C5,Units!D2:D5)/SUMPRODUCT(Units!B2:B5,Units!C2:C5,Units!E2:E5)-1",P)])
wb.save(OUT+'berlin_eur.xlsx')

# ---------- Sao Paulo mall, BRL ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Premissas';ws['A1']='SHOPPING VILA NOVA – ACQUISITION (BRL)'
I=inputs(ws,[('price','Purchase price (R$)',620000000,N),('noi','NOI year 1 (R$)',61000000,N),('ipca','IPCA',0.045,P),('real','Real rent growth',0.01,P),('selic','SELIC',0.1125,P),('spread','Debt spread over SELIC',0.018,P),
  ('ltv','LTV',0.45,P),('exit','Exit cap rate',0.095,P),('hold','Hold (years)',7,'0'),('ocup','Occupancy',0.94,P)])
cf=wb.create_sheet('Fluxo');cf['A3']='Year';cf['A4']='NOI';cf['A5']='Sale';cf['A6']='Purchase';cf['A7']='Unlevered';cf['A8']='Debt';cf['A9']='Levered'
for y in range(0,9):
    c=L(y+2);cf[f'{c}3']=y;on=f"({c}$3>=1)*({c}$3<={I['hold']})"
    cf[f'{c}4']=f"={on}*{I['noi']}*{I['ocup']}/0.94*((1+{I['ipca']})*(1+{I['real']}))^MAX(0,{c}$3-1)"
    cf[f'{c}5']=f"=IF({c}$3={I['hold']},{c}4*(1+{I['ipca']})/{I['exit']},0)";cf[f'{c}6']=f"=IF({c}$3=0,-{I['price']},0)";cf[f'{c}7']=f"=SUM({c}4:{c}6)"
    cf[f'{c}8']=f"=IF({c}$3=0,{I['price']}*{I['ltv']},0)-{on}*{I['price']}*{I['ltv']}*({I['selic']}+{I['spread']})-IF({c}$3={I['hold']},{I['price']}*{I['ltv']},0)";cf[f'{c}9']=f"={c}7+{c}8"
summary(wb,'Resultados',[('Unlevered IRR',"=IRR(Fluxo!B7:J7)",P),('Levered IRR',"=IRR(Fluxo!B9:J9)",P),('Net profit',"=SUM(Fluxo!B7:J7)",N),('Real unlevered IRR (over IPCA)',f"=(1+B3)/(1+{I['ipca']})-1",P),('Going-in cap rate',f"={I['noi']}/{I['price']}",P),('Equity multiple',"=SUMIF(Fluxo!B9:J9,\">0\")/-SUMIF(Fluxo!B9:J9,\"<0\")",'0.00"x"')])
wb.save(OUT+'mall_brl.xlsx')

# ---------- Riyadh off-plan with Wafi escrow milestones ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='AL YASMIN HEIGHTS – OFF-PLAN (WAFI) – SAR'
I=inputs(ws,[('units','Apartments',220,N),('price','Average price (SAR)',1150000,N),('sold','Sold off-plan',0.7,P),('cost','Construction cost (SAR)',148000000,N),('land','Land (SAR)',62000000,N),
  ('quarters','Construction (quarters)',10,'0'),('rate','Bank finance rate',0.07,P)])
ms=wb.create_sheet('Milestones');ms.append(['Construction complete','Cumulative release from escrow'])
for r in [(0,0.1),(0.2,0.2),(0.4,0.4),(0.6,0.6),(0.8,0.8),(1,0.95)]:ms.append(list(r))
for r in range(2,8):ms.cell(r,1).number_format=P;ms.cell(r,2).number_format=P
cf=wb.create_sheet('Quarterly');cf['A3']='Quarter';cf['A4']='Completion';cf['A5']='Escrow released (cumulative %)';cf['A6']='Escrow release';cf['A7']='Unsold sales at handover';cf['A8']='Construction';cf['A9']='Land';cf['A10']='Net cash flow'
last=L(13)
for q in range(0,12):
    c=L(q+2);p=L(q+1);cf[f'{c}3']=q
    cf[f'{c}4']=f"=MIN(1,{c}$3/{I['quarters']})"
    cf[f'{c}5']=f"=INDEX(Milestones!$B$2:$B$7,MATCH({c}4,Milestones!$A$2:$A$7,1))+IF({c}$3>{I['quarters']},0.05,0)"
    cf[f'{c}6']=(f"=({c}5-{p}5)*{I['units']}*{I['sold']}*{I['price']}" if q>0 else f"={c}5*{I['units']}*{I['sold']}*{I['price']}")
    cf[f'{c}7']=f"=IF({c}$3={I['quarters']}+1,{I['units']}*(1-{I['sold']})*{I['price']},0)"
    cf[f'{c}8']=f"=-IF(AND({c}$3>=1,{c}$3<={I['quarters']}),{I['cost']}/{I['quarters']},0)";cf[f'{c}9']=f"=IF({c}$3=0,-{I['land']},0)";cf[f'{c}10']=f"=SUM({c}6:{c}9)"
summary(wb,'Summary',[('Project IRR (annual)',f"=(1+IRR(Quarterly!B10:M10))^4-1",P),('Net profit',"=SUM(Quarterly!B10:M10)",N),('Gross revenue',"=SUM(Quarterly!B6:M7)",N),('Total development cost',"=-SUM(Quarterly!B8:M9)",N),('Profit on cost','=B4/B6',P),('Peak funding need',"=-MIN(Quarterly!B10:M10)",N)])
wb.save(OUT+'wafi_sar.xlsx')

# ---------- Virginia data hall, monthly with EOMONTH and SUMIFS by year ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='ASHBURN HALL 3 – DATA HALL LEASE-UP (USD)'
I=inputs(ws,[('start','First month',datetime.date(2027,1,31),D),('mw','IT capacity (MW)',24,N2),('rate','Rate ($/kW/month)',165,N2),('ramp','MW leased per month',1.5,N2),('capex','Capex ($m/MW)',11.5,N2),
  ('opex','Opex ($/kW/month)',22,N2),('exit','Exit multiple of annual EBITDA',18,'0.0'),('hold','Exit after (months)',60,'0')])
cf=wb.create_sheet('Monthly');cf['A3']='Month';cf['A4']='Date';cf['A5']='Year';cf['A6']='MW leased';cf['A7']='Revenue';cf['A8']='Opex';cf['A9']='Capex';cf['A10']='Sale';cf['A11']='Net cash flow'
M=60;last=L(M+2)
for m in range(0,M+1):
    c=L(m+2);cf[f'{c}3']=m;cf[f'{c}4']=f"=EOMONTH({I['start']},{c}$3)";cf[f'{c}4'].number_format=D;cf[f'{c}5']=f"=YEAR({c}4)"
    cf[f'{c}6']=f"=MIN({I['mw']},{I['ramp']}*{c}$3)";cf[f'{c}7']=f"={c}6*1000*{I['rate']}";cf[f'{c}8']=f"=-{c}6*1000*{I['opex']}"
    cf[f'{c}9']=f"=IF({c}$3=0,-{I['mw']}*{I['capex']}*1000000,0)";cf[f'{c}10']=f"=IF({c}$3={I['hold']},({c}7+{c}8)*12*{I['exit']},0)";cf[f'{c}11']=f"=SUM({c}7:{c}10)"
yr=wb.create_sheet('Annual');yr.append(['Year','Revenue','EBITDA'])
for i,y in enumerate(range(2027,2032),start=2):yr.cell(i,1,y);yr.cell(i,2,f"=SUMIFS(Monthly!$B$7:${last}$7,Monthly!$B$5:${last}$5,A{i})");yr.cell(i,3,f"=SUMIFS(Monthly!$B$7:${last}$7,Monthly!$B$5:${last}$5,A{i})+SUMIFS(Monthly!$B$8:${last}$8,Monthly!$B$5:${last}$5,A{i})")
summary(wb,'Returns',[('Unlevered IRR (annualised)',f"=(1+IRR(Monthly!B11:{last}11,0.01))^12-1",P),('Net profit',f"=SUM(Monthly!B11:{last}11)",N),('Total capex',f"=-Monthly!B9",N),('Stabilised EBITDA (2031)',"=Annual!C6",N),('Yield on cost (stabilised)',"=Annual!C6/B5",P)])
wb.save(OUT+'hall_usd.xlsx')
print('ok')
