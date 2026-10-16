# round 8: five more model types and layouts not seen yet: a lender's mezzanine model in $000s, a retail centre with
# percentages typed as whole numbers, inputs sitting on top of the cash flow on one sheet, a model with a circular
# interest calculation (iterative), and a fund-level model with four assets and LP net returns after fees and carry
import openpyxl,datetime
from openpyxl.utils import get_column_letter as L
OUT='/tmp/claude-0/-home-user-workspace/c7e8f42c-9ed0-5ba7-9852-d8a131185935/scratchpad/m8/'
P='0.00%';N='#,##0';
def put(ws,r,lab,val,fmt=None,unit=None,lc=1,vc=2,uc=3):
    ws.cell(r,lc,lab);c=ws.cell(r,vc,val)
    if fmt:c.number_format=fmt
    if unit:ws.cell(r,uc,unit)
    return f"${L(vc)}${r}"
# ---------- F. mezzanine loan, lender's view, USD thousands ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Loan Terms';ws['A1']='HARBOUR POINT – MEZZANINE FACILITY (lender model)';ws['A2']='USD 000s unless stated'
r={};row=4
for k,lab,v,f,u in [('val','Property value at close',185000,N,'$000s'),('sen','Senior loan',111000,N,'$000s'),('mez','Mezzanine commitment',27750,N,'$000s'),('cash','Cash coupon',0.09,P,'p.a.'),('pik','PIK coupon',0.035,P,'p.a., accrues'),
  ('ofee','Origination fee',0.015,P,'of commitment'),('xfee','Exit fee',0.01,P,'of balance repaid'),('term','Term',3,'0','years'),('vg','Value growth',0.03,P,'p.a.'),('dflt','Default year (0 = none)',0,'0','year'),('rec','Recovery on default',0.55,P,'of balance')]:
    r[k]="'Loan Terms'!"+put(ws,row,lab,v,f,u);row+=1
cf=wb.create_sheet('Lender CF');cf['A1']='LENDER CASH FLOWS ($000s)';cf['A3']='Year'
ls=['Opening balance','PIK accrued','Closing balance','Cash interest','Fees','Repayment','Lender cash flow','Property value','Senior + mezz LTV']
R={l:5+i for i,l in enumerate(ls)}
for l,rr in R.items():cf[f'A{rr}']=l
for y in range(0,5):
    c=L(y+2);p=L(y+1);cf[f'{c}3']=y;g=lambda l:f'{c}{R[l]}'
    alive=f"AND({c}$3>=1,{c}$3<={r['term']},OR({r['dflt']}=0,{c}$3<={r['dflt']}))"
    cf[g('Opening balance')]=f"=IF({c}$3=0,0,IF({c}$3=1,{r['mez']},{p}{R['Closing balance']}))"
    cf[g('PIK accrued')]=f"=IF({alive},{g('Opening balance')}*{r['pik']},0)"
    cf[g('Closing balance')]=f"=IF(OR({c}$3=0,NOT({alive})),0,IF(OR({c}$3={r['term']},{c}$3={r['dflt']}),0,{g('Opening balance')}+{g('PIK accrued')}))"
    cf[g('Cash interest')]=f"=IF(AND({alive},NOT({c}$3={r['dflt']})),{g('Opening balance')}*{r['cash']},0)"
    cf[g('Fees')]=f"=IF({c}$3=0,{r['mez']}*{r['ofee']},0)+IF(AND({c}$3={r['term']},{r['dflt']}=0),({g('Opening balance')}+{g('PIK accrued')})*{r['xfee']},0)"
    cf[g('Repayment')]=f"=IF({c}$3=0,-{r['mez']},IF(AND({c}$3={r['dflt']},{r['dflt']}>0),({g('Opening balance')}+{g('PIK accrued')})*{r['rec']},IF(AND({c}$3={r['term']},{r['dflt']}=0),{g('Opening balance')}+{g('PIK accrued')},0)))"
    cf[g('Lender cash flow')]=f"={g('Cash interest')}+{g('Fees')}+{g('Repayment')}"
    cf[g('Property value')]=f"={r['val']}*(1+{r['vg']})^{c}$3"
    cf[g('Senior + mezz LTV')]=f"=IF({g('Property value')}>0,({r['sen']}+{g('Opening balance')}+IF({c}$3=0,{r['mez']},0))/{g('Property value')},0)"
sm=wb.create_sheet('Returns');rg=lambda l:f"'Lender CF'!$B${R[l]}:$F${R[l]}"
for i,(lab,f,fm) in enumerate([('Lender IRR',f'=IRR({rg("Lender cash flow")})',P),('Lender MOIC',f'=SUMIF({rg("Lender cash flow")},">0")/-SUMIF({rg("Lender cash flow")},"<0")','0.00"x"'),
   ('Loan-to-value at close (through mezz)',f"=({r['sen']}+{r['mez']})/{r['val']}",P),('Debt yield on total debt',"=0.064*"+r['val']+"/("+r['sen']+"+"+r['mez']+")",P),('Lender profit ($000s)',f'=SUM({rg("Lender cash flow")})',N)]):
    sm.cell(3+i,1,lab);c=sm.cell(3+i,2,f);c.number_format=fm
wb.save(OUT+'mezz_usd000.xlsx')
# ---------- G. retail centre, EUR, percentages typed as whole numbers ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='CENTRO VERDE SHOPPING CENTRE – ACQUISITION (EUR)';ws['A2']='Percentages entered as whole numbers (6.25 = 6.25%)'
r={};row=4
for k,lab,v,u in [('price','Purchase price (EUR)',142000000,'EUR'),('ttx','Transfer tax (%)',7.5,'%'),('anr','Anchor rent (EUR p.a.)',3400000,'EUR'),('anx','Anchor lease expiry (year)',3,'year'),('void','Re-letting void (months)',9,'months'),
  ('rerent','Anchor re-let rent vs current (%)',85,'%'),('unit','Line-shop rent (EUR p.a.)',6100000,'EUR'),('occ','Line-shop occupancy (%)',93,'%'),('trn','Turnover rent (% of rent)',4,'%'),('grw','Rental growth (%)',2,'%'),
  ('nrc','Non-recoverable costs (% of income)',6,'%'),('capx','Refurbishment capex (EUR)',9000000,'EUR'),('ey','Exit yield (%)',6.25,'%'),('hold','Hold period (years)',6,'years'),('ltv','LTV (%)',50,'%'),('rate','All-in interest rate (%)',4.6,'%')]:
    r[k]='Inputs!'+put(ws,row,lab,v,N if v>1000 else '0.00',u);row+=1
cf=wb.create_sheet('CashFlow');cf['A1']='ANNUAL CASH FLOW (EUR)';cf['A3']='Year'
ls=['Anchor rent','Line-shop rent','Turnover rent','Non-recoverable costs','NOI','Capex','Acquisition','Sale','Unlevered','Loan','Interest','Repayment','Levered']
R={l:5+i for i,l in enumerate(ls)}
for l,rr in R.items():cf[f'A{rr}']=l
for y in range(0,8):
    c=L(y+2);cf[f'{c}3']=y;g=lambda l:f'{c}{R[l]}';op_=f"AND({c}$3>=1,{c}$3<={r['hold']})";gr=f"(1+{r['grw']}/100)^({c}$3-1)"
    cf[g('Anchor rent')]=f"=IF({op_},IF({c}$3<={r['anx']},{r['anr']},IF({c}$3={r['anx']}+1,{r['anr']}*{r['rerent']}/100*MAX(0,12-{r['void']})/12,{r['anr']}*{r['rerent']}/100))*{gr},0)"
    cf[g('Line-shop rent')]=f"=IF({op_},{r['unit']}*{r['occ']}/100*{gr},0)"
    cf[g('Turnover rent')]=f"=({g('Anchor rent')}+{g('Line-shop rent')})*{r['trn']}/100"
    cf[g('Non-recoverable costs')]=f"=-({g('Anchor rent')}+{g('Line-shop rent')}+{g('Turnover rent')})*{r['nrc']}/100"
    cf[g('NOI')]=f"=SUM({c}{R['Anchor rent']}:{c}{R['Non-recoverable costs']})"
    cf[g('Capex')]=f"=IF({c}$3=1,-{r['capx']},0)"
    cf[g('Acquisition')]=f"=IF({c}$3=0,-{r['price']}*(1+{r['ttx']}/100),0)"
    nx=L(y+3)
    cf[g('Sale')]=f"=IF({c}$3={r['hold']},{nx}{R['NOI']}/({r['ey']}/100),0)" if y<7 else '=0'
    cf[g('Unlevered')]=f"={g('NOI')}+{g('Capex')}+{g('Acquisition')}+{g('Sale')}"
    cf[g('Loan')]=f"=IF({c}$3=0,{r['price']}*{r['ltv']}/100,0)"
    cf[g('Interest')]=f"=IF({op_},-{r['price']}*{r['ltv']}/100*{r['rate']}/100,0)"
    cf[g('Repayment')]=f"=IF({c}$3={r['hold']},-{r['price']}*{r['ltv']}/100,0)"
    cf[g('Levered')]=f"={g('Unlevered')}+{g('Loan')}+{g('Interest')}+{g('Repayment')}"
# year 7 NOI exists so the sale in year 6 can use forward income: extend the operating test to hold+1 for NOI rows only
for l in ['Anchor rent','Line-shop rent']:
    c='I';cf[f'{c}{R[l]}']=cf[f'H{R[l]}'].value.replace('H$3','I$3').replace('AND(I$3>=1,I$3<=Inputs!$B$17)','AND(I$3>=1,I$3<=Inputs!$B$17+1)')
sm=wb.create_sheet('Summary');rg=lambda l:f"CashFlow!$B${R[l]}:$I${R[l]}"
for i,(lab,f,fm) in enumerate([('Levered IRR',f'=IRR({rg("Levered")})',P),('Unlevered IRR',f'=IRR({rg("Unlevered")})',P),('Equity multiple',f'=SUMIF({rg("Levered")},">0")/-SUMIF({rg("Levered")},"<0")','0.00"x"'),
   ('Net initial yield',f"=CashFlow!C{R['NOI']}/({r['price']}*(1+{r['ttx']}/100))",P),('Net profit',f'=SUM({rg("Levered")})',N),('Total acquisition cost',f"={r['price']}*(1+{r['ttx']}/100)+{r['capx']}",N)]):
    sm.cell(3+i,1,lab);c=sm.cell(3+i,2,f);c.number_format=fm
wb.save(OUT+'retail_eur.xlsx')
# ---------- H. senior living development, SAR, inputs on top of the monthly cash flow on one sheet ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Model';ws['A1']='AL WAHA SENIOR LIVING – DEVELOPMENT MODEL (SAR)'
r={};row=3
for k,lab,v,f in [('units','Units',180,N),('fee','Monthly fee per unit (SAR)',19500,N),('stab','Stabilised occupancy',0.92,P),('ramp','Ramp-up to stabilisation (months)',30,'0'),('staff','Staff cost per occupied unit / month',5200,N),
  ('oth','Other opex (% of revenue)',0.12,P),('dev','Development cost (SAR)',265000000,N),('build','Build period (months)',24,'0'),('cap','Exit cap rate',0.075,P),('exm','Exit month',72,'0')]:
    r[k]='Model!'+put(ws,row,lab,v,f);row+=1
top=16;ws[f'A{top}']='Month';rows=['Occupancy','Revenue','Staff costs','Other opex','NOI','Development','Exit value','Net cash flow']
RR={l:top+1+i for i,l in enumerate(rows)}
for l,rr in RR.items():ws[f'A{rr}']=l
for m in range(1,73):
    c=L(m+1);ws[f'{c}{top}']=m;g=lambda l:f'{c}{RR[l]}'
    ws[g('Occupancy')]=f"=IF({c}${top}<={r['build']},0,MIN({r['stab']},{r['stab']}*({c}${top}-{r['build']})/{r['ramp']}))"
    ws[g('Revenue')]=f"={r['units']}*{g('Occupancy')}*{r['fee']}"
    ws[g('Staff costs')]=f"=-{r['units']}*{g('Occupancy')}*{r['staff']}"
    ws[g('Other opex')]=f"=-{g('Revenue')}*{r['oth']}"
    ws[g('NOI')]=f"={g('Revenue')}+{g('Staff costs')}+{g('Other opex')}"
    ws[g('Development')]=f"=IF({c}${top}<={r['build']},-{r['dev']}/{r['build']},0)"
    ws[g('Exit value')]=f"=IF({c}${top}={r['exm']},{r['units']}*{r['stab']}*({r['fee']}*(1-{r['oth']})-{r['staff']})*12/{r['cap']},0)"
    ws[g('Net cash flow')]=f"={g('NOI')}+{g('Development')}+{g('Exit value')}"
s2=wb.create_sheet('Outputs');rg=lambda l:f"Model!$B${RR[l]}:$BU${RR[l]}"
for i,(lab,f,fm) in enumerate([('Project IRR',f'=(1+IRR({rg("Net cash flow")},0.01))^12-1',P),('Net profit',f'=SUM({rg("Net cash flow")})',N),('Profit on cost',f"=B4/{r['dev']}",P),('Stabilised yield on cost',f"={r['units']}*{r['stab']}*({r['fee']}*(1-{r['oth']})-{r['staff']})*12/{r['dev']}",P),('Total development cost',f"={r['dev']}",N)]):
    s2.cell(3+i,1,lab);c=s2.cell(3+i,2,f);c.number_format=fm
wb.save(OUT+'seniorliving_sar.xlsx')
# ---------- I. self-storage with a circular interest calculation (iterative calc on) ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='STOREWAY SELF STORAGE – DEVELOPMENT (USD)';ws['A2']='Interest is funded by the loan (circular: iterative calculation on)'
r={};row=4
for k,lab,v,f in [('nra','Net rentable area (sf)',82000,N),('rent','Rent per sf per year',19.5,'0.00'),('occ','Stabilised occupancy',0.88,P),('opx','Operating expense ratio',0.36,P),('hard','Hard cost',9800000,N),('land','Land',2600000,N),
  ('ltc','Loan to cost',0.65,P),('ir','Interest rate',0.0725,P),('cap','Exit cap rate',0.0575,P)]:
    r[k]='Inputs!'+put(ws,row,lab,v,f);row+=1
cs=wb.create_sheet('Budget');cs['A1']='BUDGET';cs['A3']='Land';cs['B3']=f"={r['land']}";cs['A4']='Hard cost';cs['B4']=f"={r['hard']}";cs['A5']='Interest reserve';cs['B5']=f"=B7*{r['ir']}*1.5/2"
cs['A6']='Total cost';cs['B6']='=SUM(B3:B5)';cs['A7']='Loan';cs['B7']=f"=B6*{r['ltc']}";cs['A8']='Equity';cs['B8']='=B6-B7'
cs['A10']='Stabilised NOI';cs['B10']=f"={r['nra']}*{r['rent']}*{r['occ']}*(1-{r['opx']})";cs['A11']='Value at exit';cs['B11']=f"=B10/{r['cap']}"
cs['A12']='Profit';cs['B12']='=B11-B6';cs['A13']='Profit on cost';cs['B13']='=B12/B6';cs['A14']='Yield on cost';cs['B14']='=B10/B6';cs['A15']='Equity multiple';cs['B15']='=(B11-B7)/B8'
wb.calculation.iterate=True;wb.calculation.iterateCount=100
wb.save(OUT+'storage_circular_usd.xlsx')
# ---------- J. fund model: four assets, fees and carry, LP net returns (EUR) ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Assets';ws['A1']='EUROPEAN LOGISTICS FUND I – FUND MODEL (EUR m)'
ws.append([]);ws.append(['Asset','Equity in (EUR m)','Entry year','Exit year','Exit multiple on equity','Annual yield to equity'])
assets=[('Lyon DC',42,0,5,1.65,0.05),('Rotterdam XXL',58,0,6,1.80,0.045),('Milan last-mile',35,1,5,1.55,0.055),('Poznan park',28,1,6,1.70,0.06)]
for a in assets:ws.append(list(a))
f=wb.create_sheet('Fund');f['A1']='FUND CASH FLOWS (EUR m)';f['A3']='Year';rr=['Gross equity calls','Gross distributions','Gross cash flow','Management fee','Net before carry','Carry','LP net cash flow']
R={l:5+i for i,l in enumerate(rr)};
for l,x in R.items():f[f'A{x}']=l
f['A14']='Management fee (% of committed p.a.)';f['B14']=0.015;f['A15']='Carried interest';f['B15']=0.20;f['A16']='Hurdle (simple, on contributions)';f['B16']=0.08;f['A17']='Commitments (EUR m)';f['B17']='=SUM(Assets!B4:B7)'
for y in range(0,8):
    c=L(y+2);f[f'{c}3']=y;g=lambda l:f'{c}{R[l]}'
    f[g('Gross equity calls')]=f'=-SUMIF(Assets!$C$4:$C$7,{c}$3,Assets!$B$4:$B$7)'
    f[g('Gross distributions')]=f'=SUMPRODUCT((Assets!$C$4:$C$7<{c}$3)*(Assets!$D$4:$D$7>={c}$3)*Assets!$B$4:$B$7*Assets!$F$4:$F$7)+SUMPRODUCT((Assets!$D$4:$D$7={c}$3)*Assets!$B$4:$B$7*Assets!$E$4:$E$7)'
    f[g('Gross cash flow')]=f"={g('Gross equity calls')}+{g('Gross distributions')}"
    f[g('Management fee')]=f'=IF({c}$3<=6,-$B$17*$B$14,0)'
    f[g('Net before carry')]=f"={g('Gross cash flow')}+{g('Management fee')}"
    f[g('Carry')]=f"=IF({c}$3=7,0,IF({c}$3=6,-MAX(0,SUM($B${R['Net before carry']}:{c}{R['Net before carry']})-$B$16*$B$17*4)*$B$15,0))"
    f[g('LP net cash flow')]=f"={g('Net before carry')}+{g('Carry')}"
sm=wb.create_sheet('Fund Returns');rg=lambda l:f"Fund!$B${R[l]}:$I${R[l]}"
for i,(lab,fm_,fmt) in enumerate([('Gross IRR',f'=IRR({rg("Gross cash flow")})',P),('Net IRR to LPs',f'=IRR({rg("LP net cash flow")})',P),('Gross MOIC',f'=SUMIF({rg("Gross cash flow")},">0")/-SUMIF({rg("Gross cash flow")},"<0")','0.00"x"'),
   ('Net TVPI',f'=SUMIF({rg("LP net cash flow")},">0")/-SUMIF({rg("LP net cash flow")},"<0")','0.00"x"'),('Total carry (EUR m)',f'=-SUM({rg("Carry")})','0.0'),('Net profit to LPs (EUR m)',f'=SUM({rg("LP net cash flow")})','0.0')]):
    sm.cell(3+i,1,lab);c=sm.cell(3+i,2,fm_);c.number_format=fmt
wb.save(OUT+'fund_eur.xlsx')
print('ok')
