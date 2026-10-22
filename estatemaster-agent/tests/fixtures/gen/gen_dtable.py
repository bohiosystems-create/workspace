"""Synthetic workbook with Excel Data Tables (What-If Analysis → Data Table): a two-way table (exit yield × rent growth) and a
one-way table (exit yield), their cached values as Excel computes them. Run: python3 gen_dtable.py <out.xlsx>"""
import sys, zipfile, re, shutil, tempfile, os
from openpyxl import Workbook
OUT = sys.argv[1]
wb = Workbook(); ws = wb.active; ws.title = 'Sens'
ws['A3'], ws['B3'] = 'Exit yield', 0.07
ws['A4'], ws['B4'] = 'Rent growth', 0.03
ws['A5'], ws['B5'] = 'Rent (year 1)', 1000
ws['A6'], ws['B6'] = 'Value', '=B5*(1+B4)^5/B3'
def val(y, g, rent=1000): return rent * (1 + g) ** 5 / y
ys, gs = [0.065, 0.07, 0.075], [0.02, 0.03, 0.04]
ws['E8'] = '=B6'
for j, y in enumerate(ys): ws.cell(8, 6 + j, y)
for i, g in enumerate(gs):
    ws.cell(9 + i, 5, g)
    for j, y in enumerate(ys): ws.cell(9 + i, 6 + j, val(y, g))
ws['K8'] = '=B6'
for i, y in enumerate(ys):
    ws.cell(9 + i, 10, y); ws.cell(9 + i, 11, val(y, 0.03))
ws['A12'], ws['B12'] = 'Top-left of the table', '=F9'
tmp = tempfile.mktemp(suffix='.xlsx'); wb.save(tmp)
zin = zipfile.ZipFile(tmp); zout = zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    data = zin.read(it.filename)
    if it.filename == 'xl/worksheets/sheet1.xml':
        x = data.decode()
        # two-way: values across row 8 feed the row input cell B3, values down column E feed the column input cell B4
        x = re.sub(r'(<c r="F9"[^>]*>)', r'\1<f t="dataTable" ref="F9:H11" dt2D="1" dtr="1" r1="B3" r2="B4"/>', x, 1)
        # one-way, column-oriented: values down column J feed B3, the formula in K8
        x = re.sub(r'(<c r="K9"[^>]*>)', r'\1<f t="dataTable" ref="K9:K11" dt2D="0" dtr="0" r1="B3"/>', x, 1)
        # the results Excel saves with every formula
        for ref, v in (('B6', val(0.07, 0.03)), ('E8', val(0.07, 0.03)), ('K8', val(0.07, 0.03)), ('B12', val(ys[0], gs[0]))):
            x = re.sub(r'(<c r="%s"[^>]*>)(<f>[^<]*</f>)(<v>[^<]*</v>|<v\s*/>)?' % ref, lambda m: m.group(1) + m.group(2) + '<v>%r</v>' % v, x, 1)
        data = x.encode()
    zout.writestr(it, data)
zout.close(); os.remove(tmp)
print('written', OUT)
