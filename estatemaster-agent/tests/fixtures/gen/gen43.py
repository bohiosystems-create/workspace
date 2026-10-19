# Round 43 fixture (synthetic): a build-to-rent residential scheme, with made-up figures, laid out the way BTR models often are.
#  r43_btr_sar: units by type in a mix table (studio, 1-bed, 2-bed, 3-bed) with count, size and monthly rent; construction over six
#  quarters; a lease-up curve after completion; stabilised occupancy, rent growth, operating costs per unit a year; a refinance at
#  stabilisation that repays the construction loan with an investment loan; a sale at the exit cap; quarterly; SAR in full units;
#  returns on a "Dashboard" whose rows read "Unlevered IRR", "Levered IRR", "Yield on cost (stabilised)" and "Equity multiple".
import openpyxl,sys,datetime,calendar
from openpyxl.utils import get_column_letter as L
out=sys.argv[1].rstrip('/')
wb=openpyxl.Workbook();D=wb.active;D.title='Dashboard';A=wb.create_sheet('Assumptions');Q=wb.create_sheet('CF')
N=32   # quarters 0..32
A['B2']='AL RAWDAH RESIDENCES (BTR) · ASSUMPTIONS (SAR)'
A['B4']='Unit mix';A['C4']='Units';A['D4']='Size (sqm)';A['E4']='Rent (a month)'
mix=[('Studio',60,42,4200),('1-bed',140,68,6100),('2-bed',110,105,8600),('3-bed',40,150,11800)]
for i,(t,u,sz,r) in enumerate(mix):
    rr=5+i;A.cell(rr,2,t);A.cell(rr,3,u);A.cell(rr,4,sz);A.cell(rr,5,r).number_format='#,##0'
A['B9']='Total';A['C9']='=SUM(C5:C8)';A['E9']='=SUMPRODUCT(C5:C8,E5:E8)';A['E9'].number_format='#,##0'
rows=[(11,'Land',38000000,'#,##0'),(12,'Construction cost per sqm',7500,'#,##0'),(13,'Soft costs (% of construction)',0.11,'0.0%'),
 (14,'Construction start (quarter)',1,'0'),(15,'Construction length (quarters)',6,'0'),(16,'Lease-up (quarters to stabilise)',4,'0'),
 (17,'Stabilised occupancy',0.94,'0.0%'),(18,'Rent growth (a year)',0.03,'0.0%'),(19,'Operating costs per unit (a year)',9500,'#,##0'),
 (20,'Exit quarter',28,'0'),(21,'Exit cap rate',0.065,'0.00%'),(22,'Selling costs',0.015,'0.0%'),
 (23,'Construction loan (share of cost)',0.6,'0.0%'),(24,'Construction loan rate',0.075,'0.00%'),(25,'Investment loan (share of value at refinance)',0.55,'0.0%'),
 (26,'Investment loan rate',0.06,'0.00%'),(27,'Refinance cap rate',0.065,'0.00%')]
for r,l,v,f in rows:A.cell(r,2,l);A.cell(r,3,v).number_format=f
Q['B2']='QUARTERLY CASH FLOW (SAR)';Q['B3']='Quarter ending';Q['B4']='Quarter'
lab={6:'Construction progress',7:'Occupancy',8:'Rent',9:'Operating costs',10:'NOI',12:'Land',13:'Construction',14:'Soft costs',15:'Sale (net)',
 16:'Unlevered cash flow',18:'Construction loan draw',19:'Construction loan interest',20:'Construction loan repayment',21:'Construction loan balance',
 22:'Investment loan draw',23:'Investment loan interest',24:'Investment loan repayment',25:'Investment loan balance',27:'Levered cash flow'}
for r,l in lab.items():Q.cell(r,2,l)
def qend(k):y,m=divmod(3*k+2,12);y+=2026;m+=1;return datetime.datetime(y,m,calendar.monthrange(y,m)[1])
cs='ROUND(Assumptions!$C$14,0)';cl='ROUND(Assumptions!$C$15,0)';lu='ROUND(Assumptions!$C$16,0)';ex='ROUND(Assumptions!$C$20,0)';done=f'({cs}+{cl})';stab=f'({done}+{lu})'
cost='(SUMPRODUCT(Assumptions!$C$5:$C$8,Assumptions!$D$5:$D$8)*Assumptions!$C$12)'
for k in range(N+1):
    c=L(3+k);p=L(2+k);n=L(4+k);qq=f'{c}$4'
    Q[f'{c}3']=qend(k);Q[f'{c}3'].number_format='mmm-yy';Q[f'{c}4']=k
    Q[f'{c}6']=f'=IF(AND({qq}>={cs},{qq}<{done}),1/{cl},0)'
    Q[f'{c}7']=f'=IF({qq}<{done},0,Assumptions!$C$17*MIN(1,({qq}-{done}+1)/{lu}))'
    Q[f'{c}8']=f'=Assumptions!$E$9*3*{c}7*(1+Assumptions!$C$18)^(MAX(0,{qq}-{done})/4)'
    Q[f'{c}9']=f'=IF({qq}<{done},0,-Assumptions!$C$9*Assumptions!$C$19/4*(1+Assumptions!$C$18)^(MAX(0,{qq}-{done})/4))'
    Q[f'{c}10']=f'={c}8+{c}9'
    Q[f'{c}12']=f'=IF({qq}=0,-Assumptions!$C$11,0)';Q[f'{c}13']=f'=-{cost}*{c}6';Q[f'{c}14']=f'={c}13*Assumptions!$C$13'
    Q[f'{c}15']=f'=IF({qq}={ex},{n}10*4/Assumptions!$C$21*(1-Assumptions!$C$22),0)' if k<N else '=0'
    Q[f'{c}16']=f'=IF({qq}>{ex},0,{c}10+{c}12+{c}13+{c}14+{c}15)'
    pc='0' if k==0 else f'{p}21';pi='0' if k==0 else f'{p}25'
    Q[f'{c}18']=f'=-({c}12+{c}13+{c}14)*Assumptions!$C$23'
    Q[f'{c}19']=f'=-{pc}*Assumptions!$C$24/4'
    Q[f'{c}22']=f'=IF({qq}={stab},{n}10*4/Assumptions!$C$27*Assumptions!$C$25,0)' if k<N else '=0'
    Q[f'{c}20']=f'=-IF({qq}={stab},{pc}+{c}18,0)'
    Q[f'{c}21']=f'={pc}+{c}18+{c}20'
    Q[f'{c}23']=f'=-{pi}*Assumptions!$C$26/4'
    Q[f'{c}24']=f'=-IF({qq}={ex},{pi}+{c}22,0)'
    Q[f'{c}25']=f'={pi}+{c}22+{c}24'
    Q[f'{c}27']=f'=IF({qq}>{ex},0,{c}16+{c}18+{c}19+{c}20+{c}22+{c}23+{c}24)'
E=L(3+N)
D['B2']='AL RAWDAH RESIDENCES (BTR) · DASHBOARD (SAR)'
dash=[(5,'Unlevered IRR',f'=(1+IRR(CF!C16:{E}16))^4-1','0.00%'),(6,'Levered IRR',f'=(1+IRR(CF!C27:{E}27))^4-1','0.00%'),
 (7,'Yield on cost (stabilised)',f"=INDEX(CF!C10:{E}10,1,ROUND(Assumptions!C14,0)+ROUND(Assumptions!C15,0)+ROUND(Assumptions!C16,0)+1)*4/-SUM(CF!C12:{E}14)",'0.00%'),
 (8,'Equity multiple',f'=SUMIF(CF!C27:{E}27,">0")/-SUMIF(CF!C27:{E}27,"<0")','0.00"x"'),(9,'Total development cost',f'=-SUM(CF!C12:{E}14)','#,##0'),
 (10,'Levered profit',f'=SUM(CF!C27:{E}27)','#,##0')]
for r,l,f,fm in dash:D.cell(r,2,l);D.cell(r,3,f).number_format=fm
wb.save(f'{out}/r43_btr_sar.xlsx')
