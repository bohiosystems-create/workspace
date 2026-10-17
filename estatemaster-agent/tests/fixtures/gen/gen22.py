# round 22: a CIO's desk on new ground
#  dc        - USD data-centre development: three phases of MW with starts and build lengths, pre-let share by case, lease-up, rent per kW
#              escalated yearly, power pass-through (PUE, tariff, utilisation, recovery), capex per MW indexed by start year (INT), a
#              construction loan converting to a mini-perm when occupancy passes a threshold (first quarter found by MATCH), cash sweep
#              above a DSCR, exit on forward NOI; Base / Downside / Upside through CHOOSE
#  branded   - AED hotel and branded residences: off-plan sales with a payment plan tied to construction milestones (quarter each is
#              reached found with COUNTIF on the S-curve), RERA escrow released against works certified, hotel ramp with seasonality by
#              calendar quarter, management fees with an incentive fee above the owner's priority, FF&E reserve by year, branding fee,
#              equity first then a facility with a cash sweep, XIRR on dates
#  ff        - GBP student housing forward funding: S-curve by VLOOKUP on the elapsed share of the build, funder's coupon rolled up,
#              delay penalties past a long-stop, purchase price on NIY less SDLT by bands (VLOOKUP) and fees, developer's profit at
#              practical completion, the investor's 10-year hold and exit, fiscal-year summary by labels made with TEXT
#  sukuk     - SAR off-plan community with an ijara sukuk: buyer payments lagged by a quarter count read from a cell (OFFSET),
#              Wafi escrow, real-estate transaction tax absorbed on part of the units, zakat at each year end, sukuk coverage
#  second    - EUR secondaries fund: six LP interests bought at discounts to NAV on dates, unfunded commitments called, distributions
#              and NAV growth, fees on NAV, a subscription line bridging calls two quarters, European waterfall with catch-up, XIRR,
#              TVPI / DPI / RVPI at a reporting date found by MATCH
import openpyxl,os,datetime
from openpyxl.utils import get_column_letter as L
OUT='/tmp/claude-0/-home-user-workspace/c7e8f42c-9ed0-5ba7-9852-d8a131185935/scratchpad/m22/'
P='0.00%';N='#,##0';X='0.00"x"';D='dd-mmm-yyyy'
def ref(ws,col,r):return (f"'{ws.title}'" if ' ' in ws.title else ws.title)+f"!${L(col)}${r}"
def inputs(ws,rows,col=2,start=3):
    I={};r=start
    for k,lab,v,f in rows:
        ws.cell(r,1,lab);c=ws.cell(r,col,v);c.number_format=f;I[k]=ref(ws,col,r);r+=1
    return I
def table(ws,rows,start=3,col=2):
    for i,(lab,f,fm) in enumerate(rows):ws.cell(start+i,1,lab);c=ws.cell(start+i,col,f);c.number_format=fm
def timeline(ws,labels,T,fill,first_row=3):
    R={l:first_row+i for i,l in enumerate(labels)}
    for l,r in R.items():ws.cell(r,1,l)
    for t in range(T+1):
        c=L(t+2);p=L(t+1)
        for l,f in fill(t,c,p,R).items():ws[f'{c}{R[l]}']=f if not isinstance(f,str) or f.startswith('=') else '='+f
    return R

# ================= 22A data centre =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='NORTHGATE DATA CENTRE CAMPUS – DEVELOPMENT AND HOLD (USD)'
I=inputs(ws,[('case','Case (1 Base, 2 Downside, 3 Upside)',1,'0'),('start','Model start',datetime.date(2026,12,31),D),('capmw','Capex per MW ($, today)',11500000,N),('cpi','Cost inflation a year',0.03,P),
 ('r1','Rent per kW a month ($) · Base',118,N),('r2','Rent per kW a month ($) · Downside',102,N),('r3','Rent per kW a month ($) · Upside',130,N),('esc','Rent escalation a year',0.025,P),
 ('pue','PUE',1.35,'0.00'),('tariff','Power tariff ($ per MWh)',85,N),('util','IT load utilisation',0.7,P),('rec','Power cost recovered from tenants',0.9,P),('opex','Other operating cost (% of rent)',0.08,P),
 ('ltc','Construction loan to cost',0.65,P),('rc','Construction loan rate',0.085,P),('occ','Conversion trigger (occupancy)',0.85,P),('rp','Mini-perm rate',0.0625,P),('amort','Mini-perm amortisation a year',0.02,P),
 ('swd','Cash sweep above DSCR',1.4,X),('sw','Cash sweep share',0.5,P),('exitq','Exit quarter',36,'0'),('c1','Exit cap · Base',0.0625,P),('c2','Exit cap · Downside',0.07,P),('c3','Exit cap · Upside',0.0575,P)])
ws['A28']='Rent per kW (live case)';ws['B28']=f"=CHOOSE({I['case']},{I['r1']},{I['r2']},{I['r3']})";ws['A29']='Exit cap (live case)';ws['B29']=f"=CHOOSE({I['case']},{I['c1']},{I['c2']},{I['c3']})";ws['B29'].number_format=P
RENT='Inputs!$B$28';CAP='Inputs!$B$29'
ph=wb.create_sheet('Phases');ph.append(['Phase','MW','Start quarter','Build (quarters)','Pre-let · Base','Pre-let · Downside','Pre-let · Upside','Lease-up of the rest (quarters)','Pre-let (live case)'])
for i,(mw,s,b,p1,p2,p3,lu) in enumerate([(24,1,7,0.6,0.4,0.75,6),(32,6,8,0.5,0.3,0.65,8),(40,13,8,0.35,0.2,0.5,10)],2):
    for j,v in enumerate([f'Phase {i-1}',mw,s,b,p1,p2,p3,lu],1):ph.cell(i,j,v)
    ph.cell(i,9,f"=CHOOSE({I['case']},E{i},F{i},G{i})")
cf=wb.create_sheet('Cash Flow');T=40
labs=['Quarter','Date']+[f'P{k} {x}' for k in (1,2,3) for x in ('capex','leased MW','rent')]+['Capex','MW built','Leased MW','Occupancy','Rent','Power cost','Power recovered','Other opex','NOI','Converted','Opening loan','Loan draw','Interest capitalised','Interest paid','Amortisation','DSCR','Cash sweep','Loan balance','Sale','Loan repaid at sale','Equity in','Equity cash flow','Project cash flow']
def dcf(t,c,p,R):
    q=f"{c}$3";first=t==0;d={'Quarter':t,'Date':f"EOMONTH({I['start']},3*{q})"}
    for k in (1,2,3):
        r=k+1;mw=f"Phases!$B${r}";s=f"Phases!$C${r}";b=f"Phases!$D${r}";pl=f"Phases!$I${r}";lu=f"Phases!$H${r}"
        d[f'P{k} capex']=f"IF(AND({q}>={s},{q}<{s}+{b}),{mw}*{I['capmw']}*(1+{I['cpi']})^INT(({s}-1)/4)/{b},0)"
        d[f'P{k} leased MW']=f"IF({q}>={s}+{b},MIN({mw},{mw}*{pl}+{mw}*(1-{pl})*MIN(1,({q}-{s}-{b}+1)/{lu})),0)"
        d[f'P{k} rent']=f"IF({q}>={s}+{b},{c}{R[f'P{k} leased MW']}*1000*{RENT}*3*(1+{I['esc']})^INT(({q}-{s}-{b})/4),0)"
    d['Capex']=f"IF({q}<={I['exitq']},{c}{R['P1 capex']}+{c}{R['P2 capex']}+{c}{R['P3 capex']},0)"
    d['MW built']="+".join(f"IF({q}>=Phases!$C${k+1}+Phases!$D${k+1},Phases!$B${k+1},0)" for k in (1,2,3))
    d['Leased MW']=f"{c}{R['P1 leased MW']}+{c}{R['P2 leased MW']}+{c}{R['P3 leased MW']}"
    d['Occupancy']=f"IF({c}{R['MW built']}>0,{c}{R['Leased MW']}/{c}{R['MW built']},0)"
    d['Rent']=f"{c}{R['P1 rent']}+{c}{R['P2 rent']}+{c}{R['P3 rent']}"
    d['Power cost']=f"{c}{R['Leased MW']}*{I['util']}*{I['pue']}*2190*{I['tariff']}"
    d['Power recovered']=f"{c}{R['Power cost']}*{I['rec']}";d['Other opex']=f"{c}{R['Rent']}*{I['opex']}"
    d['NOI']=f"{c}{R['Rent']}+{c}{R['Power recovered']}-{c}{R['Power cost']}-{c}{R['Other opex']}"
    d['Converted']='0' if first else f"IF({p}{R['Converted']}=1,1,IF(AND({c}{R['MW built']}>0,{c}{R['Occupancy']}>={I['occ']}),1,0))"
    d['Opening loan']='0' if first else f"{p}{R['Loan balance']}"
    d['Loan draw']=f"{c}{R['Capex']}*{I['ltc']}"
    d['Interest capitalised']=f"IF({c}{R['Converted']}=0,{c}{R['Opening loan']}*{I['rc']}/4,0)"
    d['Interest paid']=f"IF({c}{R['Converted']}=1,{c}{R['Opening loan']}*{I['rp']}/4,0)"
    d['Amortisation']=f"IF({c}{R['Converted']}=1,MIN({c}{R['Opening loan']},{c}{R['Opening loan']}*{I['amort']}/4),0)"
    d['DSCR']=f"IF({c}{R['Interest paid']}+{c}{R['Amortisation']}>0,{c}{R['NOI']}/({c}{R['Interest paid']}+{c}{R['Amortisation']}),0)"
    d['Cash sweep']=f"IF(AND({c}{R['Converted']}=1,{q}<{I['exitq']}),MIN({c}{R['Opening loan']}-{c}{R['Amortisation']},MAX(0,{c}{R['NOI']}-({c}{R['Interest paid']}+{c}{R['Amortisation']})*{I['swd']})*{I['sw']}),0)"
    d['Loan balance']=f"IF({q}>={I['exitq']},0,{c}{R['Opening loan']}+{c}{R['Loan draw']}+{c}{R['Interest capitalised']}-{c}{R['Amortisation']}-{c}{R['Cash sweep']})"
    nx=L(t+3) if t<T else c
    d['Sale']=f"IF({q}={I['exitq']},{nx}{R['NOI']}*4/{CAP},0)"
    d['Loan repaid at sale']=f"IF({q}={I['exitq']},{c}{R['Opening loan']}+{c}{R['Loan draw']}+{c}{R['Interest capitalised']}-{c}{R['Amortisation']},0)"
    d['Equity in']=f"{c}{R['Capex']}*(1-{I['ltc']})"
    d['Equity cash flow']=f"IF({q}>{I['exitq']},0,-{c}{R['Equity in']}+{c}{R['NOI']}-{c}{R['Interest paid']}-{c}{R['Amortisation']}-{c}{R['Cash sweep']}+{c}{R['Sale']}-{c}{R['Loan repaid at sale']})"
    d['Project cash flow']=f"IF({q}>{I['exitq']},0,-{c}{R['Capex']}+{c}{R['NOI']}+{c}{R['Sale']})"
    return d
R=timeline(cf,labs,T,dcf);E=L(T+2);row=lambda n:f"'Cash Flow'!$B${R[n]}:${E}${R[n]}"
for t in range(T+1):cf[f'{L(t+2)}{R["Date"]}'].number_format=D
table(wb.create_sheet('Returns'),[('Levered IRR',f"=(1+IRR({row('Equity cash flow')},0.03))^4-1",P),('Project IRR',f"=(1+IRR({row('Project cash flow')},0.02))^4-1",P),
 ('Equity multiple',f"=SUMIF({row('Equity cash flow')},\">0\")/-SUMIF({row('Equity cash flow')},\"<0\")",X),('Peak debt ($)',f"=MAX({row('Loan balance')})",N),
 ('Minimum DSCR (after conversion)',f"=_xlfn.MINIFS({row('DSCR')},{row('DSCR')},\">0\")",X),('Conversion quarter',f"=IFERROR(MATCH(1,{row('Converted')},0)-1,\"not reached\")",'0'),
 ('Total development cost ($)',f"=SUM({row('Capex')})",N),('Yield on cost (exit-year NOI)',f"=INDEX({row('NOI')},Inputs!$B$23+1)*4/B9",P),('Exit value ($)',f"=SUM({row('Sale')})",N),
 ('Net profit ($)',f"=SUM({row('Equity cash flow')})",N),('Cash swept ($)',f"=SUM({row('Cash sweep')})",N)])
wb.save(OUT+'dc_usd.xlsx')

# ================= 22B hotel + branded residences =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='MARSA BRANDED RESIDENCES AND HOTEL – DUBAI (AED)'
I=inputs(ws,[('start','Model start',datetime.date(2026,9,30),D),('units','Residences (units)',220,N),('size','Average unit size (sq ft)',1850,N),('price','Sale price (AED per sq ft)',3400,N),('brand','Branding fee on sales',0.05,P),
 ('book','Booking payment',0.1,P),('inst','Each construction instalment',0.1,P),('hand','Paid on handover',0.5,P),('bcost','Residences build cost (AED per sq ft)',1150,N),('land','Land (AED)',260000000,N),('build','Build (quarters)',12,'0'),
 ('keys','Hotel keys',250,N),('hcap','Hotel capex (AED per key)',1900000,N),('adr','ADR year 1 (AED)',1450,N),('adrg','ADR growth a year',0.04,P),('other','Other revenue (% of rooms)',0.45,P),('gop','GOP margin',0.38,P),
 ('base','Base management fee (% of revenue)',0.03,P),('prio','Owner\'s priority (% of hotel capex a year)',0.1,P),('inc','Incentive fee (% of GOP above the priority)',0.08,P),('eq','Equity commitment (AED)',420000000,N),('rate','Facility rate',0.075,P),('hexit','Hotel exit cap',0.07,P),('exitq','Exit quarter',32,'0')])
ws['A30']='Instalment milestones (% complete)';ms=[0.2,0.4,0.6,0.8]
for k,m in enumerate(ms):ws.cell(31+k,1,f'Milestone {k+1}');ws.cell(31+k,2,m).number_format=P
sc=wb.create_sheet('Curves');sc.append(['Build quarter','Cumulative % complete','Sales (% of units in the quarter)','Hotel occupancy by year'])
cum=[0.03,0.08,0.15,0.24,0.34,0.45,0.56,0.67,0.77,0.86,0.94,1.0];sales=[0.18,0.12,0.09,0.08,0.07,0.07,0.06,0.06,0.05,0.05,0.04,0.04,0.03,0.02,0.02,0.02];occ=[0.58,0.66,0.72,0.74]
for i in range(16):sc.cell(i+2,1,i+1);sc.cell(i+2,2,cum[i] if i<12 else 1.0);sc.cell(i+2,3,sales[i]);sc.cell(i+2,4,occ[i] if i<4 else None)
ws['A36']='Seasonality (occupancy index) Q1, Q2, Q3, Q4';
for k,v in enumerate([1.15,0.95,0.7,1.2]):ws.cell(36,2+k,v)
cf=wb.create_sheet('Cash Flow');T=32;H=f"{I['build']}"
labs=['Quarter','Date','Units sold','Sales value','Cumulative sales value','Booking collected']+[f'Milestone {k+1} collected' for k in range(4)]+['Handover collected','Collections','Escrow released','Escrow balance','Branding fee','Residence build cost','Hotel capex','Land','Hotel open','Year of operation','Occupancy','ADR','Rooms revenue','Total hotel revenue','GOP','Base fee','FF&E reserve','Owner\'s priority','Incentive fee','Hotel EBITDA','Pre-finance cash flow','Opening facility','Interest','Equity drawn','Facility drawn','Facility repaid','Facility balance','Hotel sale','Distributions','Equity cash flow']
def bcf(t,c,p,R):
    q=f"{c}$3";first=t==0;cumS=f"SUM($B${R['Sales value']}:{c}{R['Sales value']})";d={'Quarter':t,'Date':f"EOMONTH({I['start']},3*{q})"}
    d['Units sold']=f"IF(AND({q}>=1,{q}<=16),{I['units']}*INDEX(Curves!$C$2:$C$17,MAX(1,{q})),0)"
    d['Sales value']=f"{c}{R['Units sold']}*{I['size']}*{I['price']}";d['Cumulative sales value']=cumS
    d['Booking collected']=f"IF({q}<{H},{c}{R['Sales value']}*{I['book']},0)"
    for k in range(4):
        mq=f"(COUNTIF(Curves!$B$2:$B$13,\"<\"&Inputs!$B${31+k})+1)"
        d[f'Milestone {k+1} collected']=f"IF({q}>={H},0,IF({q}={mq},{c}{R['Cumulative sales value']}*{I['inst']},IF({q}>{mq},{c}{R['Sales value']}*{I['inst']},0)))"
    d['Handover collected']=f"IF({q}={H},{c}{R['Cumulative sales value']}*{I['hand']}+{c}{R['Sales value']}*({I['book']}+4*{I['inst']}),IF({q}>{H},{c}{R['Sales value']},0))"
    d['Collections']=f"SUM({c}{R['Booking collected']}:{c}{R['Handover collected']})"
    d['Escrow released']=f"IF({q}>={H},{c}{R['Collections']}+{('0' if first else p+str(R['Escrow balance']))},MIN({('0' if first else p+str(R['Escrow balance']))}+{c}{R['Collections']},{c}{R['Residence build cost']}))"
    d['Escrow balance']=f"{('0' if first else p+str(R['Escrow balance']))}+{c}{R['Collections']}-{c}{R['Escrow released']}"
    d['Branding fee']=f"{c}{R['Collections']}*{I['brand']}"
    d['Residence build cost']=f"IF(AND({q}>=1,{q}<={H}),{I['units']}*{I['size']}*{I['bcost']}*(INDEX(Curves!$B$2:$B$13,MAX(1,MIN(12,{q})))-IF({q}>1,INDEX(Curves!$B$2:$B$13,MAX(1,MIN(12,{q}-1))),0)),0)"
    d['Hotel capex']=f"IF(AND({q}>=1,{q}<={H}),{I['keys']}*{I['hcap']}/{H},0)";d['Land']=f"IF({q}=0,{I['land']},0)"
    d['Hotel open']=f"IF(AND({q}>{H},{q}<={I['exitq']}),1,0)";d['Year of operation']=f"IF({c}{R['Hotel open']}=1,INT(({q}-{H}-1)/4)+1,0)"
    d['Occupancy']=f"IF({c}{R['Hotel open']}=1,MIN(0.95,INDEX(Curves!$D$2:$D$5,MIN(4,{c}{R['Year of operation']}))*INDEX(Inputs!$B$36:$E$36,INT((MONTH({c}{R['Date']})-1)/3)+1)),0)"
    d['ADR']=f"IF({c}{R['Hotel open']}=1,{I['adr']}*(1+{I['adrg']})^({c}{R['Year of operation']}-1),0)"
    d['Rooms revenue']=f"{I['keys']}*91.25*{c}{R['Occupancy']}*{c}{R['ADR']}";d['Total hotel revenue']=f"{c}{R['Rooms revenue']}*(1+{I['other']})"
    d['GOP']=f"{c}{R['Total hotel revenue']}*{I['gop']}";d['Base fee']=f"{c}{R['Total hotel revenue']}*{I['base']}"
    d['FF&E reserve']=f"IF({c}{R['Hotel open']}=1,{c}{R['Total hotel revenue']}*CHOOSE(MIN(3,{c}{R['Year of operation']}),0.02,0.03,0.04),0)"
    d["Owner's priority"]=f"{c}{R['Hotel open']}*{I['keys']}*{I['hcap']}*{I['prio']}/4"
    d['Incentive fee']=f"MAX(0,{c}{R['GOP']}-{c}{R['Base fee']}-{c}{R['FF&E reserve']}-{c}{R[chr(79)+'wner'+chr(39)+'s priority']})*{I['inc']}"
    d['Hotel EBITDA']=f"{c}{R['GOP']}-{c}{R['Base fee']}-{c}{R['FF&E reserve']}-{c}{R['Incentive fee']}"
    d['Pre-finance cash flow']=f"{c}{R['Escrow released']}-{c}{R['Branding fee']}-{c}{R['Residence build cost']}-{c}{R['Hotel capex']}-{c}{R['Land']}+{c}{R['Hotel EBITDA']}"
    d['Opening facility']='0' if first else f"{p}{R['Facility balance']}"
    d['Interest']=f"{c}{R['Opening facility']}*{I['rate']}/4"
    need=f"MAX(0,-({c}{R['Pre-finance cash flow']}-{c}{R['Interest']}))";cumE='0' if first else f"SUM($B${R['Equity drawn']}:{p}{R['Equity drawn']})"
    d['Equity drawn']=f"MIN({need},MAX(0,{I['eq']}-{cumE}))";d['Facility drawn']=f"{need}-{c}{R['Equity drawn']}"
    d['Facility repaid']=f"MIN({c}{R['Opening facility']},MAX(0,{c}{R['Pre-finance cash flow']}-{c}{R['Interest']}))"
    d['Facility balance']=f"IF({q}>={I['exitq']},0,{c}{R['Opening facility']}+{c}{R['Facility drawn']}-{c}{R['Facility repaid']})"
    nx=L(t+3) if t<T else c
    d['Hotel sale']=f"IF({q}={I['exitq']},SUM({L(max(2,t-1))}{R['Hotel EBITDA']}:{c}{R['Hotel EBITDA']})*4/MIN(4,{q}-{H})/{I['hexit']},0)"
    d['Distributions']=f"MAX(0,{c}{R['Pre-finance cash flow']}-{c}{R['Interest']}-{c}{R['Facility repaid']})+{c}{R['Hotel sale']}-IF({q}={I['exitq']},{c}{R['Opening facility']}+{c}{R['Facility drawn']}-{c}{R['Facility repaid']},0)"
    d['Equity cash flow']=f"IF({q}>{I['exitq']},0,{c}{R['Distributions']}-{c}{R['Equity drawn']})"
    return d
R=timeline(cf,labs,T,bcf);E=L(T+2);row=lambda n:f"'Cash Flow'!$B${R[n]}:${E}${R[n]}"
for t in range(T+1):cf[f'{L(t+2)}{R["Date"]}'].number_format=D
table(wb.create_sheet('Returns'),[('Levered IRR (XIRR)',f"=XIRR({row('Equity cash flow')},{row('Date')})",P),('Equity multiple',f"=SUMIF({row('Equity cash flow')},\">0\")/-SUMIF({row('Equity cash flow')},\"<0\")",X),
 ('Residences profit (AED)',f"=SUM({row('Collections')})-SUM({row('Branding fee')})-SUM({row('Residence build cost')})-{I['land']}*0.55",N),('Hotel EBITDA, third year (AED)',f"=SUMIF({row('Year of operation')},3,{row('Hotel EBITDA')})",N),
 ('Sales collected by handover',f"=SUMIF({row('Quarter')},\"<=\"&{I['build']},{row('Collections')})/SUM({row('Sales value')})",P),('Operator incentive fee (AED)',f"=SUM({row('Incentive fee')})",N),
 ('Peak facility (AED)',f"=MAX({row('Facility balance')})",N),('Peak escrow balance (AED)',f"=MAX({row('Escrow balance')})",N),('Equity drawn (AED)',f"=SUM({row('Equity drawn')})",N),('Net profit (AED)',f"=SUM({row('Equity cash flow')})",N)])
wb.save(OUT+'branded_aed.xlsx')

# ================= 22C forward funding =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='CASTLE QUAY STUDENT – FORWARD FUNDING (GBP)'
I=inputs(ws,[('start','Funding date',datetime.date(2026,7,31),D),('land','Land (£)',18000000,N),('beds','Beds',650,N),('rent','Rent per bed a week (£)',289,N),('weeks','Weeks let a year',51,N),('g','Rent growth a year',0.03,P),
 ('void','Voids',0.03,P),('opex','Operating cost (% of income)',0.28,P),('cost','Construction cost (£)',86000000,N),('build','Build (quarters)',9,'0'),('delay','Delay (quarters)',0,'0'),('longstop','Long-stop (quarters)',10,'0'),
 ('pen','Delay penalty a quarter past long-stop (£)',400000,N),('coupon','Funder\'s coupon a year',0.065,P),('niy','Purchase NIY',0.0475,P),('fees','Purchaser\'s fees (% of price)',0.018,P),('hold','Investor hold after completion (quarters)',40,'0'),('xcap','Investor exit yield',0.05,P)])
cv=wb.create_sheet('Curve');cv.append(['Share of build elapsed','Cumulative cost'])
for i,(a,b) in enumerate([(0,0),(0.1,0.04),(0.2,0.11),(0.3,0.21),(0.4,0.33),(0.5,0.46),(0.6,0.59),(0.7,0.71),(0.8,0.82),(0.9,0.92),(1.0,1.0)],2):cv.cell(i,1,a);cv.cell(i,2,b)
sd=wb.create_sheet('SDLT');sd.append(['Band from (£)','Rate','Tax below band (£)'])
for i,(a,r,t) in enumerate([(0,0,0),(150000,0.02,0),(250000,0.05,2000)],2):sd.cell(i,1,a);sd.cell(i,2,r);sd.cell(i,3,t)
cf=wb.create_sheet('Cash Flow');T=52;PCQ=f"({I['build']}+{I['delay']})"
labs=['Quarter','Date','Fiscal year','Elapsed share','Cumulative cost share','Construction draw','Land','Funds advanced','Coupon','Penalty','Operating','Gross income','Net operating income','Investor exit','Investor cash flow']
def fcf(t,c,p,R):
    q=f"{c}$3";first=t==0;d={'Quarter':t,'Date':f"EOMONTH({I['start']},3*{q})"}
    d['Fiscal year']=f"\"FY\"&TEXT(EDATE({c}{R['Date']},-3),\"yy\")"
    d['Elapsed share']=f"MIN(1,{q}/{PCQ})"
    d['Cumulative cost share']=f"VLOOKUP({c}{R['Elapsed share']},Curve!$A$2:$B$12,2)"
    d['Construction draw']='0' if first else f"IF({q}<={PCQ},{I['cost']}*({c}{R['Cumulative cost share']}-{p}{R['Cumulative cost share']}),0)"
    d['Land']=f"IF({q}=0,{I['land']},0)"
    d['Funds advanced']=f"SUM($B${R['Construction draw']}:{c}{R['Construction draw']})+{I['land']}"
    d['Coupon']='0' if first else f"IF({q}<={PCQ},{p}{R['Funds advanced']}*{I['coupon']}/4,0)"
    d['Penalty']=f"IF(AND({q}>{I['longstop']},{q}<={PCQ}),{I['pen']},0)"
    d['Operating']=f"IF(AND({q}>{PCQ},{q}<={PCQ}+{I['hold']}),1,0)"
    d['Gross income']=f"{c}{R['Operating']}*{I['beds']}*{I['rent']}*{I['weeks']}/4*(1-{I['void']})*(1+{I['g']})^INT(({q}-{PCQ}-1)/4)"
    d['Net operating income']=f"{c}{R['Gross income']}*(1-{I['opex']})"
    nx=L(t+3) if t<T else c
    d['Investor exit']=f"IF({q}={PCQ}+{I['hold']},{c}{R['Net operating income']}*4*(1+{I['g']})/{I['xcap']},0)"
    d['Investor cash flow']=f"-{c}{R['Land']}-{c}{R['Construction draw']}-IF({q}={PCQ},Returns!$B$8+Returns!$B$7+Returns!$B$6*{I['fees']},0)+{c}{R['Net operating income']}+{c}{R['Investor exit']}"
    return d
R=timeline(cf,labs,T,fcf);E=L(T+2);row=lambda n:f"'Cash Flow'!$B${R[n]}:${E}${R[n]}"
for t in range(T+1):cf[f'{L(t+2)}{R["Date"]}'].number_format=D
rt=wb.create_sheet('Returns')
table(rt,[('Practical completion (quarter)',f"={PCQ}",'0'),('Year-1 net rent (£)',f"={I['beds']}*{I['rent']}*{I['weeks']}*(1-{I['void']})*(1-{I['opex']})",N),('Gross value at NIY (£)',f"=B4/{I['niy']}",N),
 ('Purchase price (£)',f"=B5/(1+{I['fees']})-B7",N),('SDLT (£)',f"=VLOOKUP(B5,SDLT!$A$2:$C$4,3)+(B5-VLOOKUP(B5,SDLT!$A$2:$C$4,1))*VLOOKUP(B5,SDLT!$A$2:$C$4,2)",N),
 ('Developer profit (£)',f"=B6-{I['land']}-{I['cost']}-SUM({row('Coupon')})-SUM({row('Penalty')})",N),('Funder\'s coupon (£)',f"=SUM({row('Coupon')})",N),('Delay penalties (£)',f"=SUM({row('Penalty')})",N),
 ('Developer margin on price',f"=B8/B6",P),('Investor yield on cost',f"=B4/(B6+B7+B6*{I['fees']})",P),('Investor IRR (XIRR)',f"=XIRR({row('Investor cash flow')},{row('Date')})",P),('Net profit to investor (£)',f"=SUM({row('Investor cash flow')})",N)])
fy=wb.create_sheet('FY Summary');fy.append(['Fiscal year','Net operating income (£)','Gross income (£)'])
for i,y in enumerate(range(27,40),2):fy.cell(i,1,f'FY{y}');fy.cell(i,2,f"=SUMIF({row('Fiscal year')},A{i},{row('Net operating income')})");fy.cell(i,3,f"=SUMIF({row('Fiscal year')},A{i},{row('Gross income')})")
rt.cell(16,1,'NOI in FY30 (£)');rt.cell(16,2,"=INDEX('FY Summary'!$B$2:$B$14,MATCH(\"FY30\",'FY Summary'!$A$2:$A$14,0))").number_format=N
wb.save(OUT+'ff_gbp.xlsx')

# ================= 22D off-plan community with a sukuk =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='RAWDAT AL KHAIR COMMUNITY – OFF-PLAN WITH IJARA SUKUK (SAR)'
I=inputs(ws,[('start','Model start',datetime.date(2026,12,31),D),('villas','Villas',540,N),('vprice','Villa price (SAR)',2750000,N),('apts','Apartments',880,N),('aprice','Apartment price (SAR)',1380000,N),
 ('lag','Buyer payment lag (quarters)',2,'0'),('down','Down payment',0.2,P),('build','Build (quarters)',14,'0'),('vcost','Villa build cost (SAR)',1350000,N),('acost','Apartment build cost (SAR)',640000,N),('infra','Infrastructure (SAR)',310000000,N),
 ('land','Land (SAR)',450000000,N),('rett','Transaction tax absorbed (5% on this share of units)',0.3,P),('suk','Sukuk issued (SAR)',900000000,N),('rate','Sukuk periodic distribution a year',0.0675,P),('mat','Sukuk maturity (quarter)',20,'0'),
 ('zak','Zakat rate',0.025,P),('eq','Equity (SAR)',700000000,N)])
sp=wb.create_sheet('Sales Plan');sp.append(['Quarter','Villas sold','Apartments sold'])
vs=[30,40,45,45,45,40,40,35,35,30,30,30,25,25,25,20];as_=[60,70,70,70,65,65,60,60,55,55,50,50,45,40,35,30]
for i in range(16):sp.cell(i+2,1,i+1);sp.cell(i+2,2,vs[i]);sp.cell(i+2,3,as_[i])
cf=wb.create_sheet('Cash Flow');T=24
labs=['Quarter','Date','Villas sold','Apartments sold','Contracted sales','Down payments','Balance due (lagged)','Handover payments','Collections','Construction','Infrastructure','Land','Transaction tax absorbed','Escrow released','Escrow balance','Sukuk proceeds','Periodic distribution','Sukuk redemption','Coverage ratio','Cumulative profit','Zakat','Cash before equity','Equity in','Distributions','Cash balance','Equity cash flow']
def scf(t,c,p,R):
    q=f"{c}$3";first=t==0;d={'Quarter':t,'Date':f"EOMONTH({I['start']},3*{q})"}
    d['Villas sold']=f"IF(AND({q}>=1,{q}<=16),INDEX('Sales Plan'!$B$2:$B$17,MAX(1,MIN(16,{q}))),0)";d['Apartments sold']=f"IF(AND({q}>=1,{q}<=16),INDEX('Sales Plan'!$C$2:$C$17,MAX(1,MIN(16,{q}))),0)"
    d['Contracted sales']=f"{c}{R['Villas sold']}*{I['vprice']}+{c}{R['Apartments sold']}*{I['aprice']}"
    d['Down payments']=f"{c}{R['Contracted sales']}*{I['down']}"
    d['Balance due (lagged)']=f"IF(AND({q}<{I['build']},{q}>{I['lag']}),OFFSET({c}{R['Contracted sales']},0,-{I['lag']})*0.4,0)" if t>=1 else '0'
    d['Handover payments']=f"IF({q}={I['build']},SUM($B${R['Contracted sales']}:{c}{R['Contracted sales']})*(1-{I['down']})-SUM($B${R['Balance due (lagged)']}:{p}{R['Balance due (lagged)']}),IF({q}>{I['build']},{c}{R['Contracted sales']}*(1-{I['down']}),0))" if t>=1 else '0'
    d['Collections']=f"{c}{R['Down payments']}+{c}{R['Balance due (lagged)']}+{c}{R['Handover payments']}"
    d['Construction']=f"IF(AND({q}>=1,{q}<={I['build']}),({I['villas']}*{I['vcost']}+{I['apts']}*{I['acost']})/{I['build']},0)"
    d['Infrastructure']=f"IF(AND({q}>=1,{q}<=6),{I['infra']}/6,0)";d['Land']=f"IF({q}=0,{I['land']},0)"
    d['Transaction tax absorbed']=f"{c}{R['Contracted sales']}*0.05*{I['rett']}"
    pe='0' if first else f"{p}{R['Escrow balance']}"
    d['Escrow released']=f"IF({q}>={I['build']},{pe}+{c}{R['Collections']},MIN({pe}+{c}{R['Collections']},{c}{R['Construction']}+{c}{R['Infrastructure']}))"
    d['Escrow balance']=f"{pe}+{c}{R['Collections']}-{c}{R['Escrow released']}"
    d['Sukuk proceeds']=f"IF({q}=1,{I['suk']},0)";d['Periodic distribution']=f"IF(AND({q}>=2,{q}<={I['mat']}),{I['suk']}*{I['rate']}/4,0)";d['Sukuk redemption']=f"IF({q}={I['mat']},{I['suk']},0)"
    cfads=f"({c}{R['Escrow released']}-{c}{R['Construction']}-{c}{R['Infrastructure']}-{c}{R['Transaction tax absorbed']})"
    pc='0' if first else f"{p}{R['Cash balance']}"
    d['Coverage ratio']=f"IF({c}{R['Periodic distribution']}>0,MAX(0,{pc}+{cfads})/{c}{R['Periodic distribution']},0)"
    d['Cumulative profit']=f"SUM($B${R['Escrow released']}:{c}{R['Escrow released']})-SUM($B${R['Construction']}:{c}{R['Construction']})-SUM($B${R['Infrastructure']}:{c}{R['Infrastructure']})-SUM($B${R['Transaction tax absorbed']}:{c}{R['Transaction tax absorbed']})-SUM($B${R['Periodic distribution']}:{c}{R['Periodic distribution']})-{I['land']}"
    d['Zakat']=f"IF(MONTH({c}{R['Date']})=12,MAX(0,{I['eq']}+{c}{R['Cumulative profit']})*{I['zak']},0)"
    d['Cash before equity']=f"{pc}+{c}{R['Escrow released']}-{c}{R['Construction']}-{c}{R['Infrastructure']}-{c}{R['Land']}-{c}{R['Transaction tax absorbed']}+{c}{R['Sukuk proceeds']}-{c}{R['Periodic distribution']}-{c}{R['Sukuk redemption']}-{c}{R['Zakat']}"
    d['Equity in']=f"MAX(0,-{c}{R['Cash before equity']})"
    d['Distributions']=f"IF({q}>={I['mat']},MAX(0,{c}{R['Cash before equity']}),0)"
    d['Cash balance']=f"{c}{R['Cash before equity']}+{c}{R['Equity in']}-{c}{R['Distributions']}"
    d['Equity cash flow']=f"{c}{R['Distributions']}-{c}{R['Equity in']}"
    return d
R=timeline(cf,labs,T,scf);E=L(T+2);row=lambda n:f"'Cash Flow'!$B${R[n]}:${E}${R[n]}"
for t in range(T+1):cf[f'{L(t+2)}{R["Date"]}'].number_format=D
table(wb.create_sheet('Returns'),[('Levered IRR',f"=(1+IRR({row('Equity cash flow')},0.02))^4-1",P),('Net profit (SAR)',f"=SUM({row('Equity cash flow')})",N),('Minimum sukuk coverage',f"=_xlfn.MINIFS({row('Coverage ratio')},{row('Periodic distribution')},\">0\")",X),('Equity injected (SAR)',f"=SUM({row('Equity in')})",N),
 ('Zakat paid (SAR)',f"=SUM({row('Zakat')})",N),('Peak escrow balance (SAR)',f"=MAX({row('Escrow balance')})",N),('Units sold by handover',f"=SUMIF({row('Quarter')},\"<=\"&{I['build']},{row('Villas sold')})+SUMIF({row('Quarter')},\"<=\"&{I['build']},{row('Apartments sold')})",N),
 ('Transaction tax absorbed (SAR)',f"=SUM({row('Transaction tax absorbed')})",N),('Total revenue (SAR)',f"=SUM({row('Contracted sales')})",N),('Sukuk distributions (SAR)',f"=SUM({row('Periodic distribution')})",N)])
wb.save(OUT+'sukuk_sar.xlsx')

# ================= 22E secondaries fund =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Fund';ws['A1']='HELIX SECONDARIES FUND II – EUR'
I=inputs(ws,[('start','First close',datetime.date(2026,6,30),D),('fee','Management fee (% of NAV a year)',0.01,P),('line','Subscription line bridge (quarters)',2,'0'),('lrate','Line rate a year',0.045,P),
 ('pref','Preferred return',0.08,P),('carry','Carried interest',0.125,P),('cu','GP catch-up',1.0,P),('rep','Reporting date',datetime.date(2031,6,30),D)])
bk=wb.create_sheet('Interests');bk.append(['Interest','Purchase quarter','NAV at purchase (€)','Discount to NAV','Unfunded at purchase (€)','Calls a year (% of unfunded)','Distributions a year (% of NAV)','NAV growth a year','Fund end (quarter)'])
for i,row_ in enumerate([('Buyout 2018',1,62e6,0.12,8e6,0.3,0.18,0.06,20),('Growth 2019',2,48e6,0.18,14e6,0.35,0.12,0.08,24),('Infra 2017',2,55e6,0.07,3e6,0.4,0.1,0.04,22),('Venture 2020',4,30e6,0.3,12e6,0.3,0.06,0.1,28),('RE debt 2019',5,40e6,0.09,6e6,0.5,0.2,0.03,16),('Buyout 2021',7,70e6,0.15,25e6,0.3,0.08,0.07,30)],2):
    for j,v in enumerate(row_,1):bk.cell(i,j,v)
cf=wb.create_sheet('Cash Flow');T=32
labs=['Quarter','Date']+[f'{x} {k}' for k in range(1,7) for x in ('NAV','Unfunded','Call','Distribution','Purchase')]+['Purchases','Calls','Distributions','NAV','Fees','Investments','Line drawn','Line repaid','Line balance','Line interest','LP called','Net to LPs','ROC','Pref account','Pref paid','Catch-up','Split LP','Split GP','LP cash flow','Gross cash flow']
def ecf(t,c,p,R):
    q=f"{c}$3";first=t==0;d={'Quarter':t,'Date':f"EOMONTH({I['start']},3*{q})"}
    for k in range(1,7):
        r=k+1;pq=f"Interests!$B${r}";nav0=f"Interests!$C${r}";disc=f"Interests!$D${r}";un0=f"Interests!$E${r}";cr=f"Interests!$F${r}";dy=f"Interests!$G${r}";gr=f"Interests!$H${r}";end=f"Interests!$I${r}"
        pn='0' if first else f"{p}{R[f'NAV {k}']}";pu='0' if first else f"{p}{R[f'Unfunded {k}']}"
        d[f'Call {k}']=f"IF(AND({q}>{pq},{q}<{end}),{pu}*{cr}/4,0)"
        d[f'Distribution {k}']=f"IF({q}={end},{pn}*(1+{gr}/4)+{c}{R[f'Call {k}']},IF(AND({q}>{pq},{q}<{end}),{pn}*{dy}/4,0))"
        d[f'NAV {k}']=f"IF({q}={pq},{nav0},IF(AND({q}>{pq},{q}<{end}),{pn}*(1+{gr}/4)+{c}{R[f'Call {k}']}-{c}{R[f'Distribution {k}']},0))"
        d[f'Unfunded {k}']=f"IF({q}={pq},{un0},IF(AND({q}>{pq},{q}<{end}),{pu}-{c}{R[f'Call {k}']},0))"
        d[f'Purchase {k}']=f"IF({q}={pq},{nav0}*(1-{disc}),0)"
    S=lambda x:"+".join(f"{c}{R[f'{x} {k}']}" for k in range(1,7))
    d['Purchases']=S('Purchase');d['Calls']=S('Call');d['Distributions']=S('Distribution');d['NAV']=S('NAV')
    d['Fees']='0' if first else f"{p}{R['NAV']}*{I['fee']}/4"
    d['Investments']=f"{c}{R['Purchases']}+{c}{R['Calls']}+{c}{R['Fees']}"
    d['Line drawn']=f"{c}{R['Investments']}"
    d['Line repaid']='0' if first else f"IF({q}>{I['line']},OFFSET({c}{R['Line drawn']},0,-{I['line']}),0)"
    d['Line balance']=f"SUM($B${R['Line drawn']}:{c}{R['Line drawn']})-SUM($B${R['Line repaid']}:{c}{R['Line repaid']})"
    d['Line interest']='0' if first else f"{p}{R['Line balance']}*{I['lrate']}/4"
    d['LP called']=f"{c}{R['Line repaid']}+{c}{R['Line interest']}+IF({q}={T},{c}{R['Line balance']},0)"
    d['Net to LPs']=f"{c}{R['Distributions']}"
    called=f"SUM($B${R['LP called']}:{c}{R['LP called']})";roc0='0' if first else f"SUM($B${R['ROC']}:{p}{R['ROC']})"
    d['ROC']=f"MIN({c}{R['Net to LPs']},{called}-{roc0})"
    d['Pref account']='0' if first else f"({p}{R['Pref account']}-{p}{R['Pref paid']})*(1+{I['pref']}/4)+({called}-{c}{R['LP called']}-{roc0})*{I['pref']}/4"
    d['Pref paid']=f"MIN({c}{R['Net to LPs']}-{c}{R['ROC']},{c}{R['Pref account']})"
    cuacc=f"SUM($B${R['Pref paid']}:{c}{R['Pref paid']})*{I['carry']}/(1-{I['carry']})-{('0' if first else 'SUM($B'+str(R['Catch-up'])+':'+p+str(R['Catch-up'])+')')}"
    d['Catch-up']=f"MIN(({c}{R['Net to LPs']}-{c}{R['ROC']}-{c}{R['Pref paid']})*{I['cu']},MAX(0,{cuacc}))"
    rest=f"({c}{R['Net to LPs']}-{c}{R['ROC']}-{c}{R['Pref paid']}-{c}{R['Catch-up']})"
    d['Split LP']=f"{rest}*(1-{I['carry']})";d['Split GP']=f"{rest}*{I['carry']}"
    d['LP cash flow']=f"{c}{R['ROC']}+{c}{R['Pref paid']}+{c}{R['Split LP']}-{c}{R['LP called']}"
    d['Gross cash flow']=f"{c}{R['Distributions']}-{c}{R['Purchases']}-{c}{R['Calls']}"
    return d
R=timeline(cf,labs,T,ecf);E=L(T+2);row=lambda n:f"'Cash Flow'!$B${R[n]}:${E}${R[n]}"
for t in range(T+1):cf[f'{L(t+2)}{R["Date"]}'].number_format=D
rq=f"MATCH({I['rep']},{row('Date')},0)"
table(wb.create_sheet('Returns'),[('Net IRR to LPs (XIRR)',f"=XIRR({row('LP cash flow')},{row('Date')})",P),('Gross IRR (XIRR)',f"=XIRR({row('Gross cash flow')},{row('Date')})",P),
 ('TVPI (net)',f"=SUMIF({row('LP cash flow')},\">0\")/SUM({row('LP called')})",X),('DPI at the reporting date',f"=SUM(OFFSET('Cash Flow'!$B${R['ROC']},0,0,1,{rq}))/SUM(OFFSET('Cash Flow'!$B${R['LP called']},0,0,1,{rq}))+SUM(OFFSET('Cash Flow'!$B${R['Pref paid']},0,0,1,{rq}))/SUM(OFFSET('Cash Flow'!$B${R['LP called']},0,0,1,{rq}))+SUM(OFFSET('Cash Flow'!$B${R['Split LP']},0,0,1,{rq}))/SUM(OFFSET('Cash Flow'!$B${R['LP called']},0,0,1,{rq}))",X),
 ('RVPI at the reporting date',f"=INDEX({row('NAV')},{rq})/SUM(OFFSET('Cash Flow'!$B${R['LP called']},0,0,1,{rq}))",X),('Discount captured (€)',"=SUMPRODUCT(Interests!$C$2:$C$7,Interests!$D$2:$D$7)",N),
 ('Peak unfunded exposure (€)',f"=MAX({row('Unfunded 1')})+0*1",N),('GP carry (€)',f"=SUM({row('Catch-up')})+SUM({row('Split GP')})",N),('Line interest (€)',f"=SUM({row('Line interest')})",N),('LP net profit (€)',f"=SUM({row('LP cash flow')})",N)])
# the true peak unfunded exposure: across the six interests, per quarter
rt=wb['Returns'];cf.cell(R['Gross cash flow']+2,1,'Unfunded exposure')
for t in range(T+1):c=L(t+2);cf[f'{c}{R["Gross cash flow"]+2}']="="+"+".join(f"{c}{R[f'Unfunded {k}']}" for k in range(1,7))
rt['B9']=f"=MAX('Cash Flow'!$B${R['Gross cash flow']+2}:${E}${R['Gross cash flow']+2})"
wb.save(OUT+'second_eur.xlsx')
print('ok')
