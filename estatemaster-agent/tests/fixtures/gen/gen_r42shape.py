# Round 42 (synthetic): the shapes a figure takes. python3 -I gen/gen_r42shape.py → models/r42_shape_sar.xlsx
#  unit words naming figures (GFA, NLA, Years); an occupancy row that is 0 for its first 9 periods (beyond the 8 cells a row used to be read
#  from); a levered IRR whose cash flow never turns positive (#NUM!) beside a working unlevered IRR
import openpyxl, os
from openpyxl.utils import get_column_letter as L
D = os.path.join(os.path.dirname(__file__), '..', 'models')
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Inputs'
for r in [('SAHAB OFFICES – SHAPES (SAR)', None), ('GFA', 50000), ('NLA', 40000), ('Years', 10), ('Rent per sqm', 1200), ('Exit yield', 0.07), ('Cost per sqm GFA', 4000)]:
    ws.append(list(r))
cf = wb.create_sheet('Cash Flow'); n = 12
cf.cell(1, 1, 'Period')
occ = [0] * 9 + [0.6, 0.9, 0.95]
cf.cell(2, 1, 'Occupancy'); cf.cell(3, 1, 'Rent'); cf.cell(4, 1, 'Cost'); cf.cell(5, 1, 'Exit'); cf.cell(6, 1, 'Net cash flow'); cf.cell(7, 1, 'Fee to partner')
for i in range(n):
    c = i + 2; cl = L(c)
    cf.cell(1, c, i + 1)
    cf.cell(2, c, occ[i])
    cf.cell(3, c, f'=Inputs!$B$3*Inputs!$B$5*{cl}2')
    cf.cell(4, c, '=-Inputs!$B$2*Inputs!$B$7/3' if i < 3 else 0)
    cf.cell(5, c, f'=IF({cl}1={n},{cl}3/Inputs!$B$6,0)')
    cf.cell(6, c, f'={cl}3+{cl}4+{cl}5')
    cf.cell(7, c, -100000)
last = L(n + 1)
r = wb.create_sheet('Returns')
r.append(['Unlevered IRR', f"=IRR('Cash Flow'!B6:{last}6)"])
r.append(['Levered IRR', f"=IRR('Cash Flow'!B7:{last}7)"])
r.append(['Net profit', f"=SUM('Cash Flow'!B6:{last}6)"])
wb.save(os.path.join(D, 'r42_shape_sar.xlsx'))
