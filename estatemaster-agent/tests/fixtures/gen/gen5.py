# five synthetic development / investment models of types the app has not seen; formulas only (LibreOffice stores the results)
import openpyxl,datetime
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
OUT='/tmp/claude-0/-home-user-workspace/c7e8f42c-9ed0-5ba7-9852-d8a131185935/scratchpad/m5/'
def inputs(ws,rows,lc=1,vc=3,uc=4,start=3):
    r=start;ref={}
    for row in rows:
        if isinstance(row,str): ws.cell(r,lc,row).font=openpyxl.styles.Font(bold=True);r+=1;continue
        k,lab,val,unit,fmt=row;ws.cell(r,lc,lab);c=ws.cell(r,vc,val);c.number_format=fmt
        if unit:ws.cell(r,uc,unit)
        ref[k]=f"{ws.title if ' ' not in ws.title else chr(39)+ws.title+chr(39)}!${L(vc)}${r}";r+=1
    return ref
P='0.00%';N='#,##0';N2='#,##0.00'
# ---------------- A. logistics park, AED, quarterly ----------------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='AL QUOZ LOGISTICS PARK – DEVELOPMENT APPRAISAL';ws['A2']='Currency: AED'
I=inputs(ws,['SITE & AREA',('site','Site area (sqm)',120000,'sqm',N),('cov','Site coverage',0.55,'%',P),('gla','Leasable area (sqm GLA)','=C4*C5','sqm',N),
  'COSTS',('land','Land price',125000000,'AED',N),('cc','Construction cost',2100,'AED/sqm',N),('fees','Professional fees',0.08,'% of construction',P),('cont','Contingency',0.05,'% of construction',P),('per','Construction period',6,'quarters','0'),
  'REVENUE',('rent','Market rent',420,'AED/sqm/yr',N),('lu','Lease-up period',4,'quarters after completion','0'),('opx','Non-recoverable opex',0.10,'% of rent',P),('ey','Exit yield',0.0725,'',P),('exq','Exit quarter',14,'quarter','0'),('sc','Sales costs',0.015,'% of GDV',P),
  'FINANCE',('ltc','Loan to cost',0.60,'',P),('ir','Interest rate',0.07,'p.a.',P),('af','Arrangement fee',0.01,'of facility',P)])
cf=wb.create_sheet('CF');Q=16;cf['A1']='QUARTERLY CASH FLOW (AED)';cf['A3']='Quarter';cf['A4']='Quarter ending'
for q in range(1,Q+1):
    c=L(q+2);cf[f'{c}3']=q;cf[f'{c}4']=datetime.date(2026,3,31) if q==1 else f'=EDATE({L(q+1)}4,3)';cf[f'{c}4'].number_format='mmm-yy'
labels=['Land','Construction','Professional fees','Contingency','Total development cost','Occupancy','Gross rent','Non-recoverable opex','Net operating income','Gross sale value','Sales costs','Net sale proceeds','Unlevered cash flow','Debt drawdown','Interest','Interest capitalised','Interest paid','Arrangement fee','Debt repayment','Loan balance','Levered cash flow','Cumulative levered']
R={l:6+i for i,l in enumerate(labels)}
for l,r in R.items():cf[f'A{r}']=l
for q in range(1,Q+1):
    c=L(q+2);p=L(q+1);g=lambda l:f'{c}{R[l]}'
    cf[g('Land')]=f'=IF({c}$3=1,-{I["land"]},0)'
    cf[g('Construction')]=f'=IF({c}$3<={I["per"]},-{I["gla"]}*{I["cc"]}/{I["per"]},0)'
    cf[g('Professional fees')]=f'={g("Construction")}*{I["fees"]}';cf[g('Contingency')]=f'={g("Construction")}*{I["cont"]}'
    cf[g('Total development cost')]=f'=SUM({c}{R["Land"]}:{c}{R["Contingency"]})'
    cf[g('Occupancy')]=f'=IF({c}$3>{I["exq"]},0,MIN(1,MAX(0,({c}$3-{I["per"]})/{I["lu"]})))'
    cf[g('Gross rent')]=f'={I["gla"]}*{I["rent"]}/4*{g("Occupancy")}'
    cf[g('Non-recoverable opex')]=f'=-{g("Gross rent")}*{I["opx"]}'
    cf[g('Net operating income')]=f'={g("Gross rent")}+{g("Non-recoverable opex")}'
    cf[g('Gross sale value')]=f'=IF({c}$3={I["exq"]},{I["gla"]}*{I["rent"]}*(1-{I["opx"]})/{I["ey"]},0)'
    cf[g('Sales costs')]=f'=-{g("Gross sale value")}*{I["sc"]}'
    cf[g('Net sale proceeds')]=f'={g("Gross sale value")}+{g("Sales costs")}'
    cf[g('Unlevered cash flow')]=f'={g("Total development cost")}+{g("Net operating income")}+{g("Net sale proceeds")}'
    cf[g('Debt drawdown')]=f'=-{g("Total development cost")}*{I["ltc"]}'
    prevbal=f'{p}{R["Loan balance"]}' if q>1 else '0'
    cf[g('Interest')]=f'={prevbal}*{I["ir"]}/4'
    cf[g('Interest capitalised')]=f'=IF({c}$3<={I["per"]},{g("Interest")},0)'
    cf[g('Interest paid')]=f'=-({g("Interest")}-{g("Interest capitalised")})'
    cf[g('Arrangement fee')]=f'=IF({c}$3=1,-{I["af"]}*{I["ltc"]}*-SUM($C${R["Total development cost"]}:$R${R["Total development cost"]}),0)'
    cf[g('Debt repayment')]=f'=IF({c}$3={I["exq"]},-({prevbal}+{g("Debt drawdown")}+{g("Interest capitalised")}),0)'
    cf[g('Loan balance')]=f'={prevbal}+{g("Debt drawdown")}+{g("Interest capitalised")}+{g("Debt repayment")}'
    cf[g('Levered cash flow')]=f'={g("Unlevered cash flow")}+{g("Debt drawdown")}+{g("Interest paid")}+{g("Arrangement fee")}+{g("Debt repayment")}'
    cf[g('Cumulative levered')]=f'={g("Levered cash flow")}' if q==1 else f'={p}{R["Cumulative levered"]}+{g("Levered cash flow")}'
rs=wb.create_sheet('Returns');rr=lambda l:f"CF!$C${R[l]}:$R${R[l]}"
for i,(lab,f,fmt) in enumerate([('Unlevered IRR',f'=(1+IRR({rr("Unlevered cash flow")}))^4-1',P),('Levered IRR',f'=(1+IRR({rr("Levered cash flow")}))^4-1',P),
   ('Total development cost',f'=-SUM({rr("Total development cost")})',N),('Gross development value',f'=SUM({rr("Gross sale value")})',N),
   ('Net profit',f'=SUM({rr("Levered cash flow")})',N),('Profit on cost',f'=C7/C5',P),('Development yield (yield on cost)',f'={I["gla"]}*{I["rent"]}*(1-{I["opx"]})/C5',P),
   ('Peak equity requirement',f'=-MIN({rr("Cumulative levered")})',N),('Equity multiple',f'=SUMIF({rr("Levered cash flow")},">0")/-SUMIF({rr("Levered cash flow")},"<0")','0.00"x"')]):
    rs.cell(3+i,1,lab);c=rs.cell(3+i,3,f);c.number_format=fmt
rs['A1']='RETURNS SUMMARY'
wb.save(OUT+'logistics_aed.xlsx')
# ---------------- B. student housing, GBP, case switch ----------------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Control';ws['A1']='RIVERSIDE STUDENT LIVING – INVESTMENT MODEL';ws['A3']='Active case (1 = Base, 2 = Downside, 3 = Upside)';ws['C3']=1;ws['A4']='Currency';ws['C4']='GBP'
a=wb.create_sheet('Assumptions');a.append(['Assumption','Unit','Base','Downside','Upside','Live'])
rows=[('Beds','beds',650,650,650),('Weekly rent','£/bed/week',215,200,225),('Let weeks','weeks/yr',51,51,51),('Occupancy','%',0.97,0.92,0.98),('Rental growth','% p.a.',0.03,0.02,0.035),
 ('Operating costs','£/bed/yr',2900,3100,2800),('Cost inflation','% p.a.',0.03,0.035,0.03),('Total development cost','£',78000000,78000000,78000000),('Exit yield','%',0.05,0.055,0.0475),
 ('Senior loan to cost','%',0.55,0.55,0.55),('Interest rate','%',0.0625,0.0725,0.0575),('Refinance year','year',3,3,3),('Refinance LTV','%',0.60,0.55,0.60),('Hold period','years',7,7,7)]
AR={}
for i,(lab,u,b,d,up) in enumerate(rows):
    r=i+2;a.append([lab,u,b,d,up,f'=CHOOSE(Control!$C$3,C{r},D{r},E{r})']);AR[lab]=f'Assumptions!$F${r}'
    for col in 'CDEF':a[f'{col}{r}'].number_format=P if u.startswith('%') else N
op=wb.create_sheet('Operations');op['A1']='ANNUAL CASH FLOW (£)';op['A3']='Year';yrs=range(0,8)
ol=['Development cost','Gross rental income','Operating costs','Net operating income','Exit value','Unlevered cash flow','Opening debt','Refinance proceeds','Closing debt','Interest','DSCR','Debt flows','Levered cash flow']
OR={l:5+i for i,l in enumerate(ol)}
for l,r in OR.items():op[f'A{r}']=l
for y in yrs:
    c=L(y+3);p=L(y+2);op[f'{c}3']=y;g=lambda l:f'{c}{OR[l]}';A=AR
    op[g('Development cost')]=f'=IF({c}$3=0,-{A["Total development cost"]},0)'
    op[g('Gross rental income')]=f'=IF(AND({c}$3>=1,{c}$3<={A["Hold period"]}),{A["Beds"]}*{A["Weekly rent"]}*{A["Let weeks"]}*{A["Occupancy"]}*(1+{A["Rental growth"]})^({c}$3-1),0)'
    op[g('Operating costs')]=f'=IF({g("Gross rental income")}>0,-{A["Beds"]}*{A["Operating costs"]}*(1+{A["Cost inflation"]})^({c}$3-1),0)'
    op[g('Net operating income')]=f'={g("Gross rental income")}+{g("Operating costs")}'
    op[g('Exit value')]=f'=IF({c}$3={A["Hold period"]},{g("Net operating income")}*(1+{A["Rental growth"]})/{A["Exit yield"]},0)'
    op[g('Unlevered cash flow')]=f'={g("Development cost")}+{g("Net operating income")}+{g("Exit value")}'
    op[g('Opening debt')]=f'=IF({c}$3=0,{A["Total development cost"]}*{A["Senior loan to cost"]},{p}{OR["Closing debt"]})'
    op[g('Refinance proceeds')]=f'=IF({c}$3={A["Refinance year"]},MAX(0,{A["Refinance LTV"]}*{g("Net operating income")}*(1+{A["Rental growth"]})/{A["Exit yield"]}-{g("Opening debt")}),0)'
    op[g('Closing debt')]=f'=IF({c}$3={A["Hold period"]},0,{g("Opening debt")}+{g("Refinance proceeds")})'
    op[g('Interest')]=f'=IF({c}$3=0,0,-{g("Opening debt")}*{A["Interest rate"]})'
    op[g('DSCR')]=f'=IF({g("Interest")}<0,{g("Net operating income")}/-{g("Interest")},"")'
    op[g('Debt flows')]=f'=IF({c}$3=0,{g("Opening debt")},{g("Refinance proceeds")}+{g("Interest")}-IF({c}$3={A["Hold period"]},{g("Opening debt")},0))'
    op[g('Levered cash flow')]=f'={g("Unlevered cash flow")}+{g("Debt flows")}'
sm=wb.create_sheet('Summary');sm['A1']='RETURNS';rg=lambda l:f'Operations!$C${OR[l]}:$J${OR[l]}'
for i,(lab,f,fmt) in enumerate([('Levered IRR',f'=IRR({rg("Levered cash flow")})',P),('Unlevered IRR',f'=IRR({rg("Unlevered cash flow")})',P),('Equity multiple',f'=SUMIF({rg("Levered cash flow")},">0")/-SUMIF({rg("Levered cash flow")},"<0")','0.00"x"'),
   ('Yield on cost',f'=Operations!D{OR["Net operating income"]}/{AR["Total development cost"]}',P),('Minimum DSCR',f'=MIN({rg("DSCR")})','0.00"x"'),('Net profit',f'=SUM({rg("Levered cash flow")})',N),
   ('Total development cost',f'={AR["Total development cost"]}',N),('Cash-on-cash yield (year 2)',f'=Operations!E{OR["Levered cash flow"]}/-Operations!C{OR["Levered cash flow"]}',P)]):
    sm.cell(3+i,1,lab);c=sm.cell(3+i,3,f);c.number_format=fmt
wb.save(OUT+'student_gbp.xlsx')
# ---------------- C. value-add multifamily, USD, LP/GP ----------------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Assumptions';ws['A1']='OAK HOLLOW APARTMENTS – VALUE-ADD ACQUISITION';ws['A2']='All figures in USD'
I=inputs(ws,['PROPERTY',('u','Units',240,'units',N),('pp','Purchase price',48000000,'$',N),('ccx','Closing costs',0.02,'% of price',P),
 'REVENUE',('ir_','In-place rent',1450,'$/unit/month',N),('rpy','Renovations per year',60,'units/yr',N),('rc','Renovation cost per unit',18000,'$',N),('rp','Rent premium (renovated)',225,'$/unit/month',N),
   ('vac','Vacancy & credit loss',0.06,'',P),('oi','Other income',75,'$/unit/month',N),('g','Rent growth',0.035,'p.a.',P),('opr','Operating expense ratio',0.38,'of EGI',P),
 'EXIT',('ec','Exit cap rate',0.055,'',P),('sc','Selling costs',0.02,'of sale price',P),
 'DEBT',('ltv','Loan to value',0.65,'',P),('ir','Interest rate',0.06,'',P),('am','Amortization',30,'years','0'),
 'EQUITY WATERFALL',('pref','Preferred return',0.08,'',P),('lps','LP share of equity',0.90,'',P),('prom','GP promote above pref',0.20,'',P)],lc=1,vc=3,uc=4)
cf=wb.create_sheet('Pro Forma');cf['A1']='ANNUAL PRO FORMA (USD)';cf['A3']='Year'
pl=['Renovated units (cumulative)','Gross potential rent','Vacancy','Other income','Effective gross income','Operating expenses','Net operating income','Renovation capex','Acquisition','Sale price','Selling costs','Unlevered cash flow',
    'Loan proceeds','Debt service','Loan balance (end)','Loan payoff','Levered cash flow','Equity contributed','Distributions','Hurdle balance','Tier 1 (return of capital + pref)','Tier 2 (residual)','LP cash flow','GP cash flow']
R={l:5+i for i,l in enumerate(pl)}
for l,r in R.items():cf[f'A{r}']=l
for y in range(0,6):
    c=L(y+3);p=L(y+2);cf[f'{c}3']=y;g=lambda l:f'{c}{R[l]}'
    op_=f'AND({c}$3>=1,{c}$3<=5)'
    cf[g('Renovated units (cumulative)')]=f'=MIN({I["u"]},{c}$3*{I["rpy"]})'
    cf[g('Gross potential rent')]=f'=IF({op_},(({I["u"]}-{g("Renovated units (cumulative)")})*{I["ir_"]}+{g("Renovated units (cumulative)")}*({I["ir_"]}+{I["rp"]}))*12*(1+{I["g"]})^({c}$3-1),0)'
    cf[g('Vacancy')]=f'=-{g("Gross potential rent")}*{I["vac"]}'
    cf[g('Other income')]=f'=IF({op_},{I["u"]}*{I["oi"]}*12,0)'
    cf[g('Effective gross income')]=f'={g("Gross potential rent")}+{g("Vacancy")}+{g("Other income")}'
    cf[g('Operating expenses')]=f'=-{g("Effective gross income")}*{I["opr"]}'
    cf[g('Net operating income')]=f'={g("Effective gross income")}+{g("Operating expenses")}'
    prevren=f'{p}{R["Renovated units (cumulative)"]}' if y>0 else '0'
    cf[g('Renovation capex')]=f'=-({g("Renovated units (cumulative)")}-{prevren})*{I["rc"]}'
    cf[g('Acquisition')]=f'=IF({c}$3=0,-{I["pp"]}*(1+{I["ccx"]}),0)'
    cf[g('Sale price')]=f'=IF({c}$3=5,{g("Net operating income")}*(1+{I["g"]})/{I["ec"]},0)'
    cf[g('Selling costs')]=f'=-{g("Sale price")}*{I["sc"]}'
    cf[g('Unlevered cash flow')]=f'={g("Net operating income")}+{g("Renovation capex")}+{g("Acquisition")}+{g("Sale price")}+{g("Selling costs")}'
    cf[g('Loan proceeds')]=f'=IF({c}$3=0,{I["pp"]}*{I["ltv"]},0)'
    cf[g('Debt service')]=f'=IF({op_},PMT({I["ir"]},{I["am"]},$C${R["Loan proceeds"]}),0)'
    prevbal=f'{p}{R["Loan balance (end)"]}' if y>0 else '0'
    cf[g('Loan balance (end)')]=f'=IF({c}$3=0,{g("Loan proceeds")},IF({c}$3>5,0,{prevbal}*(1+{I["ir"]})+{g("Debt service")}))'
    cf[g('Loan payoff')]=f'=IF({c}$3=5,-{g("Loan balance (end)")},0)'
    cf[g('Levered cash flow')]=f'={g("Unlevered cash flow")}+{g("Loan proceeds")}+{g("Debt service")}+{g("Loan payoff")}'
    cf[g('Equity contributed')]=f'=MAX(0,-{g("Levered cash flow")})';cf[g('Distributions')]=f'=MAX(0,{g("Levered cash flow")})'
    prevh=f'{p}{R["Hurdle balance"]}' if y>0 else '0'
    cf[g('Tier 1 (return of capital + pref)')]=f'=MIN({g("Distributions")},{prevh}*(1+{I["pref"]})+{g("Equity contributed")})'
    cf[g('Hurdle balance')]=f'={prevh}*(1+{I["pref"]})+{g("Equity contributed")}-{g("Tier 1 (return of capital + pref)")}'
    cf[g('Tier 2 (residual)')]=f'={g("Distributions")}-{g("Tier 1 (return of capital + pref)")}'
    cf[g('LP cash flow')]=f'=-{g("Equity contributed")}*{I["lps"]}+{g("Tier 1 (return of capital + pref)")}*{I["lps"]}+{g("Tier 2 (residual)")}*(1-{I["prom"]})*{I["lps"]}'
    cf[g('GP cash flow')]=f'={g("Levered cash flow")}-{g("LP cash flow")}'
sm=wb.create_sheet('Returns');rg=lambda l:f"'Pro Forma'!$C${R[l]}:$H${R[l]}"
for i,(lab,f,fmt) in enumerate([('Unlevered IRR',f'=IRR({rg("Unlevered cash flow")})',P),('Levered IRR',f'=IRR({rg("Levered cash flow")})',P),('Equity multiple',f'=SUMIF({rg("Levered cash flow")},">0")/-SUMIF({rg("Levered cash flow")},"<0")','0.00"x"'),
  ('LP IRR',f'=IRR({rg("LP cash flow")})',P),('LP equity multiple',f'=SUMIF({rg("LP cash flow")},">0")/-SUMIF({rg("LP cash flow")},"<0")','0.00"x"'),('GP IRR',f'=IRR({rg("GP cash flow")})',P),
  ('Going-in cap rate',f"='Pro Forma'!D{R['Net operating income']}/{I['pp']}",P),('Net profit',f'=SUM({rg("Levered cash flow")})',N),('Total project cost',f'=-SUM({rg("Acquisition")})-SUM({rg("Renovation capex")})',N)]):
    sm.cell(3+i,1,lab);c=sm.cell(3+i,3,f);c.number_format=fmt
wb.save(OUT+'multifamily_usd.xlsx')
# ---------------- D. land subdivision, AUD, monthly, XIRR, two input blocks side by side ----------------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Project Inputs';ws['B1']='WATTLE GROVE ESTATE – RESIDENTIAL SUBDIVISION FEASIBILITY (AUD)'
I=inputs(ws,['REVENUE',('lots','Number of lots',180,'lots',N),('lp','Average lot price',385000,'AUD/lot',N),('esc','Price escalation',0.04,'p.a.',P),('rate','Sales rate',8,'lots/month',N),('pre','Pre-sales at first settlement',25,'lots',N),
  ('com','Sales commission',0.025,'of revenue',P),('mkt','Marketing',0.02,'of revenue',P)],lc=2,vc=3,uc=4)
I.update(inputs(ws,['COSTS & FUNDING',('land','Land purchase',30000000,'AUD',N),('sd','Stamp duty',0.055,'of land',P),('civ','Civil works per lot',165000,'AUD/lot',N),('cm','Civil works period',9,'months','0'),
  ('eq','Equity contribution',12000000,'AUD',N),('ir','Interest rate',0.085,'p.a.',P)],lc=6,vc=7,uc=8))
cf=wb.create_sheet('Cash Flow');M=36;cf['A1']='MONTHLY CASH FLOW (AUD)';cf['A3']='Month';cf['A4']='Date'
dl=['Lots settled','Cumulative lots','Sales revenue','Commission & marketing','Land & duty','Civil works','Project cash flow','Equity draw','Cumulative equity','Funding need','Interest','Loan balance','Surplus to equity','Equity cash flow']
R={l:6+i for i,l in enumerate(dl)}
for l,r in R.items():cf[f'A{r}']=l
for m in range(1,M+1):
    c=L(m+1);p=L(m);cf[f'{c}3']=m;cf[f'{c}4']=datetime.date(2026,1,31) if m==1 else f'=EOMONTH({p}4,1)';cf[f'{c}4'].number_format='mmm-yy';g=lambda l:f'{c}{R[l]}'
    prevcum=f'{p}{R["Cumulative lots"]}' if m>1 else '0'
    cf[g('Lots settled')]=f'=IF({c}$3<={I["cm"]},0,MIN({I["lots"]}-{prevcum},{I["rate"]}+IF({c}$3={I["cm"]}+1,{I["pre"]},0)))'
    cf[g('Cumulative lots')]=f'={prevcum}+{g("Lots settled")}'
    cf[g('Sales revenue')]=f'={g("Lots settled")}*{I["lp"]}*(1+{I["esc"]})^(({c}$3-1)/12)'
    cf[g('Commission & marketing')]=f'=-{g("Sales revenue")}*({I["com"]}+{I["mkt"]})'
    cf[g('Land & duty')]=f'=IF({c}$3=1,-{I["land"]}*(1+{I["sd"]}),0)'
    cf[g('Civil works')]=f'=IF({c}$3<={I["cm"]},-{I["lots"]}*{I["civ"]}/{I["cm"]},0)'
    cf[g('Project cash flow')]=f'=SUM({c}{R["Sales revenue"]}:{c}{R["Civil works"]})'
    prevq=f'{p}{R["Cumulative equity"]}' if m>1 else '0'
    cf[g('Equity draw')]=f'=MIN(MAX(0,-{g("Project cash flow")}),{I["eq"]}-{prevq})'
    cf[g('Cumulative equity')]=f'={prevq}+{g("Equity draw")}'
    cf[g('Funding need')]=f'=-{g("Project cash flow")}-{g("Equity draw")}'
    prevl=f'{p}{R["Loan balance"]}' if m>1 else '0'
    cf[g('Interest')]=f'={prevl}*{I["ir"]}/12'
    cf[g('Loan balance')]=f'=MAX(0,{prevl}+{g("Interest")}+{g("Funding need")})'
    cf[g('Surplus to equity')]=f'=MAX(0,-({prevl}+{g("Interest")}+{g("Funding need")}))'
    cf[g('Equity cash flow')]=f'=-{g("Equity draw")}+{g("Surplus to equity")}'
sm=wb.create_sheet('Feasibility Summary');rg=lambda l:f"'Cash Flow'!$B${R[l]}:$AK${R[l]}";dts="'Cash Flow'!$B$4:$AK$4"
for i,(lab,f,fmt) in enumerate([('Gross realisation',f'=SUM({rg("Sales revenue")})',N),('Total development cost',f'=-SUM({rg("Commission & marketing")})-SUM({rg("Land & duty")})-SUM({rg("Civil works")})+SUM({rg("Interest")})',N),
  ('Net profit',f'=B3-B4',N),('Development margin',f'=B5/B4',P),('Project IRR',f'=XIRR({rg("Project cash flow")},{dts})',P),('Equity IRR',f'=XIRR({rg("Equity cash flow")},{dts})',P),('Peak debt',f'=MAX({rg("Loan balance")})',N),
  ('Equity multiple',f'=SUM({rg("Surplus to equity")})/SUM({rg("Equity draw")})','0.00"x"')]):
    sm.cell(3+i,1,lab);c=sm.cell(3+i,2,f);c.number_format=fmt
wb.save(OUT+'subdivision_aud.xlsx')
# ---------------- E. data centre, USD, named ranges, EV/EBITDA exit ----------------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='NORTHGATE DATA CENTRE CAMPUS – USD'
rows=[('IT_MW','IT capacity',24,'MW','0.0'),('Capex_MW','Development capex per MW',11500000,'$/MW',N),('Land_DC','Land and power connection',38000000,'$',N),('Leaseup_MW','Lease-up',8,'MW per year','0.0'),
 ('Rent_kW','Rent',128,'$/kW/month',N),('Rent_Esc','Rent escalation',0.025,'p.a.',P),('PUE','PUE',1.35,'','0.00'),('Power_Price','Power cost',0.085,'$/kWh','0.000'),('Passthrough','Power pass-through to tenants',0.85,'',P),
 ('Fixed_Opex','Fixed opex per MW',380000,'$/MW/yr',N),('Exit_Mult','Exit EV/EBITDA multiple',17,'x','0.0"x"'),('Exit_Yr','Exit year',8,'year','0'),('Debt_LTC','Loan to cost',0.50,'',P),('Debt_Rate','Debt interest rate',0.0675,'',P)]
r=3
for k,lab,v,u,fmt in rows:
    ws.cell(r,1,lab);c=ws.cell(r,2,v);c.number_format=fmt;ws.cell(r,3,u);wb.defined_names[k]=DefinedName(k,attr_text=f'Inputs!$B${r}');r+=1
cf=wb.create_sheet('Model');cf['A1']='ANNUAL MODEL (USD)';cf['A3']='Year'
ml=['Leased MW','Rental revenue','Power cost (net of pass-through)','Fixed opex','EBITDA','Capex','Exit value','Unlevered cash flow','Debt drawn','Interest','Debt repaid','Levered cash flow']
R={l:5+i for i,l in enumerate(ml)}
for l,rr_ in R.items():cf[f'A{rr_}']=l
for y in range(0,11):
    c=L(y+2);cf[f'{c}3']=y;g=lambda l:f'{c}{R[l]}'
    cf[g('Leased MW')]=f'=IF(OR({c}$3<2,{c}$3>Exit_Yr),0,MIN(IT_MW,({c}$3-1)*Leaseup_MW))'
    cf[g('Rental revenue')]=f'={g("Leased MW")}*1000*Rent_kW*12*(1+Rent_Esc)^MAX(0,{c}$3-2)'
    cf[g('Power cost (net of pass-through)')]=f'=-{g("Leased MW")}*1000*PUE*8760*0.7*Power_Price*(1-Passthrough)'
    cf[g('Fixed opex')]=f'=IF({g("Leased MW")}>0,-IT_MW*Fixed_Opex,0)'
    cf[g('EBITDA')]=f'=SUM({c}{R["Rental revenue"]}:{c}{R["Fixed opex"]})'
    cf[g('Capex')]=f'=IF({c}$3=0,-Land_DC-IT_MW*Capex_MW*0.6,IF({c}$3=1,-IT_MW*Capex_MW*0.4,0))'
    nxt=L(y+3)
    cf[g('Exit value')]=f'=IF({c}$3=Exit_Yr,{g("EBITDA")}*(1+Rent_Esc)*Exit_Mult,0)'
    cf[g('Unlevered cash flow')]=f'={g("EBITDA")}+{g("Capex")}+{g("Exit value")}'
    cf[g('Debt drawn')]=f'=-{g("Capex")}*Debt_LTC'
    cf[g('Interest')]=f'=-SUM($B${R["Debt drawn"]}:{c}{R["Debt drawn"]})*Debt_Rate*({c}$3<=Exit_Yr)*({c}$3>0)'
    cf[g('Debt repaid')]=f'=IF({c}$3=Exit_Yr,-SUM($B${R["Debt drawn"]}:{c}{R["Debt drawn"]}),0)'
    cf[g('Levered cash flow')]=f'={g("Unlevered cash flow")}+{g("Debt drawn")}+{g("Interest")}+{g("Debt repaid")}'
sm=wb.create_sheet('Dashboard');rg=lambda l:f'Model!$B${R[l]}:$L${R[l]}'
for i,(lab,f,fmt) in enumerate([('Project IRR (unlevered)',f'=IRR({rg("Unlevered cash flow")})',P),('Equity IRR (levered)',f'=IRR({rg("Levered cash flow")})',P),('Equity multiple',f'=SUMIF({rg("Levered cash flow")},">0")/-SUMIF({rg("Levered cash flow")},"<0")','0.00"x"'),
  ('Total capex',f'=-SUM({rg("Capex")})',N),('Stabilised EBITDA',f'=MAX({rg("EBITDA")})',N),('Yield on cost',f'=B7/B6',P),('Net profit',f'=SUM({rg("Levered cash flow")})',N)]):
    sm.cell(3+i,1,lab);c=sm.cell(3+i,2,f);c.number_format=fmt
wb.save(OUT+'datacentre_usd.xlsx')
print('ok')
