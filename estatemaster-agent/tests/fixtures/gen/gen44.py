# Round 44 fixture (synthetic): an open-ended property fund reported by NAV, with made-up figures.
#  r44_nav_fund_sar: a portfolio of five income assets valued each year at their own cap rates; fund debt; NAV and NAV per unit;
#  subscriptions and redemptions at NAV; distributions; a Fund Report whose figures are "NAV per unit", "Distribution yield",
#  "Total return (a year)", "Gearing (LTV)" and "Income return" with no IRR anywhere: the reader must not make one up.
import openpyxl,sys
from openpyxl.utils import get_column_letter as L
out=sys.argv[1].rstrip('/')
wb=openpyxl.Workbook();R=wb.active;R.title='Fund Report';A=wb.create_sheet('Portfolio');N=wb.create_sheet('NAV')
Y=5
A['B2']='AL NOOR INCOME FUND · PORTFOLIO (SAR)';A['B4']='Asset';A['C4']='NOI year 1';A['D4']='Cap rate';A['E4']='NOI growth'
assets=[('Riyadh logistics park',42000000,0.0675,0.025),('Jeddah retail centre',36500000,0.0750,0.015),('Khobar offices',21000000,0.0800,0.010),
 ('Riyadh residential compound',28500000,0.0700,0.030),('Dammam warehouses',17500000,0.0725,0.020)]
for i,(n,noi,cap,g) in enumerate(assets):
    r=5+i;A.cell(r,2,n);A.cell(r,3,noi).number_format='#,##0';A.cell(r,4,cap).number_format='0.00%';A.cell(r,5,g).number_format='0.00%'
A['B11']='Fund debt';A['C11']=650000000;A['C11'].number_format='#,##0';A['B12']='Debt cost';A['C12']=0.058;A['C12'].number_format='0.00%'
A['B13']='Fund costs (% of GAV)';A['C13']=0.0085;A['C13'].number_format='0.00%';A['B14']='Payout (% of net income)';A['C14']=0.9;A['C14'].number_format='0%'
A['B15']='Units in issue (start)';A['C15']=120000000;A['C15'].number_format='#,##0';A['B16']='Net subscriptions (SAR a year)';A['C16']=80000000;A['C16'].number_format='#,##0'
N['B2']='NAV BRIDGE (SAR)';N['B4']='Year'
lab={6:'Gross asset value',7:'Net operating income',8:'Debt cost',9:'Fund costs',10:'Net income',11:'Distributions',12:'Net subscriptions',13:'NAV',14:'Units in issue',15:'NAV per unit',16:'Cash'}
for r,l in lab.items():N.cell(r,2,l)
for y in range(Y+1):
    c=L(3+y);p=L(2+y);N[f'{c}4']=y
    N[f'{c}6']='='+'+'.join(f'Portfolio!$C${5+i}*(1+Portfolio!$E${5+i})^{c}$4/Portfolio!$D${5+i}' for i in range(len(assets)))
    N[f'{c}7']='=0' if y==0 else '='+'+'.join(f'Portfolio!$C${5+i}*(1+Portfolio!$E${5+i})^({c}$4-1)' for i in range(len(assets)))
    N[f'{c}8']='=0' if y==0 else '=-Portfolio!$C$11*Portfolio!$C$12';N[f'{c}9']='=0' if y==0 else f'=-{p}6*Portfolio!$C$13'
    N[f'{c}10']=f'=SUM({c}7:{c}9)';N[f'{c}11']=f'=-{c}10*Portfolio!$C$14';N[f'{c}12']='=0' if y==0 else '=Portfolio!$C$16'
    N[f'{c}16']=f'={c}10+{c}11+{c}12' if y==0 else f'={p}16+{c}10+{c}11+{c}12'
    N[f'{c}13']=f'={c}6-Portfolio!$C$11+{c}16'
    N[f'{c}14']='=Portfolio!$C$15' if y==0 else f'={p}14+{c}12/{p}15'
    N[f'{c}15']=f'={c}13/{c}14'
R['B2']='AL NOOR INCOME FUND · FUND REPORT (SAR)'
rep=[(5,'NAV per unit (year 1)','=NAV!D15','0.0000'),(6,'NAV per unit (start)','=NAV!C15','0.0000'),(7,'Distribution yield','=-NAV!D11/NAV!D14/NAV!C15','0.00%'),
 (8,'Income return','=NAV!D10/NAV!C13','0.00%'),(9,'Total return (a year)','=(NAV!D15-NAV!C15-NAV!D11/NAV!D14)/NAV!C15','0.00%'),(10,'Gearing (LTV)','=Portfolio!C11/NAV!D6','0.00%'),
 (11,'Gross asset value','=NAV!D6','#,##0'),(12,'NAV','=NAV!D13','#,##0')]
for r,l,f,fm in rep:R.cell(r,2,l);R.cell(r,3,f).number_format=fm
wb.save(f'{out}/r44_nav_fund_sar.xlsx')
