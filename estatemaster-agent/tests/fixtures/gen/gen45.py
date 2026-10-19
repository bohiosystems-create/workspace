# Round 45 fixture (synthetic): a data-centre development laid out the way data-centre underwriting templates often are, with made-up figures.
#  r45_datacenter_usd: one Underwriting sheet with the inputs on the left and a monthly timeline from column S; a KEY METRICS matrix
#  (Unlevered / Levered rows by EQUITY MULTIPLE · IRR · NET PROFIT columns, no "levered IRR" label anywhere); a uses table phased by
#  start/end month and method (Straight-Line or S-Curve), power and cooling costed per kW of IT load; tenants leased by kW with a rent
#  per kW per month, a ramp-up and annual growth; operating costs recovered pro rata to each tenant's share of the IT load; a
#  construction loan drawn pari passu and refinanced by a permanent loan at stabilisation (LTV-sized on stabilised NOI); a sale at a
#  typed month on forward NOI; yield on cost untrended / trended / at sale and a development spread in bps; returns by XIRR over
#  OFFSET ranges that end at the sale month; defined names (IT_Load, Sale_Month, Stabilization_1st_Month, Project_Cost_Total, ...).
#  An Annual CF sheet sums the timeline by year.
import openpyxl,sys,datetime
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
out=sys.argv[1].rstrip('/')
wb=openpyxl.Workbook();U=wb.active;U.title='Underwriting';A=wb.create_sheet('Annual CF')
def name(n,ref):wb.defined_names[n]=DefinedName(n,attr_text=ref)
N=84;C0=19   # months 0..N on the timeline, month 0 in column S
U['B2']='DATA CENTER DEVELOPMENT | UNDERWRITING';U['B3']='All amounts in USD'
# ---- key metrics ----
U['H3']='KEY METRICS';U['I4']='EQUITY MULTIPLE';U['J4']='IRR';U['K4']='NET PROFIT';U['H5']='Unlevered';U['H6']='Levered'
E=L(C0+N)
for r,row in ((5,104),(6,112)):
    U[f'J{r}']=f'=XIRR(OFFSET($S${row},0,0,1,Sale_Month+1),OFFSET($S$51,0,0,1,Sale_Month+1))'
    U[f'I{r}']=f'=SUMIF($S${row}:${E}${row},">0")/-SUMIF($S${row}:${E}${row},"<0")';U[f'K{r}']=f'=SUM($S${row}:${E}${row})'
    U[f'I{r}'].number_format='0.00"x"';U[f'J{r}'].number_format='0.00%';U[f'K{r}'].number_format='#,##0'
U['H8']='YIELD ON COST';U['H9']='Untrended';U['J9']='=Stabilized_NOI_Untrended/Project_Cost_Total';U['H10']='Trended';U['J10']='=Stabilized_NOI/Project_Cost_Total'
U['H11']='At Sale';U['J11']='=Exit_NOI/Project_Cost_Total';U['H12']='Development Spread (bps)';U['J12']='=(J10-Exit_Cap)*10000'
for a in ('J9','J10','J11'):U[a].number_format='0.00%'
U['J12'].number_format='0'
# ---- timing and site ----
U['B6']='TIMING';U['B7']='Analysis Start';U['E7']=datetime.datetime(2027,1,1);U['E7'].number_format='mmm-yy'
U['B8']='Sale Month';U['E8']=72;U['B9']='Stabilization (1st month)';U['E9']='=MAX(F29+G29,F30+G30,F31+G31,F32+G32)'
U['B10']='Total IT Load (kW)';U['E10']='=SUM(C29:C32)';U['E10'].number_format='#,##0'
U['B11']='Land Area (SF)';U['E11']=420000;U['B12']='Gross Building Area (SF)';U['E12']=185000
for a in ('E11','E12'):U[a].number_format='#,##0'
# ---- uses ----
U['B14']='USES';U['C14']='Per kW';U['D14']='Per GSF';U['E14']='Total';U['F14']='Start Month';U['G14']='End Month';U['H14']='Method'
uses=((15,'Land Acquisition',None,None,'=E11*42',0,0,'Straight-Line'),
      (16,'Shell & Core',None,410,'=D16*E12',1,16,'S-Curve'),
      (17,'Power & Cooling (MEP)',9800,None,'=C17*IT_Load',4,20,'S-Curve'),
      (18,'Fit-Out & IT Infrastructure',1400,None,'=C18*IT_Load',12,22,'Straight-Line'),
      (19,'Soft Costs',None,None,'=0.11*SUM(E16:E18)',1,22,'Straight-Line'),
      (20,'Developer Fee',None,None,'=0.03*SUM(E16:E18)',1,22,'Straight-Line'),
      (21,'Contingency',None,None,'=0.05*SUM(E16:E18)',4,22,'S-Curve'))
for r,l,pk,pg,tot,s,e,m in uses:
    U[f'B{r}']=l;U[f'E{r}']=tot;U[f'F{r}']=s;U[f'G{r}']=e;U[f'H{r}']=m;U[f'E{r}'].number_format='#,##0'
    if pk is not None:U[f'C{r}']=pk
    if pg is not None:U[f'D{r}']=pg
U['B22']='Total Project Cost';U['E22']='=SUM(E15:E21)';U['E22'].number_format='#,##0'
# ---- financing ----
U['B24']='FINANCING';U['B25']='Construction Loan to Cost';U['E25']=0.6;U['F25']='Rate';U['G25']=0.075
U['B26']='Permanent Loan LTV';U['E26']=0.5;U['F26']='Rate';U['G26']=0.0625
for a in ('E25','G25','E26','G26'):U[a].number_format='0.00%'
# ---- tenants ----
U['B28']='TENANT';U['C28']='IT Load (kW)';U['D28']='Rent / kW / Month';U['E28']='Growth (% /Yr.)';U['F28']='Lease Start Month';U['G28']='Ramp-Up (months)'
for r,(t,kw,rent,g,s,ramp) in enumerate((('Hyperscaler A',12000,128,0.025,20,12),('Hyperscaler B',8000,134,0.025,24,9),('Enterprise Colocation',2500,195,0.03,22,18),('Cloud On-Ramp',1500,170,0.03,26,6)),29):
    U[f'B{r}']=t;U[f'C{r}']=kw;U[f'D{r}']=rent;U[f'E{r}']=g;U[f'F{r}']=s;U[f'G{r}']=ramp;U[f'C{r}'].number_format='#,##0';U[f'E{r}'].number_format='0.00%'
U['B33']='Total';U['C33']='=SUM(C29:C32)';U['C33'].number_format='#,##0'
# ---- operating costs, year 1 dollars ----
U['B35']='OPERATING EXPENSES';U['C35']='Year 1';U['D35']='Basis'
for r,l,v,b in ((36,'Utilities',1150,'per leased kW / yr'),(37,'Maintenance',2600000,'fixed / yr'),(38,'Staffing',3300000,'fixed / yr'),(39,'Insurance',950000,'fixed / yr'),
                (40,'G&A',650000,'fixed / yr'),(41,'Property Taxes',1900000,'fixed / yr'),(42,'Management Fee',0.02,'% of rent')):
    U[f'B{r}']=l;U[f'C{r}']=v;U[f'D{r}']=b;U[f'C{r}'].number_format='0.00%' if r==42 else '#,##0'
U['B43']='Expense Growth (% /Yr.)';U['C43']=0.03;U['B44']='Expense Recovery %';U['C44']=1;U['B45']='Opex Start Month';U['C45']='=MIN(F29:F32)'
for a in ('C43','C44'):U[a].number_format='0.00%'
# ---- reversion ----
U['B47']='REVERSION';U['B48']='Exit Cap Rate';U['C48']=0.0625;U['B49']='Selling Costs';U['C49']=0.02;U['B50']='NOI Growth to Forward Year';U['C50']=0.025
for a in ('C48','C49','C50'):U[a].number_format='0.00%'
for n,ref in (('IT_Load','Underwriting!$E$10'),('Sale_Month','Underwriting!$E$8'),('Stabilization_1st_Month','Underwriting!$E$9'),('Project_Cost_Total','Underwriting!$E$22'),
              ('Timing_Start','Underwriting!$E$7'),('Exit_Cap','Underwriting!$C$48'),('Perm_Debt','Underwriting!$E$54'),('Exit_NOI','Underwriting!$E$53'),
              ('Stabilized_NOI','Underwriting!$E$55'),('Stabilized_NOI_Untrended','Underwriting!$E$56')):name(n,ref)
U['D53']='Exit NOI (forward)';U['E53']=f'=SUMPRODUCT(($S$50:${E}$50>Sale_Month-12)*($S$50:${E}$50<=Sale_Month)*$S$83:${E}$83)*(1+C50)'
U['D54']='Permanent Loan';U['E54']=f'=E26*E55/C48';U['D55']='Stabilized NOI (annualised)';U['E55']=f'=12*SUMPRODUCT(($S$50:${E}$50=Stabilization_1st_Month)*$S$83:${E}$83)'
U['D56']='Stabilized NOI (untrended)';U['E56']='=SUMPRODUCT(C29:C32,D29:D32)*12*(1-C42)-SUM(C37:C41)-C36*IT_Load+C44*(SUM(C37:C41)+C36*IT_Load)'
for a in ('E53','E54','E55','E56'):U[a].number_format='#,##0'
# ---- monthly timeline ----
lab={50:'Month',51:'Month Ending',52:'Year',54:'Land Acquisition',55:'Shell & Core',56:'Power & Cooling (MEP)',57:'Fit-Out & IT Infrastructure',58:'Soft Costs',59:'Developer Fee',
     60:'Contingency',61:'Total Development Cost',63:'Leased IT Load (kW)',64:'Rent - Hyperscaler A',65:'Rent - Hyperscaler B',66:'Rent - Enterprise Colocation',67:'Rent - Cloud On-Ramp',
     68:'Total Rental Revenue',70:'Recovery - Hyperscaler A',71:'Recovery - Hyperscaler B',72:'Recovery - Enterprise Colocation',73:'Recovery - Cloud On-Ramp',74:'Total Recovery Income',
     75:'Effective Gross Revenue',76:'Utilities',77:'Maintenance',78:'Staffing',79:'Insurance',80:'G&A',81:'Property Taxes',82:'Management Fee',83:'Net Operating Income',84:'Total Operating Expenses',
     86:'Gross Sale Price',87:'Selling Costs',88:'Net Sale Proceeds',90:'Construction Loan Draw',91:'Construction Loan Interest',92:'Construction Loan Repayment',93:'Construction Loan Balance',
     95:'Permanent Loan Proceeds',96:'Permanent Loan Interest',97:'Permanent Loan Repayment',104:'Unlevered Cash Flow',112:'Levered Cash Flow'}
for r,l in lab.items():U.cell(r,2,l)
for k in range(N+1):
    c=L(C0+k);p=L(C0+k-1);live=f'({c}$50<=Sale_Month)'
    U[f'{c}50']=k;U[f'{c}51']=f'=EOMONTH(Timing_Start,{c}50-1)';U[f'{c}51'].number_format='mmm-yy';U[f'{c}52']=f'=ROUNDUP({c}50/12,0)'
    for i,r in enumerate(range(54,61)):
        u=15+i;n=f'($G${u}-$F${u}+1)';t=f'({c}$50-$F${u})'
        U[f'{c}{r}']=(f'=IF(AND({c}$50>=$F${u},{c}$50<=$G${u}),$E${u}*IF($H${u}="S-Curve",(COS(PI()*{t}/{n})-COS(PI()*({t}+1)/{n}))/2,1/{n}),0)')
    U[f'{c}61']=f'=SUM({c}54:{c}60)'
    for i in range(4):
        t=29+i;ramp=f'MAX(0,MIN(1,({c}$50-$F${t}+1)/$G${t}))'
        U[f'{c}{64+i}']=f'={live}*$C${t}*{ramp}*$D${t}*(1+$E${t})^MAX(0,ROUNDUP(({c}$50-$F${t}+1)/12,0)-1)'
        U[f'{c}{70+i}']=f'={live}*$C${t}/IT_Load*{ramp}*$C$44*(SUM({c}77:{c}81)+$C$36/12*IT_Load*(1+$C$43)^({c}$52-1))'
    U[f'{c}63']=f'='+'+'.join(f'$C${29+i}*MAX(0,MIN(1,({c}$50-$F${29+i}+1)/$G${29+i}))' for i in range(4))
    U[f'{c}68']=f'=SUM({c}64:{c}67)';U[f'{c}74']=f'=SUM({c}70:{c}73)';U[f'{c}75']=f'={c}68+{c}74'
    U[f'{c}76']=f'={live}*$C$36/12*{c}63*(1+$C$43)^({c}$52-1)'
    for r,src in ((77,37),(78,38),(79,39),(80,40),(81,41)):U[f'{c}{r}']=f'={live}*({c}$50>=$C$45)*$C${src}/12*(1+$C$43)^({c}$52-1)'
    U[f'{c}82']=f'=$C$42*{c}68';U[f'{c}84']=f'=SUM({c}76:{c}82)';U[f'{c}83']=f'={c}75-{c}84'
    U[f'{c}86']=f'=({c}$50=Sale_Month)*Exit_NOI/Exit_Cap';U[f'{c}87']=f'=-{c}86*$C$49';U[f'{c}88']=f'={c}86+{c}87'
    U[f'{c}90']=f'={c}61*$E$25'
    U[f'{c}91']='=0' if k==0 else f'=IF({c}$50<=Stabilization_1st_Month,{p}93*$G$25/12,0)'
    U[f'{c}92']=f'=IF({c}$50=Stabilization_1st_Month,{"0" if k==0 else p+"93"}+{c}90,0)'
    U[f'{c}93']=f'={"0" if k==0 else p+"93"}+{c}90-{c}92'
    U[f'{c}95']=f'=({c}$50=Stabilization_1st_Month)*Perm_Debt'
    U[f'{c}96']=f'=AND({c}$50>Stabilization_1st_Month,{live})*Perm_Debt*$G$26/12'
    U[f'{c}97']=f'=({c}$50=Sale_Month)*Perm_Debt'
    U[f'{c}104']=f'={live}*(-{c}61+{c}83+{c}88)'
    U[f'{c}112']=f'={live}*({c}104+{c}90-{c}91-{c}92+{c}95-{c}96-{c}97)'
    for r in list(range(54,62))+list(range(64,89))+[90,91,92,93,95,96,97,104,112]:U[f'{c}{r}'].number_format='#,##0;(#,##0)'
# ---- annual cash flow ----
A['B2']='DATA CENTER DEVELOPMENT | ANNUAL CASH FLOW';A['B4']='Year'
for r,(l,src) in enumerate((('Total Development Cost',61),('Total Rental Revenue',68),('Total Recovery Income',74),('Total Operating Expenses',84),('Net Operating Income',83),
                            ('Net Sale Proceeds',88),('Unlevered Cash Flow',104),('Levered Cash Flow',112)),5):
    A[f'B{r}']=l
    for y in range(8):
        c=L(3+y);A[f'{c}4']=y;A[f'{c}{r}']=f"=SUMIF(Underwriting!$S$52:${E}$52,{c}$4,Underwriting!$S${src}:${E}${src})";A[f'{c}{r}'].number_format='#,##0;(#,##0)'
wb.save(f'{out}/r45_datacenter_usd.xlsx')
