# Round 39 fixture (synthetic): a closed-ended fund laid out the way fund administrators' models often are, with made-up figures.
#  r39_fund_waterfall_eur: three assets bought, held and sold on an annual timeline; capital called as needed against commitments;
#  a management fee on commitments, then on invested capital; a European waterfall (return of capital, an 8% preferred return,
#  a full GP catch-up, then 80/20); a Fund Summary whose returns read just "Gross IRR", "Net IRR", "TVPI" and "DPI", the gross
#  figure being the fund's before fees and carry and the net figure the LPs' after both, plus the GP's carried interest.
import openpyxl,sys
from openpyxl.utils import get_column_letter as L
out=sys.argv[1].rstrip('/')
wb=openpyxl.Workbook();S=wb.active;S.title='Fund Summary';A=wb.create_sheet('Assumptions');C=wb.create_sheet('Fund Cash Flow');W=wb.create_sheet('Waterfall')
Y=10   # years 0..Y
# ---- Assumptions ----
A['B2']='AURORA REAL ESTATE FUND I · ASSUMPTIONS (EUR)'
rows=[(4,'Total commitments',150000000,'#,##0'),(5,'GP commitment (share of total)',0.02,'0.00%'),(6,'Management fee (investment period, on commitments)',0.015,'0.00%'),
 (7,'Management fee (after, on invested capital)',0.0125,'0.00%'),(8,'Investment period (years)',4,'0'),(9,'Fund expenses (a year)',350000,'#,##0'),
 (10,'Preferred return',0.08,'0.00%'),(11,'GP catch-up',1.0,'0%'),(12,'Carried interest',0.2,'0.00%')]
for r,l,v,f in rows:A.cell(r,2,l);A.cell(r,3,v).number_format=f
A['B14']='Assets';A['C14']='Equity in';A['D14']='Year in';A['E14']='Net income yield';A['F14']='Exit multiple of cost';A['G14']='Year out'
for i,(n,e,yi,yl,m,yo) in enumerate([('Rhine Logistics',42000000,1,0.055,1.55,6),('Porto Residences',38000000,2,0.045,1.70,8),('Milan Offices',51000000,3,0.05,1.35,9)]):
    r=15+i;A.cell(r,2,n);A.cell(r,3,e).number_format='#,##0';A.cell(r,4,yi);A.cell(r,5,yl).number_format='0.00%';A.cell(r,6,m).number_format='0.00"x"';A.cell(r,7,yo)
# ---- Fund Cash Flow: year 0 in column D ----
C['B2']='FUND CASH FLOW (EUR)';C['B4']='Year'
lab={6:'Acquisitions',7:'Net income',8:'Sale proceeds',9:'Gross asset cash flow',11:'Management fee',12:'Fund expenses',13:'Fund cash flow before carry',15:'Capital called',16:'Distributable cash'}
for r,l in lab.items():C.cell(r,2,l)
for y in range(Y+1):
    c=L(4+y);C[f'{c}4']=y;yy=f'{c}$4'
    acq='+'.join(f'IF({yy}=Assumptions!$D${r},-Assumptions!$C${r},0)' for r in (15,16,17))
    inc='+'.join(f'IF(AND({yy}>Assumptions!$D${r},{yy}<=Assumptions!$G${r}),Assumptions!$C${r}*Assumptions!$E${r},0)' for r in (15,16,17))
    sal='+'.join(f'IF({yy}=Assumptions!$G${r},Assumptions!$C${r}*Assumptions!$F${r},0)' for r in (15,16,17))
    C[f'{c}6']=f'={acq}';C[f'{c}7']=f'={inc}';C[f'{c}8']=f'={sal}';C[f'{c}9']=f'=SUM({c}6:{c}8)'
    inv='+'.join(f'IF(AND({yy}>=Assumptions!$D${r},{yy}<Assumptions!$G${r}),Assumptions!$C${r},0)' for r in (15,16,17))
    C[f'{c}11']='=0' if y==0 else f'=-IF({yy}<=Assumptions!$C$8,Assumptions!$C$4*Assumptions!$C$6,({inv})*Assumptions!$C$7)'
    C[f'{c}12']='=0' if y==0 else f'=-Assumptions!$C$9'
    C[f'{c}13']=f'={c}9+{c}11+{c}12'
    C[f'{c}15']=f'=-MIN(0,{c}13)';C[f'{c}16']=f'=MAX(0,{c}13)'
E=L(4+Y)
# ---- Waterfall (European, whole fund) ----
W['B2']='DISTRIBUTION WATERFALL (EUR)';W['B4']='Year'
wl={6:'Contributions (all partners)',7:'Distributable cash',9:'Unreturned capital, opening',10:'1. Return of capital',11:'Unreturned capital, closing',
 13:'Preferred return accrued, opening',14:'Preferred return for the year',15:'2. Preferred return paid',16:'Preferred return accrued, closing',
 18:'3. GP catch-up',19:'Cumulative LP profit distributed',20:'Cumulative GP catch-up',21:'4. Split: to LPs',22:'4. Split: to GP',
 24:'Total to all partners (pro rata to commitments)',25:'Carried interest to GP',27:'LP net cash flow',28:'GP carry cash flow'}
for r,l in wl.items():W.cell(r,2,l)
for y in range(Y+1):
    c=L(4+y);p=L(3+y);W[f'{c}4']=y
    W[f'{c}6']=f"='Fund Cash Flow'!{c}15";W[f'{c}7']=f"='Fund Cash Flow'!{c}16"
    W[f'{c}9']='=0' if y==0 else f'={p}11';W[f'{c}10']=f'=MIN({c}7,{c}9+{c}6)';W[f'{c}11']=f'={c}9+{c}6-{c}10'
    W[f'{c}13']='=0' if y==0 else f'={p}16';W[f'{c}14']=f'=({c}9+{c}13)*Assumptions!$C$10';W[f'{c}15']=f'=MIN({c}7-{c}10,{c}13+{c}14)';W[f'{c}16']=f'={c}13+{c}14-{c}15'
    cumLP=f'SUM($D$15:{c}15)' ;prevcu='0' if y==0 else f'SUM($D$18:{p}18)'
    W[f'{c}18']=f'=MIN({c}7-{c}10-{c}15,MAX(0,Assumptions!$C$12/(1-Assumptions!$C$12)*{cumLP}-{prevcu})*Assumptions!$C$11)'
    W[f'{c}19']=f'=SUM($D$15:{c}15)+SUM($D$21:{c}21)';W[f'{c}20']=f'=SUM($D$18:{c}18)'
    W[f'{c}21']=f'=({c}7-{c}10-{c}15-{c}18)*(1-Assumptions!$C$12)';W[f'{c}22']=f'=({c}7-{c}10-{c}15-{c}18)*Assumptions!$C$12'
    W[f'{c}24']=f'={c}10+{c}15+{c}21';W[f'{c}25']=f'={c}18+{c}22'
    W[f'{c}27']=f'=(-{c}6+{c}24)*(1-Assumptions!$C$5)';W[f'{c}28']=f'={c}25'
# ---- Fund Summary ----
S['B2']='AURORA REAL ESTATE FUND I · FUND SUMMARY';S['B3']='All amounts in EUR'
S['B5']='Fund returns';S['B6']='Gross IRR';S['C6']=f"=IRR('Fund Cash Flow'!D9:{E}9)";S['B7']='Net IRR';S['C7']=f'=IRR(Waterfall!D27:{E}27)'
S['B8']='TVPI';S['C8']=f'=SUM(Waterfall!D24:{E}24)*(1-Assumptions!C5)/SUM(Waterfall!D6:{E}6)/(1-Assumptions!C5)'
S['B9']='DPI';S['C9']='=C8'
S['B11']='Fund economics';S['B12']='Total capital called';S['C12']=f"=SUM('Fund Cash Flow'!D15:{E}15)";S['B13']='Total distributions';S['C13']=f'=SUM(Waterfall!D24:{E}24)'
S['B14']='Gross profit (assets)';S['C14']=f"=SUM('Fund Cash Flow'!D9:{E}9)";S['B15']='Carried interest (GP)';S['C15']=f'=SUM(Waterfall!D25:{E}25)';S['B16']='Management fees';S['C16']=f"=-SUM('Fund Cash Flow'!D11:{E}11)"
for a in ('C6','C7'):S[a].number_format='0.00%'
for a in ('C8','C9'):S[a].number_format='0.00"x"'
for a in ('C12','C13','C14','C15','C16'):S[a].number_format='#,##0'
wb.save(f'{out}/r39_fund_waterfall_eur.xlsx')
