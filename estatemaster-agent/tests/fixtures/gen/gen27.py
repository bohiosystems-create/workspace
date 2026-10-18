# round 27: a CIO's desk on new ground again
#  park     - USD retail park: tenants in rows, rent steps chosen with IFS and SWITCH, MAXIFS/MINIFS over the rent roll, XNPV on
#             irregular payment dates, NETWORKDAYS for the fit-out, DATEDIF for the lease term left
#  landbank - SAR land bank: plots sold over years, the sales window summed with INDIRECT("C"&a&":C"&b) (an address built as text),
#             a hand-built sensitivity grid (land cost × sale price) laid out as a table of formulas
#  zurich   - CHF office in kCHF: a negative SARON floored at zero for the margin, apostrophe thousands (#'##0), amounts in thousands
#  pune     - INR build-to-sell in ₹ crore: Indian digit grouping (##,##,##0), amounts in crore (1 crore = 10 million)
#  jv       - EUR joint venture: a dashboard with each value to the LEFT of its label, merged headings, a catch-up waterfall
import openpyxl,os,datetime
from openpyxl.utils import get_column_letter as L
OUT='/tmp/claude-0/-home-user-workspace/c7e8f42c-9ed0-5ba7-9852-d8a131185935/scratchpad/m27/'
os.makedirs(OUT,exist_ok=True)
P='0.00%';N='#,##0';X='0.00"x"';D='dd-mmm-yyyy'
def ref(ws,col,r):return (f"'{ws.title}'" if ' ' in ws.title else ws.title)+f"!${L(col)}${r}"
def inputs(ws,rows,col=2,start=3):
    I={};r=start
    for k,lab,v,f in rows:
        ws.cell(r,1,lab);c=ws.cell(r,col,v);c.number_format=f;I[k]=ref(ws,col,r);r+=1
    return I
def table(ws,rows,start=3,col=2):
    for i,(lab,f,fm) in enumerate(rows):ws.cell(start+i,1,lab);c=ws.cell(start+i,col,f);c.number_format=fm
def timeline(ws,labels,T,fill,first_row=3):
    R={l:first_row+i for i,l in enumerate(labels)}
    for l,r in R.items():ws.cell(r,1,l)
    for t in range(T+1):
        c=L(t+2);p=L(t+1)
        for l,f in fill(t,c,p,R).items():ws[f'{c}{R[l]}']=f if not isinstance(f,str) or f.startswith('=') else '='+f
    return R

# ================= 27A retail park =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='CEDAR CROSSING RETAIL PARK (USD)'
I=inputs(ws,[('val','Valuation date',datetime.date(2026,12,31),D),('price','Purchase price (USD)',38000000,N),('disc','Discount rate',0.075,P),('xy','Exit cap rate',0.0675,P),('grow','Market rent growth',0.025,P),
 ('fit','Fit-out start',datetime.date(2027,2,1),D),('fitend','Fit-out end',datetime.date(2027,5,28),D),('opex','Non-recoverable costs (USD a year)',650000,N)])
tn=wb.create_sheet('Rent Roll');tn.append(['Tenant','Type','Area (sq ft)','Rent (USD/sq ft)','Lease end','Grade'])
for row_ in [('Anchor grocer','Anchor',62000,14.5,datetime.date(2036,1,31),'A'),('Home store','Big box',38000,16.0,datetime.date(2031,6,30),'B'),('Pharmacy','Shop',12000,28.0,datetime.date(2029,12,31),'A'),
             ('Pet store','Shop',9000,24.5,datetime.date(2030,9,30),'C'),('Restaurant pad','Pad',6500,42.0,datetime.date(2033,3,31),'B'),('Gym','Big box',24000,15.5,datetime.date(2028,8,31),'C')]:tn.append(list(row_))
for r in range(2,8):tn.cell(r,5).number_format=D
tn['H1']='Years left';tn['I1']='Credit haircut';tn['J1']='Step-up'
for r in range(2,8):
    tn[f'H{r}']=f'=DATEDIF({I["val"]},E{r},"y")'
    tn[f'I{r}']=f'=_xlfn.SWITCH(F{r},"A",0,"B",0.02,"C",0.05,0.1)'
    tn[f'J{r}']=f'=_xlfn.IFS(B{r}="Anchor",0.01,B{r}="Big box",0.015,B{r}="Pad",0.025,TRUE,0.02)'
cf=wb.create_sheet('Cash Flow');T=10
labs=['Year','Date','Rent','Costs','Exit','Cash flow']
def pf(t,c,p,R):
    y=f"{c}$3";first=t==0
    rent="+".join(f"IF({c}{R['Date']}<='Rent Roll'!$E${r},'Rent Roll'!$C${r}*'Rent Roll'!$D${r}*(1+'Rent Roll'!$J${r})^({y}-1)*(1-'Rent Roll'!$I${r}),'Rent Roll'!$C${r}*'Rent Roll'!$D${r}*(1+{I['grow']})^{y}*0.9)" for r in range(2,8))
    d={'Year':t,'Date':f"EDATE({I['val']},12*{y}+IF({y}>0,MOD({y},3),0))",'Rent':'0' if first else rent,'Costs':'0' if first else f"{I['opex']}*(1.03)^({y}-1)",
       'Exit':f"IF({y}={T},{c}{R['Rent']}*1.025/{I['xy']},0)"}
    d['Cash flow']=f"-{I['price']}-NETWORKDAYS({I['fit']},{I['fitend']})*4500" if first else f"{c}{R['Rent']}-{c}{R['Costs']}+{c}{R['Exit']}"
    return d
R=timeline(cf,labs,T,pf);E=L(T+2);row=lambda n:f"'Cash Flow'!$B${R[n]}:${E}${R[n]}"
for t in range(T+1):cf[f'{L(t+2)}{R["Date"]}'].number_format=D
table(wb.create_sheet('Returns'),[('Unlevered IRR',f"=XIRR({row('Cash flow')},{row('Date')})",P),('NPV at the discount rate (USD)',f"=XNPV({I['disc']},{row('Cash flow')},{row('Date')})",N),
 ('Highest rent (USD/sq ft)',"=_xlfn.MAXIFS('Rent Roll'!D2:D7,'Rent Roll'!B2:B7,\"<>Anchor\")",'0.00'),('Lowest big-box rent (USD/sq ft)',"=_xlfn.MINIFS('Rent Roll'!D2:D7,'Rent Roll'!B2:B7,\"Big box\")",'0.00'),
 ('Fit-out working days',f"=NETWORKDAYS({I['fit']},{I['fitend']})",'0'),('WALE (years)',"=SUMPRODUCT('Rent Roll'!C2:C7,'Rent Roll'!H2:H7)/SUM('Rent Roll'!C2:C7)",'0.0'),('Net profit (USD)',f"=SUM({row('Cash flow')})",N)])
wb.save(OUT+'park_usd.xlsx')

# ================= 27B land bank: INDIRECT address strings, a sensitivity grid =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='WADI HANIFA LAND BANK (SAR)'
I=inputs(ws,[('land','Land cost (SAR)',240000000,N),('plots','Plots',400,N),('price','Sale price per plot (SAR)',1150000,N),('infra','Infrastructure (SAR)',95000000,N),('first','First year of sales',2,'0'),
 ('last','Last year of sales',6,'0'),('disc','Discount rate',0.1,P)])
cf=wb.create_sheet('Cash Flow');cf['A1']='Year';cf['B1']='Plots sold';cf['C1']='Sales';cf['D1']='Costs';cf['E1']='Cash flow'
for y in range(0,9):
    r=y+2;cf[f'A{r}']=y
    cf[f'B{r}']=f"=IF(AND(A{r}>={I['first']},A{r}<={I['last']}),{I['plots']}/({I['last']}-{I['first']}+1),0)"
    cf[f'C{r}']=f"=B{r}*{I['price']}*1.03^A{r}"
    cf[f'D{r}']=f"=IF(A{r}=0,{I['land']},0)+IF(AND(A{r}>=1,A{r}<=3),{I['infra']}/3,0)"
    cf[f'E{r}']=f"=C{r}-D{r}"
rs=wb.create_sheet('Returns')
table(rs,[('IRR',"=IRR('Cash Flow'!E2:E10)",P),('NPV at the discount rate (SAR)',f"=NPV({I['disc']},'Cash Flow'!E3:E10)+'Cash Flow'!E2",N),
 ('Sales in the selling window (SAR)',f"=SUM(INDIRECT(\"'Cash Flow'!C\"&({I['first']}+2)&\":C\"&({I['last']}+2)))",N),('Net profit (SAR)',"=SUM('Cash Flow'!E2:E10)",N)])
sg=wb.create_sheet('Sensitivity');sg['A1']='Net profit (SAR m): land cost (rows) × sale price per plot (columns)'
prices=[1000000,1150000,1300000];lands=[200000000,240000000,280000000]
for j,pr in enumerate(prices):sg.cell(2,j+2,pr).number_format=N
for i,ld in enumerate(lands):
    sg.cell(i+3,1,ld).number_format=N
    for j in range(3):
        c=L(j+2);sg.cell(i+3,j+2,f"=(SUMPRODUCT('Cash Flow'!$B$2:$B$10,1.03^'Cash Flow'!$A$2:$A$10)*{c}$2-$A{i+3}-{I['infra']})/1000000").number_format='0.0'
wb.save(OUT+'landbank_sar.xlsx')

# ================= 27C Zurich office in kCHF =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='SEEFELD BÜROHAUS – OFFICE, ZURICH (kCHF)'
APO="#'##0"
I=inputs(ws,[('price','Purchase price (kCHF)',68500,APO),('rent','Rent a year (kCHF)',2950,APO),('idx','CPI indexation',0.009,P),('saron','SARON',-0.0025,P),('margin','Loan margin',0.0095,P),
 ('ltv','LTV',0.55,P),('xy','Exit yield',0.0385,P),('costs','Operating costs (kCHF a year)',410,APO)])
cf=wb.create_sheet('Cash Flow');T=10
labs=['Year','Rent','Costs','Interest','Loan','Exit','Cash flow']
def zf(t,c,p,R):
    y=f"{c}$3";first=t==0
    d={'Year':t,'Rent':'0' if first else f"{I['rent']}*(1+{I['idx']})^({y}-1)",'Costs':'0' if first else f"{I['costs']}*(1.01)^({y}-1)",
       'Interest':'0' if first else f"{I['price']}*{I['ltv']}*(MAX(0,{I['saron']})+{I['margin']})",'Loan':f"IF({y}=0,{I['price']}*{I['ltv']},IF({y}={T},-{I['price']}*{I['ltv']},0))",
       'Exit':f"IF({y}={T},{c}{R['Rent']}*(1+{I['idx']})/{I['xy']},0)"}
    d['Cash flow']=f"-{I['price']}+{c}{R['Loan']}" if first else f"{c}{R['Rent']}-{c}{R['Costs']}-{c}{R['Interest']}+{c}{R['Loan']}+{c}{R['Exit']}"
    return d
R=timeline(cf,labs,T,zf);E=L(T+2);row=lambda n:f"'Cash Flow'!$B${R[n]}:${E}${R[n]}"
for n in labs[1:]:
    for t in range(T+1):cf[f'{L(t+2)}{R[n]}'].number_format=APO
table(wb.create_sheet('Returns'),[('Levered IRR',f"=IRR({row('Cash flow')})",P),('All-in loan rate',f"=MAX(0,{I['saron']})+{I['margin']}",P),('Exit value (kCHF)',f"=SUM({row('Exit')})",APO),
 ('Net profit (kCHF)',f"=SUM({row('Cash flow')})",APO),('Cash-on-cash, year 1',f"=('Cash Flow'!C{R['Rent']}-'Cash Flow'!C{R['Costs']}-'Cash Flow'!C{R['Interest']})/-'Cash Flow'!B{R['Cash flow']}",P)])
wb.save(OUT+'zurich_chf.xlsx')

# ================= 27D Pune build-to-sell in ₹ crore =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='RIVERSIDE GREENS, PUNE – BUILD TO SELL (₹ crore)'
IN='[>=10000000]##\\,##\\,##\\,##0;[>=100000]##\\,##\\,##0;##,##0'
I=inputs(ws,[('units','Apartments',820,N),('area','Saleable area per apartment (sq ft)',1150,N),('rate','Sale rate (₹ per sq ft)',8200,IN),('cost','Construction cost (₹ crore)',412,'0.00'),('land','Land (₹ crore)',165,'0.00'),
 ('gst','GST on sales',0.05,P),('sell','Sales period (quarters)',10,'0'),('disc','Discount rate',0.14,P)])
cf=wb.create_sheet('Cash Flow');T=14
labs=['Quarter','Sales (₹ crore)','GST (₹ crore)','Construction (₹ crore)','Land (₹ crore)','Cash flow (₹ crore)']
def nf(t,c,p,R):
    q=f"{c}$3"
    d={'Quarter':t,'Sales (₹ crore)':f"IF(AND({q}>=2,{q}<2+{I['sell']}),{I['units']}*{I['area']}*{I['rate']}*1.01^{q}/{I['sell']}/10000000,0)",
       'GST (₹ crore)':f"{c}{R['Sales (₹ crore)']}*{I['gst']}/(1+{I['gst']})",'Construction (₹ crore)':f"IF(AND({q}>=1,{q}<=10),{I['cost']}/10,0)",'Land (₹ crore)':f"IF({q}=0,{I['land']},0)"}
    d['Cash flow (₹ crore)']=f"{c}{R['Sales (₹ crore)']}-{c}{R['GST (₹ crore)']}-{c}{R['Construction (₹ crore)']}-{c}{R['Land (₹ crore)']}"
    return d
R=timeline(cf,labs,T,nf);E=L(T+2);row=lambda n:f"'Cash Flow'!$B${R[n]}:${E}${R[n]}"
table(wb.create_sheet('Returns'),[('Project IRR',f"=(1+IRR({row('Cash flow (₹ crore)')}))^4-1",P),('Revenue (₹ crore)',f"=SUM({row('Sales (₹ crore)')})",'0.00'),('Net profit (₹ crore)',f"=SUM({row('Cash flow (₹ crore)')})",'0.00'),
 ('Profit margin',"=B5/B4",P),('Revenue per apartment (₹)',f"=B4*10000000/{I['units']}",IN)])
wb.save(OUT+'pune_inr.xlsx')

# ================= 27E JV: a dashboard with values left of their labels =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='HAFENCITY JV – RESIDENTIAL DEVELOPMENT (EUR)'
I=inputs(ws,[('cost','Development cost (EUR)',148000000,N),('gdv','Sales (EUR)',196000000,N),('opeq','Operating partner equity share',0.1,P),('pref','Preferred return',0.06,P),('catch','Catch-up to the operating partner',0.5,P),
 ('promote','Promote above the catch-up',0.25,P),('years','Years to sell out',4,'0')])
cf=wb.create_sheet('Cash Flow');T=4
labs=['Year','Project cash flow','Capital partner','Operating partner']
def jf(t,c,p,R):
    y=f"{c}$3";first=t==0
    d={'Year':t,'Project cash flow':f"-{I['cost']}" if first else f"{I['gdv']}/{I['years']}"}
    d['Capital partner']=f"{c}{R['Project cash flow']}*(1-{I['opeq']})" if first else "0"
    d['Operating partner']=f"{c}{R['Project cash flow']}*{I['opeq']}" if first else "0"
    return d
R=timeline(cf,labs,T,jf);E=L(T+2)
wf=wb.create_sheet('Waterfall')
for i,(lab,f) in enumerate([('Distributable',f"=SUMIF('Cash Flow'!C{R['Project cash flow']}:{E}{R['Project cash flow']},\">0\")"),('Equity in',f"={I['cost']}"),
  ('Pref to all',f"=MIN(B3,B4*((1+{I['pref']})^{I['years']}-1))"),('Return of equity',f"=MIN(B3-B5,B4)"),('Catch-up',f"=MIN(B3-B5-B6,B5*{I['catch']}/(1-{I['catch']})*{I['promote']}/{I['catch']})"),
  ('Remaining',"=B3-B5-B6-B7"),('Operating partner total',f"=(B5+B6)*{I['opeq']}+B7+B8*{I['promote']}+B8*(1-{I['promote']})*{I['opeq']}"),('Capital partner total',"=B3-B9")],3):
    wf.cell(i,1,lab);wf.cell(i,2,f).number_format=N
db=wb.create_sheet('Dashboard');db.merge_cells('A1:D1');db['A1']='HAFENCITY JV – DASHBOARD'
db.merge_cells('A3:B3');db['A3']='Returns'
rows=[("=Waterfall!B10/(Inputs!B3*(1-Inputs!B5))",'Capital partner multiple',X),("=Waterfall!B9/(Inputs!B3*Inputs!B5)",'Operating partner multiple',X),
      ("=Waterfall!B3-Inputs!B3",'Net profit (EUR)',N),("=(Waterfall!B3/Inputs!B3)^(1/Inputs!B9)-1",'Project return a year',P),("=Waterfall!B7+Waterfall!B8*Inputs!B8",'Promote paid (EUR)',N)]
for i,(f,lab,fm) in enumerate(rows,4):c=db.cell(i,1,f);c.number_format=fm;db.cell(i,2,lab)
wb.save(OUT+'jv_eur.xlsx')
print('ok')

# ======================================== FUNDS ========================================
# ================= 27F mutual fund: monthly NAV, flows, fees, performance fee over a high-water mark =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='GULF INCOME & GROWTH FUND – OPEN-ENDED MUTUAL FUND (USD)'
I=inputs(ws,[('nav0','Opening NAV (USD)',250000000,N),('units0','Units outstanding',25000000,N),('ret','Gross return a month',0.0065,P),('vol','Return swing (+/-) every other month',0.012,P),
 ('mgmt','Management fee a year',0.0125,P),('perf','Performance fee over the high-water mark',0.15,P),('other','Other expenses a year',0.0035,P),('subs','Subscriptions a month (USD)',4500000,N),('reds','Redemptions a month (USD)',3200000,N)])
cf=wb.create_sheet('NAV');T=24
labs=['Month','Gross return','NAV before fees','Management fee','Other expenses','NAV per unit before performance fee','High-water mark','Performance fee','NAV after fees','NAV per unit','Units subscribed','Units redeemed','Units outstanding','NAV (USD)']
def mf_(t,c,p,R):
    m=f"{c}$3";first=t==0
    d={'Month':t}
    if first:
        d.update({'Gross return':'0','NAV before fees':f"{I['nav0']}",'Management fee':'0','Other expenses':'0','NAV per unit before performance fee':f"{I['nav0']}/{I['units0']}",'High-water mark':f"{I['nav0']}/{I['units0']}",
                  'Performance fee':'0','NAV after fees':f"{I['nav0']}",'NAV per unit':f"{I['nav0']}/{I['units0']}",'Units subscribed':'0','Units redeemed':'0','Units outstanding':f"{I['units0']}",'NAV (USD)':f"{I['nav0']}"})
        return d
    d['Gross return']=f"{I['ret']}+IF(MOD({m},2)=0,{I['vol']},-{I['vol']})"
    d['NAV before fees']=f"{p}{R['NAV (USD)']}*(1+{c}{R['Gross return']})"
    d['Management fee']=f"{c}{R['NAV before fees']}*{I['mgmt']}/12"
    d['Other expenses']=f"{c}{R['NAV before fees']}*{I['other']}/12"
    d['NAV per unit before performance fee']=f"({c}{R['NAV before fees']}-{c}{R['Management fee']}-{c}{R['Other expenses']})/{p}{R['Units outstanding']}"
    d['Performance fee']=f"MAX(0,{c}{R['NAV per unit before performance fee']}-{p}{R['High-water mark']})*{I['perf']}*{p}{R['Units outstanding']}"
    d['High-water mark']=f"MAX({p}{R['High-water mark']},{c}{R['NAV per unit before performance fee']}-{c}{R['Performance fee']}/{p}{R['Units outstanding']})"
    d['NAV after fees']=f"{c}{R['NAV before fees']}-{c}{R['Management fee']}-{c}{R['Other expenses']}-{c}{R['Performance fee']}"
    d['NAV per unit']=f"{c}{R['NAV after fees']}/{p}{R['Units outstanding']}"
    d['Units subscribed']=f"{I['subs']}/{c}{R['NAV per unit']}"
    d['Units redeemed']=f"MIN({I['reds']}/{c}{R['NAV per unit']},{p}{R['Units outstanding']})"
    d['Units outstanding']=f"{p}{R['Units outstanding']}+{c}{R['Units subscribed']}-{c}{R['Units redeemed']}"
    d['NAV (USD)']=f"{c}{R['Units outstanding']}*{c}{R['NAV per unit']}"
    return d
R=timeline(cf,labs,T,mf_);E=L(T+2);row=lambda n:f"NAV!$B${R[n]}:${E}${R[n]}"
table(wb.create_sheet('Returns'),[('Net return a year (NAV per unit)',f"=(NAV!{E}{R['NAV per unit']}/NAV!B{R['NAV per unit']})^(12/{T})-1",P),('NAV at month 24 (USD)',f"=NAV!{E}{R['NAV (USD)']}",N),
 ('NAV per unit at month 24 (USD)',f"=NAV!{E}{R['NAV per unit']}",'0.0000'),('Total expense ratio',f"=(SUM({row('Management fee')})+SUM({row('Other expenses')}))/AVERAGE({row('NAV (USD)')})*12/{T}",P),
 ('Performance fees paid (USD)',f"=SUM({row('Performance fee')})",N),('Net flows (USD)',f"=({I['subs']}-{I['reds']})*{T}",N)])
wb.save(OUT+'mutual_usd.xlsx')

# ================= 27G listed REIT: FFO, AFFO, 90% payout, NAV per unit vs price =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='AL MANAZIL REIT – LISTED SAUDI REIT (SAR)'
I=inputs(ws,[('units','Units in issue',150000000,N),('price','Unit price on the exchange (SAR)',9.2,'0.00'),('mgmt','Fund manager fee a year (of NAV)',0.0075,P),('payout','Distribution of net income',0.9,P),
 ('rate','Murabaha rate',0.062,P),('debt','Murabaha facility (SAR)',520000000,N),('capexr','Maintenance capex share of NOI',0.06,P),('zak','Zakat rate',0.025,P),('growth','Rent growth a year',0.02,P),('cap','Valuation cap rate',0.075,P)])
pp=wb.create_sheet('Properties');pp.append(['Property','Sector','NOI a year (SAR)','Occupancy'])
for row_ in [('Olaya Towers','Office',61000000,0.93),('Jeddah Logistics Park','Logistics',38500000,1.0),('Al Khobar Mall','Retail',47200000,0.88),('Riyadh Hotel Apartments','Hospitality',22800000,0.79),('Dammam Schools','Education',18400000,1.0)]:pp.append(list(row_))
for r in range(2,7):pp.cell(r,4).number_format=P
cf=wb.create_sheet('Projection');T=5
labs=['Year','NOI','Fund manager fee','Finance cost','Depreciation','Net income','FFO','Maintenance capex','AFFO','Zakat','Distribution','Distribution per unit','Property value','NAV','NAV per unit']
def rf_(t,c,p,R):
    y=f"{c}$3";first=t==0
    d={'Year':t,'NOI':'0' if first else f"SUMPRODUCT(Properties!$C$2:$C$6,Properties!$D$2:$D$6/MAX(Properties!$D$2:$D$6))*(1+{I['growth']})^({y}-1)"}
    d['Property value']=f"SUMPRODUCT(Properties!$C$2:$C$6,Properties!$D$2:$D$6/MAX(Properties!$D$2:$D$6))*(1+{I['growth']})^{y}/{I['cap']}"
    d['NAV']=f"{c}{R['Property value']}-{I['debt']}"
    d['NAV per unit']=f"{c}{R['NAV']}/{I['units']}"
    d['Fund manager fee']='0' if first else f"{p}{R['NAV']}*{I['mgmt']}"
    d['Finance cost']='0' if first else f"{I['debt']}*{I['rate']}"
    d['Depreciation']='0' if first else f"{c}{R['Property value']}*0.012"
    d['Net income']=f"{c}{R['NOI']}-{c}{R['Fund manager fee']}-{c}{R['Finance cost']}-{c}{R['Depreciation']}"
    d['FFO']=f"{c}{R['Net income']}+{c}{R['Depreciation']}"
    d['Maintenance capex']=f"{c}{R['NOI']}*{I['capexr']}"
    d['AFFO']=f"{c}{R['FFO']}-{c}{R['Maintenance capex']}"
    d['Zakat']=f"MAX(0,{c}{R['Net income']})*{I['zak']}"
    d['Distribution']=f"MAX(0,{c}{R['Net income']}-{c}{R['Zakat']})*{I['payout']}"
    d['Distribution per unit']=f"{c}{R['Distribution']}/{I['units']}"
    return d
R=timeline(cf,labs,T,rf_);E=L(T+2)
table(wb.create_sheet('Returns'),[('Dividend yield on the unit price',f"=Projection!C{R['Distribution per unit']}/{I['price']}",P),('FFO per unit, year 1 (SAR)',f"=Projection!C{R['FFO']}/{I['units']}",'0.000'),
 ('AFFO payout ratio',f"=Projection!C{R['Distribution']}/Projection!C{R['AFFO']}",P),('NAV per unit (SAR)',f"=Projection!B{R['NAV per unit']}",'0.00'),('Premium or discount to NAV',f"={I['price']}/Projection!B{R['NAV per unit']}-1",P),
 ('Gearing (debt to property value)',f"={I['debt']}/Projection!B{R['Property value']}",P),('Total return to unitholders a year',f"=((Projection!{E}{R['NAV per unit']}+SUM(Projection!C{R['Distribution per unit']}:{E}{R['Distribution per unit']}))/{I['price']})^(1/{T})-1",P)])
wb.save(OUT+'reit_sar.xlsx')

# ================= 27H closed-ended fund: commitments, calls, distributions, European waterfall with clawback =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Terms';ws['A1']='MERIDIAN REAL ESTATE FUND III – CLOSED-ENDED (EUR)'
I=inputs(ws,[('commit','Commitments (EUR)',400000000,N),('fee','Management fee on commitments, years 1–4',0.015,P),('feei','Management fee on invested capital after year 4',0.01,P),('pref','Preferred return',0.08,P),
 ('carry','Carried interest',0.2,P),('catch','GP catch-up',1.0,P),('mult','Gross exit multiple on deals',1.75,X)])
dl=wb.create_sheet('Deals');dl.append(['Deal','Equity (EUR)','Invest year','Exit year','Yield to equity'])
for row_ in [('Madrid offices',62000000,1,5,0.045),('Lyon logistics',48000000,1,6,0.055),('Milan resi',71000000,2,6,0.04),('Berlin student',39000000,2,7,0.05),('Warsaw retail',44000000,3,7,0.065),('Lisbon hotel',56000000,3,8,0.03)]:dl.append(list(row_))
for r in range(2,8):dl.cell(r,5).number_format=P
cf=wb.create_sheet('Fund');T=10
labs=['Year','Invested','Income','Exit proceeds','Management fee','Capital called','Distributions','Net cash to LPs','Cumulative called','Cumulative distributed','Unrealised value','DPI','RVPI','TVPI']
def ff(t,c,p,R):
    y=f"{c}$3"
    d={'Year':t,'Invested':f"SUMIF(Deals!$C$2:$C$7,{y},Deals!$B$2:$B$7)",
       'Income':f"SUMPRODUCT((Deals!$C$2:$C$7<{y})*(Deals!$D$2:$D$7>={y})*Deals!$B$2:$B$7*Deals!$E$2:$E$7)",
       'Exit proceeds':f"SUMIF(Deals!$D$2:$D$7,{y},Deals!$B$2:$B$7)*{I['mult']}",
       'Management fee':f"IF({y}=0,0,IF({y}<=4,{I['commit']}*{I['fee']},SUMPRODUCT((Deals!$C$2:$C$7<{y})*(Deals!$D$2:$D$7>={y})*Deals!$B$2:$B$7)*{I['feei']}))"}
    d['Capital called']=f"{c}{R['Invested']}+{c}{R['Management fee']}"
    d['Distributions']=f"{c}{R['Income']}+{c}{R['Exit proceeds']}"
    d['Net cash to LPs']=f"{c}{R['Distributions']}-{c}{R['Capital called']}-IF({y}={T},Waterfall!$B$9-Waterfall!$B$11,0)"
    d['Cumulative called']=f"SUM($B${R['Capital called']}:{c}{R['Capital called']})"
    d['Cumulative distributed']=f"SUM($B${R['Distributions']}:{c}{R['Distributions']})"
    d['Unrealised value']=f"SUMPRODUCT((Deals!$C$2:$C$7<={y})*(Deals!$D$2:$D$7>{y})*Deals!$B$2:$B$7)*(1+({I['mult']}-1)*0.5)"
    d['DPI']=f"IF({c}{R['Cumulative called']}>0,{c}{R['Cumulative distributed']}/{c}{R['Cumulative called']},0)"
    d['RVPI']=f"IF({c}{R['Cumulative called']}>0,{c}{R['Unrealised value']}/{c}{R['Cumulative called']},0)"
    d['TVPI']=f"{c}{R['DPI']}+{c}{R['RVPI']}"
    return d
R=timeline(cf,labs,T,ff);E=L(T+2);row=lambda n:f"Fund!$B${R[n]}:${E}${R[n]}"
wf=wb.create_sheet('Waterfall')
for i,(lab,f) in enumerate([('Total called',f"=SUM({row('Capital called')})"),('Total distributed',f"=SUM({row('Distributions')})"),('Profit',"=B4-B3"),
  ('Preferred return',f"=B3*((1+{I['pref']})^5-1)"),('GP catch-up',f"=MIN(MAX(0,B5-B6),B6*{I['carry']}/(1-{I['carry']})*{I['catch']})"),('Split above',"=MAX(0,B5-B6-B7)"),
  ('Carried interest',f"=B7+B8*{I['carry']}"),('Carry if it were a fifth of profit',f"=B5*{I['carry']}"),('Clawback',"=MAX(0,B9-B10)"),('Net to LPs',"=B4-B9+B11")],3):
    wf.cell(i,1,lab);wf.cell(i,2,f).number_format=N
table(wb.create_sheet('Returns'),[('Net IRR to LPs',f"=IRR({row('Net cash to LPs')})",P),('TVPI at year 10',f"=Fund!{E}{R['TVPI']}",X),('DPI at year 5',f"=Fund!G{R['DPI']}",X),('RVPI at year 5',f"=Fund!G{R['RVPI']}",X),
 ('Carried interest (EUR)',"=Waterfall!B9",N),('Clawback (EUR)',"=Waterfall!B11",N),('Peak capital called (EUR)',f"=MAX({row('Cumulative called')})",N)])
wb.save(OUT+'closed_eur.xlsx')

# ================= 27I open-ended core fund: quarterly NAV, a redemption queue with a 10% gate, liquidity sleeve =================
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='NORDIC CORE PROPERTY FUND – OPEN-ENDED, QUARTERLY (EUR)'
I=inputs(ws,[('gav','Property value (EUR)',1800000000,N),('cash0','Cash (EUR)',90000000,N),('debt','Debt (EUR)',420000000,N),('noiy','NOI yield on value',0.048,P),('cg','Capital growth a year',0.01,P),
 ('subs','Subscriptions a quarter (EUR)',20000000,N),('reds','Redemption requests a quarter (EUR)',160000000,N),('gate','Gate (share of NAV a quarter)',0.1,P),('levy','Dilution levy on redemptions',0.01,P),
 ('fee','Management fee a year (of NAV)',0.0065,P),('units0','Units in issue',13000000,N),('mincash','Minimum cash (share of NAV)',0.03,P)])
cf=wb.create_sheet('NAV');T=12
labs=['Quarter','Property value','Income','Fee','Requests','Paid out','Queue','Property sold','Cash','NAV','Units','NAV per unit']
def of(t,c,p,R):
    q=f"{c}$3";first=t==0
    if first:return{'Quarter':0,'Property value':f"{I['gav']}",'Income':'0','Fee':'0','Requests':'0','Paid out':'0','Queue':'0','Property sold':'0','Cash':f"{I['cash0']}",
        'NAV':f"{I['gav']}+{I['cash0']}-{I['debt']}",'Units':f"{I['units0']}",'NAV per unit':f"({I['gav']}+{I['cash0']}-{I['debt']})/{I['units0']}"}
    d={'Quarter':t}
    d['Income']=f"{p}{R['Property value']}*{I['noiy']}/4"
    d['Fee']=f"{p}{R['NAV']}*{I['fee']}/4"
    d['Requests']=f"{I['reds']}+{p}{R['Queue']}"
    d['Paid out']=f"MIN({c}{R['Requests']},{p}{R['NAV']}*{I['gate']})*(1-{I['levy']})"
    d['Queue']=f"{c}{R['Requests']}-{c}{R['Paid out']}/(1-{I['levy']})"
    d['Property sold']=f"MAX(0,{p}{R['NAV']}*{I['mincash']}-({p}{R['Cash']}+{c}{R['Income']}-{c}{R['Fee']}+{I['subs']}-{c}{R['Paid out']}))"
    d['Property value']=f"{p}{R['Property value']}*(1+{I['cg']})^(1/4)-{c}{R['Property sold']}"
    d['Cash']=f"{p}{R['Cash']}+{c}{R['Income']}-{c}{R['Fee']}+{I['subs']}-{c}{R['Paid out']}+{c}{R['Property sold']}"
    d['NAV']=f"{c}{R['Property value']}+{c}{R['Cash']}-{I['debt']}"
    d['Units']=f"{p}{R['Units']}+{I['subs']}/{p}{R['NAV per unit']}-{c}{R['Paid out']}/(1-{I['levy']})/{p}{R['NAV per unit']}"
    d['NAV per unit']=f"{c}{R['NAV']}/{c}{R['Units']}"
    return d
R=timeline(cf,labs,T,of);E=L(T+2);row=lambda n:f"NAV!$B${R[n]}:${E}${R[n]}"
table(wb.create_sheet('Returns'),[('Total return a year (NAV per unit)',f"=(NAV!{E}{R['NAV per unit']}/NAV!B{R['NAV per unit']})^(4/{T})-1",P),('Redemption queue at the end (EUR)',f"=NAV!{E}{R['Queue']}",N),
 ('Quarters gated',f"=COUNTIF({row('Queue')},\">0\")",'0'),('Property sold to pay redemptions (EUR)',f"=SUM({row('Property sold')})",N),('Loan to value at the end',f"={I['debt']}/NAV!{E}{R['Property value']}",P),
 ('NAV at the end (EUR)',f"=NAV!{E}{R['NAV']}",N)])
wb.save(OUT+'openend_eur.xlsx')
print('funds ok')
