# A fourth model, unseen by the tuning: Saudi mixed-use in SAR. Labels in A, units in B, values in C; percentages typed as whole
# numbers with a "%" unit column (no % format); two phases with the same "Construction cost" label; a start date; annual cash flow with XIRR.
import openpyxl, subprocess, sys, os, shutil, datetime
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as L
wb=openpyxl.Workbook();I=wb.active;I.title="Inputs"
I["A1"]="AL WAHA MIXED-USE · inputs";I["A1"].font=Font(bold=True)
rows=[("General",None,None),("Project start date","date",datetime.datetime(2026,1,1)),("Land","SAR",60000000),("Build period","years",2),("Hold after completion","years",5),
("Phase 1 · Residential",None,None),("Residential GLA","sqm",24000),("Residential rent","SAR/sqm/yr",1100),("Construction cost","SAR/sqm",4200),
("Phase 2 · Retail",None,None),("Retail GLA","sqm",9000),("Retail rent","SAR/sqm/yr",1800),("Construction cost","SAR/sqm",5000),
("Operations",None,None),("Vacancy","%",5),("Operating costs","% of rent",20),("Rent growth","% p.a.",3),("Exit yield","%",7.5),
("Funding",None,None),("Debt (loan to cost)","%",50),("Interest rate","%",6.5)]
ref={};r=3
for lab,u,v in rows:
    I.cell(r,1,lab)
    if v is None:I.cell(r,1).font=Font(bold=True)
    else:
        I.cell(r,2,u);c=I.cell(r,3,v)
        c.number_format='dd/mm/yyyy' if u=='date' else '#,##0' if isinstance(v,int) and v>=1000 else '0.0' if u in('%','% p.a.','% of rent') else '#,##0'
        ref.setdefault(lab,[]).append(f"Inputs!$C${r}")
    r+=1
g=lambda k,i=0:ref[k][i]
C=wb.create_sheet("CashFlow");C["A1"]="Annual cash flow (SAR)";C["A1"].font=Font(bold=True)
lab=["Year","Date","Land","Construction","Rent (net)","Exit value","Unlevered cash flow","Debt draw","Interest","Debt repayment","Levered cash flow"]
for i,t in enumerate(lab):C.cell(3+i,1,t)
N=10;cost=f"({g('Residential GLA')}*{g('Construction cost',0)}+{g('Retail GLA')}*{g('Construction cost',1)})"
gross=f"({g('Residential GLA')}*{g('Residential rent')}+{g('Retail GLA')}*{g('Retail rent')})"
bp=g('Build period');hd=g('Hold after completion')
for n in range(N):
    col=L(2+n);pc=L(1+n)
    C[f"{col}3"]=n
    C[f"{col}4"]=f"=EDATE({g('Project start date')},12*{col}3)";C[f"{col}4"].number_format='dd/mm/yyyy'
    C[f"{col}5"]=f"=IF({col}3=0,-{g('Land')},0)"
    C[f"{col}6"]=f"=IF({col}3<{bp},-{cost}/{bp},0)"
    C[f"{col}7"]=f"=IF(AND({col}3>={bp},{col}3<{bp}+{hd}),{gross}*(1-{g('Vacancy')}/100)*(1-{g('Operating costs')}/100)*(1+{g('Rent growth')}/100)^({col}3-{bp}),0)"
    C[f"{col}8"]=f"=IF({col}3={bp}+{hd}-1,{gross}*(1-{g('Vacancy')}/100)*(1-{g('Operating costs')}/100)*(1+{g('Rent growth')}/100)^({col}3-{bp}+1)/({g('Exit yield')}/100),0)"
    C[f"{col}9"]=f"=SUM({col}5:{col}8)"
    C[f"{col}10"]=f"=IF({col}3<{bp},-{col}6*{g('Debt (loan to cost)')}/100,0)"
    C[f"{col}11"]=f"=-({'0' if n==0 else 'SUM($B$10:'+pc+'10)-SUM($B$12:'+pc+'12)'})*{g('Interest rate')}/100"
    C[f"{col}12"]=f"=IF({col}3={bp}+{hd}-1,SUM($B$10:{col}10)-{'0' if n==0 else 'SUM($B$12:'+pc+'12)'},0)"
    C[f"{col}13"]=f"={col}9+{col}10+{col}11-{col}12"
R=wb.create_sheet("Results");R["A1"]="AL WAHA MIXED-USE";R["A1"].font=Font(bold=True)
res=[("Total development cost","=-SUM(CashFlow!B5:K6)","#,##0"),("Net profit","=SUM(CashFlow!B13:K13)","#,##0"),("Profit on cost","=B4/B3","0.0%"),
("Project IRR (unlevered)","=XIRR(CashFlow!B9:K9,CashFlow!B4:K4)","0.00%"),("Equity IRR (levered)","=XIRR(CashFlow!B13:K13,CashFlow!B4:K4)","0.00%"),
("Peak debt","=SUM(CashFlow!B10:K10)","#,##0")]
for i,(l,f,fm) in enumerate(res):R.cell(3+i,1,l);c=R.cell(3+i,2,f);c.number_format=fm
tmp=os.path.dirname(os.path.abspath(__file__))+"/waha_src";os.makedirs(tmp,exist_ok=True);src=tmp+"/waha_sar.xlsx";wb.save(src)
out=tmp+"/out";shutil.rmtree(out,ignore_errors=True);os.makedirs(out)
subprocess.run(["soffice","--headless","--calc","--convert-to","xlsx:Calc MS Excel 2007 XML","--outdir",out,src],check=True,capture_output=True)
shutil.copy(out+"/waha_sar.xlsx",sys.argv[1]);print("written")
