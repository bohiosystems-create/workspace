# Round 42 fixture (synthetic): a multi-tenant office acquisition with a lease-by-lease rent roll, with made-up figures.
#  r42_office_gbp: a Rent Roll table (tenant, area, rent per sq ft, lease expiry year, renewal probability, downtime months, TI and LC);
#  each tenant's rent by year from the table: in place to expiry, then the market rent weighted by the renewal probability, less the
#  downtime on the share that leaves; TI and leasing commissions at each expiry; the cash flow summed over tenants; a sale at the
#  exit cap on the next year's NOI; returns on a Summary whose figures read "IRR (ungeared)", "IRR (geared)", "Profit on cost" is
#  absent and "Equity multiple" is given; amounts in GBP; rent per sq ft a year.
import openpyxl,sys
from openpyxl.utils import get_column_letter as L
out=sys.argv[1].rstrip('/')
wb=openpyxl.Workbook();S=wb.active;S.title='Summary';A=wb.create_sheet('Assumptions');RR=wb.create_sheet('Rent Roll');T=wb.create_sheet('Tenant CF');C=wb.create_sheet('Cash Flow')
Y=11   # years 0..11; exit at the end of year 10 on year 11's NOI
A['B2']='ST MARY AXE HOUSE · ASSUMPTIONS (GBP)'
rows=[(4,'Purchase price',150000000,'#,##0'),(5,'Purchase costs (SDLT and fees)',0.068,'0.00%'),(6,'Market rent (per sq ft a year)',72.5,'0.00'),
 (7,'Market rent growth',0.025,'0.00%'),(8,'Non-recoverable costs (% of rent)',0.035,'0.00%'),(9,'Exit cap rate',0.0525,'0.00%'),(10,'Selling costs',0.015,'0.00%'),
 (11,'Hold (years)',10,'0'),(12,'Loan to value',0.5,'0.00%'),(13,'Interest rate',0.058,'0.00%'),(14,'Arrangement fee',0.01,'0.00%')]
for r,l,v,f in rows:A.cell(r,2,l);A.cell(r,3,v).number_format=f
RR['B2']='RENT ROLL';hd=['Tenant','Area (sq ft)','Rent (per sq ft)','Lease expiry (year)','Renewal probability','Downtime (months)','TI (per sq ft)','LC (months of rent)']
for j,h in enumerate(hd):RR.cell(4,2+j,h)
ten=[('Harcourt Legal LLP',32000,64.0,3,0.65,9,45,2.0),('Northgate Capital',21500,70.0,5,0.70,6,40,2.0),('Ferro Digital',18000,58.5,2,0.50,12,50,2.5),
 ('Albion Insurance',40000,66.0,8,0.75,9,45,2.0),('Quill & Co',9500,74.0,6,0.60,6,35,1.5),('Tessera Labs',14000,61.0,4,0.55,9,50,2.0)]
for i,t in enumerate(ten):
    r=5+i
    for j,v in enumerate(t):
        c=RR.cell(r,2+j,v)
        if j==1:c.number_format='#,##0'
        if j in (2,6):c.number_format='0.00'
        if j==4:c.number_format='0%'
NT=len(ten);RR[f'B{5+NT}']='Total';RR[f'C{5+NT}']=f'=SUM(C5:C{4+NT})';RR[f'C{5+NT}'].number_format='#,##0'
# tenant cash flow: one block of three rows per tenant (rent, TI and LC), year 0 in column D
T['B2']='TENANT CASH FLOW (GBP)';T['B4']='Year'
for y in range(Y+1):T.cell(4,4+y,y)
for i in range(NT):
    rr=5+i;base=6+i*4;T.cell(base,2,f"='Rent Roll'!B{rr}");T.cell(base,3,'Rent');T.cell(base+1,3,'TI');T.cell(base+2,3,'Leasing commission')
    for y in range(Y+1):
        c=L(4+y);yy=f'{c}$4';exp=f"'Rent Roll'!$E${rr}";area=f"'Rent Roll'!$C${rr}";p=f"'Rent Roll'!$F${rr}";dt=f"ROUND('Rent Roll'!$G${rr},0)"
        mkt=f'Assumptions!$C$6*(1+Assumptions!$C$7)^({yy}-1)'
        if y==0:
            for k in range(3):T[f'{c}{base+k}']=0
            continue
        T[f'{c}{base}']=(f"=IF({yy}<={exp},{area}*'Rent Roll'!$D${rr},{area}*{mkt}*({p}+(1-{p})*IF({yy}={exp}+1,MAX(0,12-{dt})/12,1)))")
        T[f'{c}{base+1}']=f"=IF({yy}={exp}+1,-{area}*'Rent Roll'!$H${rr}*(1-{p}*0.5),0)"
        T[f'{c}{base+2}']=f"=IF({yy}={exp}+1,-{area}*{mkt}/12*'Rent Roll'!$I${rr},0)"
E=L(4+Y)
C['B2']='CASH FLOW (GBP)';C['B4']='Year'
cl={6:'Gross rent',7:'Non-recoverable costs',8:'Net operating income',9:'TI and leasing costs',10:'Purchase',11:'Sale (net)',12:'Ungeared cash flow',
 14:'Loan',15:'Interest',16:'Repayment',17:'Geared cash flow'}
for r,l in cl.items():C.cell(r,2,l)
rent_rows=[str(6+i*4) for i in range(NT)];cap_rows=[str(7+i*4) for i in range(NT)]+[str(8+i*4) for i in range(NT)]
for y in range(Y+1):
    c=L(4+y);n=L(5+y);yy=f'{c}$4';H='ROUND(Assumptions!$C$11,0)'
    C[f'{c}4']=y
    C[f'{c}6']='='+'+'.join(f"'Tenant CF'!{c}{r}" for r in rent_rows)
    C[f'{c}7']=f'=-{c}6*Assumptions!$C$8';C[f'{c}8']=f'={c}6+{c}7'
    C[f'{c}9']='='+'+'.join(f"'Tenant CF'!{c}{r}" for r in cap_rows)
    C[f'{c}10']=f'=IF({yy}=0,-Assumptions!$C$4*(1+Assumptions!$C$5),0)'
    C[f'{c}11']=f'=IF({yy}={H},{n}8/Assumptions!$C$9*(1-Assumptions!$C$10),0)' if y<Y else '=0'
    C[f'{c}12']=f'=IF({yy}>{H},0,SUM({c}8:{c}11))'
    C[f'{c}14']=f'=IF({yy}=0,Assumptions!$C$4*Assumptions!$C$12*(1-Assumptions!$C$14),0)'
    C[f'{c}15']=f'=IF(AND({yy}>=1,{yy}<={H}),-Assumptions!$C$4*Assumptions!$C$12*Assumptions!$C$13,0)'
    C[f'{c}16']=f'=IF({yy}={H},-Assumptions!$C$4*Assumptions!$C$12,0)'
    C[f'{c}17']=f'={c}12+{c}14+{c}15+{c}16'
S['B2']='ST MARY AXE HOUSE · SUMMARY (GBP)'
sm=[(5,'Net initial yield',f"='Cash Flow'!E8/(Assumptions!C4*(1+Assumptions!C5))",'0.00%'),(6,'IRR (ungeared)',f"=IRR('Cash Flow'!D12:{E}12)",'0.00%'),
 (7,'IRR (geared)',f"=IRR('Cash Flow'!D17:{E}17)",'0.00%'),(8,'Equity multiple',f"=SUMIF('Cash Flow'!D17:{E}17,\">0\")/-SUMIF('Cash Flow'!D17:{E}17,\"<0\")",'0.00"x"'),
 (9,'WAULT (years)',f"=SUMPRODUCT('Rent Roll'!C5:C{4+NT},'Rent Roll'!D5:D{4+NT},'Rent Roll'!E5:E{4+NT})/SUMPRODUCT('Rent Roll'!C5:C{4+NT},'Rent Roll'!D5:D{4+NT})",'0.0'),
 (10,'Geared profit',f"=SUM('Cash Flow'!D17:{E}17)",'#,##0')]
for r,l,f,fm in sm:S.cell(r,2,l);S.cell(r,3,f).number_format=fm
wb.save(f'{out}/r42_office_gbp.xlsx')
