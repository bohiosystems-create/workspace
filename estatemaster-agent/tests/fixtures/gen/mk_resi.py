#!/usr/bin/env python3
# A second, deliberately different model: UK residential scheme in GBP, inputs on one sheet (labels col B, values col D,
# percentages stored as fractions), a monthly cash flow with debt drawn to a loan-to-cost cap, and a returns sheet with
# IRR() annualised, NPV, profit on cost, peak debt. Saved via LibreOffice so every formula carries its calculated value.
import openpyxl, subprocess, sys, os, shutil
from openpyxl.styles import Font
wb = openpyxl.Workbook()
A = wb.active; A.title = "Assumptions"
A["B2"] = "RIVERSIDE GARDENS · residential scheme"; A["B2"].font = Font(bold=True)
rows = [
 ("Site", None, None),
 ("Site area (sqm)", 6500, "#,##0"),
 ("Gross internal area (sqm)", 12000, "#,##0"),
 ("Net sales area (sqm)", 10200, "#,##0"),
 ("Revenue", None, None),
 ("Sales price (£/sqm)", 8500, "£#,##0"),
 ("Sales period (months)", 12, "0"),
 ("Costs", None, None),
 ("Land cost", 25000000, "£#,##0"),
 ("Build cost (£/sqm GIA)", 2800, "£#,##0"),
 ("Professional fees", 0.10, "0.0%"),
 ("Contingency", 0.05, "0.0%"),
 ("Build period (months)", 24, "0"),
 ("Funding", None, None),
 ("Loan to cost", 0.65, "0%"),
 ("Interest rate", 0.075, "0.00%"),
 ("Discount rate", 0.10, "0.0%"),
]
r = 4
ref = {}
for lab, v, fmt in rows:
    A.cell(r, 2, lab)
    if v is None: A.cell(r, 2).font = Font(bold=True)
    else:
        c = A.cell(r, 4, v); c.number_format = fmt; ref[lab] = f"Assumptions!$D${r}"
    r += 1
C = wb.create_sheet("CashFlow")
N = 40
C["B2"] = "Monthly cash flow (£)"; C["B2"].font = Font(bold=True)
C["B4"] = "Month"; C["B5"] = "Land"; C["B6"] = "Build"; C["B7"] = "Fees & contingency"; C["B8"] = "Sales revenue"
C["B9"] = "Net cash flow (unlevered)"; C["B10"] = "Cumulative cost"; C["B11"] = "Debt draw"; C["B12"] = "Interest"; C["B13"] = "Debt balance"; C["B14"] = "Debt repayment"; C["B15"] = "Net cash flow to equity"
gdv = f"({ref['Net sales area (sqm)']}*{ref['Sales price (£/sqm)']})"
build = f"({ref['Gross internal area (sqm)']}*{ref['Build cost (£/sqm GIA)']})"
bp = ref["Build period (months)"]; sp = ref["Sales period (months)"]
for m in range(1, N + 1):
    col = openpyxl.utils.get_column_letter(3 + m); pc = openpyxl.utils.get_column_letter(2 + m)
    C[f"{col}4"] = m
    C[f"{col}5"] = f"=IF({col}4=1,-{ref['Land cost']},0)"
    C[f"{col}6"] = f"=IF({col}4<={bp},-{build}/{bp},0)"
    C[f"{col}7"] = f"={col}6*({ref['Professional fees']}+{ref['Contingency']})"
    C[f"{col}8"] = f"=IF(AND({col}4>{bp},{col}4<={bp}+{sp}),{gdv}/{sp},0)"
    C[f"{col}9"] = f"=SUM({col}5:{col}8)"
    C[f"{col}10"] = f"=-SUM($D$5:{col}7)"
    # debt: costs are funded by debt once equity has covered (1-LTC) of total cost; interest rolls up
    C[f"{col}11"] = f"=MAX(0,MIN(-({col}5+{col}6+{col}7),{col}10-(1-{ref['Loan to cost']})*(-SUM($D$5:$AQ$7))))"
    C[f"{col}12"] = f"={'0' if m==1 else pc+'13'}*{ref['Interest rate']}/12"
    C[f"{col}14"] = f"=IF({col}8>0,MIN({col}8,{'0' if m==1 else pc+'13'}+{col}11+{col}12),0)"
    C[f"{col}13"] = f"={'0' if m==1 else pc+'13'}+{col}11+{col}12-{col}14"
    C[f"{col}15"] = f"={col}9+{col}11-{col}14"
    for rr in range(5, 16):
        C[f"{col}{rr}"].number_format = "#,##0;(#,##0)"
R = wb.create_sheet("Returns")
R["B2"] = "RIVERSIDE GARDENS · returns"; R["B2"].font = Font(bold=True)
out = [
 ("Gross development value", f"={gdv}", "£#,##0"),
 ("Total development cost", f"=-SUM(CashFlow!D5:AQ7)+SUM(CashFlow!D12:AQ12)", "£#,##0"),
 ("Net profit", "=D4-D5", "£#,##0"),
 ("Profit on cost", "=D6/D5", "0.00%"),
 ("Unlevered IRR", "=(1+IRR(CashFlow!D9:AQ9,0.02))^12-1", "0.00%"),
 ("Levered IRR", "=(1+IRR(CashFlow!D15:AQ15,0.02))^12-1", "0.00%"),
 ("Peak debt", "=MAX(CashFlow!D13:AQ13)", "£#,##0"),
 ("NPV at discount rate", f"=NPV({ref['Discount rate']}/12,CashFlow!D9:AQ9)", "£#,##0"),
 ("Equity multiple", "=SUMIF(CashFlow!D15:AQ15,\">0\")/-SUMIF(CashFlow!D15:AQ15,\"<0\")", "0.00\"x\""),
]
r = 4
for lab, f, fmt in out:
    R.cell(r, 2, lab); c = R.cell(r, 4, f); c.number_format = fmt; r += 1
K = wb.create_sheet("Checks")
K["B2"] = "Checks and dates"; K["B2"].font = Font(bold=True)
chk = [
 ("Analysis date", "=TODAY()", "dd/mm/yyyy"),
 ("Completion date", "=EDATE(D4,Assumptions!D16)", "dd/mm/yyyy"),
 ("Average monthly sales after build", '=AVERAGEIFS(CashFlow!D8:AQ8,CashFlow!D4:AQ4,">"&Assumptions!D16)', "£#,##0"),
 ("Scheme", '=_xlfn.CONCAT("Scheme: ",Returns!B2)', "@"),
 ("Debt balance at month 30", "=LOOKUP(30,CashFlow!D4:AQ4,CashFlow!D13:AQ13)", "£#,##0"),
 ("Rank of levered IRR", "=RANK(Returns!D9,Returns!D7:D9)", "0"),
 ("Debt trend at month 41", "=FORECAST(41,CashFlow!D13:AQ13,CashFlow!D4:AQ4)", "£#,##0"),
 ("Strength", '=_xlfn.IFS(Returns!D7>0.2,"strong",Returns!D7>0.1,"ok",TRUE,"weak")', "@"),
 ("Sales after build", "=SUMPRODUCT((CashFlow!D4:AQ4>Assumptions!D16)*CashFlow!D8:AQ8)", "£#,##0"),
 ("Days to completion", "=D5-D4", "0"),
]
r = 4
for lab, f, fmt in chk:
    K.cell(r, 2, lab); c = K.cell(r, 4, f); c.number_format = fmt; r += 1
wb.move_sheet("Returns", offset=-2)   # returns first: a different order from the hotel model
tmp = __import__("tempfile").mkdtemp()
os.makedirs(tmp, exist_ok=True); src = tmp + "/resi_uk.xlsx"; wb.save(src)
outdir = tmp + "/out"; shutil.rmtree(outdir, ignore_errors=True); os.makedirs(outdir)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", outdir, src], check=True, capture_output=True)
shutil.copy(outdir + "/resi_uk.xlsx", sys.argv[1]); print("written", sys.argv[1])
