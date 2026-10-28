"""Adel Income Fund I: a synthetic three-asset income fund, annual 2024-2031, every result a formula.
   python3 gen_fund.py budget|actual OUT.xlsx   (the actual file replaces the typed 2025 figures with the year's actuals)"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

kind, out = sys.argv[1], sys.argv[2]
A = kind == 'actual'
wb = Workbook()
B = Font(bold=True); H = Font(bold=True, size=13); Y = PatternFill('solid', fgColor='FFF2CC')
PCT, X, M = '0.00%', '0.00"x"', '#,##0'

ws = wb.active; ws.title = 'Inputs'
ws['A1'] = 'ADEL INCOME FUND I – ' + ('ACTUAL 2025 AND REFORECAST' if A else 'BUDGET 2025'); ws['A1'].font = H
ws['A2'] = 'Yellow cells are inputs. ' + ('2025 figures are the year\'s actuals (audited management accounts, valuer\'s report of 31 Dec 2025).' if A else '2025 figures are the approved budget.')
fund = [
    ('Fund', None),
    ('Acquisition year (assets bought on 31 December)', 2024),
    ('Unit price at launch (SAR)', 10),
    ('Debt as share of acquisition price', 0.35),
    ('Real estate transaction tax and acquisition costs', 0.055),
    ('SAIBOR', 0.0585 if A else 0.055),
    ('Margin (bps)', 200),
    ('Management fee (share of NAV)', 0.0085),
    ('Fund administration and custody fees (SAR a year)', 4_900_000 if A else 4_500_000),
    ('Cost inflation a year', 0.03),
    ('Distribution policy (share of cash available)', 0.9),
    ('Exit year (assets sold on 31 December)', 2031),
    ('Disposal costs', 0.015),
]
r = 4
for lab, v in fund:
    ws.cell(r, 1, lab)
    if v is None: ws.cell(r, 1).font = B
    else:
        c = ws.cell(r, 3, v); c.fill = Y
        c.number_format = PCT if isinstance(v, float) and v < 1 else M
    r += 1
F = {lab: f'Inputs!$C${4 + i}' for i, (lab, v) in enumerate(fund)}
r += 1
hdr = r; ws.cell(r, 1, 'Assets').font = B
assets = ['Al Olaya Office Tower, Riyadh', 'Corniche Retail Centre, Jeddah', 'Dammam Logistics Park']
for j, n in enumerate(assets): ws.cell(r, 3 + j, n).font = B; ws.cell(r, 3 + j).alignment = Alignment(wrap_text=True)
# (label, budget values, actual values, format)
rows = [
    ('Purchase price (SAR)', [620e6, 480e6, 380e6], None, M),
    ('Leasable area (sqm)', [48_000, 36_000, 140_000], None, M),
    ('Rent in 2025 (SAR per sqm a year)', [1450, 1350, 330], [1450, 1300, 345], M),
    ('Occupancy in 2025', [0.90, 0.88, 0.95], [0.86, 0.84, 0.98], PCT),
    ('Stabilised occupancy (from 2026)', [0.93, 0.92, 0.96], [0.92, 0.90, 0.97], PCT),
    ('Rent growth a year', [0.03, 0.025, 0.03], None, PCT),
    ('Operating costs (share of rent)', [0.18, 0.22, 0.08], [0.195, 0.22, 0.08], PCT),
    ('Capital expenditure in 2025 (SAR)', [12e6, 8e6, 3e6], [15.5e6, 7.2e6, 3e6], M),
    ('Maintenance capex from 2026 (share of purchase price)', [0.005, 0.005, 0.003], None, PCT),
    ('Valuation cap rate', [0.075, 0.08, 0.0725], [0.075, 0.0825, 0.07], PCT),
]
AR = {}
for lab, bv, av, fmt in rows:
    r += 1; ws.cell(r, 1, lab); AR[lab] = r
    for j in range(3):
        v = (av if (A and av) else bv)[j]
        c = ws.cell(r, 3 + j, v); c.fill = Y; c.number_format = fmt
ws.column_dimensions['A'].width = 52
for j in range(3): ws.column_dimensions[L(3 + j)].width = 18
ac = lambda lab, j: f"Inputs!${L(3 + j)}${AR[lab]}"

# Assets: annual 2024 (acquisition) to 2031 (exit), columns C..J
YRS = list(range(2024, 2032)); C0 = 3
col = lambda y: L(C0 + y - 2024)
wa = wb.create_sheet('Assets')
wa['A1'] = 'ASSETS – ANNUAL (SAR)'; wa['A1'].font = H
wa.cell(3, 1, 'Year').font = B
for y in YRS: wa.cell(3, C0 + y - 2024, y).font = B
wa.cell(4, 1, 'Operating year (0 = acquisition)')
for y in YRS: wa.cell(4, C0 + y - 2024, f'={col(y)}3-{F["Acquisition year (assets bought on 31 December)"]}')
wa.cell(5, 1, 'Exit flag')
for y in YRS: wa.cell(5, C0 + y - 2024, f'=IF({col(y)}3={F["Exit year (assets sold on 31 December)"]},1,0)')
AS = {}
r = 7
for j, n in enumerate(assets):
    wa.cell(r, 1, n).font = B; base = r; AS[j] = {}
    lines = ['Occupancy', 'Rent (SAR per sqm a year)', 'Gross rent', 'Operating costs', 'Net operating income', 'Capital expenditure', 'Valuation at year end', 'Exit value net of disposal costs']
    for k, ln in enumerate(lines):
        rr = base + 1 + k; wa.cell(rr, 1, ln); AS[j][ln] = rr
    for y in YRS:
        c = col(y); t = f'{c}$4'
        g = lambda ln: f'{c}{AS[j][ln]}'
        wa[f'{c}{AS[j]["Occupancy"]}'] = f'=IF({t}<1,0,IF({t}=1,{ac("Occupancy in 2025", j)},{ac("Stabilised occupancy (from 2026)", j)}))'
        wa[f'{c}{AS[j]["Rent (SAR per sqm a year)"]}'] = f'=IF({t}<1,0,{ac("Rent in 2025 (SAR per sqm a year)", j)}*(1+{ac("Rent growth a year", j)})^({t}-1))'
        wa[f'{c}{AS[j]["Gross rent"]}'] = f'={ac("Leasable area (sqm)", j)}*{g("Rent (SAR per sqm a year)")}*{g("Occupancy")}'
        wa[f'{c}{AS[j]["Operating costs"]}'] = f'={g("Gross rent")}*{ac("Operating costs (share of rent)", j)}'
        wa[f'{c}{AS[j]["Net operating income"]}'] = f'={g("Gross rent")}-{g("Operating costs")}'
        wa[f'{c}{AS[j]["Capital expenditure"]}'] = f'=IF({t}<1,0,IF({t}=1,{ac("Capital expenditure in 2025 (SAR)", j)},{ac("Purchase price (SAR)", j)}*{ac("Maintenance capex from 2026 (share of purchase price)", j)}))'
        wa[f'{c}{AS[j]["Valuation at year end"]}'] = f'=IF({t}<1,{ac("Purchase price (SAR)", j)},{g("Net operating income")}*(1+{ac("Rent growth a year", j)})/{ac("Valuation cap rate", j)})'
        wa[f'{c}{AS[j]["Exit value net of disposal costs"]}'] = f'={c}$5*{g("Valuation at year end")}*(1-{F["Disposal costs"]})'
        for ln in lines:
            wa[f'{c}{AS[j][ln]}'].number_format = PCT if ln == 'Occupancy' else M
    r = base + len(lines) + 2
wa.column_dimensions['A'].width = 40
for y in YRS: wa.column_dimensions[col(y)].width = 15
S = lambda c, ln: '+'.join(f'Assets!{c}{AS[j][ln]}' for j in range(3))

# Fund: annual
wf = wb.create_sheet('Fund')
wf['A1'] = 'FUND – ANNUAL (SAR)'; wf['A1'].font = H
wf.cell(3, 1, 'Year').font = B
for y in YRS: wf.cell(3, C0 + y - 2024, y).font = B
fl = ['Acquisition price', 'Acquisition costs', 'Debt drawn', 'Equity raised', 'Units in issue',
      'Net operating income', 'Capital expenditure', 'Gross asset value', 'Debt outstanding', 'Net asset value', 'Management fee',
      'Fund administration and custody fees', 'Interest', 'Cash available for distribution', 'Distributions to unitholders',
      'Exit proceeds net of debt', 'Cash flow to unitholders', 'Distribution yield on equity', 'NAV per unit (SAR)', 'Loan to value', 'Interest cover (NOI / interest)']
FR = {ln: 4 + k for k, ln in enumerate(fl)}
for ln, rr in FR.items(): wf.cell(rr, 1, ln)
price = '+'.join(ac('Purchase price (SAR)', j) for j in range(3))
for y in YRS:
    c = col(y); t = f'Assets!{c}$4'; g = lambda ln: f'{c}{FR[ln]}'
    wf[g('Acquisition price')] = f'=IF({t}=0,{price},0)'
    wf[g('Acquisition costs')] = f'={g("Acquisition price")}*{F["Real estate transaction tax and acquisition costs"]}'
    wf[g('Debt drawn')] = f'={g("Acquisition price")}*{F["Debt as share of acquisition price"]}'
    wf[g('Equity raised')] = f'={g("Acquisition price")}+{g("Acquisition costs")}-{g("Debt drawn")}'
    wf[g('Units in issue')] = f'=$C${FR["Equity raised"]}/{F["Unit price at launch (SAR)"]}'
    wf[g('Net operating income')] = f'={S(c, "Net operating income")}'
    wf[g('Capital expenditure')] = f'={S(c, "Capital expenditure")}'
    wf[g('Gross asset value')] = f'={S(c, "Valuation at year end")}'
    wf[g('Debt outstanding')] = f'=IF(Assets!{c}$5=1,0,$C${FR["Debt drawn"]})'
    wf[g('Net asset value')] = f'={g("Gross asset value")}-$C${FR["Debt drawn"]}'
    wf[g('Management fee')] = f'=IF({t}<1,0,{g("Net asset value")}*{F["Management fee (share of NAV)"]})'
    wf[g('Fund administration and custody fees')] = f'=IF({t}<1,0,{F["Fund administration and custody fees (SAR a year)"]}*(1+{F["Cost inflation a year"]})^({t}-1))'
    wf[g('Interest')] = f'=IF({t}<1,0,$C${FR["Debt drawn"]}*({F["SAIBOR"]}+{F["Margin (bps)"]}/10000))'
    wf[g('Cash available for distribution')] = f'={g("Net operating income")}-{g("Capital expenditure")}-{g("Management fee")}-{g("Fund administration and custody fees")}-{g("Interest")}'
    wf[g('Distributions to unitholders')] = f'=MAX(0,{g("Cash available for distribution")}*{F["Distribution policy (share of cash available)"]})'
    wf[g('Exit proceeds net of debt')] = f'=IF(Assets!{c}$5=1,{S(c, "Exit value net of disposal costs")}-$C${FR["Debt drawn"]}+SUM($C{FR["Cash available for distribution"]}:{c}{FR["Cash available for distribution"]})-SUM($C{FR["Distributions to unitholders"]}:{c}{FR["Distributions to unitholders"]}),0)'
    wf[g('Cash flow to unitholders')] = f'=-{g("Equity raised")}+{g("Distributions to unitholders")}+{g("Exit proceeds net of debt")}'
    wf[g('Distribution yield on equity')] = f'=IF({t}<1,0,{g("Distributions to unitholders")}/$C${FR["Equity raised"]})'
    wf[g('NAV per unit (SAR)')] = f'={g("Net asset value")}/{g("Units in issue")}'
    wf[g('Loan to value')] = f'=IF({g("Gross asset value")}=0,0,$C${FR["Debt drawn"]}/{g("Gross asset value")})'
    wf[g('Interest cover (NOI / interest)')] = f'=IF({g("Interest")}=0,0,{g("Net operating income")}/{g("Interest")})'
    for ln in fl:
        wf[g(ln)].number_format = PCT if ln in ('Distribution yield on equity', 'Loan to value') else X if ln.startswith('Interest cover') else '#,##0.00' if ln.startswith('NAV per unit') else M
wf.column_dimensions['A'].width = 40
for y in YRS: wf.column_dimensions[col(y)].width = 15
fr = lambda ln, y=None: f"Fund!{col(y)}{FR[ln]}" if y else f"Fund!$C${FR[ln]}:$J${FR[ln]}"

wd = wb.create_sheet('Dashboard', 0)
wd['A1'] = 'ADEL INCOME FUND I – DASHBOARD (SAR)'; wd['A1'].font = H
wd['A2'] = 'Version'; wd['B2'] = 'Actual 2025 and reforecast' if A else 'Budget 2025'
D = [
    ('Fund IRR (net, to unitholders)', f'=IRR({fr("Cash flow to unitholders")})', PCT),
    ('Equity multiple', f'=SUMIF({fr("Cash flow to unitholders")},">0")/-SUMIF({fr("Cash flow to unitholders")},"<0")', X),
    ('Equity raised (SAR)', f'={fr("Equity raised", 2024)}', M),
    ('Net operating income 2025 (SAR)', f'={fr("Net operating income", 2025)}', M),
    ('Distributions to unitholders 2025 (SAR)', f'={fr("Distributions to unitholders", 2025)}', M),
    ('Distribution yield 2025', f'={fr("Distribution yield on equity", 2025)}', PCT),
    ('Net asset value 31 Dec 2025 (SAR)', f'={fr("Net asset value", 2025)}', M),
    ('NAV per unit 31 Dec 2025 (SAR)', f'={fr("NAV per unit (SAR)", 2025)}', '#,##0.00'),
    ('Loan to value 31 Dec 2025', f'={fr("Loan to value", 2025)}', PCT),
    ('Interest cover 2025', f'={fr("Interest cover (NOI / interest)", 2025)}', X),
]
for i, (lab, f, fmt) in enumerate(D):
    wd.cell(4 + i, 1, lab); c = wd.cell(4 + i, 2, f); c.number_format = fmt
wd.column_dimensions['A'].width = 44; wd.column_dimensions['B'].width = 20
wb.save(out)
