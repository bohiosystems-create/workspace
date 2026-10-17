# round 19: financially complex models, built from scratch
#  capstack - US office recap, quarterly: senior on a SOFR forward curve with amortisation and a DSCR cash trap that sweeps cash to
#             the senior, mezz with a PIK toggle, accruing preferred equity, common equity cures shortfalls; exit pays the stack in order
#  fund     - European value-add fund on dated quarters: five deals, subscription line with a call lag and its interest, management fee
#             on commitments then on invested cost, European waterfall (return of capital, 8% pref, 100% GP catch-up, 80/20), XIRR
#  tenancy  - London office, tenant by tenant: lease expiries as dates, renewal probability, voids, rent free and TI/LC weighted by
#             renewal, ERV growth, DCF at a discount rate with an exit on year 11 NOI
#  hedge    - US investor buying a London asset in GBP: GBP loan on SONIA, GBP/USD path and forward points, a rolled annual hedge on
#             part of the equity, GBP and USD (hedged / unhedged) IRRs
#  aftertax - US multifamily after tax: building straight-line over 27.5 years, cost segregation with bonus depreciation, interest-only
#             then amortising loan (PMT/IPMT), losses carried forward, recapture at 25% / ordinary and capital gains at sale
import openpyxl,os,math,datetime
from openpyxl.utils import get_column_letter as L
OUT='/tmp/claude-0/-home-user-workspace/c7e8f42c-9ed0-5ba7-9852-d8a131185935/scratchpad/m19/';os.makedirs(OUT+'calc',exist_ok=True)
P='0.00%';N='#,##0';N2='#,##0.00';X='0.00"x"';D='dd-mmm-yyyy'
def ref(ws,col,r):return (f"'{ws.title}'" if ' ' in ws.title else ws.title)+f"!${L(col)}${r}"
def inputs(ws,rows,col=2,start=3):
    I={};r=start
    for k,lab,v,f in rows:
        ws.cell(r,1,lab);c=ws.cell(r,col,v);c.number_format=f;I[k]=ref(ws,col,r);r+=1
    return I
def table(ws,rows,start=3,col=2):
    for i,(lab,f,fm) in enumerate(rows):ws.cell(start+i,1,lab);c=ws.cell(start+i,col,f);c.number_format=fm
def timeline(ws,labels,T,fill,first_row=3):
    R={l:first_row+i for i,l in enumerate(labels)}
    for l,r in R.items():ws.cell(r,1,l)
    for t in range(T+1):
        c=L(t+2);p=L(t+1)
        for l,f in fill(t,c,p,R).items():
            v=ws[f'{c}{R[l]}'];v.value=f if not isinstance(f,str) or f.startswith('=') else '='+f
    return R

# ---------- 19A capital stack with cash trap ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='LAKESHORE PLAZA – RECAPITALISATION, CAPITAL STACK AND CASH TRAP (USD)'
I=inputs(ws,[('price','Purchase price ($)',250000000,N),('close','Closing costs',0.02,P),('noi','NOI, year 1 ($)',15500000,N),('g','NOI growth a year',0.03,P),('capex','Capital reserve a quarter ($)',500000,N),
 ('hold','Hold (quarters)',28,'0'),('exit','Exit cap rate',0.0625,P),('sell','Selling costs',0.015,P),
 ('sltv','Senior loan to value',0.55,P),('marg','Senior margin over SOFR',0.0225,P),('amort','Senior amortisation a year (of original)',0.01,P),('trig','Cash trap: senior DSCR below',1.25,X),
 ('mltv','Mezzanine (% of price)',0.15,P),('mrate','Mezzanine coupon',0.11,P),('pltv','Preferred equity (% of price)',0.10,P),('prate','Preferred return (accrues)',0.12,P)])
cv=wb.create_sheet('Curves');cv['A1']='Quarter';cv['B1']='SOFR forward'
for q in range(1,41):cv.cell(q+1,1,q);c=cv.cell(q+1,2,round(0.043-0.009*(1-math.exp(-q/8)),6));c.number_format=P
cf=wb.create_sheet('Cash Flow');T=40
labs=['Quarter','SOFR','NOI','Capital reserve','CFADS','Senior opening','Senior interest','Senior amortisation','Senior debt service','Senior DSCR','Cash after senior','Equity cure',
 'Mezz opening','Mezz interest due','Mezz interest paid','Mezz PIK','Cash after mezz','Pref opening','Pref accrual','Pref paid','Residual cash','Cash trap on','Cash swept','Distributions',
 'Senior closing','Mezz closing','Pref closing','Gross sale','Net sale','Common at exit','Common cash flow','Senior cash flow','Mezz cash flow','Pref cash flow']
def f19a(t,c,p,R):
    first=t==0;on=f"AND({c}$3>=1,{c}$3<={I['hold']})";ex=f"({c}$3={I['hold']})"
    S0=f"{I['price']}*{I['sltv']}";M0=f"{I['price']}*{I['mltv']}";P0=f"{I['price']}*{I['pltv']}"
    d={'Quarter':t,'SOFR':'0' if first else f"INDEX(Curves!$B$2:$B$41,{c}$3)",
     'NOI':f"IF({on},{I['noi']}/4*(1+{I['g']})^(({c}$3-1)/4),0)",'Capital reserve':f"IF({on},{I['capex']},0)",'CFADS':f"{c}{R['NOI']}-{c}{R['Capital reserve']}",
     'Senior opening':'0' if first else f"{p}{R['Senior closing']}",'Senior interest':f"{c}{R['Senior opening']}*({c}{R['SOFR']}+{I['marg']})/4",
     'Senior amortisation':f"IF({on},MIN({c}{R['Senior opening']},{S0}*{I['amort']}/4),0)",'Senior debt service':f"{c}{R['Senior interest']}+{c}{R['Senior amortisation']}",
     'Senior DSCR':f"IF({c}{R['Senior debt service']}>0,{c}{R['CFADS']}/{c}{R['Senior debt service']},99)",'Cash after senior':f"{c}{R['CFADS']}-{c}{R['Senior debt service']}",
     'Equity cure':f"MAX(0,-{c}{R['Cash after senior']})",'Mezz opening':'0' if first else f"{p}{R['Mezz closing']}",'Mezz interest due':f"{c}{R['Mezz opening']}*{I['mrate']}/4",
     'Mezz interest paid':f"MIN(MAX(0,{c}{R['Cash after senior']}),{c}{R['Mezz interest due']})",'Mezz PIK':f"{c}{R['Mezz interest due']}-{c}{R['Mezz interest paid']}",
     'Cash after mezz':f"MAX(0,{c}{R['Cash after senior']})-{c}{R['Mezz interest paid']}",'Pref opening':'0' if first else f"{p}{R['Pref closing']}",
     'Pref accrual':f"{c}{R['Pref opening']}*{I['prate']}/4",'Pref paid':f"MIN({c}{R['Cash after mezz']},{c}{R['Pref opening']}+{c}{R['Pref accrual']})",
     'Residual cash':f"{c}{R['Cash after mezz']}-{c}{R['Pref paid']}",'Cash trap on':f"IF(AND({on},{c}{R['Senior DSCR']}<{I['trig']}),1,0)",
     'Cash swept':f"MIN({c}{R['Cash trap on']}*{c}{R['Residual cash']},{c}{R['Senior opening']}-{c}{R['Senior amortisation']})",'Distributions':f"{c}{R['Residual cash']}-{c}{R['Cash swept']}",
     'Senior closing':S0 if first else f"IF({c}$3>={I['hold']},0,{c}{R['Senior opening']}-{c}{R['Senior amortisation']}-{c}{R['Cash swept']})",
     'Mezz closing':M0 if first else f"IF({c}$3>={I['hold']},0,{c}{R['Mezz opening']}+{c}{R['Mezz PIK']})",
     'Pref closing':P0 if first else f"IF({c}$3>={I['hold']},0,{c}{R['Pref opening']}+{c}{R['Pref accrual']}-{c}{R['Pref paid']})",
     'Gross sale':f"IF({ex},{c}{R['NOI']}*4*(1+{I['g']})/{I['exit']},0)",'Net sale':f"{c}{R['Gross sale']}*(1-{I['sell']})",
     'Common at exit':f"IF({ex},{c}{R['Net sale']}-({c}{R['Senior opening']}-{c}{R['Senior amortisation']}-{c}{R['Cash swept']})-({c}{R['Mezz opening']}+{c}{R['Mezz PIK']})-({c}{R['Pref opening']}+{c}{R['Pref accrual']}-{c}{R['Pref paid']}),0)",
     'Common cash flow':f"-({I['price']}*(1+{I['close']})-{S0}-{M0}-{P0})" if first else f"{c}{R['Distributions']}-{c}{R['Equity cure']}+{c}{R['Common at exit']}",
     'Senior cash flow':f"-{S0}" if first else f"{c}{R['Senior debt service']}+{c}{R['Cash swept']}+IF({ex},{c}{R['Senior opening']}-{c}{R['Senior amortisation']}-{c}{R['Cash swept']},0)",
     'Mezz cash flow':f"-{M0}" if first else f"{c}{R['Mezz interest paid']}+IF({ex},{c}{R['Mezz opening']}+{c}{R['Mezz PIK']},0)",
     'Pref cash flow':f"-{P0}" if first else f"{c}{R['Pref paid']}+IF({ex},{c}{R['Pref opening']}+{c}{R['Pref accrual']}-{c}{R['Pref paid']},0)"}
    return d
R=timeline(cf,labs,T,f19a);E=L(T+2);row=lambda n:f"'Cash Flow'!B{R[n]}:{E}{R[n]}";rowc=lambda n:f"'Cash Flow'!C{R[n]}:{E}{R[n]}"
table(wb.create_sheet('Returns'),[('Common equity IRR',f"=(1+IRR({row('Common cash flow')},0.03))^4-1",P),('Common equity multiple',f"=SUMIF({row('Common cash flow')},\">0\")/-SUMIF({row('Common cash flow')},\"<0\")",X),
 ('Mezzanine IRR',f"=(1+IRR({row('Mezz cash flow')},0.02))^4-1",P),('Preferred equity IRR',f"=(1+IRR({row('Pref cash flow')},0.02))^4-1",P),('Senior IRR',f"=(1+IRR({row('Senior cash flow')},0.01))^4-1",P),
 ('Minimum senior DSCR',f"=MIN({rowc('Senior DSCR')})",X),('Quarters in cash trap',f"=SUM({row('Cash trap on')})",'0'),('Cash swept to senior ($)',f"=SUM({row('Cash swept')})",N),
 ('Peak mezzanine balance with PIK ($)',f"=MAX({row('Mezz closing')})",N),('Equity cures ($)',f"=SUM({row('Equity cure')})",N),('Net sale proceeds ($)',f"=SUM({row('Net sale')})",N)])
wb.save(OUT+'capstack_usd.xlsx')

# ---------- 19B fund with dated quarters, sub line and European waterfall ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Fund';ws['A1']='EUROPA VALUE-ADD FUND II – FUND MODEL (EUR)'
I=inputs(ws,[('commit','Commitments (€)',500000000,N),('start','First close',datetime.date(2025,3,31),D),('invend','Investment period ends',datetime.date(2028,3,31),D),
 ('fee1','Management fee on commitments (investment period)',0.015,P),('fee2','Management fee on invested cost (after)',0.0125,P),('exp','Fund expenses a year (% of commitments)',0.002,P),
 ('lag','Capital call lag on the subscription line (quarters)',2,'0'),('sub','Subscription line rate',0.055,P),('pref','Preferred return',0.08,P),('carry','Carried interest',0.2,P),('gp','GP commitment',0.02,P)])
dl=wb.create_sheet('Deals');dl.append(['Deal','Acquisition date','Equity (€)','Exit date','Gross multiple'])
deals=[('Milan offices',datetime.date(2025,6,30),95000000,datetime.date(2029,12,31),1.75),('Lyon logistics',datetime.date(2025,12,31),70000000,datetime.date(2029,6,30),1.9),
 ('Madrid resi',datetime.date(2026,9,30),85000000,datetime.date(2030,9,30),1.6),('Warsaw retail park',datetime.date(2027,3,31),60000000,datetime.date(2030,3,31),1.45),('Lisbon hotel',datetime.date(2027,12,31),90000000,datetime.date(2031,6,30),2.1)]
for i,d in enumerate(deals,start=2):
    for j,v in enumerate(d,start=1):c=dl.cell(i,j,v);c.number_format=D if isinstance(v,datetime.date) else (N if j==3 else ('0.00"x"' if j==5 else 'General'))
cf=wb.create_sheet('Cash Flow');T=28
labs=['Quarter','Date','Deal equity','Deal proceeds','Invested cost','Management fee','Fund expenses','Sub line drawn','Capital called','Sub line interest','Sub line balance',
 'Contributions','Distributable','LP contributions','ROC account','Return of capital','Pref account','Pref paid','Catch-up account','Catch-up to GP','Split 80/20 LP','Split 80/20 GP',
 'LP distributions','GP carry','LP net cash flow','Gross deal cash flow']
def f19b(t,c,p,R):
    first=t==0;DA="Deals!$B$2:$B$6";EQ="Deals!$C$2:$C$6";EX="Deals!$D$2:$D$6";MU="Deals!$E$2:$E$6"
    prevd=f"{p}{R['Date']}" if not first else f"({c}{R['Date']}-1)"
    lagc=f"OFFSET({c}{R['Sub line drawn']},0,-{I['lag']})"
    d={'Quarter':t,'Date':f"EOMONTH({I['start']},3*{c}$3)",
     'Deal equity':f"SUMPRODUCT(({DA}>{prevd})*({DA}<={c}{R['Date']})*{EQ})",'Deal proceeds':f"SUMPRODUCT(({EX}>{prevd})*({EX}<={c}{R['Date']})*{EQ}*{MU})",
     'Invested cost':f"SUMPRODUCT(({DA}<={c}{R['Date']})*({EX}>{c}{R['Date']})*{EQ})",
     'Management fee':f"IF({c}{R['Date']}<={I['invend']},{I['commit']}*{I['fee1']},{c}{R['Invested cost']}*{I['fee2']})/4",'Fund expenses':f"{I['commit']}*{I['exp']}/4",
     'Sub line drawn':f"{c}{R['Deal equity']}",
     'Capital called':f"IF({c}$3>={I['lag']},{lagc},0)+{c}{R['Management fee']}+{c}{R['Fund expenses']}+{c}{R['Sub line interest']}",
     'Sub line interest':'0' if first else f"{p}{R['Sub line balance']}*{I['sub']}/4",
     'Sub line balance':f"{c}{R['Sub line drawn']}-IF({c}$3>={I['lag']},{lagc},0)" if first else f"{p}{R['Sub line balance']}+{c}{R['Sub line drawn']}-IF({c}$3>={I['lag']},{lagc},0)",
     'Contributions':f"{c}{R['Capital called']}",'Distributable':f"{c}{R['Deal proceeds']}",'LP contributions':f"{c}{R['Contributions']}",
     'ROC account':f"{c}{R['Contributions']}" if first else f"{p}{R['ROC account']}-{p}{R['Return of capital']}+{c}{R['Contributions']}",
     'Return of capital':f"MIN({c}{R['Distributable']},{c}{R['ROC account']})",
     'Pref account':'0' if first else f"{p}{R['Pref account']}-{p}{R['Pref paid']}+({p}{R['ROC account']}-{p}{R['Return of capital']}+{p}{R['Pref account']}-{p}{R['Pref paid']})*((1+{I['pref']})^(({c}{R['Date']}-{p}{R['Date']})/365)-1)",
     'Pref paid':f"MIN({c}{R['Distributable']}-{c}{R['Return of capital']},{c}{R['Pref account']})",
     'Catch-up account':f"{c}{R['Pref paid']}*{I['carry']}/(1-{I['carry']})" if first else f"{p}{R['Catch-up account']}-{p}{R['Catch-up to GP']}+{c}{R['Pref paid']}*{I['carry']}/(1-{I['carry']})",
     'Catch-up to GP':f"MIN({c}{R['Distributable']}-{c}{R['Return of capital']}-{c}{R['Pref paid']},{c}{R['Catch-up account']})",
     'Split 80/20 LP':f"({c}{R['Distributable']}-{c}{R['Return of capital']}-{c}{R['Pref paid']}-{c}{R['Catch-up to GP']})*(1-{I['carry']})",
     'Split 80/20 GP':f"({c}{R['Distributable']}-{c}{R['Return of capital']}-{c}{R['Pref paid']}-{c}{R['Catch-up to GP']})*{I['carry']}",
     'LP distributions':f"{c}{R['Return of capital']}+{c}{R['Pref paid']}+{c}{R['Split 80/20 LP']}",'GP carry':f"{c}{R['Catch-up to GP']}+{c}{R['Split 80/20 GP']}",
     'LP net cash flow':f"{c}{R['LP distributions']}-{c}{R['LP contributions']}",'Gross deal cash flow':f"{c}{R['Deal proceeds']}-{c}{R['Deal equity']}"}
    return d
R=timeline(cf,labs,T,f19b);E=L(T+2);row=lambda n:f"'Cash Flow'!B{R[n]}:{E}{R[n]}"
for t in range(T+1):cf[f'{L(t+2)}{R["Date"]}'].number_format=D
table(wb.create_sheet('Returns'),[('Gross IRR (deals)',f"=XIRR({row('Gross deal cash flow')},{row('Date')})",P),('LP net IRR',f"=XIRR({row('LP net cash flow')},{row('Date')})",P),
 ('LP TVPI',f"=SUM({row('LP distributions')})/SUM({row('LP contributions')})",X),('GP carried interest (€)',f"=SUM({row('GP carry')})",N),
 ('Management fees (€)',f"=SUM({row('Management fee')})",N),('Sub line interest (€)',f"=SUM({row('Sub line interest')})",N),('Peak sub line balance (€)',f"=MAX({row('Sub line balance')})",N),
 ('Total called (€)',f"=SUM({row('Contributions')})",N),('Carry as share of profit',f"=B6/(SUM({row('Deal proceeds')})-B10)",P)])
wb.save(OUT+'fund_eur.xlsx')

# ---------- 19C tenant-by-tenant office DCF ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='70 FENCHURCH AVENUE – MULTI-LET OFFICE DCF (GBP)'
I=inputs(ws,[('val','Valuation date',datetime.date(2026,1,1),D),('price','Purchase price (£)',245000000,N),('costs','Purchase costs',0.068,P),('ervg','ERV growth a year',0.025,P),
 ('void','Void on a new letting (months)',12,'0'),('rfn','Rent free, new letting (months)',18,'0'),('rfr','Rent free, renewal (months)',6,'0'),('tin','Fit-out contribution, new (£/sq ft)',60,N),
 ('tir','Fit-out contribution, renewal (£/sq ft)',20,N),('lc','Letting fees (% of first-year rent)',0.15,P),('nonrec','Non-recoverable costs (% of rent)',0.04,P),('disc','Discount rate',0.075,P),('exit','Exit yield',0.0525,P),('sell','Sale costs',0.018,P)])
tn=wb.create_sheet('Tenancy');tn.append(['Floor','Tenant','Area (sq ft)','Passing rent (£/sq ft)','Lease expiry','ERV (£/sq ft)','Renewal probability'])
import random;random.seed(19)
ten=[('Ground','Fenchurch Retail',12000,72.0,datetime.date(2031,6,24),70,0.6),('1','Halden Insurance',21000,61.5,datetime.date(2027,3,24),68,0.55),('2','Halden Insurance',21000,61.5,datetime.date(2027,3,24),68,0.55),
 ('3','Marlow Legal',21000,64.0,datetime.date(2029,9,28),68,0.7),('4','Vacant',21000,0,datetime.date(2026,1,1),68,0),('5','Coral Re',20500,66.0,datetime.date(2032,12,24),69,0.65),
 ('6','Tessel Bank',20500,63.0,datetime.date(2028,6,23),69,0.5),('7','Tessel Bank',20500,63.0,datetime.date(2028,6,23),69,0.5),('8','Arkwright Partners',19000,70.0,datetime.date(2034,3,24),72,0.6),
 ('9','Arkwright Partners',19000,70.0,datetime.date(2034,3,24),72,0.6),('10','Nimbus Capital',17000,78.0,datetime.date(2030,12,24),76,0.45),('11','Nimbus Capital',17000,78.0,datetime.date(2030,12,24),76,0.45)]
for i,t in enumerate(ten,start=2):
    for j,v in enumerate(t,start=1):c=tn.cell(i,j,v);c.number_format=D if isinstance(v,datetime.date) else (P if j==7 else (N2 if j in (4,6) else 'General'))
rs=wb.create_sheet('Rent');Y=11;rs['A2']='Year';rs['B2']='Year start'
for y in range(1,Y+1):rs.cell(2,y+2,y);c=rs.cell(3,y+2,f"=EDATE({I['val']},12*{y-1})");c.number_format=D
rs['A3']='Year start date';rs['A4']='Tenant rows: rent received (probability-weighted)';n=len(ten)
for k in range(n):
    r=5+k;tr=k+2;rs.cell(r,1,f"=Tenancy!B{tr}&\" · \"&Tenancy!A{tr}")
    for y in range(1,Y+1):
        c=L(y+2);ys=f"{c}$3";ye=f"EDATE({c}$3,12)"
        exp=f"Tenancy!$E${tr}";area=f"Tenancy!$C${tr}";pas=f"Tenancy!$D${tr}";erv=f"Tenancy!$F${tr}*(1+{I['ervg']})^({c}$2-1)";p_=f"Tenancy!$G${tr}"
        before=f"MAX(0,MIN(1,YEARFRAC({ys},MIN({exp},{ye}),1)*({exp}>{ys})))"
        # after expiry in this year: renewal (rent free rfr) with p, new letting after void+rfn with 1-p; months from expiry to year end
        mo=f"MAX(0,MIN(12,DATEDIF(MAX({exp},{ys}),{ye},\"m\")))"
        aft=f"IF({exp}<{ye},{p_}*MAX(0,{mo}-MAX(0,{I['rfr']}-IF({exp}<{ys},DATEDIF({exp},{ys},\"m\"),0)))/12+(1-{p_})*MAX(0,{mo}-MAX(0,{I['void']}+{I['rfn']}-IF({exp}<{ys},DATEDIF({exp},{ys},\"m\"),0)))/12,0)"
        rs[f'{c}{r}']=f"={area}*({pas}*{before}+{erv}*{aft})"
    rs[f'{c}{r}'].number_format=N
cr=5+n;rs.cell(cr,1,'Rent received');rs.cell(cr+1,1,'Leasing costs (probability-weighted)');rs.cell(cr+2,1,'Non-recoverable costs');rs.cell(cr+3,1,'Net operating income');rs.cell(cr+4,1,'Exit value');rs.cell(cr+5,1,'Net cash flow')
for y in range(1,Y+1):
    c=L(y+2);rs[f'{c}{cr}']=f"=SUM({c}5:{c}{cr-1})"
    # leasing costs land in the year the lease expires: fit-out and letting fees, weighted by renewal probability
    rs[f'{c}{cr+1}']=f"=SUMPRODUCT((Tenancy!$E$2:$E${n+1}>={c}$3)*(Tenancy!$E$2:$E${n+1}<EDATE({c}$3,12))*Tenancy!$C$2:$C${n+1}*(Tenancy!$G$2:$G${n+1}*({I['tir']}+{I['lc']}*Tenancy!$F$2:$F${n+1}*(1+{I['ervg']})^({c}$2-1))+(1-Tenancy!$G$2:$G${n+1})*({I['tin']}+{I['lc']}*Tenancy!$F$2:$F${n+1}*(1+{I['ervg']})^({c}$2-1))))"
    rs[f'{c}{cr+2}']=f"={c}{cr}*{I['nonrec']}";rs[f'{c}{cr+3}']=f"={c}{cr}-{c}{cr+2}"
    rs[f'{c}{cr+4}']=f"=IF({c}$2=10,{L(y+3)}{cr+3}/{I['exit']}*(1-{I['sell']}),0)" if y<Y else "=0"
    rs[f'{c}{cr+5}']=f"=IF({c}$2<=10,{c}{cr+3}-{c}{cr+1}+{c}{cr+4},0)"
    for k in range(6):rs[f'{c}{cr+k}'].number_format=N
cfr=f"Rent!C{cr+5}:L{cr+5}"
table(wb.create_sheet('Valuation'),[('Net present value (£)',f"=NPV({I['disc']},{cfr})",N),('Value less price and costs (£)',f"=B3-{I['price']}*(1+{I['costs']})",N),
 ('Unlevered IRR',f"=IRR(Valuation!B12:L12)",P),('Year 1 net operating income (£)',f"=Rent!C{cr+3}",N),('Net initial yield',f"=B6/({I['price']}*(1+{I['costs']}))",P),
 ('WAULT to expiry (years)',f"=SUMPRODUCT(Tenancy!C2:C{n+1},Tenancy!D2:D{n+1},(Tenancy!E2:E{n+1}-{I['val']})/365.25)/SUMPRODUCT(Tenancy!C2:C{n+1},Tenancy!D2:D{n+1})",N2),('Exit value (£)',f"=SUM(Rent!C{cr+4}:L{cr+4})",N)])
vs=wb['Valuation'];vs['A12']='IRR cash flow';vs['B12']=f"=-{I['price']}*(1+{I['costs']})"
for y in range(1,11):vs.cell(12,y+2,f"=Rent!{L(y+2)}{cr+5}")
wb.save(OUT+'tenancy_gbp.xlsx')

# ---------- 19D cross-border hedged acquisition ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='ALDGATE YARD – US INVESTOR, GBP ASSET WITH FX HEDGE'
I=inputs(ws,[('price','Purchase price (£)',180000000,N),('costs','Purchase costs',0.068,P),('noi','NOI year 1 (£)',10200000,N),('g','NOI growth',0.025,P),('exit','Exit yield',0.0575,P),('hold','Hold (years)',7,'0'),
 ('ltv','GBP loan to value',0.5,P),('sonia','SONIA (flat)',0.04,P),('marg','Loan margin',0.018,P),('spot','GBP/USD spot',1.27,'0.0000'),('drift','GBP/USD drift a year',-0.01,P),('hedge','Hedge ratio (of GBP equity)',0.75,P),('hcost','Hedge execution cost (bps a year)',15,'0')])
cvs=wb.create_sheet('Curves');cvs['A1']='Year';cvs['B1']='Forward points (annual, % of spot)'
for y in range(1,11):cvs.cell(y+1,1,y);c=cvs.cell(y+1,2,round(-0.0045-0.0005*y,5));c.number_format='0.000%'
cf=wb.create_sheet('Cash Flow');T=10
labs=['Year','NOI (£)','Sale (£)','Purchase (£)','Loan (£)','Interest (£)','Repayment (£)','Equity cash flow (£)','GBP/USD spot','Equity (£) hedged notional','Forward rate locked','Hedge settlement ($)','Equity cash flow unhedged ($)','Equity cash flow hedged ($)']
def f19d(t,c,p,R):
    first=t==0;on=f"AND({c}$3>=1,{c}$3<={I['hold']})";ex=f"({c}$3={I['hold']})";LN=f"{I['price']}*{I['ltv']}"
    eq0=f"{I['price']}*(1+{I['costs']})-{LN}"
    d={'Year':t,'NOI (£)':f"IF({on},{I['noi']}*(1+{I['g']})^({c}$3-1),0)",'Sale (£)':f"IF({ex},{I['noi']}*(1+{I['g']})^{c}$3/{I['exit']},0)",
     'Purchase (£)':f"IF({c}$3=0,-{I['price']}*(1+{I['costs']}),0)",'Loan (£)':f"IF({c}$3=0,{LN},0)",'Interest (£)':f"IF({on},-{LN}*({I['sonia']}+{I['marg']}),0)",'Repayment (£)':f"IF({ex},-{LN},0)",
     'Equity cash flow (£)':f"SUM({c}{R['NOI (£)']}:{c}{R['Repayment (£)']})",'GBP/USD spot':f"{I['spot']}*(1+{I['drift']})^{c}$3",
     'Equity (£) hedged notional':f"IF({on},{eq0}*{I['hedge']},0)",'Forward rate locked':'0' if first else f"IF({on},{p}{R['GBP/USD spot']}*(1+INDEX(Curves!$B$2:$B$11,{c}$3)),0)",
     'Hedge settlement ($)':f"{c}{R['Equity (£) hedged notional']}*({c}{R['Forward rate locked']}-{c}{R['GBP/USD spot']})-{c}{R['Equity (£) hedged notional']}*{c}{R['GBP/USD spot']}*{I['hcost']}/10000",
     'Equity cash flow unhedged ($)':f"{c}{R['Equity cash flow (£)']}*{c}{R['GBP/USD spot']}",'Equity cash flow hedged ($)':f"{c}{R['Equity cash flow unhedged ($)']}+{c}{R['Hedge settlement ($)']}"}
    return d
R=timeline(cf,labs,T,f19d);E=L(T+2);row=lambda n:f"'Cash Flow'!B{R[n]}:{E}{R[n]}"
table(wb.create_sheet('Returns'),[('Levered IRR (GBP)',f"=IRR({row('Equity cash flow (£)')})",P),('Levered IRR (USD, unhedged)',f"=IRR({row('Equity cash flow unhedged ($)')})",P),('Levered IRR (USD, hedged)',f"=IRR({row('Equity cash flow hedged ($)')})",P),
 ('Equity multiple (USD, hedged)',f"=SUMIF({row('Equity cash flow hedged ($)')},\">0\")/-SUMIF({row('Equity cash flow hedged ($)')},\"<0\")",X),('Hedge result ($)',f"=SUM({row('Hedge settlement ($)')})",N),('Net profit (USD, hedged)',f"=SUM({row('Equity cash flow hedged ($)')})",N)])
wb.save(OUT+'hedge_gbp.xlsx')

# ---------- 19E after-tax US multifamily ----------
wb=openpyxl.Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='CEDAR HOLLOW APARTMENTS – AFTER-TAX INVESTMENT MODEL (USD)'
I=inputs(ws,[('price','Purchase price ($)',85000000,N),('land','Land share of price',0.2,P),('seg','Cost segregation (share of building, 5-year property)',0.2,P),('bonus','Bonus depreciation',0.6,P),
 ('noi','NOI year 1 ($)',4900000,N),('g','NOI growth',0.03,P),('ltv','Loan to value',0.65,P),('rate','Interest rate',0.061,P),('io','Interest-only (years)',3,'0'),('amyrs','Amortisation (years)',30,'0'),
 ('hold','Hold (years)',7,'0'),('exit','Exit cap rate',0.055,P),('sell','Selling costs',0.02,P),('ord','Ordinary income tax rate',0.37,P),('recap','Unrecaptured §1250 rate',0.25,P),('cg','Capital gains rate (incl. NIIT)',0.238,P)])
cf=wb.create_sheet('Cash Flow');T=10
labs=['Year','NOI','Loan opening','Interest','Principal','Loan closing','Debt service','Building depreciation','Cost seg depreciation','Taxable income','Loss carried in','Loss used','Loss carried out','Tax on operations','Pre-tax equity cash flow','Sale price','Sale costs','Tax at sale','After-tax equity cash flow']
def f19e(t,c,p,R):
    first=t==0;on=f"AND({c}$3>=1,{c}$3<={I['hold']})";ex=f"({c}$3={I['hold']})";LN=f"{I['price']}*{I['ltv']}";B=f"{I['price']}*(1-{I['land']})"
    pay=f"PMT({I['rate']},{I['amyrs']},-{LN})"
    d={'Year':t,'NOI':f"IF({on},{I['noi']}*(1+{I['g']})^({c}$3-1),0)",'Loan opening':'0' if first else f"{p}{R['Loan closing']}",
     'Interest':f"IF({on},{c}{R['Loan opening']}*{I['rate']},0)",'Principal':f"IF(AND({on},{c}$3>{I['io']}),{pay}-{c}{R['Interest']},0)",
     'Loan closing':LN if first else f"IF({c}$3>={I['hold']},0,{c}{R['Loan opening']}-{c}{R['Principal']})",'Debt service':f"{c}{R['Interest']}+{c}{R['Principal']}",
     'Building depreciation':f"IF({on},{B}*(1-{I['seg']})/27.5*IF({c}$3=1,11.5/12,1)*IF({ex},11.5/12,1),0)",
     'Cost seg depreciation':f"IF({on},IF({c}$3=1,{B}*{I['seg']}*{I['bonus']},0)+IF({c}$3<=5,{B}*{I['seg']}*(1-{I['bonus']})/5,0),0)",
     'Taxable income':f"{c}{R['NOI']}-{c}{R['Interest']}-{c}{R['Building depreciation']}-{c}{R['Cost seg depreciation']}",
     'Loss carried in':'0' if first else f"{p}{R['Loss carried out']}",'Loss used':f"IF({c}{R['Taxable income']}>0,MIN({c}{R['Taxable income']},{c}{R['Loss carried in']}),0)",
     'Loss carried out':f"{c}{R['Loss carried in']}-{c}{R['Loss used']}+MAX(0,-{c}{R['Taxable income']})",
     'Tax on operations':f"MAX(0,{c}{R['Taxable income']}-{c}{R['Loss used']})*{I['ord']}",
     'Pre-tax equity cash flow':f"-({I['price']}-{LN})" if first else f"{c}{R['NOI']}-{c}{R['Debt service']}+IF({ex},{c}{R['Sale price']}-{c}{R['Sale costs']}-{c}{R['Loan opening']}+{c}{R['Principal']},0)",
     'Sale price':f"IF({ex},{I['noi']}*(1+{I['g']})^{c}$3/{I['exit']},0)",'Sale costs':f"{c}{R['Sale price']}*{I['sell']}",
     # at sale: cost seg recaptured at ordinary rates, building depreciation at 25%, the rest of the gain at capital gains; losses still carried offset ordinary first
     'Tax at sale':f"IF({ex},MAX(0,MIN(SUM($C{R['Cost seg depreciation']}:{c}{R['Cost seg depreciation']}),{c}{R['Sale price']}-{c}{R['Sale costs']}-({I['price']}-SUM($C{R['Building depreciation']}:{c}{R['Building depreciation']})-SUM($C{R['Cost seg depreciation']}:{c}{R['Cost seg depreciation']})))-{c}{R['Loss carried out']})*{I['ord']}"
                   f"+MIN(SUM($C{R['Building depreciation']}:{c}{R['Building depreciation']}),MAX(0,{c}{R['Sale price']}-{c}{R['Sale costs']}-({I['price']}-SUM($C{R['Building depreciation']}:{c}{R['Building depreciation']}))))*{I['recap']}"
                   f"+MAX(0,{c}{R['Sale price']}-{c}{R['Sale costs']}-{I['price']})*{I['cg']}-MAX(0,{c}{R['Loss carried out']}-SUM($C{R['Cost seg depreciation']}:{c}{R['Cost seg depreciation']}))*{I['ord']},0)",
     'After-tax equity cash flow':f"{c}{R['Pre-tax equity cash flow']}-{c}{R['Tax on operations']}-{c}{R['Tax at sale']}"}
    return d
R=timeline(cf,labs,T,f19e);E=L(T+2);row=lambda n:f"'Cash Flow'!B{R[n]}:{E}{R[n]}"
table(wb.create_sheet('Returns'),[('Levered IRR (pre-tax)',f"=IRR({row('Pre-tax equity cash flow')})",P),('After-tax IRR',f"=IRR({row('After-tax equity cash flow')})",P),
 ('After-tax equity multiple',f"=SUMIF({row('After-tax equity cash flow')},\">0\")/-SUMIF({row('After-tax equity cash flow')},\"<0\")",X),('Tax at sale ($)',f"=SUM({row('Tax at sale')})",N),
 ('Depreciation taken ($)',f"=SUM({row('Building depreciation')})+SUM({row('Cost seg depreciation')})",N),('Net profit after tax ($)',f"=SUM({row('After-tax equity cash flow')})",N),('Year 1 tax shield ($)',f"=-MIN(0,'Cash Flow'!C{R['Taxable income']})*{I['ord']}",N)])
wb.save(OUT+'aftertax_usd.xlsx')
print('ok')
