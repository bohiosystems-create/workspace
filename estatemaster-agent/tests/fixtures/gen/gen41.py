# Round 41 fixture (synthetic): a hotel development and operation, with made-up figures, laid out the way hotel models often are.
#  r41_hotel_aed: keys, ADR and occupancy ramping up by year; rooms, F&B and other revenue; departmental costs, undistributed costs
#  and GOP; the operator's base fee (% of revenue) and incentive fee (% of GOP); an FF&E reserve; EBITDA; construction over two
#  years; an exit at a multiple of the next year's EBITDA; a bank loan with interest-only then amortisation; returns on a
#  "Returns" sheet in a matrix: rows "Unleveraged" and "Leveraged", columns "IRR", "MOIC" and "Profit". Annual, AED.
import openpyxl,sys
from openpyxl.utils import get_column_letter as L
out=sys.argv[1].rstrip('/')
wb=openpyxl.Workbook();R=wb.active;R.title='Returns';A=wb.create_sheet('Assumptions');P=wb.create_sheet('Operations');F=wb.create_sheet('Cash Flow')
Y=12   # years 0..Y; construction years 1-2, opening year 3, exit at the end of year 10 on year 11's EBITDA
A['B2']='MARINA BAY HOTEL · ASSUMPTIONS (AED)'
rows=[(4,'Keys',320,'0'),(5,'Land',95000000,'#,##0'),(6,'Construction cost per key',1150000,'#,##0'),(7,'Pre-opening and FF&E (total)',48000000,'#,##0'),
 (8,'Opening year',3,'0'),(9,'Stabilised ADR',780,'#,##0'),(10,'Stabilised occupancy',0.74,'0.0%'),(11,'ADR growth (a year)',0.03,'0.0%'),
 (12,'F&B revenue (% of rooms)',0.42,'0.0%'),(13,'Other revenue (% of rooms)',0.06,'0.0%'),(14,'Departmental costs (% of revenue)',0.38,'0.0%'),
 (15,'Undistributed costs (% of revenue)',0.17,'0.0%'),(16,'Base management fee (% of revenue)',0.03,'0.0%'),(17,'Incentive fee (% of GOP)',0.08,'0.0%'),
 (18,'FF&E reserve (% of revenue)',0.04,'0.0%'),(19,'Exit year',10,'0'),(20,'Exit EBITDA multiple',13.5,'0.0"x"'),(21,'Selling costs',0.02,'0.0%'),
 (22,'Loan to cost',0.55,'0.0%'),(23,'Interest rate',0.065,'0.00%'),(24,'Amortisation (years after opening)',20,'0')]
for r,l,v,f in rows:A.cell(r,2,l);A.cell(r,3,v).number_format=f
A['B26']='Ramp-up (share of stabilised occupancy and ADR, by operating year)';A['B27']='Operating year';A['B28']='Occupancy ramp';A['B29']='ADR ramp'
for i,(o,a) in enumerate([(0.70,0.88),(0.86,0.95),(1.0,1.0)]):c=L(3+i);A[f'{c}27']=i+1;A[f'{c}28']=o;A[f'{c}29']=a
P['B2']='OPERATIONS (AED)';P['B4']='Year';P['B5']='Operating year'
lab={7:'Occupancy',8:'ADR',9:'RevPAR',10:'Rooms revenue',11:'F&B revenue',12:'Other revenue',13:'Total revenue',14:'Departmental costs',15:'Undistributed costs',
 16:'GOP',17:'Base management fee',18:'Incentive fee',19:'FF&E reserve',20:'EBITDA'}
for r,l in lab.items():P.cell(r,2,l)
for y in range(Y+1):
    c=L(3+y);P[f'{c}4']=y;P[f'{c}5']=f'=MAX(0,{c}4-ROUND(Assumptions!$C$8,0)+1)';op=f'{c}5'
    ramp=lambda row:f'IF({op}=0,0,IF({op}>=3,1,INDEX(Assumptions!$C${row}:$E${row},1,{op})))'
    P[f'{c}7']=f'=Assumptions!$C$10*{ramp(28)}';P[f'{c}8']=f'=IF({op}=0,0,Assumptions!$C$9*{ramp(29)}*(1+Assumptions!$C$11)^({c}4-ROUND(Assumptions!$C$8,0)))'
    P[f'{c}9']=f'={c}7*{c}8';P[f'{c}10']=f'={c}9*Assumptions!$C$4*365';P[f'{c}11']=f'={c}10*Assumptions!$C$12';P[f'{c}12']=f'={c}10*Assumptions!$C$13'
    P[f'{c}13']=f'=SUM({c}10:{c}12)';P[f'{c}14']=f'=-{c}13*Assumptions!$C$14';P[f'{c}15']=f'=-{c}13*Assumptions!$C$15';P[f'{c}16']=f'=SUM({c}13:{c}15)'
    P[f'{c}17']=f'=-{c}13*Assumptions!$C$16';P[f'{c}18']=f'=-MAX(0,{c}16)*Assumptions!$C$17';P[f'{c}19']=f'=-{c}13*Assumptions!$C$18';P[f'{c}20']=f'=SUM({c}16:{c}19)'
    for r in (7,):P[f'{c}{r}'].number_format='0.0%'
E=L(3+Y)
F['B2']='CASH FLOW (AED)';F['B4']='Year'
fl={6:'Land',7:'Construction',8:'Pre-opening and FF&E',9:'EBITDA',10:'Sale (net)',11:'Unleveraged cash flow',13:'Loan draw',14:'Interest',15:'Repayment',16:'Loan balance',18:'Leveraged cash flow'}
for r,l in fl.items():F.cell(r,2,l)
for y in range(Y+1):
    c=L(3+y);p=L(2+y);n=L(4+y);yy=f'{c}$4';ex='ROUND(Assumptions!$C$19,0)';ob='ROUND(Assumptions!$C$8,0)'
    F[f'{c}4']=y;F[f'{c}6']=f'=IF({yy}=0,-Assumptions!$C$5,0)'
    F[f'{c}7']=f'=IF(AND({yy}>=1,{yy}<{ob}),-Assumptions!$C$4*Assumptions!$C$6/({ob}-1),0)'
    F[f'{c}8']=f'=IF({yy}={ob}-1,-Assumptions!$C$7,0)'
    F[f'{c}9']=f"=IF({yy}<={ex},Operations!{c}20,0)"
    F[f'{c}10']=f"=IF({yy}={ex},Operations!{n}20*Assumptions!$C$20*(1-Assumptions!$C$21),0)" if y<Y else '=0'
    F[f'{c}11']=f'=SUM({c}6:{c}10)'
    F[f'{c}13']=f'=-SUM({c}6:{c}8)*Assumptions!$C$22'
    F[f'{c}14']='=0' if y==0 else f'=-{p}16*Assumptions!$C$23'
    am=f'IF(AND({yy}>={ob},{yy}<{ex}),{p}16/MAX(1,Assumptions!$C$24-({yy}-{ob})),0)'
    F[f'{c}15']='=0' if y==0 else f'=-IF({yy}={ex},{p}16,IF({yy}>{ex},0,{am}))'
    F[f'{c}16']=f'={c}13' if y==0 else f'={p}16+{c}13+{c}15'
    F[f'{c}18']=f'={c}11+{c}13+{c}14+{c}15'
R['B2']='MARINA BAY HOTEL · RETURNS (AED)';R['C5']='IRR';R['D5']='MOIC';R['E5']='Profit'
R['B6']='Unleveraged';R['C6']=f"=IRR('Cash Flow'!C11:{E}11)";R['D6']=f"=SUMIF('Cash Flow'!C11:{E}11,\">0\")/-SUMIF('Cash Flow'!C11:{E}11,\"<0\")";R['E6']=f"=SUM('Cash Flow'!C11:{E}11)"
R['B7']='Leveraged';R['C7']=f"=IRR('Cash Flow'!C18:{E}18)";R['D7']=f"=SUMIF('Cash Flow'!C18:{E}18,\">0\")/-SUMIF('Cash Flow'!C18:{E}18,\"<0\")";R['E7']=f"=SUM('Cash Flow'!C18:{E}18)"
for r in (6,7):R[f'C{r}'].number_format='0.00%';R[f'D{r}'].number_format='0.00"x"';R[f'E{r}'].number_format='#,##0'
R['B9']='Stabilised RevPAR (year 5)';R['C9']='=Operations!H9';R['B10']='Yield on cost (year 5 EBITDA)';R['C10']=f"=Operations!H20/-SUM('Cash Flow'!C6:{E}8)";R['C10'].number_format='0.00%'
R['B11']='Total development cost';R['C11']=f"=-SUM('Cash Flow'!C6:{E}8)";R['C11'].number_format='#,##0'
wb.save(f'{out}/r41_hotel_aed.xlsx')
