# Round 37 fixture (synthetic): a single-asset acquisition laid out the way US acquisition templates often are, with made-up figures.
#  r37_acq_monthly_usd: the inputs (labels in C, values in D) on the same sheet as a monthly DCF (labels in I, month 0 in J); the start
#  typed as a month name and a year; a fiscal-year row; the hold period in years masking the cash flow; the sale at the forward twelve
#  months' NOI over the exit cap, in the month whose name and calendar year match the exit; taxes and insurance paid in a named month;
#  typed leasing costs on the timeline; an unlevered and a levered block on one sheet, each with a row labelled just "IRR" (the
#  unlevered cash-flow row keeps the template's typo "Unlvered"); a loan from PMT with its repayment read from an amortisation sheet;
#  and an annual returns sheet whose "Unlevered Return Summary" and "Levered Return Summary" both hold an "IRR" row.
#  Month counts and the hold are rounded where the timeline reads them, so a fractional input still gives a sale.
import openpyxl,sys
from openpyxl.utils import get_column_letter as L
out=sys.argv[1].rstrip('/')
M='Inputs & Monthly CF';Q="'Inputs & Monthly CF'!"
MONTHS=['January','February','March','April','May','June','July','August','September','October','November','December']
wb=openpyxl.Workbook();s=wb.active;s.title=M
s['B2']='INPUTS AND MONTHLY CASH FLOW';s['C10']='Inputs';s['H10']='Monthly DCF';s['C12']='Property Inputs';s['H12']='Annual Growth Rate';s['I12']='Unlevered DCF'
for r,m in enumerate(MONTHS,87):s.cell(r,4,m);s.cell(r,5,r-86)
inp=[(19,'Analysis Start','=DATE(D21,MATCH(D20,$D$87:$D$98,0),1)','mmm-yy'),(20,'Start Month','January',None),(21,'Start Year',2026,'0'),
 (22,'Property Name','Lakeside Logistics Park',None),(23,'City, State','Columbus, OH',None),(25,'Property SF',145000,'#,##0'),
 (26,'Purchase Price',18500000,'#,##0'),(27,'Going in Cap Rate','=SUM($K$28:$V$28)/D26','0.00%'),(28,'$/PSF','=D26/D25','#,##0.00'),
 (29,'Closing costs',0.02,'0.00%'),(30,'All-in Purchase Cost ','=D26*(1+$D$29)','#,##0'),(31,'Hold Period',5,'0'),
 (32,'Exit Month','=D20',None),(33,'Exit Year','=YEAR(D19)+ROUND(D31,0)','0'),
 (34,'Gross Potential Rent - Year 1',1480000,'#,##0'),(35,'Reimbursable Expenses Year 1',395000,'#,##0'),(36,'Misc. Income',22000,'#,##0'),
 (37,'Gross Potential Income Year 1','=SUM(D34:D36)','#,##0'),(38,'Expected Vacancy Loss (% of GPI)',0.05,'0.00%'),
 (39,'Tax in Year 1',214000,'#,##0'),(40,'Tax Payable in Month','October',None),(41,'Annual Insurance premium',51000,'#,##0'),
 (42,'Insurance Payable in Month','March',None),(43,'CAM (Annual Expense)',168000,'#,##0'),(44,'Property Mgmt Fee ',0.03,'0.00%'),
 (45,'Ave Annual Cap ex & Misc',29000,'#,##0'),(51,'Going Out Cap Rate',0.0725,'0.00%'),(52,'Selling Expenses',0.015,'0.00%'),
 (55,'Debt Inputs',None,None),(56,'LTV',0.65,'0.00%'),(57,'Int Rate',0.0625,'0.00%'),(58,'Amort. Period (Max 30 years)',30,'0'),(59,'Points',1,'0.00'),
 (60,'Pmt','=-PMT(D57/12,D58*12,D26*D56)','#,##0'),(61,'Loan Proceeds','=D26*D56*(1-D59/100)','#,##0'),
 (62,'Loan Repayment','=INDEX(AmortSched!$I$14:$I$373,12*ROUND(D31,0))','#,##0')]
for r,l,v,f in inp:
    s.cell(r,3,l)
    if v is not None:s.cell(r,4,v)
    if f:s.cell(r,4).number_format=f
for r,g in ((20,0.025),(24,0.03),(25,0.03),(26,0.03),(31,0.02)):s.cell(r,8,g).number_format='0.00%'
lab={13:'Date ',14:'Calendar Year',15:'Fiscal Year',16:'Month',18:'Fiscal Month',19:'Purchase Price',20:'Gross Potential Income',21:'Expected Vacancy',
 22:'Additional Vacancy - (Input as Negative)',23:'Effective Gross Income',24:'Tax',25:'Insurance',26:'CAM',27:'Mgmt Fee',28:'NOI',
 29:'Leasing Commissions - (Input as Negative)',30:'Tenant Improvements - (Input as Negative)',31:'Cap Ex & Misc',33:'CFO',35:'Sale',36:'Selling Expenses',
 37:'Proceeds',39:'Unlvered Cash Flow',40:'IRR',41:'Equity Multiple',42:'Net Gain (loss)',46:'Levered DCF',49:'Purchase Price',50:'Loan Proceeds',
 59:'NOI',64:'CFO',65:'Debt Payments',66:'CFAF',68:'Sale',69:'Selling Expenses',70:'Loan Repayment',71:'Proceeds',73:'Levered Cash Flow',
 74:'IRR',75:'Equity Multiple',76:'Net Gain (loss)',78:'Year',79:'DSCR'}
for r,l in lab.items():s.cell(r,9,l)
N=132;H='ROUND($D$31,0)'
for k in range(N+1):
    c=L(10+k);p=L(9+k);g=lambda gr:f'(1+$H${gr})^({c}15-1)'
    if k==0:
        s['J13']='=$D$19';s['J14']='=$D$21';s['J15']=0;s['J16']='=$D$20';s['J18']=0;s['J19']='=-$D$30';s['J39']='=J19';s['J49']='=J19';s['J50']='=$D$61';s['J73']='=J49+J50'
        s['J13'].number_format='mmm-yy';continue
    s[f'{c}13']=f'=EDATE({p}13,1)';s[f'{c}13'].number_format='mmm-yy';s[f'{c}14']=f'=YEAR({c}13)';s[f'{c}15']=f'=ROUNDUP({c}18/12,0)'
    s[f'{c}16']=f'=INDEX($D$87:$D$98,MONTH({c}13))';s[f'{c}18']=f'={p}18+1'
    s[f'{c}20']=f'=$D$37/12*{g(20)}';s[f'{c}21']=f'=-{c}20*$D$38';s[f'{c}22']=-8000 if k<=6 else 0;s[f'{c}23']=f'=SUM({c}20:{c}22)'
    s[f'{c}24']=f'=($D$40={c}16)*-$D$39*{g(24)}';s[f'{c}25']=f'=($D$42={c}16)*-$D$41*{g(25)}';s[f'{c}26']=f'=-$D$43/12*{g(26)}'
    s[f'{c}27']=f'=-{c}23*$D$44';s[f'{c}28']=f'=SUM({c}23:{c}27)'
    if k==30:s[f'{c}29']=-64000;s[f'{c}30']=-120000
    s[f'{c}31']=f'=-$D$45/12*{g(31)}'
    s[f'{c}33']=f'=({c}15<={H})*SUM({c}28:{c}31)'
    s[f'{c}35']=f'=AND($D$32={c}16,$D$33={c}14)*SUM({L(11+k)}$28:{L(22+k)}$28)/$D$51';s[f'{c}36']=f'={c}35*$D$52';s[f'{c}37']=f'={c}35-{c}36'
    s[f'{c}39']=f'=({c}15<={H})*SUM({c}33,{c}37)'
    s[f'{c}59']=f'={c}28';s[f'{c}64']=f'={c}33';s[f'{c}65']=f'=-({c}15<={H})*$D$60';s[f'{c}66']=f'=SUM({c}64:{c}65)'
    s[f'{c}68']=f'={c}35';s[f'{c}69']=f'={c}36';s[f'{c}70']=f'=({c}68>0)*$D$62';s[f'{c}71']=f'={c}68-SUM({c}69:{c}70)';s[f'{c}73']=f'=SUM({c}66,{c}71)'
E=L(10+N)
s['J40']=f'=XIRR($J$39:${E}$39,$J$13:${E}$13)';s['J41']=f'=SUMIF($J$39:${E}$39,">0")/-SUMIF($J$39:${E}$39,"<0")';s['J42']=f'=SUM(J39:{E}39)'
s['J74']=f'=XIRR($J$73:${E}$73,$J$13:${E}$13)';s['J75']=f'=SUMIF($J$73:${E}$73,">0")/-SUMIF($J$73:${E}$73,"<0")';s['J76']=f'=SUM(J73:{E}73)'
for y in range(1,11):
    c=L(9+y);s[f'{c}78']=y
    s[f'{c}79']=f'=IF({c}78>{H},"",IFERROR(SUMIF($K$15:${E}$15,{c}78,$K$59:${E}$59)/-SUMIF($K$15:${E}$15,{c}78,$K$65:${E}$65),""))';s[f'{c}79'].number_format='0.00"x"'
for a in ('J40','J74'):s[a].number_format='0.00%'
for a in ('J41','J75'):s[a].number_format='0.00"x"'
# amortisation schedule
a=wb.create_sheet('AmortSched');a['B2']='AMORTIZATION SCHEDULE';a['C12']='Loan Amortization Schedule'
for c,h in zip('DEFGHI',['Period','Balance','Pmt','Int','Princ','End']):a[f'{c}13']=h
for i in range(360):
    r=14+i;a[f'D{r}']=i+1;a[f'E{r}']=f'={Q}$D$26*{Q}$D$56' if i==0 else f'=I{r-1}'
    a[f'F{r}']=f'=-PMT({Q}$D$57/12,{Q}$D$58*12,$E$14)';a[f'G{r}']=f'=E{r}*{Q}$D$57/12';a[f'H{r}']=f'=F{r}-G{r}';a[f'I{r}']=f'=E{r}-H{r}'
# annual returns
n=wb.create_sheet('AnnualCF & Returns');n['B2']='RETURNS AND ANNUAL CASH FLOW';n['C10']='Return Metrics & Sensitivity';n['H10']='Annual Cash Flow';n['C12']='Return Summary';n['H12']='Unlevered DCF'
sm=[(16,'Hold Period',f'={Q}D31','0'),(17,'Purchase Price',f'={Q}D26','#,##0'),(18,'Purchase Price PSF',f'={Q}D28','#,##0.00'),(19,'Going-In Cap Rate','=J25/-I16','0.00%'),
 (20,'Going-Out Cap Rate',f'={Q}D51','0.00%'),(21,'Exit Price',f'=SUM({Q}K35:{E}35)','#,##0'),(22,'Unlevered Return Summary',None,None),
 (23,'IRR',f'={Q}J40','0.00%'),(24,'Equity Multiple',f'={Q}J41','0.00"x"'),(25,'Net Gain (loss)',f'={Q}J42','#,##0'),(27,'NPV Discount Rate',0.08,'0.00%'),
 (28,'NPV','=NPV($D$27,$J$36:$S$36)+$I$36','#,##0'),(29,'Levered Return Summary',None,None),(30,'IRR',f'={Q}J74','0.00%'),(31,'Equity Multiple',f'={Q}J75','0.00"x"'),
 (32,'Equity Needed at Acquisition','=-(I46+I47)','#,##0'),(33,'Net Gain (loss)',f'={Q}J76','#,##0'),(35,'Min DSCR',f'=MIN({Q}J79:S79)','0.00"x"')]
for r,l,v,f in sm:
    n.cell(r,3,l)
    if v is not None:n.cell(r,4,v)
    if f:n.cell(r,4).number_format=f
n['H14']='Calendar Year';n['H15']='Fiscal Year';n['H16']='Purchase Price';n['I16']=f'={Q}J19'
for y in range(11):
    c=L(9+y);n[f'{c}15']=y;n[f'{c}14']=f'={Q}$D$21+{c}15'
for r,src,l in ((25,28,'NOI'),(36,39,'Unlvered Cash Flow'),(70,73,'Levered Cash Flow')):
    n[f'H{r}']=l
    for y in range(1,11):c=L(9+y);n[f'{c}{r}']=f'=SUMIF({Q}$K$15:${E}$15,{c}$15,{Q}$K{src}:${E}{src})'
n['I36']=f'={Q}J39';n['H43']='Levered DCF';n['H46']='Purchase Price';n['I46']=f'={Q}J49';n['H47']='Loan Proceeds';n['I47']=f'={Q}J50';n['I70']=f'={Q}J73'
wb.save(f'{out}/r37_acq_monthly_usd.xlsx')
