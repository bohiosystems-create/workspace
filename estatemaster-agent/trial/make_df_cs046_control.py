"""
Control workbook for the Altus sample model DF_CS046 Industrial Logistics Options.emdf (EstateMaster DF trial),
built from a screenshot of the Input sheet's financing section. Values were read from the screenshot: check each
one against the model and put Y in column H. Link column E of each row in EstateMaster, then put Y in column J.

    python trial/make_df_cs046_control.py      ->  trial/DF_CS046_control.xlsx
"""
import csv, shutil
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "trial" / "DF_CS046_control.xlsx"
shutil.copy(ROOT / "setup" / "KINAN_control_workbook_template.xlsx", OUT)
wb = load_workbook(OUT)

# Project
ps = wb["Project"]
for r, v in enumerate(["DF_CS046 Industrial Logistics Options (EstateMaster trial)", "narjis",
                       r"C:\Bohio\Models\DF_CS046 Industrial Logistics Options.emdf", "Bohio (from Input sheet screenshot)",
                       "", "", "draft"], start=2):
    ps.cell(row=r, column=2, value=v)

# Inputs: the 33 core lines were designed for the Al Narjis demo; none maps cleanly to this model -> N/A
ws = wb["Inputs"]
for r in csv.DictReader(open(ROOT / "runner" / "register_map.csv", encoding="utf-8")):
    n = int(r["cell"][1:])
    ws[f"E{n}"] = None
    ws[f"H{n}"] = "N/A"
    ws[f"I{n}"] = "Not in DF_CS046 (industrial logistics); use the Lines sheet"

# Lines: what the screenshot shows (Input sheet, financing section)
EQ, L1, L2, SL, FC = ("Finance › Equity", "Finance › Loan 1 (Construction Loan)", "Finance › Loan 2 (Investment Loan 1)",
                      "Finance › Senior Loan", "Finance › Financing costs")
LINES = [
    (EQ, "Developer's equity contribution (% of net cash flow to be funded)", "%", 30, "Equity › Percentage"),
    (EQ, "Interest charged on equity", "% p.a.", 0, "Equity › Interest Charged on Equity"),
    (EQ, "Interest received on surplus cash", "% p.a.", 0, "Equity › Interest received on Surplus Cash"),
    (EQ, "% of available funds to repay equity before debt", "%", 0, "Equity › % of Available Funds to Repay Equity"),
    (L1, "Loan 1 facility limit (% of net cash flow to be funded)", "%", 70, "Loan 1 › Facility Limit › Percentage"),
    (L1, "Loan 1 maturity month", "month", 24, "Loan 1 › Maturity Month (Manual)"),
    (L1, "Loan 1 interest rate", "% p.a.", 7.5, "Loan 1 › Interest Rate"),
    (L1, "Loan 1 application fee", "%", 0, "Loan 1 › Fees › Application Fee"),
    (L1, "Loan 1 annual line fee", "%", 0, "Loan 1 › Fees › Annual Line Fee"),
    (L1, "Loan 1 standby fee", "%", 0, "Loan 1 › Fees › Standby Fee"),
    (L1, "Loan 1 profit split to lender", "%", 0, "Loan 1 › Profit Split to Lender"),
    (L2, "Loan 2 maturity month", "month", 140, "Loan 2 › Maturity Month (Manual)"),
    (L2, "Loan 2 interest rate", "% p.a.", 7, "Loan 2 › Interest Rate"),
    (L2, "Loan 2 term of P&I loan", "months", 120, "Loan 2 › Term of P & I Loan"),
    (L2, "Loan 2 application fee (amount)", "currency", 50000, "Loan 2 › Fees › Application Fee › Amount"),
    (L2, "Loan 2 application fee: month paid", "month", 35, "Loan 2 › Fees › Application Fee › Month Paid"),
    (L2, "Loan 2 annual line fee", "%", 0, "Loan 2 › Fees › Annual Line Fee"),
    (L2, "Loan 2 standby fee", "%", 0, "Loan 2 › Fees › Standby Fee"),
    (L2, "Loan 2 profit split to lender", "%", 0, "Loan 2 › Profit Split to Lender"),
    (SL, "Senior loan interest rate", "% p.a.", 0, "Senior Loan › Interest Rate"),
    (SL, "Senior loan: maintain leverage", "%", 0, "Senior Loan › Maintain Leverage"),
    (FC, "Non Utilisation Fee: base rate", "currency", 20000, "Financing Costs › Non Utilisation Fee › Base Rate/Unit"),
    (FC, "Non Utilisation Fee: month start", "month", 68, "Financing Costs › Non Utilisation Fee › Month Start"),
    (FC, "Commitment Fee: base rate", "currency", 50000, "Financing Costs › Commitment Fee › Base Rate/Unit"),
    (FC, "Commitment Fee: month start", "month", 68, "Financing Costs › Commitment Fee › Month Start"),
    (FC, "Monitoring Surveyor Fees: base rate", "currency", 24000, "Financing Costs › Monitoring Surveyor Fees › Base Rate/Unit"),
    (FC, "Monitoring Surveyor Fees: month start", "month", 56, "Financing Costs › Monitoring Surveyor Fees › Month Start"),
    (FC, "Monitoring Surveyor Fees: month span", "months", 12, "Financing Costs › Monitoring Surveyor Fees › Month Span"),
    (FC, "Investment Loan Fee: base rate", "currency", 50000, "Financing Costs › Investment Loan Fee › Base Rate/Unit"),
    (FC, "Investment Loan Fee: month start", "month", 75, "Financing Costs › Investment Loan Fee › Month Start"),
]
ls = wb["Lines"]
for i, (sec, label, unit, val, src) in enumerate(LINES):
    n = 4 + i
    for col, v in zip("ABCDEFGHIJ", [f"L{i + 1:04d}", sec, label, unit, val, "Input sheet: " + src, "Bohio (screenshot)",
                                      None, "Read from a screenshot: check against the model, then Y in column H", "N"]):
        ls[f"{col}{n}"] = v

rm = wb["ReadMe"]
rm["A13"] = "DF_CS046 trial workbook: 30 model lines from the Input sheet (financing section), read from a screenshot."
rm["A14"] = "1) Check each value against the model and put Y in column H.  2) Link column E of the rows you want in EstateMaster, then Y in column J."
rm["A13"].font = Font(bold=True)
wb.save(OUT)
print(OUT.relative_to(ROOT), len(LINES), "model lines")
