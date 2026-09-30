"""
Presses EstateMaster's buttons the way an analyst would (Windows UI Automation via pywinauto).
Shared by em_runner.py (unattended runs) and em_ui_trial.py (the trial). NOT yet run against EstateMaster.

Sequence, from ui_steps.json:  open model -> refresh_links -> recalculate -> export_excel -> save_copy -> close
The button names in ui_steps.json are placeholders until em_ui_probe.py has recorded the real ones.
Any failure raises UIFailed with a screenshot path; nothing is guessed.
"""
import json, os, re, time
from pathlib import Path

HERE = Path(__file__).parent
TITLE_RE = os.environ.get("EM_TITLE_RE", r".*EstateMaster.*")
GROUPS = ("refresh_links", "recalculate", "export_excel", "save_copy", "close")


class UIFailed(Exception):
    def __init__(self, msg, screenshot=None):
        super().__init__(msg)
        self.screenshot = screenshot


def steps():
    return json.loads((HERE / "ui_steps.json").read_text(encoding="utf-8"))


def _pw():
    # imported lazily so the rest of the runner (and its tests) work on machines without pywinauto
    from pywinauto import Desktop, keyboard
    from pywinauto.application import Application
    return Desktop, keyboard, Application


def main_window(timeout=30):
    Desktop, _, Application = _pw()
    deadline = time.time() + timeout
    while time.time() < deadline:
        for w in Desktop(backend="uia").windows():
            if w.is_visible() and re.match(TITLE_RE, w.window_text() or ""):
                return Application(backend="uia").connect(handle=w.handle).window(handle=w.handle)
        time.sleep(1)
    raise UIFailed("EstateMaster window not found")


def _find(win, step):
    crit = {k: step[k] for k in ("title", "title_re", "control_type", "auto_id") if k in step}
    ctrl = win.child_window(**crit)
    ctrl.wait("exists enabled", timeout=step.get("timeout", 20))
    return ctrl


def _do(win, step, ctx):
    Desktop, keyboard, _ = _pw()
    kind = step["do"]
    if kind == "click":
        c = _find(win, step)
        try:
            c.invoke()
        except Exception:
            c.click_input()
    elif kind == "keys":
        win.set_focus()
        keyboard.send_keys(step["keys"])
    elif kind == "type":
        _find(win, step).set_edit_text(step["text"].format(**ctx))
    elif kind == "wait_dialog":
        Desktop(backend="uia").window(title_re=step["title_re"]).wait("visible", timeout=step.get("timeout", 30))
    elif kind == "wait_idle":
        time.sleep(step.get("seconds", 2))
    else:
        raise UIFailed(f"unknown step type {kind}")


def round_trip(emdf, export_dir, stamp, timings=None, shots_dir=None):
    """Open the model, refresh the Excel links, recalculate, export, save a copy, close.
    Returns the export files produced. Appends (step, seconds) to timings."""
    timings = timings if timings is not None else []
    export_dir, shots_dir = Path(export_dir), Path(shots_dir or export_dir)
    export_dir.mkdir(parents=True, exist_ok=True)
    cfg, ctx, win = steps(), {"export_dir": str(export_dir), "stamp": stamp}, None
    try:
        t0 = time.time()
        os.startfile(str(emdf))  # Windows only: opens the model through the .emdf file association
        win = main_window(timeout=float(os.environ.get("EM_OPEN_WAIT_S", "60")))
        timings.append(("open model", round(time.time() - t0, 1)))
        for g in GROUPS:
            t0 = time.time()
            for s in cfg[g]:
                _do(win, s, ctx)
            timings.append((g, round(time.time() - t0, 1)))
    except Exception as e:
        shot = None
        try:
            shot = shots_dir / f"failed_{stamp}.png"
            (win or main_window(5)).capture_as_image().save(shot)
        except Exception:
            shot = None
        raise UIFailed(f"{type(e).__name__}: {e}", str(shot) if shot else None) from e
    files = sorted(export_dir.glob(f"*{stamp}*.xls*"))
    if not files:
        raise UIFailed(f"EstateMaster ran but no export named *{stamp}* appeared in {export_dir}")
    return files


def latest_exports(export_dir, since):
    """Exports saved by a person (fallback mode) after `since` (epoch seconds)."""
    return sorted((f for f in Path(export_dir).glob("*.xls*") if f.stat().st_mtime >= since), key=os.path.getmtime)


def read_export(files, labels):
    """Find each output by its row label in EstateMaster's export (not by cell address).
    labels: {name: label}. Takes the last number on the matching row; confirm that rule in the trial."""
    from openpyxl import load_workbook
    want = {v.strip().lower(): k for k, v in labels.items()}
    found = {}
    for f in files:
        for ws in load_workbook(f, data_only=True, read_only=True).worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [c for c in row if c is not None]
                for c in cells:
                    key = c.strip().lower() if isinstance(c, str) else None
                    if key in want and want[key] not in found:
                        nums = [v for v in cells if isinstance(v, (int, float)) and not isinstance(v, bool)]
                        if nums:
                            found[want[key]] = nums[-1]
    return found, [k for k in labels if k not in found]
